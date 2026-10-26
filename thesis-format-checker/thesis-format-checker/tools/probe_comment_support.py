"""Probe native Word comment stories and collapsed OOXML anchors on disposable files."""
from pathlib import Path
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from thesis_format_checker.core.word_session import WordSession


def main():
    folder = Path("samples/_revision_20261012/word_probe")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "collapsed.docx"
    doc = Document()
    for text in ("", "Top text", "Bottom text"):
        doc.add_paragraph(text)
    doc.sections[0].header.paragraphs[0].text = "Header"
    doc.sections[0].footer.paragraphs[0].text = "Footer"
    for index, offset in ((0, 0), (1, 0), (2, 1)):
        p = doc.paragraphs[index]._p
        c = doc.comments.add_comment(f"point {index}", author="probe")
        start, end, run, ref = [OxmlElement(t) for t in (
            "w:commentRangeStart", "w:commentRangeEnd", "w:r", "w:commentReference")]
        for node in (start, end, ref):
            node.set(qn("w:id"), str(c.comment_id))
        run.append(ref)
        if offset:
            p.extend((start, end, run))
        else:
            for n, node in enumerate((start, end, run)):
                p.insert(n, node)
    doc.save(path)
    result = {}
    with WordSession() as session:
        wd = session.open_readonly(path.resolve())
        result["word_version"] = session.app.Version
        result["comment_count"] = wd.Comments.Count
        result["opened_comments"] = []
        for i in range(1, wd.Comments.Count + 1):
            c = wd.Comments(i)
            result["opened_comments"].append(
                {"id": i, "text": c.Range.Text, "start": c.Scope.Start,
                 "end": c.Scope.End, "scope": c.Scope.Text, "page": c.Scope.Information(3)})
        for name, story in (("header", wd.Sections(1).Headers(1)),
                            ("footer", wd.Sections(1).Footers(1))):
            try:
                c = wd.Comments.Add(Range=story.Range, Text=name)
                result[name] = {"accepted": True, "scope_story": c.Scope.StoryType}
            except Exception as exc:
                result[name] = {"accepted": False, "error": str(exc)}
    (folder / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
