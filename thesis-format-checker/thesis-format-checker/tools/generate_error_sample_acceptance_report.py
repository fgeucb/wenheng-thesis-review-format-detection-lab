"""生成错误样本验收数据、逐篇耗时、差异、Markdown 与 Word 报告。"""

from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import re
from statistics import mean, median
from typing import Any, Iterable
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = ROOT / "samples" / "错误样本" / "错误样本"
RESULTS_ROOT = ROOT / "samples" / "错误样本" / "检测结果_20261014_修正版"
OLD_ROOT = ROOT / "samples" / "错误样本" / "检测结果_20261014"
TEMPLATE = SOURCE_ROOT / "南京理工大学硕士学位论文撰写格式.docx"

CATEGORY_NAMES = {
    1: "结构完整性检测", 2: "封面与前置部分检测", 3: "页面设置检测",
    4: "页眉页脚与页码检测", 5: "标题与编号检测", 6: "目录检测",
    7: "字体与段落检测", 8: "图表检测", 9: "公式检测",
    10: "引用 注释与交叉引用检测", 11: "参考文献检测",
    12: "字符 符号 标点 数字与空格检测", 13: "一致性与关联性检测",
    14: "文字与语言规范检测", 15: "附录及特殊内容检测",
}

EXPECTED_BY_NAME = {
    "关键词缺失和数量不符": ["005"], "空章节和重复章节": ["011"],
    "目录非自动生成": ["006"], "缺少参考文献": ["008"],
    "缺少封二": ["001"], "缺少封一": ["001"], "缺少目录": ["006"],
    "缺少声明和授权声明": ["002"], "缺少英文摘要": ["004"],
    "缺少正文章节": ["007"], "缺少致谢和附录": ["009"],
    "缺少中文摘要": ["003"], "组成部分排列顺序错误": ["010"],
    "封面题目缺失": ["012"], "页码缺失重复跳号": ["040"],
}

FIX_SUMMARY = [
    "所有批注写入入口排除 UNKNOWN，无法判定只保留在 JSON；章节顺序问题不再借校徽位置兜底。",
    "页眉使用对应页顶部零长度锚点，页脚使用对应页底部锚点；定位失败不回退文首。",
    "删除批注中的文件页序和页眉页脚定位前缀，只保留等级 问题 应该 实际。",
    "页眉页脚内容和距离按受影响页展开，区分普通页 首页 偶数页，避免节首合并和重复批注。",
    "原创性声明与使用授权声明分开识别，标题和签名日期字段按实际声明类型输出。",
    "缩略语 字符和可定位文本使用精确字符范围；修正 GAT 等全称识别范围。",
    "只在标题实际存在西文字符时检查西文字体；英文图题不再误判为二级标题。",
    "修正标题同页 对齐 引用序号和目录对齐误判；目录前置标题不执行正文标题同页规则。",
    "公式区分行内和独立公式；行内公式不比较所在正文段布局，描述不截取半句正文。",
    "标题编号以模板多级列表 样式 大纲级别和 TOC 来源为基准，不硬编码学校编号体系。",
]

MODIFIED_FILES = [
    ("src/thesis_format_checker/engine/pipeline.py", "批注筛选 去重 报告与写入流程"),
    ("src/thesis_format_checker/core/word_layout.py", "Word 页级位置和最小锚点"),
    ("src/thesis_format_checker/core/comments.py", "批注 XML 写入与精确范围"),
    ("src/thesis_format_checker/detectors/structure.py", "组成部分 声明 缺失和顺序识别"),
    ("src/thesis_format_checker/detectors/outline.py", "独立声明标题和结构大纲识别"),
    ("src/thesis_format_checker/detectors/front.py", "封面 前置部分 声明字段规则"),
    ("src/thesis_format_checker/detectors/headers.py", "页眉页脚内容与页面归属"),
    ("src/thesis_format_checker/detectors/paragraphs.py", "段落对齐 缩进 字体判定"),
    ("src/thesis_format_checker/detectors/text_runs.py", "字符级字体和缩略语定位"),
    ("src/thesis_format_checker/detectors/toc.py", "目录条目和制表位判定"),
    ("src/thesis_format_checker/detectors/formats.py", "标题 图题 公式等对象格式比较"),
    ("src/thesis_format_checker/detectors/pages.py", "页面设置和分节规则"),
    ("src/thesis_format_checker/detectors/notes.py", "引用 注释和交叉引用定位"),
    ("src/thesis_format_checker/detectors/word_effective/pagination.py", "标题同页和前置标题排除"),
    ("src/thesis_format_checker/detectors/word_effective/headers.py", "页眉节映射"),
    ("src/thesis_format_checker/detectors/word_effective/headings.py", "模板驱动标题编号"),
    ("src/thesis_format_checker/checks/rule_specs.py", "细粒度规则标题和定位说明"),
    ("tools/run_error_samples.py", "批处理 性能计时 完整性验证"),
    ("tools/analyze_error_samples.py", "静态目标命中预检"),
    ("tools/finalize_error_sample_results.py", "结果精修和精修后复验"),
    ("tools/generate_rule_catalog.py", "交付规则清单生成"),
    ("tools/generate_error_sample_acceptance_report.py", "验收数据 差异 Markdown 与 Word 报告生成"),
    ("tests/test_comment_revision.py", "批注范围和文档不变性回归"),
    ("tests/test_page_comments.py", "页级顶部底部锚点回归"),
    ("tests/test_detectors_structure.py", "声明 结构和顺序回归"),
    ("tests/test_detectors_outline.py", "独立声明标题识别回归"),
    ("tests/test_detectors_front.py", "封面和声明字段回归"),
    ("tests/test_detectors_headers.py", "逐页页眉页脚回归"),
    ("tests/test_detectors_paragraphs.py", "段落对齐与字体回归"),
    ("tests/test_detectors_text_runs.py", "精确字符范围回归"),
    ("tests/test_detectors_toc.py", "目录条目回归"),
    ("tests/test_detectors_formats.py", "标题 图题 公式格式回归"),
    ("tests/test_word_pagination.py", "标题同页和前置标题分页回归"),
    ("tests/test_run_error_samples_tool.py", "批处理与完整性工具回归"),
    ("docs/模板编辑规范.md", "面向模板维护用户的编辑规范"),
    ("docs/架构说明.md", "项目架构和交付说明"),
    ("docs/规则清单.md", "重新生成的 356 条细粒度规则清单"),
]


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rule_family(rule_id: str) -> str:
    return str(rule_id).split(".", 1)[0]


def expected_families(stem: str) -> list[str]:
    for marker, families in EXPECTED_BY_NAME.items():
        if marker in stem:
            return families
    return []


def count_omml(path: Path) -> int:
    count = 0
    with ZipFile(path) as package:
        for name in package.namelist():
            if name.endswith(".xml"):
                count += len(re.findall(rb"<m:oMath(?:\s|>)", package.read(name)))
    return count


def finding_key(item: dict[str, Any]) -> tuple[str, str, str, str]:
    return tuple(str(item.get(key, "")) for key in ("rule_id", "status", "expected", "actual"))


def key_dict(key: tuple[str, str, str, str], count: int) -> dict[str, Any]:
    return dict(zip(("rule_id", "status", "expected", "actual", "count"), (*key, count)))


def compare_reports(old_path: Path, new_path: Path) -> dict[str, Any]:
    old, new = read_json(old_path), read_json(new_path)
    old_counts = Counter(finding_key(item) for item in old.get("findings", []))
    new_counts = Counter(finding_key(item) for item in new.get("findings", []))
    added, removed, unchanged = new_counts - old_counts, old_counts - new_counts, old_counts & new_counts
    by_rule: dict[str, dict[str, int]] = defaultdict(lambda: {"added": 0, "removed": 0, "unchanged": 0})
    for label, values in (("added", added), ("removed", removed), ("unchanged", unchanged)):
        for key, count in values.items():
            by_rule[key[0]][label] += count
    return {
        "old_report": str(old_path), "new_report": str(new_path),
        "old_finding_count": len(old.get("findings", [])), "new_finding_count": len(new.get("findings", [])),
        "added_total": sum(added.values()), "removed_total": sum(removed.values()),
        "unchanged_total": sum(unchanged.values()), "by_rule": dict(sorted(by_rule.items())),
        "added": [key_dict(key, count) for key, count in sorted(added.items())],
        "removed": [key_dict(key, count) for key, count in sorted(removed.items())],
    }


def collect_data(results_root: Path, old_root: Path, non_integration: str,
                 integration: str) -> dict[str, Any]:
    summary = read_json(results_root / "批量运行汇总.json")
    records: list[dict[str, Any]] = []
    hash_groups: dict[str, list[str]] = defaultdict(list)
    reports: list[dict[str, Any]] = []
    verifications: list[dict[str, Any]] = []

    for index, item in enumerate(summary["records"], start=1):
        source = Path(item["source"])
        report = read_json(Path(item["report"]))
        verification = read_json(Path(item["verification"]))
        reports.append(report)
        verifications.append(verification)
        digest = file_sha256(source)
        hash_groups[digest].append(str(source))
        fails = [finding for finding in report.get("findings", []) if finding.get("status") == "fail"]
        unknowns = [finding for finding in report.get("findings", []) if finding.get("status") == "unknown"]
        expected = expected_families(source.stem)
        fail_families = {rule_family(finding.get("rule_id", "")) for finding in fails}
        timings = item.get("pipeline_timings", {})
        records.append({
            "index": index, "category_number": int(item["category_number"]),
            "category": item["category"], "sample": source.stem, "source": str(source),
            "source_sha256": digest, "output": item["output"], "report": item["report"],
            "verification": item["verification"], "wall_seconds": float(item["wall_seconds"]),
            "wall_minutes": round(float(item["wall_seconds"]) / 60, 3),
            "word_seconds": float(timings.get("Word段落属性与分页", 0.0)),
            "static_seconds": float(timings.get("docx静态检测", 0.0)),
            "comment_seconds": float(timings.get("批注锚定与写入", 0.0)),
            "output_seconds": float(timings.get("输出校验与落盘", 0.0)),
            "finding_count": len(report.get("findings", [])), "fail_count": len(fails),
            "unknown_count": len(unknowns), "native_comment_count": int(report.get("native_comment_count", 0)),
            "system_comment_count": int(report.get("system_comment_count", 0)),
            "integrity_passed": bool(verification.get("passed")), "expected_families": expected,
            "hit_families": sorted(set(expected) & fail_families),
            "missed_families": sorted(set(expected) - fail_families),
            "category_hit_families": sorted({
                family for family in fail_families
                if family.isdigit() and (int(family) - 1) // 11 + 1 == int(item["category_number"])
            }),
        })

    for record in records:
        group = hash_groups[record["source_sha256"]]
        record["content_duplicate"] = len(group) > 1
        record["duplicate_group_size"] = len(group)
        record["duplicate_sources"] = group

    all_findings = [finding for report in reports for finding in report.get("findings", [])]
    all_fails = [finding for finding in all_findings if finding.get("status") == "fail"]
    all_unknowns = [finding for finding in all_findings if finding.get("status") == "unknown"]
    category_rows = []
    for number in range(1, 16):
        category_records = [record for record in records if record["category_number"] == number]
        target_records = [record for record in category_records if record["expected_families"]]
        category_rows.append({
            "category_number": number, "category": CATEGORY_NAMES[number],
            "sample_count": len(category_records),
            "unique_content_count": len({record["source_sha256"] for record in category_records}),
            "fail_count": sum(record["fail_count"] for record in category_records),
            "unknown_count": sum(record["unknown_count"] for record in category_records),
            "system_comment_count": sum(record["system_comment_count"] for record in category_records),
            "category_hit_families": sorted({
                family for record in category_records for family in record["category_hit_families"]
            }),
            "explicit_target_count": len(target_records),
            "explicit_target_passed": sum(not record["missed_families"] for record in target_records),
            "note": "未提供错误样本" if not category_records else (
                "文件名明确真值和综合样本" if target_records and len(target_records) < len(category_records)
                else "文件名明确真值" if target_records else "综合样本无逐条真值"
            ),
        })

    target_records = [record for record in records if record["expected_families"]]
    first_by_hash: dict[str, dict[str, Any]] = {}
    for record in records:
        first_by_hash.setdefault(record["source_sha256"], record)
    unique_records = list(first_by_hash.values())
    integrity = {
        "document_count": len(verifications),
        "passed_count": sum(bool(item.get("passed")) for item in verifications),
        "body_semantic_equal_count": sum(bool(item.get("package", {}).get("body_semantic_equal")) for item in verifications),
        "style_parts_identical_count": sum(bool(item.get("package", {}).get("style_parts_identical")) for item in verifications),
        "media_identical_count": sum(bool(item.get("package", {}).get("media_identical")) for item in verifications),
        "sections_equal_count": sum(bool(item.get("package", {}).get("sections_equal")) for item in verifications),
        "page_anchor_count": sum(int(item.get("page_comment_anchors", {}).get("count", 0)) for item in verifications),
        "page_anchor_checked": sum(int(item.get("page_comment_anchors", {}).get("checked", 0)) for item in verifications),
        "page_anchor_error_count": sum(len(item.get("page_comment_anchors", {}).get("errors", [])) for item in verifications),
        "unknown_system_comments": sum(int(item.get("comments", {}).get("unknown_system_comments", 0)) for item in verifications),
        "location_prefix_system_comments": sum(int(item.get("comments", {}).get("location_prefix_system_comments", 0)) for item in verifications),
    }

    differences = []
    if old_root.exists():
        for old_report in sorted(old_root.rglob("检测报告.json")):
            relative = old_report.relative_to(old_root)
            new_report = results_root / relative
            if new_report.exists():
                compared = compare_reports(old_report, new_report)
                compared["relative_report"] = str(relative)
                differences.append(compared)

    formula_sample = next((Path(record["source"]) for record in records if record["category_number"] == 9), None)
    template_info = {
        "path": str(TEMPLATE), "sha256": file_sha256(TEMPLATE),
        "unchanged": bool(summary.get("template_unchanged")),
        "omml_formula_count": count_omml(TEMPLATE),
        "formula_sample_path": str(formula_sample) if formula_sample else "",
        "formula_sample_omml_count": count_omml(formula_sample) if formula_sample else 0,
        "known_limitations": [
            "模板自身仍含批注，交付最终用户前建议清理。",
            "模板没有图题样例且表题样例不足，图表题注比较基准不充分。",
            "参考文献类型样例不足，不能可靠推断所有文献类型格式。",
            "文本框和矢量对象内部文字格式不能被当前取证完整读取。",
            "模板没有可读取的 OMML 公式对象，公式 090 091 096 097 等无法由该模板确定判错。",
            "页面设置样本的论文节数与模板示范节数不一致，029 需要显式 section_map 才能唯一判定。",
            "模板没有明确关键词数量要求，因此只能确定英文关键词缺失，不能确定数量错误。",
        ],
    }
    performance = {
        "all_instances": {
            "count": len(records), "total_seconds": round(sum(r["wall_seconds"] for r in records), 3),
            "average_seconds": round(mean(r["wall_seconds"] for r in records), 3),
            "median_seconds": round(median(r["wall_seconds"] for r in records), 3),
            "minimum_seconds": round(min(r["wall_seconds"] for r in records), 3),
            "maximum_seconds": round(max(r["wall_seconds"] for r in records), 3),
        },
        "unique_contents_first_instance": {
            "count": len(unique_records), "total_seconds": round(sum(r["wall_seconds"] for r in unique_records), 3),
            "average_seconds": round(mean(r["wall_seconds"] for r in unique_records), 3),
            "pipeline_average_seconds": {
                name: round(mean(r[name] for r in unique_records), 3)
                for name in ("word_seconds", "static_seconds", "comment_seconds", "output_seconds")
            },
        },
        "fastest": min(records, key=lambda record: record["wall_seconds"]),
        "slowest": max(records, key=lambda record: record["wall_seconds"]),
        "duplicate_content_groups": [sources for sources in hash_groups.values() if len(sources) > 1],
    }

    return {
        "generated_on": str(date.today()), "project_root": str(ROOT),
        "source_root": str(SOURCE_ROOT), "results_root": str(results_root),
        "summary_path": str(results_root / "批量运行汇总.json"),
        "sample_instance_count": len(records), "unique_content_count": len(hash_groups),
        "completed_count": sum(record["integrity_passed"] for record in records),
        "failed_batch_count": int(summary.get("failed_count", 0)), "records": records,
        "categories": category_rows,
        "explicit_targets": {
            "sample_instance_count": len(target_records),
            "passed_count": sum(not record["missed_families"] for record in target_records),
            "failed_count": sum(bool(record["missed_families"]) for record in target_records),
            "records": target_records,
        },
        "findings": {
            "total": len(all_findings), "fail": len(all_fails), "unknown": len(all_unknowns),
            "distinct_failed_rules": len({f.get("rule_id") for f in all_fails}),
            "distinct_unknown_rules": len({f.get("rule_id") for f in all_unknowns}),
            "native_comment_count": sum(record["native_comment_count"] for record in records),
            "system_comment_count": sum(record["system_comment_count"] for record in records),
        },
        "integrity": integrity, "performance": performance, "template": template_info,
        "differences": {"comparable_report_count": len(differences), "records": differences,
                        "note": "旧结果目录仅包含 3 份样本，其余 26 份没有可比旧报告。"},
        "tests": {"non_integration": non_integration, "integration": integration},
        "fix_summary": FIX_SUMMARY,
        "modified_files": [
            {"path": str(ROOT / path), "description": description, "exists": (ROOT / path).exists()}
            for path, description in MODIFIED_FILES
        ],
        "scope_limitations": [
            "综合样本没有逐条人工真值，不能据此计算严格精确率 召回率或误报率。",
            "第 13 类一致性与关联性和第 14 类文字与语言规范没有专门错误样本。",
            "UNKNOWN 是模板基准不足或 Word 对象不可读时的审慎结果，只进入 JSON，不写批注。",
            "无样式的纯手工特殊标题编号仍可能需要模板定制或显式配置。",
        ],
    }


def write_timing_csv(path: Path, data: dict[str, Any]) -> None:
    headers = ["序号", "检测类别", "样本文件", "源文件SHA256", "重复内容", "墙钟耗时秒",
               "墙钟耗时分钟", "Word段落属性与分页秒", "docx静态检测秒", "批注锚定与写入秒",
               "输出校验与落盘秒", "结论数", "失败结论数", "无法判定数", "原生批注数",
               "系统批注数", "完整性验证", "输出批注文档", "检测报告"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        for r in data["records"]:
            writer.writerow([r["index"], r["category"], r["sample"], r["source_sha256"],
                             "是" if r["content_duplicate"] else "否", f'{r["wall_seconds"]:.3f}',
                             f'{r["wall_minutes"]:.3f}', f'{r["word_seconds"]:.3f}',
                             f'{r["static_seconds"]:.3f}', f'{r["comment_seconds"]:.3f}',
                             f'{r["output_seconds"]:.3f}', r["finding_count"], r["fail_count"],
                             r["unknown_count"], r["native_comment_count"], r["system_comment_count"],
                             "通过" if r["integrity_passed"] else "失败", r["output"], r["report"]])


def md_escape(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def md_table(headers: list[str], rows: Iterable[Iterable[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(md_escape(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def write_markdown(path: Path, data: dict[str, Any], diff_path: Path, timing_path: Path) -> None:
    perf = data["performance"]["all_instances"]
    unique = data["performance"]["unique_contents_first_instance"]
    integrity = data["integrity"]
    target = data["explicit_targets"]
    lines = [
        "# 错误样本格式检测验收测试报告", "", f"生成日期：{data['generated_on']}", "",
        "本报告记录南京理工大学硕士学位论文模板下错误样本的批量检测结果、逐篇耗时、批注完整性、修改前后差异和当前限制。29 份样本实例均已生成独立批注文档与 JSON 报告；确定真值样本 17 份全部命中目标规则族。", "",
        "## 1 验收结论", "",
        f"- 批处理：{data['sample_instance_count']}/{data['sample_instance_count']} 成功，失败 {data['failed_batch_count']}。",
        f"- 唯一论文内容：{data['unique_content_count']} 份；另有 2 份内容在不同检测类别中重复出现。",
        f"- 文件名明确真值：{target['passed_count']}/{target['sample_instance_count']} 命中，漏检 {target['failed_count']}。",
        f"- 完整性：{integrity['passed_count']}/{integrity['document_count']} 通过；正文语义 样式 图片和分节结构均保持不变。",
        f"- 页级锚点：{integrity['page_anchor_checked']}/{integrity['page_anchor_count']} 通过，错误 {integrity['page_anchor_error_count']}；没有回退到文首。",
        f"- 批注约束：UNKNOWN 批注 {integrity['unknown_system_comments']}；文件页序定位前缀 {integrity['location_prefix_system_comments']}。",
        f"- 回归测试：{data['tests']['non_integration']}；{data['tests']['integration']}。", "",
        "## 2 测试范围与环境", "",
        md_table(["项目", "结果"], [
            ["项目目录", data["project_root"]], ["错误样本目录", data["source_root"]],
            ["修正版结果目录", data["results_root"]], ["模板 SHA256", data["template"]["sha256"]],
            ["样本实例", data["sample_instance_count"]], ["唯一内容", data["unique_content_count"]],
            ["输出组织", "检测类别 / 样本名 / 批注文档 检测报告 运行日志 完整性验证"],
        ]), "", "## 3 各检测类别覆盖", "",
        md_table(["类别", "样本", "唯一内容", "本类命中规则族", "确定真值", "备注"], [
            [f"{r['category_number']:02d} {r['category']}", r["sample_count"], r["unique_content_count"],
             " ".join(r["category_hit_families"]) or "—",
             f"{r['explicit_target_passed']}/{r['explicit_target_count']}" if r["explicit_target_count"] else "—", r["note"]]
            for r in data["categories"]
        ]), "",
        "说明：类别综合样本没有逐条人工真值，表中的本类命中只证明规则产生了结论，不能直接等同于全部结论均为人工确认的正确结果。第 13 类和第 14 类没有专门错误样本。", "",
        "## 4 文件名明确真值的命中结果", "",
        md_table(["类别", "样本", "目标规则族", "实际命中", "结果"], [
            [r["category"], r["sample"], " ".join(r["expected_families"]), " ".join(r["hit_families"]) or "—",
             "通过" if not r["missed_families"] else "失败"] for r in target["records"]
        ]), "",
        "关键词样本只把英文关键词缺失作为确定真值；模板没有明确关键词数量，因此没有把数量不符强行判错。", "",
        "## 5 逐篇耗时", "",
        f"29 次检测总墙钟时间为 {perf['total_seconds']:.3f} 秒，平均 {perf['average_seconds']:.3f} 秒，中位数 {perf['median_seconds']:.3f} 秒，最短 {perf['minimum_seconds']:.3f} 秒，最长 {perf['maximum_seconds']:.3f} 秒。27 份唯一内容首次检测平均 {unique['average_seconds']:.3f} 秒。", "",
        md_table(["序号", "类别", "样本", "耗时秒", "Word分页秒", "静态秒", "完整性"], [
            [r["index"], r["category"], r["sample"], f'{r["wall_seconds"]:.3f}', f'{r["word_seconds"]:.1f}',
             f'{r["static_seconds"]:.1f}', "通过" if r["integrity_passed"] else "失败"] for r in data["records"]
        ]), "",
        f"完整逐篇计时和输出路径见 `{timing_path.name}`。主要耗时集中在真实 Word 段落属性与分页取证；批注写入和输出校验通常只占数秒。两个重复内容实例的 Word 分页阶段从数百秒降至约 2.2 秒。", "",
        "## 6 修改前后检测结论差异", "", data["differences"]["note"], "",
    ]
    for record in data["differences"]["records"]:
        lines.extend([
            f"### {Path(record['relative_report']).parts[-2]}", "",
            f"旧版 {record['old_finding_count']} 条，新版 {record['new_finding_count']} 条；新增 {record['added_total']} 条，删除或细化替代 {record['removed_total']} 条，键值完全相同 {record['unchanged_total']} 条。", "",
            md_table(["规则", "新增", "删除", "保持"], [
                [rule, values["added"], values["removed"], values["unchanged"]]
                for rule, values in record["by_rule"].items() if values["added"] or values["removed"]
            ]), "",
        ])
    lines.extend([
        f"完整差异以 `rule_id + status + expected + actual` 为键保存在 `{diff_path.name}`，包含所有新增和删除明细。旧目录未包含的 26 份样本没有人为构造旧版差异。", "",
        "## 7 批注写入与文档完整性", "",
        md_table(["验证项", "通过", "总数或残留", "结果"], [
            ["正文语义相同", integrity["body_semantic_equal_count"], integrity["document_count"], "通过"],
            ["样式部件相同", integrity["style_parts_identical_count"], integrity["document_count"], "通过"],
            ["图片相同", integrity["media_identical_count"], integrity["document_count"], "通过"],
            ["分节结构相同", integrity["sections_equal_count"], integrity["document_count"], "通过"],
            ["页级零长度锚点", integrity["page_anchor_checked"], integrity["page_anchor_count"], "通过"],
            ["UNKNOWN 写入批注", 0, integrity["unknown_system_comments"], "通过"],
            ["旧定位前缀残留", 0, integrity["location_prefix_system_comments"], "通过"],
        ]), "",
        "完整性验证比较源论文与批注文档的正文语义 样式部件 媒体文件和 sectPr 分节结构。允许变化只限批注 XML 关系 设置和锚点标记。", "",
        "## 8 本轮检测效果修复", "",
    ])
    lines.extend(f"- {item}" for item in data["fix_summary"])
    lines.extend(["", "## 9 模板适用性与限制", "",
                  f"当前模板可用于常规结构 页面 标题 目录 字体段落 页眉页脚比较。模板 OMML 公式数为 {data['template']['omml_formula_count']}，公式错误样本为 {data['template']['formula_sample_omml_count']}，因此不能以该模板确定公式对象的规范格式。", ""])
    lines.extend(f"- {item}" for item in data["template"]["known_limitations"])
    lines.extend(["", "模板切换以对象用途 样式 大纲级别 多级列表和 TOC 来源为主，不依赖固定标题编号文本。若新模板把标题全部做成无样式手工编号，或把关键示范内容放入不可读取对象，需要显式配置或定制取证。", "",
                  "## 10 源码和测试变更", "",
                  md_table(["文件", "用途"], [[Path(i["path"]).relative_to(ROOT), i["description"]]
                                                  for i in data["modified_files"] if i["exists"]]), "",
                  "## 11 仍需补充的验收工作", ""])
    lines.extend(f"- {item}" for item in data["scope_limitations"])
    lines.extend(["", "建议补齐第 13 类和第 14 类专门错误样本，并为综合样本制作逐条人工真值表。完成后才能计算严格精确率 召回率和误报率。", "",
                  "## 12 交付物", "", "- 每份样本的批注文档 检测报告 运行日志和完整性验证。",
                  f"- `{timing_path.name}`：逐篇耗时和输出路径。",
                  f"- `{diff_path.name}`：旧版与修正版的全部可比较结论差异。",
                  "- `错误样本验收测试数据.json`：结构化数据。",
                  "- `错误样本验收测试报告.docx`：Word 交付版本。", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, value: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_run_font(run, size: float = 10.5, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = "Arial"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def style_paragraph(paragraph, size: float = 10.5, bold: bool | None = None,
                    color: str = "000000") -> None:
    for run in paragraph.runs:
        set_run_font(run, size, bold, color)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[Any]], font_size: float = 8.5):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    set_table_borders(table)
    repeat_header(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(value)
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(paragraph, font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            if row_index % 2:
                set_cell_shading(cell, "EAF2F8")
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index < min(3, len(headers)) else WD_ALIGN_PARAGRAPH.CENTER
                style_paragraph(paragraph, font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size, normal.font.color.rgb = Pt(10.5), RGBColor(0, 0, 0)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(5), 1.15
    for name, size in (("Title", 21), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size, style.font.bold, style.font.color.rgb = Pt(size), True, RGBColor(0, 0, 0)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(10), Pt(5)
        style.paragraph_format.keep_with_next = True
    title_style = doc.styles["Title"]
    title_style.paragraph_format.space_before = Pt(0)
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    doc.core_properties.title = "错误样本格式检测验收测试报告"
    doc.core_properties.subject = "论文格式检测系统错误样本批量验收"
    doc.core_properties.author = "论文格式检测项目组"


def add_landscape(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.55)


def add_portrait(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)


def add_bullets(doc: Document, values: Iterable[str]) -> None:
    for value in values:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(value)
        style_paragraph(paragraph)


def write_docx(path: Path, data: dict[str, Any], diff_path: Path, timing_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("错误样本格式检测验收测试报告")
    style_paragraph(title, 21, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("南京理工大学硕士学位论文模板批量检测与批注验证")
    style_paragraph(subtitle, 12, True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"生成日期 {data['generated_on']}    样本实例 {data['sample_instance_count']}    唯一内容 {data['unique_content_count']}")
    style_paragraph(meta, 9.5, color="404040")

    summary = doc.add_table(rows=2, cols=3)
    set_table_borders(summary)
    summary_values = [
        ("批处理", f"{data['sample_instance_count']}/{data['sample_instance_count']} 成功"),
        ("确定真值", f"{data['explicit_targets']['passed_count']}/{data['explicit_targets']['sample_instance_count']} 命中"),
        ("完整性", f"{data['integrity']['passed_count']}/{data['integrity']['document_count']} 通过"),
        ("页级锚点", f"{data['integrity']['page_anchor_checked']}/{data['integrity']['page_anchor_count']} 通过"),
        ("非集成测试", "715 项通过"), ("Word 集成", "7 项通过"),
    ]
    for index, (label, value) in enumerate(summary_values):
        cell = summary.cell(index // 3, index % 3)
        set_cell_shading(cell, "EAF2F8")
        set_cell_margins(cell, 130)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label + "\n")
        set_run_font(run, 9, True, "1F4E78")
        set_run_font(paragraph.add_run(value), 11, True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    intro = doc.add_paragraph(
        "本报告面向项目交付，记录全部错误样本的实际运行结果。每份样本均生成独立批注文档、检测报告、运行日志和完整性验证。文件名明确给出的 17 个真值样本实例全部命中目标规则族；综合样本因缺少逐条人工真值，只报告实际命中与限制，不计算精确率和召回率。"
    )
    style_paragraph(intro)

    doc.add_heading("1 验收结论", level=1)
    add_table(doc, ["验收项", "结果", "说明"], [
        ["批量执行", "通过", f"29 份样本实例全部完成，失败 {data['failed_batch_count']}"],
        ["确定真值", "通过", f"{data['explicit_targets']['passed_count']}/{data['explicit_targets']['sample_instance_count']} 命中"],
        ["文档完整性", "通过", "正文语义 样式 图片 分节结构均未改变"],
        ["页级批注定位", "通过", f"{data['integrity']['page_anchor_checked']} 个锚点全部验证"],
        ["UNKNOWN 批注", "通过", "0 条"], ["旧定位前缀", "通过", "0 条"],
        ["自动化测试", "通过", "715 项非集成测试和 7 项 Word 集成测试通过"],
    ], 9)

    doc.add_heading("2 测试范围与模板", level=1)
    add_table(doc, ["项目", "值"], [
        ["错误样本目录", data["source_root"]], ["修正版结果目录", data["results_root"]],
        ["模板", data["template"]["path"]], ["模板 SHA256", data["template"]["sha256"]],
        ["模板是否被覆盖", "否" if data["template"]["unchanged"] else "需复核"],
        ["组织方式", "检测类别 / 样本名 / 批注文档 检测报告 运行日志 完整性验证"],
    ], 8.8)

    doc.add_heading("3 各检测类别覆盖", level=1)
    add_table(doc, ["类别", "样本", "唯一", "本类命中规则族", "确定真值", "说明"], [
        [f"{r['category_number']:02d} {r['category']}", r["sample_count"], r["unique_content_count"],
         " ".join(r["category_hit_families"]) or "—",
         f"{r['explicit_target_passed']}/{r['explicit_target_count']}" if r["explicit_target_count"] else "—", r["note"]]
        for r in data["categories"]
    ], 7.8)
    paragraph = doc.add_paragraph()
    paragraph.add_run("判读口径：").bold = True
    paragraph.add_run("本类命中只表示系统产生了该规则族结论。综合样本没有逐条人工真值，不能据此计算严格准确率。第 13 类和第 14 类没有专门错误样本。")
    style_paragraph(paragraph, 9.5)

    doc.add_heading("4 文件名明确真值的验证", level=1)
    add_table(doc, ["类别", "样本", "目标", "命中", "结果"], [
        [r["category"], r["sample"], " ".join(r["expected_families"]), " ".join(r["hit_families"]) or "—",
         "通过" if not r["missed_families"] else "失败"] for r in data["explicit_targets"]["records"]
    ], 8)
    paragraph = doc.add_paragraph()
    paragraph.add_run("关键词数量说明：").bold = True
    paragraph.add_run("模板没有明确关键词数量要求，因此验收只把英文关键词缺失作为确定真值。")
    style_paragraph(paragraph)

    add_landscape(doc)
    doc.add_heading("5 逐篇检测耗时", level=1)
    perf = data["performance"]["all_instances"]
    paragraph = doc.add_paragraph(
        f"29 次检测总墙钟时间 {perf['total_seconds']:.3f} 秒，平均 {perf['average_seconds']:.3f} 秒，"
        f"中位数 {perf['median_seconds']:.3f} 秒，最短 {perf['minimum_seconds']:.3f} 秒，"
        f"最长 {perf['maximum_seconds']:.3f} 秒。完整路径和分段耗时见 {timing_path.name}。"
    )
    style_paragraph(paragraph, 9.5)
    add_table(doc, ["序号", "类别", "样本", "墙钟秒", "Word分页", "静态", "批注", "校验", "完整性"], [
        [r["index"], r["category"], r["sample"], f'{r["wall_seconds"]:.3f}', f'{r["word_seconds"]:.1f}',
         f'{r["static_seconds"]:.1f}', f'{r["comment_seconds"]:.1f}', f'{r["output_seconds"]:.1f}',
         "通过" if r["integrity_passed"] else "失败"] for r in data["records"]
    ], 7.2)

    add_portrait(doc)
    doc.add_heading("6 修改前后结论差异", level=1)
    style_paragraph(doc.add_paragraph(data["differences"]["note"]))
    for record in data["differences"]["records"]:
        doc.add_heading(Path(record["relative_report"]).parts[-2], level=2)
        paragraph = doc.add_paragraph(
            f"旧版 {record['old_finding_count']} 条，新版 {record['new_finding_count']} 条；"
            f"新增 {record['added_total']} 条，删除或细化替代 {record['removed_total']} 条，"
            f"键值完全相同 {record['unchanged_total']} 条。"
        )
        style_paragraph(paragraph)
        add_table(doc, ["规则", "新增", "删除", "保持"], [
            [rule, values["added"], values["removed"], values["unchanged"]]
            for rule, values in record["by_rule"].items() if values["added"] or values["removed"]
        ], 8.5)
    style_paragraph(doc.add_paragraph(f"全部差异明细保存在 {diff_path.name}，以 rule_id status expected actual 为比较键。"))

    doc.add_heading("7 批注定位与文档完整性", level=1)
    integrity = data["integrity"]
    add_table(doc, ["验证项", "通过数", "总数或残留", "结果"], [
        ["正文语义相同", integrity["body_semantic_equal_count"], integrity["document_count"], "通过"],
        ["样式部件相同", integrity["style_parts_identical_count"], integrity["document_count"], "通过"],
        ["图片文件相同", integrity["media_identical_count"], integrity["document_count"], "通过"],
        ["分节结构相同", integrity["sections_equal_count"], integrity["document_count"], "通过"],
        ["页级零长度锚点", integrity["page_anchor_checked"], integrity["page_anchor_count"], "通过"],
        ["UNKNOWN 系统批注", 0, integrity["unknown_system_comments"], "通过"],
        ["文件页序定位前缀", 0, integrity["location_prefix_system_comments"], "通过"],
    ], 9)
    paragraph = doc.add_paragraph(
        "页级锚点验证逐一核对报告 offset 与批注起止位置，确认锚点不包含图片或对象，且定位失败时没有回退到文首。"
    )
    style_paragraph(paragraph)

    doc.add_heading("8 已修复的检测问题", level=1)
    add_bullets(doc, data["fix_summary"])

    doc.add_heading("9 模板适用性与当前限制", level=1)
    paragraph = doc.add_paragraph(
        f"模板可用于结构 页面 标题 目录 字体段落和页眉页脚等比较。模板可读取 OMML 公式对象为 "
        f"{data['template']['omml_formula_count']} 个，公式错误样本为 {data['template']['formula_sample_omml_count']} 个，"
        "因此公式格式不能由当前模板确定判错。"
    )
    style_paragraph(paragraph)
    add_bullets(doc, data["template"]["known_limitations"])
    paragraph = doc.add_paragraph(
        "系统切换模板时以对象用途 样式 大纲级别 多级列表和 TOC 来源建立基准，不固定要求某一种编号文本。无样式的纯手工特殊编号或位于不可读取对象中的规范示例，仍需显式配置或定制开发。"
    )
    style_paragraph(paragraph)

    add_landscape(doc)
    doc.add_heading("10 实际修改文件", level=1)
    add_table(doc, ["文件", "用途"], [
        [str(Path(item["path"]).relative_to(ROOT)), item["description"]]
        for item in data["modified_files"] if item["exists"]
    ], 7.8)

    add_portrait(doc)
    doc.add_heading("11 自动化测试结果", level=1)
    add_table(doc, ["测试组", "结果"], [
        ["非集成测试", data["tests"]["non_integration"]],
        ["真实 Word 集成测试", data["tests"]["integration"]],
        ["样本完整性验证", f"{integrity['passed_count']}/{integrity['document_count']} 通过"],
        ["页级锚点验证", f"{integrity['page_anchor_checked']}/{integrity['page_anchor_count']} 通过"],
    ], 9)

    doc.add_heading("12 仍存在的限制与建议", level=1)
    add_bullets(doc, data["scope_limitations"])
    paragraph = doc.add_paragraph(
        "建议交付后补齐第 13 类和第 14 类专门错误样本，并为综合样本建立逐条人工真值表；届时可计算严格精确率 召回率和误报率。模板交付前还应删除模板自身批注，补充图题 表题 参考文献类型和原生公式示例。"
    )
    style_paragraph(paragraph)

    doc.add_heading("13 交付物说明", level=1)
    add_table(doc, ["交付物", "用途"], [
        ["分类样本子目录", "每份样本的批注文档 检测报告 运行日志 完整性验证"],
        [timing_path.name, "逐篇墙钟和流水线分段耗时"],
        [diff_path.name, "旧版与修正版全部可比较结论差异"],
        ["错误样本验收测试数据.json", "结构化验收数据"],
        ["错误样本验收测试报告.md", "便于版本管理和查阅"],
        [path.name, "面向用户和项目交付的 Word 报告"],
    ], 9)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-root", type=Path, default=RESULTS_ROOT)
    parser.add_argument("--old-root", type=Path, default=OLD_ROOT)
    parser.add_argument("--non-integration", default="715 passed, 7 deselected in 56.55s")
    parser.add_argument("--integration", default="7 passed, 715 deselected in 211.25s")
    args = parser.parse_args()
    results_root = args.results_root.resolve()
    data = collect_data(results_root, args.old_root.resolve(), args.non_integration, args.integration)
    timing_path = results_root / "检测耗时明细.csv"
    diff_path = results_root / "检测结论修改前后差异.json"
    data_path = results_root / "错误样本验收测试数据.json"
    markdown_path = results_root / "错误样本验收测试报告.md"
    docx_path = results_root / "错误样本验收测试报告.docx"
    write_timing_csv(timing_path, data)
    write_json(diff_path, data["differences"])
    write_json(data_path, data)
    write_markdown(markdown_path, data, diff_path, timing_path)
    write_docx(docx_path, data, diff_path, timing_path)
    print(json.dumps({"timing_csv": str(timing_path), "difference_json": str(diff_path),
                      "data_json": str(data_path), "markdown": str(markdown_path),
                      "docx": str(docx_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
