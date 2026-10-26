r"""批量运行错误样本，逐文件保存批注文档、JSON 报告、日志、耗时与完整性验证。

示例::

    .venv\Scripts\python.exe tools\run_error_samples.py
    .venv\Scripts\python.exe tools\run_error_samples.py --categories 1 2 3
    .venv\Scripts\python.exe tools\run_error_samples.py --overwrite
"""

from __future__ import annotations

import argparse
from hashlib import sha256
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from thesis_format_checker.core.point_comments import content_tree
from thesis_format_checker.core.point_comments import marker_offset
from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.xmlutil import qn

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "samples" / "错误样本" / "错误样本"
DEFAULT_OUTPUT = ROOT / "samples" / "错误样本" / "检测结果_20261014_修正版"
TEMPLATE_NAME = "南京理工大学硕士学位论文撰写格式.docx"

_ALLOWED_CHANGED_PARTS = {
    "[Content_Types].xml",
    "word/_rels/document.xml.rels",
    "word/comments.xml",
    "word/document.xml",
    # --show-comments 只会改变批注显示设置，不改变论文排版。
    "word/settings.xml",
}
_COMMENTS_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)
_COMMENTS_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _matching_parts(package: ZipFile, predicate) -> dict[str, bytes]:
    return {name: package.read(name) for name in package.namelist() if predicate(name)}


def _verify_page_comment_anchors(output_doc, report: dict | None) -> dict:
    result = {"count": 0, "checked": 0, "errors": [], "passed": True}
    if not report:
        return result
    rows = report.get("page_comments", [])
    result["count"] = len(rows)
    if not rows:
        return result
    ctx = DocxContext(output_doc)
    comments = {str(item.comment_id): item.text for item in output_doc.comments}
    markers = {
        tag: {node.get(qn("w:id")): node for node in output_doc.element.body.iter(qn("w:" + tag))}
        for tag in ("commentRangeStart", "commentRangeEnd", "commentReference")
    }
    for row in rows:
        identity = str(row.get("comment_id"))
        try:
            paragraph_index = int(row["paragraph"]) - 1
            paragraph = ctx.paragraphs[paragraph_index]
            start = markers["commentRangeStart"][identity]
            end = markers["commentRangeEnd"][identity]
            actual_start, actual_end = marker_offset(paragraph._p, start), marker_offset(paragraph._p, end)
            expected = int(row["offset"])
            if actual_start != expected or actual_end != expected:
                raise ValueError(
                    f"零长度位置不一致：expected={expected}, start={actual_start}, end={actual_end}"
                )
            if paragraph._p.xpath(".//w:drawing|.//w:pict|.//w:object"):
                raise ValueError("页级批注锚点段落包含图片或嵌入对象")
            text = comments[identity]
            if "无法判定" in text:
                raise ValueError("页级批注包含只应保留在报告中的无法判定信息")
            result["checked"] += 1
        except Exception as exc:
            result["errors"].append({"comment_id": identity, "error": str(exc), "row": row})
    result["passed"] = result["checked"] == result["count"] and not result["errors"]
    return result


def verify_output_integrity(source: Path, output: Path, *, source_sha256: str | None = None,
                            report: dict | None = None) -> dict:
    """验证批注写入只改变批注相关 OOXML，不改变论文内容与版式。

    该检查不依赖 PDF 渲染，因为批注在无头渲染中通常不可见；批注本身通过
    comments.xml、关系、内容类型和正文中的三个锚点标记集合交叉核对。
    """

    result = {
        "source_sha256_before": source_sha256,
        "source_sha256_after": _digest(source),
        "output_sha256": _digest(output),
        "python_docx_reopen": False,
        "package": {},
        "comments": {},
        "errors": [],
        "passed": False,
    }
    result["source_unchanged"] = (
        source_sha256 is None or source_sha256 == result["source_sha256_after"]
    )
    try:
        source_doc = Document(source)
        output_doc = Document(output)
        result["python_docx_reopen"] = True

        with ZipFile(source) as before, ZipFile(output) as after:
            before_names, after_names = set(before.namelist()), set(after.namelist())
            common = before_names & after_names
            changed = {name for name in common if before.read(name) != after.read(name)}
            added = after_names - before_names
            removed = before_names - after_names
            touched = changed | added | removed
            package = result["package"]
            package["changed_parts"] = sorted(changed)
            package["added_parts"] = sorted(added)
            package["removed_parts"] = sorted(removed)
            package["unexpected_changes"] = sorted(touched - _ALLOWED_CHANGED_PARTS)
            package["body_semantic_equal"] = (
                content_tree(etree.fromstring(before.read("word/document.xml")))
                == content_tree(etree.fromstring(after.read("word/document.xml")))
            )
            package["style_parts_identical"] = (
                _matching_parts(before, lambda name: name.startswith("word/styles"))
                == _matching_parts(after, lambda name: name.startswith("word/styles"))
            )
            before_media = _matching_parts(before, lambda name: name.startswith("word/media/"))
            after_media = _matching_parts(after, lambda name: name.startswith("word/media/"))
            package["media_count_before"] = len(before_media)
            package["media_count_after"] = len(after_media)
            package["media_identical"] = before_media == after_media

            comments_xml = after.read("word/comments.xml") if "word/comments.xml" in after_names else None
            document_xml = etree.fromstring(after.read("word/document.xml"))
            rels_xml = etree.fromstring(after.read("word/_rels/document.xml.rels"))
            types_xml = etree.fromstring(after.read("[Content_Types].xml"))

        package["sections_equal"] = (
            [etree.tostring(section._sectPr, method="c14n") for section in source_doc.sections]
            == [etree.tostring(section._sectPr, method="c14n") for section in output_doc.sections]
        )
        package["section_count_before"] = len(source_doc.sections)
        package["section_count_after"] = len(output_doc.sections)

        comment_rows = list(output_doc.comments)
        system_comments = [item for item in comment_rows if item.author == "论文格式检测"]
        comment_ids = set()
        if comments_xml is not None:
            comments_root = etree.fromstring(comments_xml)
            comment_ids = {
                node.get(qn("w:id")) for node in comments_root.findall(qn("w:comment"))
            }
        marker_ids = {
            tag: {node.get(qn("w:id")) for node in document_xml.iter(qn("w:" + tag))}
            for tag in ("commentRangeStart", "commentRangeEnd", "commentReference")
        }
        rel_types = {
            node.get("Type") for node in rels_xml
            if node.get("Target", "").replace("\\", "/").endswith("comments.xml")
        }
        comment_overrides = {
            node.get("ContentType") for node in types_xml
            if node.get("PartName") == "/word/comments.xml"
        }
        has_comments = bool(comment_rows)
        comments = result["comments"]
        comments["count"] = len(comment_rows)
        comments["system_count"] = len(system_comments)
        comments["comments_xml_present"] = comments_xml is not None
        comments["relationship_valid"] = (
            not has_comments or _COMMENTS_RELATIONSHIP in rel_types
        )
        comments["content_type_valid"] = (
            not has_comments or _COMMENTS_CONTENT_TYPE in comment_overrides
        )
        comments["anchors_complete"] = (
            all(ids == comment_ids for ids in marker_ids.values())
            if has_comments else all(not ids for ids in marker_ids.values())
        )
        comments["unknown_system_comments"] = sum(
            "无法判定" in item.text for item in system_comments
        )
        prefix = re.compile(r"^\s*第\s*\d+\s*页(?:（文件页序）)?\s*｜")
        comments["location_prefix_system_comments"] = sum(
            bool(prefix.match(item.text)) for item in system_comments
        )
        result["page_comment_anchors"] = _verify_page_comment_anchors(output_doc, report)

        result["passed"] = all((
            result["source_unchanged"],
            result["python_docx_reopen"],
            not package["unexpected_changes"],
            package["body_semantic_equal"],
            package["style_parts_identical"],
            package["media_identical"],
            package["sections_equal"],
            comments["relationship_valid"],
            comments["content_type_valid"],
            comments["anchors_complete"],
            comments["unknown_system_comments"] == 0,
            comments["location_prefix_system_comments"] == 0,
            result["page_comment_anchors"]["passed"],
        ))
    except Exception as exc:  # 交付批处理必须保留失败证据，不能让异常吞掉汇总。
        result["errors"].append(f"{type(exc).__name__}: {exc}")
    return result


def _number(name: str) -> int:
    try:
        return int(name.split(".", 1)[0])
    except (ValueError, IndexError):
        return 999


def _samples(source: Path, categories: set[int] | None):
    for category in sorted((p for p in source.iterdir() if p.is_dir()),
                           key=lambda p: (_number(p.name), p.name)):
        number = _number(category.name)
        if categories and number not in categories:
            continue
        for path in sorted(category.glob("*.doc*"), key=lambda p: p.name):
            if path.suffix.casefold() in {".doc", ".docx"}:
                yield number, category.name, path


def _existing_record(number: int, category: str, source: Path, target: Path):
    report = target / "检测报告.json"
    output = target / "批注文档.docx"
    if not report.is_file() or not output.is_file():
        return None
    data = json.loads(report.read_text(encoding="utf-8"))
    verification_path = target / "完整性验证.json"
    verification = (
        json.loads(verification_path.read_text(encoding="utf-8"))
        if verification_path.is_file() else verify_output_integrity(source, output, report=data)
    )
    if not verification_path.is_file():
        verification_path.write_text(
            json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    return {
        "category_number": number, "category": category,
        "source": str(source), "output": str(output), "report": str(report),
        "status": "completed" if verification.get("passed", False) else "failed",
        "exit_code": 0,
        "wall_seconds": data.get("batch_wall_seconds"),
        "pipeline_timings": data.get("timings", {}),
        "finding_count": len(data.get("findings", [])),
        "fail_count": sum(x.get("status") == "fail" for x in data.get("findings", [])),
        "comment_count": data.get("native_comment_count"),
        "system_comment_count": data.get("system_comment_count"),
        "verification": str(verification_path),
        "integrity_passed": verification.get("passed", False),
        "cached": True,
    }


def run(args) -> int:
    source = args.source.resolve()
    output_root = args.output.resolve()
    template = source / TEMPLATE_NAME
    if not template.is_file():
        raise SystemExit(f"模板不存在：{template}")
    output_root.mkdir(parents=True, exist_ok=True)
    records = []
    samples = list(_samples(source, set(args.categories) if args.categories else None))
    batch_started = datetime.now()
    template_sha256 = _digest(template)
    print(f"模板：{template}\n样本：{len(samples)} 份\n输出：{output_root}", flush=True)
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    for index, (number, category, document) in enumerate(samples, 1):
        target = output_root / category / document.stem
        target.mkdir(parents=True, exist_ok=True)
        output = target / "批注文档.docx"
        report = target / "检测报告.json"
        log = target / "运行日志.txt"
        verification_path = target / "完整性验证.json"
        if not args.overwrite:
            existing = _existing_record(number, category, document, target)
            if existing:
                records.append(existing)
                print(f"[{index}/{len(samples)}] 已存在，跳过：{category} / {document.name}", flush=True)
                continue
        command = [sys.executable, "-m", "thesis_format_checker", "compare",
                   "--template", str(template), "--document", str(document),
                   "--output", str(output), "--report-json", str(report),
                   "--show-comments"]
        if args.overwrite:
            command.append("--overwrite")
        print(f"[{index}/{len(samples)}] 开始：{category} / {document.name}", flush=True)
        source_sha256 = _digest(document)
        started = time.perf_counter()
        proc = subprocess.run(command, cwd=ROOT, env=env, text=True,
                              stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                              encoding="utf-8", errors="replace")
        elapsed = round(time.perf_counter() - started, 3)
        log.write_text("命令：" + subprocess.list2cmdline(command) + "\n\n" + proc.stdout,
                       encoding="utf-8")
        data = {}
        if report.is_file():
            data = json.loads(report.read_text(encoding="utf-8"))
            data["batch_wall_seconds"] = elapsed
            report.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        comment_count = None
        system_comment_count = None
        verification = {"passed": False, "errors": ["检测命令未成功生成批注文档"]}
        if output.is_file():
            comments = list(Document(output).comments)
            comment_count = len(comments)
            system_comment_count = sum(comment.author == "论文格式检测" for comment in comments)
            verification = verify_output_integrity(
                document, output, source_sha256=source_sha256, report=data
            )
        verification_path.write_text(
            json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        data["native_comment_count"] = comment_count
        data["system_comment_count"] = system_comment_count
        data["integrity_verification"] = verification
        if report.is_file():
            report.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        completed = proc.returncode == 0 and verification.get("passed", False)
        record = {
            "category_number": number, "category": category,
            "source": str(document), "output": str(output), "report": str(report),
            "log": str(log), "verification": str(verification_path),
            "status": "completed" if completed else "failed",
            "exit_code": proc.returncode, "wall_seconds": elapsed,
            "pipeline_timings": data.get("timings", {}),
            "finding_count": len(data.get("findings", [])),
            "fail_count": sum(x.get("status") == "fail" for x in data.get("findings", [])),
            "comment_count": comment_count, "system_comment_count": system_comment_count,
            "integrity_passed": verification.get("passed", False),
            "cached": False,
        }
        records.append(record)
        (output_root / "批量运行汇总.json").write_text(
            json.dumps({"generated_at": datetime.now().isoformat(timespec="seconds"),
                        "started_at": batch_started.isoformat(timespec="seconds"),
                        "template": str(template), "template_sha256": template_sha256,
                        "records": records},
                       ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"  完成：exit={proc.returncode}，耗时={elapsed:.1f}s，"
              f"问题={record['fail_count']}，批注={record['comment_count']}，"
              f"完整性={'通过' if record['integrity_passed'] else '失败'}", flush=True)
    failed = sum(item["status"] != "completed" for item in records)
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "started_at": batch_started.isoformat(timespec="seconds"),
        "template": str(template),
        "template_sha256": template_sha256,
        "template_unchanged": _digest(template) == template_sha256,
        "sample_count": len(samples),
        "completed_count": len(records) - failed,
        "failed_count": failed,
        "total_wall_seconds": round(sum(item.get("wall_seconds") or 0 for item in records), 3),
        "records": records,
    }
    (output_root / "批量运行汇总.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"批量运行结束：完成 {len(records) - failed}，失败 {failed}")
    return 1 if failed else 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--categories", nargs="*", type=int)
    parser.add_argument("--overwrite", action="store_true")
    return run(parser.parse_args(argv))


if __name__ == "__main__":
    raise SystemExit(main())
