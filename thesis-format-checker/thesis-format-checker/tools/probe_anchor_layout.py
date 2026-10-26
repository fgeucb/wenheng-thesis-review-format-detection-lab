"""Probe whether comment-reference runs affect display math and manual-page-break spacing."""
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from thesis_format_checker.core.docx_context import DocxContext, make_finding
from thesis_format_checker.core.comments import add_word_comment
from thesis_format_checker.core.point_comments import add_boundary_comment
from thesis_format_checker.core.xmlutil import qn
from thesis_format_checker.core.word_session import WordSession

folder = Path("samples/_revision_20261012/anchor_layout_probe")
folder.mkdir(exist_ok=True)
source = Document("samples/paper.docx")
ctx = DocxContext(source)
elements = [deepcopy(ctx.paragraphs[i]._p) for i in (51, 52, 53, 262, 263, 264, 265)]
names = ("original", "normal", "hidden", "ordinary", "zero_before", "zero_after", "math_only")
for name in names:
    d = Document("samples/paper.docx")
    for child in list(d.element.body):
        if child.tag != qn("w:sectPr"):
            d.element.body.remove(child)
    for i, p in enumerate(elements):
        d.element.body.insert(i, deepcopy(p))
    if name in ("normal", "hidden", "ordinary", "math_only"):
        for p in d.paragraphs:
            if name == "math_only" and not p._p.xpath(".//m:oMathPara"):
                continue
            add_word_comment(d, make_finding("139", "格式", "一致", "不一致"), p)
    if name in ("normal", "hidden", "zero_before"):
        add_boundary_comment(d, d.paragraphs[0], 0, "【警告】页脚")
    if name in ("normal", "hidden", "zero_after"):
        add_boundary_comment(d, d.paragraphs[0], 1, "【警告】页眉")
    if name == "hidden":
        for ref in d.element.body.iter(qn("w:commentReference")):
            run = ref.getparent()
            run.get_or_add_rPr().append(OxmlElement("w:vanish"))
    d.save(folder / (name + ".docx"))
with WordSession() as session:
    for name in names:
        wd = session.open_readonly((folder / (name + ".docx")).resolve())
        session.resume_pagination()
        wd.Windows(1).View.ShowRevisionsAndComments = False
        wd.ExportAsFixedFormat(str((folder / (name + ".pdf")).resolve()), 17, Item=0)
        print(name, "pages", wd.ComputeStatistics(2), "comments", wd.Comments.Count, flush=True)
        session.close_document(wd, save=False)
    wd = session.open_readonly((folder / "original.docx").resolve())
    for i in range(1, wd.OMaths.Count + 1):
        wd.Comments.Add(Range=wd.OMaths(i).Range, Text="native math")
    wd.SaveAs2(str((folder / "native.docx").resolve()), FileFormat=16)
    wd.ExportAsFixedFormat(str((folder / "native.pdf").resolve()), 17, Item=0)
    print("native math", wd.Comments.Count, flush=True)
