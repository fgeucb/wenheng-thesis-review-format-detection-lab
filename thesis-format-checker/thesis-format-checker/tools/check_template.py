"""模板体检：按《模板编辑规范》检查一份模板是否满足格式检测的最低要求。

用法（项目根目录）::

    .\\.venv\\Scripts\\python.exe tools\\check_template.py samples\\template.docx

输出三类结论：

* ``[通过]``  —— 该项已满足；
* ``[建议]``  —— 不补齐也能跑，但会降低检测准确率或产生“无法判定”；
* ``[必须补齐]``—— 缺了它对应规则无法工作。

退出码：存在 ``必须补齐`` 项时为 1，便于自动化流水线拦截不合格模板。
"""

from __future__ import annotations

import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402  （需先设置 sys.path）
from docx.oxml.ns import qn  # noqa: E402

from thesis_format_checker.core.docx_context import DocxContext  # noqa: E402
from thesis_format_checker.detectors import bibliography as bib  # noqa: E402
from thesis_format_checker.detectors.figures import collect_caption_records  # noqa: E402
from thesis_format_checker.detectors.headings import _heading_level  # noqa: E402
from thesis_format_checker.detectors.outline import extract_outline_baseline  # noqa: E402
from thesis_format_checker.detectors.paragraphs import _body_records  # noqa: E402

PASS = "通过"
WARN = "建议"
FAIL = "必须补齐"


@dataclass(frozen=True)
class Item:
    """一条体检结论。"""

    status: str
    name: str
    detail: str


def _xml(zf: zipfile.ZipFile, name: str) -> str:
    return zf.read(name).decode("utf-8", "ignore") if name in zf.namelist() else ""


def _heading_stats(ctx) -> dict[int, Counter]:
    """各级标题的样式使用情况：``{级别: {样式名: 段数}}``。"""

    stats: dict[int, Counter] = {}
    for paragraph in ctx.paragraphs:
        try:
            level = _heading_level(paragraph)
        except Exception:
            level = None
        if not level:
            continue
        stats.setdefault(level, Counter())[paragraph.style.name or "未命名样式"] += 1
    return stats


def inspect(path: Path) -> list[Item]:
    """体检一份模板，返回结论列表。"""

    items: list[Item] = []
    try:
        document = Document(str(path))
        ctx = DocxContext(document)
    except Exception as exc:
        return [Item(FAIL, "文件可读性", f"无法打开：{type(exc).__name__}；模板必须是可正常打开的 .docx")]

    with zipfile.ZipFile(str(path)) as zf:
        names = zf.namelist()
        document_xml = _xml(zf, "word/document.xml")
        footer_xml = "".join(_xml(zf, name) for name in names if name.startswith("word/footer"))
        has_revision = any(node in document_xml for node in ("<w:ins ", "<w:del ", "<w:pPrChange"))
        has_comment = any(name.startswith("word/comments") for name in names)

    items.append(Item(PASS, "文件格式", f".docx，可正常打开（{path.stat().st_size // 1024} KB）"))
    items.append(Item(WARN if has_revision else PASS, "修订与批注",
                      "存在未接受的修订：模板必须处于“已接受全部修订”状态" if has_revision
                      else ("存在批注，建议删除后再作为模板" if has_comment else "无修订、无批注")))

    # 标题层级与样式
    stats = _heading_stats(ctx)
    if not stats:
        items.append(Item(FAIL, "标题样式", "未识别到任何标题：请使用 Word 内置“标题 1/2/3”样式，"
                                            "否则章节目录、编号、字体比对都无法工作"))
    for level in (1, 2, 3):
        styles = stats.get(level)
        if not styles:
            items.append(Item(WARN if level == 3 else FAIL, f"{level} 级标题样例",
                              f"没有 {level} 级标题样例：该级规则只能输出“无法判定”"))
            continue
        total = sum(styles.values())
        detail = "、".join(f"{name} {count} 段" for name, count in styles.most_common())
        if len(styles) > 1:
            items.append(Item(WARN, f"{level} 级标题样例",
                              f"同一层级使用了多种样式：{detail}；请统一为一种样式"))
        elif total < 2:
            items.append(Item(WARN, f"{level} 级标题样例", f"只有 1 段样例（{detail}）；建议 ≥2 段"))
        else:
            items.append(Item(PASS, f"{level} 级标题样例", f"{total} 段：{detail}"))

    # 正文样例
    body = _body_records(ctx, None)
    if len(body) < 3:
        items.append(Item(WARN, "正文样例", f"只识别到 {len(body)} 段正文；建议每类正文（普通段、"
                                            "列表段、引文段）至少放 2～3 段真实样例"))
    else:
        styles = Counter(record["role"] for record in body)
        items.append(Item(PASS, "正文样例", f"{len(body)} 段：" + "、".join(f"{k} {v}" for k, v in styles.most_common(4))))

    # 题注
    captions = collect_caption_records(ctx)
    kinds = Counter(item["kind"] for item in captions)
    caption_detail = f"图题 {kinds.get('figure', 0)} 段、表题 {kinds.get('table', 0)} 段"
    if kinds.get("figure", 0) < 1 or kinds.get("table", 0) < 1:
        items.append(Item(WARN, "图表题注样例", caption_detail + "；请各留至少 2 段并统一使用“题注”样式"))
    else:
        items.append(Item(PASS, "图表题注样例", caption_detail))

    # 目录域
    if "TOC" in document_xml:
        items.append(Item(PASS, "自动目录", "存在 TOC 域（请保持“随正文更新”后再交付）"))
    else:
        items.append(Item(WARN, "自动目录", "未发现 TOC 域；请用“引用→目录→自动目录”插入，不要手工排版目录"))

    # 参考文献样例
    try:
        baseline = bib.extract_bibliography_baseline(ctx).get("bibliography", {})
        paragraphs, _anchor = bib._bibliography_paragraphs(ctx)
        parsed = [p for p in paragraphs if bib._parse_entry(p)]
        groups = baseline.get("groups") or {}
        issues = [str(issue) for issue in (baseline.get("issues") or [])]
        if len(parsed) < 3 or issues:
            detail = f"可解析条目 {len(parsed)} 条（样例组 {len(groups)} 个）"
            if issues:
                detail += "；" + "；".join(issues[:2])
            items.append(Item(WARN, "参考文献样例",
                              detail + "；请放 3 条以上完整真实条目，含 [M]/[J] 等类型标识，中英文各 ≥1 条"))
        else:
            items.append(Item(PASS, "参考文献样例", f"可解析条目 {len(parsed)} 条，样例组 {len(groups)} 个"))
    except Exception as exc:
        items.append(Item(WARN, "参考文献样例", f"读取失败：{type(exc).__name__}；请检查“参考文献”标题与条目样式"))

    # 页眉页脚与页码
    sections = list(document.sections)
    headers = [p.text.strip() for section in sections for p in section.header.paragraphs if p.text.strip()]
    if not headers:
        items.append(Item(WARN, "页眉页脚", "所有节都没有页眉文字；正文页眉规则只能输出“无法判定”"))
    else:
        distinct = len(set(headers))
        items.append(Item(PASS, "页眉页脚", f"{len(sections)} 节、{len(set(headers))} 种页眉文字"))
        if distinct > 1:
            items.append(Item(WARN, "页眉一致性", f"存在 {distinct} 种页眉文字；若确实按区域不同，请在 mapping.yaml 配置 section_map"))
        if "PAGE" not in footer_xml:
            items.append(Item(WARN, "页码域", "页脚未发现 PAGE 域；请在页脚插入“页码”域，不要手工输入数字"))
        else:
            items.append(Item(PASS, "页码域", "页脚存在 PAGE 域"))

    # 不可读对象
    body_element = document.element.body
    textboxes = [node for node in body_element.xpath(".//w:txbxContent")
                 if "".join(t.text or "" for t in node.iter(qn("w:t"))).strip()]
    objects = len(body_element.xpath(".//w:object"))
    picts = len(body_element.xpath(".//w:pict"))
    if textboxes or objects or picts:
        items.append(Item(WARN, "可读性",
                          f"含文字的文本框 {len(textboxes)} 个、OLE 对象 {objects} 个、矢量图形 {picts} 个："
                          "文本框/OLE/矢量图形内部的格式无法检测；示范段落请写进正文，说明文字可用批注"))
    else:
        items.append(Item(PASS, "可读性", "未使用文本框/OLE/VML 等不可读对象"))

    # 直接格式化比例
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    # 只有真正指定了字体名才算“直接格式化”；仅带 w:hint 的 rFonts 不影响字体选择。
    names = tuple("w:" + name for name in ("ascii", "hAnsi", "eastAsia", "cs",
                                            "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"))
    direct = [
        run for run in runs
        if run._r.rPr is not None
        and (fonts := run._r.rPr.find(qn("w:rFonts"))) is not None
        and any(fonts.get(qn(name)) for name in names)
    ]
    if runs:
        share = len(direct) / len(runs)
        items.append(Item(WARN if share > 0.5 else PASS, "样式与直接格式化",
                          f"{share:.0%} 的文字片段直接设置了字体（{len(direct)}/{len(runs)}）；"
                          "正文请尽量用样式，直接格式化只用于封面字段"))

    # 结构组成
    try:
        components = extract_outline_baseline(ctx).get("outline", {}).get("components", [])
        labels = "、".join(" ".join(str(item["label"]).split()) for item in components[:14])
        if len(components) < 8:
            items.append(Item(WARN, "必需组成部分", f"只识别到 {len(components)} 个组成部分：{labels}；"
                                                   "模板里每个必需部分都要出现一次，系统才能要求论文包含它"))
        else:
            items.append(Item(PASS, "必需组成部分", f"{len(components)} 个：{labels}"))
    except Exception as exc:
        items.append(Item(WARN, "必需组成部分", f"提取失败：{type(exc).__name__}"))

    return items


def render(items: list[Item]) -> str:
    """把结论渲染成可读报告。"""

    lines = []
    for item in items:
        lines.append(f"[{item.status}] {item.name}：{item.detail}")
    failed = sum(1 for item in items if item.status == FAIL)
    warned = sum(1 for item in items if item.status == WARN)
    lines.append("")
    if failed:
        lines.append(f"结论：不可直接用于检测——{failed} 项必须补齐、{warned} 项建议。")
    elif warned:
        lines.append(f"结论：可用于检测，但有 {warned} 项建议会影响准确率。")
    else:
        lines.append("结论：模板符合《模板编辑规范》，可直接用于检测。")
    return "\n".join(lines)


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__)
        return 2
    path = Path(argv[0])
    if not path.exists():
        print(f"模板不存在：{path}")
        return 2
    items = inspect(path)
    print(f"模板体检报告：{path}")
    print("-" * 60)
    print(render(items))
    return 1 if any(item.status == FAIL for item in items) else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
