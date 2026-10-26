from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.headers import (
    adapt_legacy_header_findings, check_header_footer_errors, extract_headers_baseline,
)
from thesis_format_checker.core.models import Finding, Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def clone(doc):
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return Document(stream)


def field(p, instruction="PAGE", cache="1", size=11):
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Times New Roman")
    props.append(font)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(size * 2))
    props.append(size_node)
    run.append(props)
    if cache is not None:
        text = OxmlElement("w:t")
        text.text = cache
        run.append(text)
    element.append(run)
    p._p.append(element)
    return element


def number_settings(section, fmt=None, start=None):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    if fmt is not None:
        node.set(qn("w:fmt"), fmt)
    if start is not None:
        node.set(qn("w:start"), str(start))


def header_text(section, value, kind="header"):
    p = getattr(section, kind).paragraphs[0]
    p.text = value
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def border(p, value="single"):
    props = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), value)
    bottom.set(qn("w:sz"), "8")
    borders.append(bottom)
    props.append(borders)


def check(template, document, mapping=None):
    before = tuple(part.blob for d in (template, document) for part in d.part.package.parts)
    result = check_header_footer_errors(extract_headers_baseline(DocxContext(template)), DocxContext(document), mapping)
    after = tuple(part.blob for d in (template, document) for part in d.part.package.parts)
    assert before == after
    return result


def failed(findings, rule):
    return [f for f in findings if family_of(f.rule_id) == rule and f.status == Status.FAIL]


def limited(findings, rule):
    # 模板基准不唯一时现在是“不适用”（不写批注），测试中与“无法判定”一并接受。
    return [f for f in findings if family_of(f.rule_id) == rule
            and f.status in (Status.UNKNOWN, Status.NOT_APPLICABLE) and UNREADABLE in f.message]


def template(role="摘要"):
    doc = Document()
    doc.add_paragraph(role)
    doc.add_paragraph("段落内容")
    return doc


def test_missing_extra_header_and_native_body_anchor():
    a = template()
    header_text(a.sections[0], "模板页眉")
    b = clone(a)
    b.sections[0].header.paragraphs[0].text = ""
    results = check(a, b)
    assert failed(results, "034")
    assert all(f.anchor._p in list(b.element.body.iter(qn("w:p"))) for f in results)
    assert failed(check(b, a), "034")
    assert not any(family_of(f.rule_id) == "035" for f in results)


def test_disabled_first_even_stories_do_not_trigger_header_or_page_errors():
    a = template()
    b = clone(a)
    b.sections[0].first_page_header.paragraphs[0].text = "隐藏首页页眉"
    field(b.sections[0].even_page_footer.paragraphs[0])
    assert not any(f.status == Status.FAIL for f in check(a, b))


def test_header_format_reads_style_and_direct_format_and_border():
    a = template()
    p = header_text(a.sections[0], "学校 header")
    border(p)
    b = clone(a)
    p = b.sections[0].header.paragraphs[0]
    p.runs[0].font.size = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).set(qn("w:val"), "double")
    matches = failed(check(a, b), "036")
    # 字号（036.3）与下边框（036.5）各成一条，这里合起来看
    assert matches
    messages = "；".join(item.message for item in matches)
    assert "18" in messages and "下边框" in messages


def test_even_header_content_and_options():
    a = template()
    a.settings.odd_and_even_pages_header_footer = True
    header_text(a.sections[0], "奇数固定页眉")
    a.sections[0].even_page_header.paragraphs[0].text = "偶数固定页眉"
    b = clone(a)
    b.sections[0].even_page_header.paragraphs[0].text = "错误的偶数文字"
    assert failed(check(a, b), "037")
    b.settings.odd_and_even_pages_header_footer = False
    assert failed(check(a, b), "037")


def two_sections():
    a = template()
    header_text(a.sections[0], "前置页眉")
    a.add_section(WD_SECTION_START.NEW_PAGE)
    a.add_heading("正文内容", 1)
    a.sections[1].header.is_linked_to_previous = False
    header_text(a.sections[1], "正文页眉")
    return a


def test_linked_header_actual_inheritance_is_checked_without_creating_parts():
    a = two_sections()
    b = clone(a)
    b.sections[1].header.is_linked_to_previous = True
    result = check(a, b)
    assert failed(result, "038")
    assert "前置页眉" in failed(result, "038")[0].message


def test_disconnected_identical_content_is_not_inherently_an_error():
    a = two_sections()
    a.sections[1].header.is_linked_to_previous = True
    b = clone(a)
    b.sections[1].header.is_linked_to_previous = False
    header_text(b.sections[1], "前置页眉")
    assert not failed(check(a, b), "038")


def test_footer_extra_text_border_and_cache_exclusion():
    a = template()
    p = a.sections[0].footer.paragraphs[0]
    field(p, "PAGE", "1")
    b = clone(a)
    bp = b.sections[0].footer.paragraphs[0]
    bp._p.find(qn("w:fldSimple")).find(".//" + qn("w:t")).text = "666"
    assert not failed(check(a, b), "039")
    bp.add_run("多余的文字")
    border(bp)
    assert failed(check(a, b), "039")


def test_missing_and_duplicate_page_but_numpages_is_not_duplicate():
    a = template()
    field(a.sections[0].footer.paragraphs[0])
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]
    field(p, "NUMPAGES", "80")
    assert not failed(check(a, b), "040")
    field(p)
    assert failed(check(a, b), "040")
    p.clear()
    assert failed(check(a, b), "040")
    p.add_run("1")
    assert not failed(check(a, b), "040")
    assert limited(check(a, b), "040")


def test_number_format_and_start_are_template_driven_not_school_constants():
    a = template()
    field(a.sections[0].footer.paragraphs[0], "PAGE", "vii")
    number_settings(a.sections[0], "lowerRoman", 7)
    b = clone(a)
    number_settings(b.sections[0], "decimal", 3)
    results = check(a, b)
    assert failed(results, "041") and failed(results, "042")
    assert "7" in failed(results, "042")[0].message
    number_settings(b.sections[0], "lowerRoman", 7)
    assert not failed(check(a, b), "041") and not failed(check(a, b), "042")


def test_page_field_switch_overrides_section_format_and_case():
    a = template()
    field(a.sections[0].footer.paragraphs[0], "PAGE \\* ROMAN", "I")
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]._p.find(qn("w:fldSimple"))
    p.set(qn("w:instr"), "PAGE \\* roman")
    assert failed(check(a, b), "041")
    p.set(qn("w:instr"), "PAGE \\# 00")
    assert limited(check(a, b), "041")


def test_page_font_position_and_no_cache_limits():
    a = template()
    p = a.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p)
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p._p.find(".//" + qn("w:sz")).set(qn("w:val"), "40")
    assert failed(check(a, b), "043")
    p.clear()
    field(p, cache=None)
    assert limited(check(a, b), "043")


def test_page_in_prohibited_front_matter_is_config_error_not_year():
    a = template()
    b = clone(a)
    header_text(b.sections[0], "2026年", "footer")
    assert not failed(check(a, b), "044")
    b.sections[0].footer.paragraphs[0].clear()
    field(b.sections[0].footer.paragraphs[0])
    assert failed(check(a, b), "044")


def test_layout_is_unknown_not_success_and_current_page_cache_is_not_evaluated():
    a = template()
    field(a.sections[0].footer.paragraphs[0], cache="888")
    b = clone(a)
    fields = b.sections[0].footer.paragraphs[0]._p
    fields.find(".//" + qn("w:t")).text = "-999"
    result = check(a, b)
    assert limited(result, "040")
    assert not failed(result, "040")


def test_malformed_reference_does_not_raise():
    a = template()
    header_text(a.sections[0], "header")
    b = clone(a)
    b.sections[0]._sectPr.find(qn("w:headerReference")).set(qn("r:id"), "missing")
    result = check(a, b)
    assert limited(result, "034") and limited(result, "038")


def test_malformed_field_does_not_raise_or_claim_page_pass():
    a = template()
    field(a.sections[0].footer.paragraphs[0])
    b = clone(a)
    run = OxmlElement("w:r")
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), "begin")
    run.append(node)
    b.sections[0].footer.paragraphs[0]._p.append(run)
    result = check(a, b)
    assert limited(result, "039") and limited(result, "040")


def test_nested_conditional_page_not_claimed_visible():
    a = template()
    field(a.sections[0].footer.paragraphs[0])
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]
    p.clear()
    outer = field(p, 'IF 1 = 1 "" ""', "")
    child = OxmlElement("w:fldSimple")
    child.set(qn("w:instr"), "PAGE")
    outer.append(child)
    assert limited(check(a, b), "040")


def test_ambiguous_multisection_templates_need_mapping():
    a = two_sections()
    b = template("无法分类的区域")
    header_text(b.sections[0], "另一个页眉")
    result = check(a, b)
    assert limited(result, "034") and not failed(result, "034")
    assert not limited(check(a, b, {"section_map": {"1": 1}}), "034")


def test_repeated_body_section_does_not_inherit_template_restart_requirement():
    a = template("正文区域")
    a.add_heading("第一章", 1)
    field(a.sections[0].footer.paragraphs[0])
    b = clone(a)
    b.add_section(WD_SECTION_START.NEW_PAGE)
    b.add_heading("第二章", 1)
    assert not failed(check(a, b), "042")
    assert limited(check(a, b), "042")


def test_read_only_linked_missing_stories_do_not_create_parts():
    a = template()
    a.add_section(WD_SECTION_START.NEW_PAGE)
    a.add_heading("正文", 1)
    b = clone(a)
    assert not any(f.status == Status.FAIL for f in check(a, b))


def test_legacy_dynamic_content_is_downgraded_static_preserved():
    a = template()
    field(a.sections[0].header.paragraphs[0], "STYLEREF 1", "模板标题")
    b = clone(a)
    old = Finding("035", "页眉文字", Status.FAIL, "原有页眉证据", section_index=1, header_kind=1)
    other = Finding("025", "边距", Status.FAIL, "边距证据")
    baseline = extract_headers_baseline(DocxContext(a))
    result = adapt_legacy_header_findings([old, other], baseline, DocxContext(b))
    # 模板页眉含动态域 → 没有可比基准，现在是“不适用”（不再写无法判定批注）。
    assert result[0].status in (Status.UNKNOWN, Status.NOT_APPLICABLE)
    assert "原有页眉证据" in result[0].message
    assert result[1] is other and old.status == Status.FAIL
    assert any(family_of(f.rule_id) == "035" for f in adapt_legacy_header_findings([], baseline, DocxContext(b)))
    a, b = template(), template()
    header_text(a.sections[0], "学校固定文字")
    header_text(b.sections[0], "其他固定文字")
    assert adapt_legacy_header_findings([old], extract_headers_baseline(DocxContext(a)), DocxContext(b))[0] is old


def test_legacy_fixed_header_vs_dynamic_field_is_a_definite_mismatch():
    a = template()
    header_text(a.sections[0], "南京农业大学本科毕业论文（设计）")
    b = template()
    field(b.sections[0].header.paragraphs[0], "STYLEREF 1", "第一章 绪论")
    old = Finding("035", "页眉文字", Status.FAIL, "原有页眉证据",
                  section_index=1, header_kind=1)

    result = adapt_legacy_header_findings(
        [old], extract_headers_baseline(DocxContext(a)), DocxContext(b))

    assert result[0] is old


def test_footer_top_border_is_detected():
    a = template()
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]
    border(p)
    p._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).tag = qn("w:top")
    assert failed(check(a, b), "039")


def test_opaque_footer_text_is_unknown_not_claimed_extra():
    a = template()
    b = clone(a)
    root = b.sections[0].footer.paragraphs[0]._p
    box = OxmlElement("w:txbxContent")
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "文本框里的页脚"
    r.append(t)
    p.append(r)
    box.append(p)
    root.append(box)
    result = check(a, b)
    assert limited(result, "039") and not failed(result, "039")


def test_hidden_page_field_is_not_counted_as_definite_display():
    a = template()
    b = clone(a)
    node = field(b.sections[0].footer.paragraphs[0])
    hidden = OxmlElement("w:vanish")
    node.find(".//" + qn("w:rPr")).append(hidden)
    result = check(a, b)
    assert limited(result, "040") and limited(result, "044")
    assert not failed(result, "044")


def test_plain_roman_number_is_unknown_not_missing():
    a = template()
    field(a.sections[0].footer.paragraphs[0], "PAGE", "iv")
    b = clone(a)
    b.sections[0].footer.paragraphs[0].text = "iv"
    result = check(a, b)
    assert limited(result, "040") and not failed(result, "040")


def test_body_page_field_is_located_and_missing_shared_page_not_guessed():
    a = template()
    field(a.sections[0].footer.paragraphs[0])
    b = clone(a)
    b.sections[0].footer.paragraphs[0].clear()
    field(b.paragraphs[1])
    result = check(a, b)
    assert not failed(result, "040") and limited(result, "040") and limited(result, "043")
    a.sections[0].footer.paragraphs[0].clear()
    assert failed(check(a, b), "044")


def test_first_blank_page_with_default_number_is_config_not_physical_page_claim():
    a = template()
    a.sections[0].different_first_page_header_footer = True
    b = clone(a)
    field(b.sections[0].footer.paragraphs[0])
    result = check(a, b)
    findings = failed(result, "044")
    assert findings and all("普通/奇数页" in f.message and "尚未确认" in f.message for f in findings)
    assert not any("首页" in f.message for f in findings)


def test_legacy_static_header_remains_even_with_different_page_number_baselines():
    a = template("说明文字")
    header_text(a.sections[0], "学校名称")
    a.add_section(WD_SECTION_START.NEW_PAGE)
    a.add_paragraph("其他说明")
    number_settings(a.sections[1], "upperRoman", 5)
    b = template("不具名区域")
    header_text(b.sections[0], "错误学校名称")
    old = Finding("035", "页眉文字", Status.FAIL, "原证据", section_index=1, header_kind=1)
    result = adapt_legacy_header_findings([old], extract_headers_baseline(DocxContext(a)), DocxContext(b))
    assert result[0] is old


def test_extra_even_header_uses_template_default_format_as_effective_baseline():
    a = template()
    header_text(a.sections[0], "学校名称")
    a.sections[0].header.paragraphs[0].alignment = 1
    b = clone(a)
    b.settings.odd_and_even_pages_header_footer = True
    header_text(b.sections[0], "学校名称", kind="even_page_header")
    b.sections[0].even_page_header.paragraphs[0].alignment = 0

    faults = failed(check(a, b), "036")

    assert faults and any(f.page_variant == "even" for f in faults)


def test_legacy_nested_comment_text_is_flattened():
    a = template()
    field(a.sections[0].header.paragraphs[0], "STYLEREF 1", "样例")
    old = Finding("035", "页眉文字", Status.FAIL, "【旧批注】\n模板要求：旧\n当前文档实际：旧\n请核对调整。", section_index=1, header_kind=1)
    result = adapt_legacy_header_findings([old], extract_headers_baseline(DocxContext(a)), DocxContext(clone(a)))
    # 统一批注格式：抬头 + 应该 + 实际；原批注里的换行必须被压平
    assert result[0].message.count("\n") == 2


def test_deleted_footer_text_and_inherited_hidden_page_are_unknown():
    a = template()
    b = clone(a)
    p = b.sections[0].footer.paragraphs[0]
    revision = OxmlElement("w:del")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "已删除页脚"
    r.append(t)
    revision.append(r)
    p._p.append(revision)
    result = check(a, b)
    assert limited(result, "039") and not failed(result, "039")
    p.clear()
    field(p)
    b.styles["Footer"].font.hidden = True
    result = check(a, b)
    assert limited(result, "044") and not failed(result, "044")
