from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from thesis_format_checker.engine.pipeline import run_comparison
from thesis_format_checker.core.point_comments import content_tree as semantic_tree
from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE, make_finding
from thesis_format_checker.core.models import RunSummary, Status
from thesis_format_checker.checks.catalog import (
    DEFERRED_RULE_IDS,
    LEGACY_RULE_IDS,
    PENDING_RULE_IDS,
    REMOVED_RULE_IDS,
)
from thesis_format_checker.engine.pipeline import SECTION_RULES, check_sections, extract_baseline
from conftest import comment_heads, family_titles


def test_exact_number_coverage_and_append_only_baseline():
    """大类十、十一、十二（100–132）的 33 个规则族全部有归属：检测 / 待实现 / 待确认 / 已删除。"""
    sections = ("notes", "bibliography", "characters", "language_lite")
    covered = {rule for section in sections for rule in SECTION_RULES[section]
               if 100 <= int(rule) <= 132}
    assert len(covered) == 25      # 原 24 条 + language_lite 的 119
    legacy_here = {rule for rule in LEGACY_RULE_IDS if 100 <= int(rule) <= 132}
    removed_here = {rule for rule in REMOVED_RULE_IDS if 100 <= int(rule) <= 132}
    pending_here = {rule for rule in PENDING_RULE_IDS if 100 <= int(rule) <= 132}
    deferred_here = {rule for rule in DEFERRED_RULE_IDS if 100 <= int(rule) <= 132}
    assert legacy_here == {"122", "123"}  # 确定性标点检测由旧版 Word 流程提供
    # 119 的“英文文献作者大小写写法一致性”已实现（119.1），不再属于已删除。
    assert removed_here == {"100", "101", "102", "116"}
    assert pending_here == set()   # 本区间已无“整族待实现”
    # 127 由 characters 模块负责（127.2 已实现，127.1 待确认），因此整族仍是已实现。
    assert deferred_here == {"109", "120"}
    families = covered | legacy_here | removed_here | pending_here | deferred_here
    assert families == {f"{i:03}" for i in range(100, 133)}
    assert len(families) == 33
    ctx = DocxContext(Document())
    old = {"earlier": {"preserved": True}}
    data = extract_baseline(ctx, old)
    assert data["earlier"] is old["earlier"]
    assert set(old) == {"earlier"}
    assert {"notes", "bibliography", "characters"} <= set(data)
    with pytest.raises(ValueError, match="未覆盖"):
        extract_baseline(ctx, data)


def test_detector_failures_isolated_into_read_limit_comments(monkeypatch):
    """注释/参考文献/字符三类各自失败时互不影响，失败的一组降级为“无法判定”。"""
    import thesis_format_checker.detectors.bibliography as bibliography
    import thesis_format_checker.detectors.characters as characters
    import thesis_format_checker.detectors.notes as notes

    def broken(*args, **kwargs):
        raise ValueError("cannot decode object")

    monkeypatch.setattr(notes, "check_note_reference_errors", broken)
    monkeypatch.setattr(bibliography, "check_bibliography_errors", broken)
    monkeypatch.setattr(characters, "check_character_errors", broken)
    ctx = DocxContext(Document())
    sections = ("notes", "bibliography", "characters")
    baseline = extract_baseline(ctx, sections=sections)
    result = check_sections(baseline, ctx, sections=sections)
    assert len(result) == 24      # notes/bibliography/characters 现在覆盖 24 个族
    assert all(f.status == Status.UNKNOWN and UNREADABLE in f.message for f in result)
    assert not (set(REMOVED_RULE_IDS) | set(PENDING_RULE_IDS) | set(DEFERRED_RULE_IDS)) & {f.rule_id for f in result}

    def good_check(*args, **kwargs):
        return [make_finding("105", "注释类型", "模板类型", "另一类型")]

    monkeypatch.setattr(notes, "check_note_reference_errors", good_check)
    result = check_sections(baseline, ctx, sections=sections)
    assert any(f.rule_id == "105" and f.status == Status.FAIL for f in result)
    # notes 换成“可用”实现后，其余 17 个族（bibliography 8 + characters 9）仍应各自降级为读取受限。
    assert sum(f.status == Status.UNKNOWN for f in result) == 17


def test_pipeline_retains_legacy_punctuation_and_reports_only_extension_skips(runtime_dir):
    from thesis_format_checker.detectors.word_effective.punctuation import inspect_punctuation
    a, b, out = [runtime_dir / name for name in ("template.docx", "paper.docx", "out.docx")]
    for path in (a, b):
        doc = Document()
        doc.add_heading("第一章 绪论", level=1)
        doc.add_paragraph("正文中文 text 123。")
        doc.add_paragraph("中文,标点使用；另有（未闭合括号。")
        doc.add_comment(doc.paragraphs[-1].runs, "保留原人工批注", author="人工")
        doc.save(path)
    before = [p.read_bytes() for p in (a, b)]
    calls = []
    def legacy(*args, **kwargs):
        calls.append(args)
        return RunSummary(findings=inspect_punctuation(Document(b).paragraphs[-1].text, 3))
    summary = run_comparison(legacy, a, b, out)
    assert len(calls) == 1
    # 细粒度拆分后按规则族统计（122.1、123.3 …）
    assert {"122", "123"} <= {rule_id.split(".")[0] for rule_id in summary.counts_by_rule()}
    assert not (set(REMOVED_RULE_IDS) | set(PENDING_RULE_IDS) | set(DEFERRED_RULE_IDS)) & {f.rule_id for f in summary.findings}
    assert any("已删除" in note for note in summary.skipped)
    assert any("待确认（业务口径）" in note for note in summary.skipped)
    assert before == [p.read_bytes() for p in (a, b)]
    comments = list(Document(out).comments)
    assert comments[0].text == "保留原人工批注"
    # 批注不再写规则编号：按 122/123 的标题核对（抬头只保留等级与问题标题）。
    heads = comment_heads(comments)
    assert heads & family_titles("122")
    assert heads & family_titles("123")
    with ZipFile(b) as original, ZipFile(out) as annotated:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in original.namelist():
            if name not in allowed:
                assert original.read(name) == annotated.read(name), name
        assert semantic_tree(etree.fromstring(original.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))


def test_complete_fixture_detects_new_rules_preserving_notes_fields_and_characters(runtime_dir):
    from fixtures_citations import fixture
    a, b, out = [runtime_dir / name for name in ("all_template.docx", "all_paper.docx", "all_out.docx")]
    fixture(a, template=True)
    fixture(b, template=False)
    before = [p.read_bytes() for p in (a, b)]
    summary = run_comparison(lambda *args, **kwargs: RunSummary(), a, b, out)
    sections = ("notes", "bibliography", "characters")
    expected = {rule for section in sections for rule in SECTION_RULES[section]}
    failed = {f.rule_id for f in summary.findings if f.status == Status.FAIL}
    # 细粒度拆分后按规则族统计（105.1、107.3 …）
    failed_families = {rule_id.split(".")[0] for rule_id in failed}
    # 本次新增的族（引用序号格式、字符卫生、文献类型标识）不在旧夹具覆盖范围内，
    # 其余族必须仍然逐族产出结论。
    newly_optional = {"103", "104", "112", "126", "127", "130", "131", "132"}
    assert expected - newly_optional <= failed_families, sorted(expected - newly_optional - failed_families)
    assert not (set(REMOVED_RULE_IDS) | set(PENDING_RULE_IDS) | set(DEFERRED_RULE_IDS)) & failed
    assert before == [p.read_bytes() for p in (a, b)]
    with ZipFile(b) as original, ZipFile(out) as annotated:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in original.namelist():
            if name not in allowed:
                assert original.read(name) == annotated.read(name), name
        assert semantic_tree(etree.fromstring(original.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))
    comments = list(Document(out).comments)
    assert comments[0].text == "保留的人工批注"
    for comment in comments[1:]:
        assert comment.text.startswith("【") and "应该：" in comment.text and "实际：" in comment.text
        # 批注已锚定在出问题的文字上，正文里不再重复写“位置：”一行。
        assert "位置：" not in comment.text
