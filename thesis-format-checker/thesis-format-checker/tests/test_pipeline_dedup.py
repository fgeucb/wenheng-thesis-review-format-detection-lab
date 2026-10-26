"""批注去重与同类问题合并：报告条数必须与实际写出的批注一致。"""

from docx import Document

from thesis_format_checker.core.docx_context import DocxContext, make_finding, unknown
from thesis_format_checker.core.models import Finding, RunSummary, Status
from thesis_format_checker.engine.pipeline import (
    _paragraph_anchor,
    dedupe_findings,
    merge_repeated_problems,
    run_comparison,
)


def _context(*texts):
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    return DocxContext(document)


def _finding(rule, expected, actual, context, index):
    return make_finding(rule, "测试规则", expected, actual, context.paragraphs[index])


def test_same_problem_on_same_paragraph_is_deduplicated():
    context = _context("正文一", "正文二")
    first = _finding("068", "Times New Roman", "宋体", context, 0)
    second = _finding("068", "Times New Roman", "宋体", context, 0)
    second.message = "旧版措辞不同，但讲的是同一件事"

    result, dropped = dedupe_findings([first, second], context)

    assert dropped == 1
    assert result == [first]


def test_same_problem_on_different_paragraphs_is_kept_by_dedupe():
    context = _context("正文一", "正文二")
    result, dropped = dedupe_findings(
        [_finding("068", "Times New Roman", "宋体", context, 0),
         _finding("068", "Times New Roman", "宋体", context, 1)],
        context,
    )
    assert dropped == 0 and len(result) == 2


def test_same_paragraph_same_problem_ignores_text_range_for_font_rules():
    """同一段被切成多个字号片段时是同一个问题，只写一条批注。"""

    context = _context("正文一")
    first = _finding("068", "Times New Roman", "宋体", context, 0)
    first.start, first.end = 0, 2
    second = _finding("068", "Times New Roman", "宋体", context, 0)
    second.start, second.end = 2, 4
    result, dropped = dedupe_findings([first, second], context)
    assert dropped == 1 and len(result) == 1


def test_character_rules_keep_each_position_when_deduplicating():
    context = _context("正文一")
    first = _finding("126", "只保留一个半角空格", "连续 2 个空格", context, 0)
    first.start, first.end = 1, 3
    second = _finding("126", "只保留一个半角空格", "连续 2 个空格", context, 0)
    second.start, second.end = 4, 6
    result, dropped = dedupe_findings([first, second], context)
    assert dropped == 0 and len(result) == 2


def test_word_paragraph_index_offset_anchors_the_closest_match():
    """Word 段号与 docx 段号存在偏移时，按最接近的命中段落定位，而不是回到文首。"""

    context = _context("本科生毕业论文（设计）", "图2-2  模型整体流程图", "正文一", "如图2-2所示")
    finding = make_finding("069", "正文字号", "12.0", "10.5")
    finding.paragraph_index = 3          # Word 段号：docx 段号偏小
    finding.excerpt = "模型整体流程图"
    assert _paragraph_anchor(context, finding).text == "图2-2  模型整体流程图"
    assert not getattr(finding, "location_note", "")


def test_unmatched_location_is_not_commented_at_all():
    """定位不到文字时既不猜位置、也不堆到文首：本条只进报告。"""

    context = _context("本科生毕业论文（设计）", "正文一", "正文二", "正文三")
    finding = make_finding("068", "正文英文及数字字体", "Times New Roman", "宋体")
    finding.paragraph_index = 3
    finding.excerpt = "文档里不存在的片段"
    assert _paragraph_anchor(context, finding) is None
    assert "未写批注" in finding.location_note


def test_out_of_range_word_paragraph_index_is_not_commented():
    """Word 段落数多于 docx 时（含文本框），越界段号不再猜位置。"""

    context = _context("本科生毕业论文（设计）", "正文一", "结论")
    finding = make_finding("045.2", "一级标题西文字体不符合要求", "黑体", "Times New Roman")
    finding.paragraph_index = 1035
    finding.excerpt = ""
    # 摘录为空且段号越界：无法核对位置。
    assert _paragraph_anchor(context, finding) is None


def test_merged_locations_are_deduplicated():
    context = _context("标题一", "标题二", "标题三")
    findings = [_finding("068", "Times New Roman", "宋体", context, 0),
                _finding("068", "Times New Roman", "宋体", context, 0),
                _finding("068", "Times New Roman", "宋体", context, 1)]
    result, merged = merge_repeated_problems(findings, context)
    assert merged == 2
    assert result[0].merged_locations == [1, 2]
    assert result[0].merged_count == 2


def test_repeated_problem_across_paragraphs_is_merged_into_one():
    context = _context("标题一", "标题二", "标题三")
    findings = [_finding("047.1", "黑体", "宋体", context, index) for index in range(3)]

    result, merged = merge_repeated_problems(findings, context)

    assert merged == 2 and len(result) == 1
    assert result[0].merged_count == 3
    assert result[0].merged_locations == [1, 2, 3]


def test_two_occurrences_are_not_merged():
    context = _context("标题一", "标题二")
    findings = [_finding("047.1", "黑体", "宋体", context, index) for index in range(2)]
    result, merged = merge_repeated_problems(findings, context)
    assert merged == 0 and len(result) == 2


def test_finding_without_location_goes_to_report_instead_of_cover(runtime_dir):
    """定位不到位置的结论只进报告：不给它随便找个位置写批注。"""

    template_path = runtime_dir / "template.docx"
    document_path = runtime_dir / "paper.docx"
    output_path = runtime_dir / "out.docx"
    Document().save(template_path)
    document = Document()
    document.add_paragraph("本科生毕业论文（设计）")
    document.add_paragraph("正文一")
    document.save(document_path)

    def legacy(template_file, document_file, output_file, **kwargs):
        finding = make_finding("068", "正文英文及数字字体", "Times New Roman", "宋体")
        finding.paragraph_index = 999        # 段号越界、无摘录：无法核对位置
        return RunSummary(findings=[finding])

    summary = run_comparison(legacy, template_path, document_path, output_path, overwrite=True)

    assert any("未能定位" in entry for entry in summary.report_only)
    # 这条 068 不写批注；文档里其他规则（结构类）的批注不受影响。
    texts = [comment.text for comment in Document(output_path).comments]
    assert not any("Times New Roman" in text for text in texts)   # 068 的“应该”值


def test_template_missing_wording_is_treated_as_not_applicable():
    """模板基准缺失/不唯一的措辞统一按“不适用”跳过（不写批注、不计入无法判定）。"""

    reasons = (
        "模板中没有相应用途和层级的编号基准",
        "模板页眉页脚基准不唯一、节设置读取失败或模板证据不可得；请按区域核对或提供section_map",
        "首页/奇偶页选项不一致，无法完整对应生效页码区域",
        "日期缺少至少两处可确认的同用途模板示例或模板字段不可解析",
    )
    for reason in reasons:
        assert unknown("051", "标题编号", reason).status is Status.NOT_APPLICABLE


def test_render_only_rules_stay_in_report_without_comments(runtime_dir):
    """需要 Word 渲染的“无法判定”只进报告，不写批注；同一规则的确定问题仍写批注。"""

    template_path = runtime_dir / "template.docx"
    document_path = runtime_dir / "paper.docx"
    output_path = runtime_dir / "out.docx"
    Document().save(template_path)
    document = Document()
    document.add_paragraph("正文一")
    document.add_paragraph("正文二")
    document.save(document_path)

    def legacy(template_file, document_file, output_file, **kwargs):
        def item(rule, title, expected, actual, paragraph, status=Status.FAIL):
            finding = make_finding(rule, title, expected, actual)
            finding.status = status
            finding.actual = actual
            finding.paragraph_index = paragraph
            return finding

        return RunSummary(findings=[
            item("108", "交叉引用显示值", "与书签一致", "PAGEREF 显示值依赖分页", 1, Status.UNKNOWN),
            item("108", "交叉引用显示值", "与书签一致", "指向不存在的书签", 2, Status.FAIL),
        ])

    summary = run_comparison(legacy, template_path, document_path, output_path, overwrite=True)

    assert any("108" in entry for entry in summary.report_only)
    written = Document(output_path)
    texts = [comment.text for comment in written.comments]
    # 只有确定问题写成了批注（批注抬头不再带规则编号，按“实际”内容核对）。
    assert sum("指向不存在的书签" in text for text in texts) == 1


def test_run_comparison_report_matches_written_comments(runtime_dir):
    """端到端：重复结论被去掉后报告条数与批注条数一致（显式开启合并时同样一致）。"""

    template_path = runtime_dir / "template.docx"
    document_path = runtime_dir / "paper.docx"
    output_path = runtime_dir / "out.docx"
    template = Document()
    template.add_paragraph("模板正文")
    template.save(template_path)
    document = Document()
    for text in ("正文一", "正文二", "正文三"):
        document.add_paragraph(text)
    document.save(document_path)

    def legacy(template_file, document_file, output_file, **kwargs):
        # 与真实 Word 路径一致：只给段号（1 起），由流水线解析锚点。
        def item(rule, title, expected, actual, paragraph):
            finding = make_finding(rule, title, expected, actual)
            finding.paragraph_index = paragraph
            return finding

        duplicated = [
            item("068", "正文英文及数字字体", "Times New Roman", "宋体", 1),
            item("068", "正文英文及数字字体", "Times New Roman", "宋体", 1),
        ] + [item("047.1", "三级标题中文字体", "黑体", "宋体", index) for index in (1, 2, 3)]
        return RunSummary(findings=duplicated)

    # 合并是可选能力：默认不写汇总批注，这里显式开启以覆盖合并路径。
    summary = run_comparison(legacy, template_path, document_path, output_path, overwrite=True,
                             merge_repeats=3)

    # 068 在同一段落只保留一条；047.1 的三处合并成一条汇总批注。
    feed_068 = [finding for finding in summary.findings if finding.rule_id == "068"]
    assert len(feed_068) == 1
    assert summary.duplicates_removed >= 1
    assert summary.merged_repeats >= 2
    written = Document(output_path)
    commented = [
        finding for finding in summary.findings
        if finding.status == Status.FAIL
    ]
    assert len(written.comments) == len(commented)   # 报告条数与实际批注条数一致
    merged = next(finding for finding in summary.findings if finding.rule_id == "047.1")
    assert merged.merged_count == 3
    assert "共 3 处同类问题" in merged.message
