"""可选的真实 Word 端到端测试：通过 thesis-checker.cmd 公共入口调用。"""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
import re
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree
import pytest

from conftest import assert_reported_comments


_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
_EXPECTED_RULES = {"001", "005", "013", "023", "024", "025", "026", "067", "068", "069", "072", "074", "075"}


def _format_run(run, *, size, east_asia="宋体", latin="Times New Roman", bold=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)


def _fixture(path, *, template):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT if template else WD_ORIENT.LANDSCAPE
    section.page_width = Mm(210 if template else 240)
    section.page_height = Mm(297 if template else 160)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(30 if template else 20)
    section.right_margin = Mm(30)
    section.gutter = Mm(10 if template else 0)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    title = document.add_paragraph()
    title.paragraph_format.line_spacing = 1.3
    _format_run(title.add_run("论文题目：模板示例题目" if template else "论文题目：实际研究题目"), size=19 if template else 11, east_asia="仿宋", bold=True)
    document.add_paragraph("作者姓名：张三")
    if template:
        # The plain instructions are evidence, while the sample title is not a value requirement.
        document.add_paragraph("必须包含：封面、题名页。", style="Title")
        document.add_paragraph("中文关键词数量：2—3个。", style="Title")
        document.add_paragraph("题名页")
        document.add_paragraph("论文题目：模板示例题目")
    document.add_paragraph("摘要")
    document.add_paragraph("本研究用于检验模板与文档的格式属性差异。")
    document.add_paragraph("关键词：甲；乙" if template else "关键词：甲；乙；丙；丁")
    document.add_heading("第一章 绪论", level=1)
    for index in range(20):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24 if template else 0)
        paragraph.paragraph_format.line_spacing = 1.5 if template else 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6 if template else 0)
        paragraph._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200" if template else "0")
        _format_run(
            paragraph.add_run(f"这是第{index + 1}段用于格式检测的普通正文，包含连续中文内容和 English scientific research 123456。"),
            size=12 if template else 10,
            east_asia="宋体" if template else "微软雅黑",
            latin="Times New Roman" if template else "Arial",
        )
    document.save(path)


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _body_without_comment_tokens(xml):
    """Compare original text/formatting while permitting lossless run splitting."""
    from thesis_format_checker.core.point_comments import content_tree
    return content_tree(etree.fromstring(xml))


@pytest.mark.integration
def test_public_cmd_retains_legacy_checks_adds_docx_comments_and_preserves_input(runtime_dir):
    if os.environ.get("RUN_WORD_INTEGRATION") != "1":
        pytest.skip("设置 RUN_WORD_INTEGRATION=1 后运行真实 cmd / Microsoft Word 集成测试")
    if os.name != "nt":
        pytest.skip("此集成测试要求 Windows 与 Microsoft Word")
    pytest.importorskip("win32com.client")
    repository = Path(__file__).resolve().parents[1]
    template = runtime_dir / "标准模板.docx"
    document = runtime_dir / "待检测文档.docx"
    output = runtime_dir / "带原生批注.docx"
    _fixture(template, template=True)
    _fixture(document, template=False)
    input_hashes = (_sha256(template), _sha256(document))
    with ZipFile(document) as archive:
        original_parts = {name: archive.read(name) for name in archive.namelist()}
    report = runtime_dir / "report.json"
    command = [str(repository / "thesis-checker.cmd"), "compare", "--template", str(template), "--document", str(document), "--output", str(output), "--report-json", str(report)]
    stdout_log = runtime_dir / "cmd_stdout.log"
    stderr_log = runtime_dir / "cmd_stderr.log"
    try:
        # File handles avoid inherited-pipe waits if Word or a child outlives cmd.
        with stdout_log.open("wb") as stdout, stderr_log.open("wb") as stderr:
            result = subprocess.run(command, cwd=repository, env=os.environ.copy(), stdout=stdout, stderr=stderr, timeout=180, check=False)
    except subprocess.TimeoutExpired:
        pytest.fail("真实 cmd 检测超过180秒；已停止等待，不自动重试，不终止任何用户 Word 进程。子进程可能仍在运行，请检查本次CLI状态。")
    assert input_hashes == (_sha256(template), _sha256(document)), "检测必须保留两份输入文件字节不变"
    assert result.returncode == 0, f"CLI 未成功结束：{stdout_log.read_text(encoding='utf-8', errors='replace')[-2000:]}\n{stderr_log.read_text(encoding='utf-8', errors='replace')[-2000:]}"
    assert output.is_file(), "固定CLI调用必须生成新DOCX"
    with ZipFile(output) as archive:
        output_parts = {name: archive.read(name) for name in archive.namelist()}
    assert "word/comments.xml" in output_parts
    allowed_changes = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "word/comments.xml"}
    for name, contents in original_parts.items():
        if name not in allowed_changes:
            assert output_parts[name] == contents, f"非批注包部件被改变：{name}"
    assert _body_without_comment_tokens(output_parts["word/document.xml"]) == _body_without_comment_tokens(original_parts["word/document.xml"])
    comments = etree.fromstring(output_parts["word/comments.xml"])
    comment_texts = ["\n".join("".join(p.itertext()) for p in comment.findall("w:p", _NS)) for comment in comments]
    # 规则目录的标题别名与实际取证标题可能不同；由报告规则编号核对批注。
    for text in comment_texts:
        first = text.splitlines()[0] if text else ""
        if first.startswith(("【错误】", "【警告】", "【提醒】")):
            assert "应该：" in text and "实际：" in text
    assert_reported_comments(report, Document(output).comments, _EXPECTED_RULES)
    body = etree.fromstring(output_parts["word/document.xml"])
    comment_ids = {comment.get(qn("w:id")) for comment in comments}
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        assert {node.get(qn("w:id")) for node in body.findall(f".//w:{tag}", _NS)} == comment_ids
    # Reload through python-docx to ensure comments remain native, parseable objects.
    reloaded = Document(output)
    assert len(reloaded.comments) == len(comment_ids)
