from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.equations import extract_equation_baseline, check_equation_errors
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def node(tag, **attrs):
    result = OxmlElement(tag)
    prefix = tag.split(":")[0]
    for key, value in attrs.items():
        result.set(qn(prefix + ":" + key), str(value))
    return result


def math_defaults(doc, font="Cambria Math"):
    root = doc.settings.element.find(qn("m:mathPr"))
    if root is None:
        root = node("m:mathPr")
        doc.settings.element.append(root)
    old = root.find(qn("m:mathFont"))
    if old is not None:
        root.remove(old)
    root.append(node("m:mathFont", val=font))
    doc.styles["Normal"].font.size = Pt(13)
    doc.styles["Normal"].font.name = "Arial"


def math_run(text="x=1", size=None, font=None, normal=False):
    run = node("m:r")
    if normal:
        prop = node("m:rPr")
        prop.append(node("m:nor"))
        run.append(prop)
    if size is not None or font:
        prop = node("w:rPr")
        if font:
            prop.append(node("w:rFonts", ascii=font, hAnsi=font, eastAsia=font))
        if size is not None:
            prop.append(node("w:sz", val=2 * size))
        run.append(prop)
    text_node = node("m:t")
    text_node.text = text
    run.append(text_node)
    return run


def equation(doc, label="(1)", jc="center", paragraph=None, size=None, font=None, normal=False, inline=False):
    p = paragraph or doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5), WD_TAB_ALIGNMENT.RIGHT)
    formula = node("m:oMath")
    formula.append(math_run(size=size, font=font, normal=normal))
    if inline:
        p.add_run("文字说明")
        p._p.append(formula)
    else:
        display = node("m:oMathPara")
        prop = node("m:oMathParaPr")
        prop.append(node("m:jc", val=jc))
        display.append(prop)
        display.append(formula)
        p._p.append(display)
    if label:
        p.add_run("\t" + label)
    return p, formula


def documents():
    a, b = Document(), Document()
    math_defaults(a)
    math_defaults(b)
    return a, b


def check(a, b, mapping=None):
    return check_equation_errors(extract_equation_baseline(DocxContext(a)), DocxContext(b), mapping)


def select(items, rule, status=Status.FAIL):
    statuses = (Status.UNKNOWN, Status.NOT_APPLICABLE) if status == Status.UNKNOWN else (status,)
    return [item for item in items if family_of(item.rule_id) == rule and item.status in statuses]


def seq(p, instruction="SEQ Equation \\* ARABIC", cache="1"):
    field = node("w:fldSimple", instr=instruction)
    run = node("w:r")
    text = node("w:t")
    text.text = cache
    run.append(text)
    field.append(run)
    p._p.append(field)
    return field


def test_090_template_math_alignment_and_tabs_not_hardcoded():
    a, b = documents()
    equation(a, jc="left")
    bp, _ = equation(b, jc="center")
    results = check(a, b)
    assert select(results, "090")
    assert select(results, "090", Status.UNKNOWN)
    bp._p.xpath("./m:oMathPara/m:oMathParaPr/m:jc")[0].set(qn("m:val"), "left")
    assert not select(check(a, b), "090")
    bp.paragraph_format.tab_stops.add_tab_stop(Inches(4), WD_TAB_ALIGNMENT.RIGHT)
    assert select(check(a, b), "090")


def test_091_manual_number_pattern_dynamic_template():
    a, b = documents()
    equation(a, "式（2.1）")
    equation(b, "(2-1)")
    results = check(a, b)
    failure = select(results, "091")[0]
    assert "式（#.#）" in failure.message
    assert failure.anchor is not None


def test_091_bare_number_is_compared_and_seq_ids_are_local():
    a, b = documents()
    equation(a, "(1)")
    equation(b, "1")
    assert select(check(a, b), "091")
    a, b = documents()
    ap, _ = equation(a, None)
    bp, _ = equation(b, None)
    seq(ap, "SEQ Equation \\* ARABIC")
    seq(bp, "SEQ MyEq \\* ARABIC")
    assert not select(check(a, b), "091")


def test_091_seq_structure_compared_but_cache_never_number_truth():
    a, b = documents()
    ap, _ = equation(a, None)
    bp, _ = equation(b, None)
    seq(ap, "SEQ Equation \\* ARABIC", "1")
    seq(bp, "SEQ Equation \\* ROMAN", "999")
    a.add_paragraph("公式全文连续编号。")
    results = check(a, b)
    assert select(results, "091")
    assert select(results, "091", Status.UNKNOWN)
    assert not select(results, "092")
    assert select(results, "092", Status.UNKNOWN)


def test_092_continuity_repeat_and_no_invented_start():
    a, b = documents()
    a.add_paragraph("公式全文连续编号。")
    equation(a, "(8)")
    equation(b, "(8)")
    equation(b, "(10)")
    equation(b, "(10)")
    failures = select(check(a, b), "092")
    assert len(failures) == 2
    assert any("应为 9" in item.message for item in failures)
    assert any("重复编号" in item.message for item in failures)


def test_092_chapter_prefix_and_reset_scope():
    a, b = documents()
    a.add_paragraph("公式按章编号，每章重新编号。")
    a.add_heading("1 第一章", 1)
    equation(a, "(1-1)")
    b.add_heading("1 第一章", 1)
    equation(b, "(1-1)")
    b.add_heading("2 第二章", 1)
    equation(b, "(3-1)")
    failures = select(check(a, b), "092")
    assert len(failures) == 1
    assert "章前缀 3" in failures[0].message


def test_092_no_sequence_policy_is_unknown():
    a, b = documents()
    equation(a)
    equation(b, "(9)")
    results = check(a, b)
    assert not select(results, "092")
    assert select(results, "092", Status.UNKNOWN)


def test_093_explicit_internal_manual_reference_and_external_exclusion():
    a, b = documents()
    equation(a)
    a.add_paragraph("由本文式(1)可得结果。")
    equation(b)
    b.add_paragraph("由本文式(9)可得结果。")
    b.add_paragraph("由文献[2]式(99)可得结果。")
    failures = select(check(a, b), "093")
    assert len(failures) == 1
    assert "式(9)" in failures[0].message


def test_093_ref_missing_bookmark_and_cached_value_uncertain():
    a, b = documents()
    equation(a)
    a.add_paragraph("由本文式(1)可得结果。")
    bp, _ = equation(b)
    bp._p.insert(0, node("w:bookmarkStart", id=40, name="eq_one"))
    bp._p.append(node("w:bookmarkEnd", id=40))
    good = b.add_paragraph("本文公式")
    seq(good, "REF eq_one \\h", "(999)")
    missing = b.add_paragraph("本文公式")
    seq(missing, "REF missing_eq \\h", "(1)")
    results = check(a, b)
    assert len(select(results, "093")) == 1
    assert "missing_eq" in select(results, "093")[0].message
    assert select(results, "093", Status.UNKNOWN)


def test_096_math_default_font_and_effective_style_size():
    a, b = documents()
    math_defaults(a, "STIX Two Math")
    math_defaults(b, "Cambria Math")
    equation(a)
    bp, _ = equation(b)
    custom = b.styles.add_style("Derived Equation", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = b.styles["Normal"]
    custom.font.size = Pt(16)
    bp.style = custom
    failures = select(check(a, b), "096")
    assert any("STIX Two Math" in item.message for item in failures)
    assert any("16" in item.message and "基础字号" in item.message for item in failures)


def test_096_direct_override_and_normal_math_text_inherits_style():
    a, b = documents()
    equation(a, normal=True, font="Arial", size=13)
    equation(b, normal=True, font="Courier New", size=18)
    failures = select(check(a, b), "096")
    assert any("Courier New" in item.message for item in failures)
    assert any("18" in item.message for item in failures)


def test_096_script_base_size_role_does_not_compare_to_base_run():
    a, b = documents()
    for doc in (a, b):
        _, formula = equation(doc)
        sub = node("m:sub")
        sub.append(math_run("2", size=8))
        formula.append(sub)
    assert not select(check(a, b), "096")


def test_097_only_formula_paragraph_baseline_not_body():
    a, b = documents()
    a.add_paragraph("正文").paragraph_format.line_spacing = 2
    equation(a)
    bp, _ = equation(b)
    bp.paragraph_format.space_after = Pt(17)
    failures = select(check(a, b), "097")
    assert len(failures) == 1
    assert "段后间距 6" in failures[0].message
    assert "段后间距 17" in failures[0].message


def test_formula_table_identifies_neighbor_number_and_template_layout():
    a, b = documents()
    for doc in (a, b):
        table = doc.add_table(rows=1, cols=2)
        equation(doc, None, paragraph=table.cell(0, 0).paragraphs[0])
        table.cell(0, 1).paragraphs[0].text = "(1)"
        table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    b.tables[0].cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    assert select(check(a, b), "090")
    assert not select(check(a, b), "091")


def test_opaque_mathtype_degrades_without_crash_or_semantic_rules():
    a, b = documents()
    equation(a)
    a.add_paragraph("公式全文连续编号。")
    p = b.add_paragraph("(1)")
    obj = node("w:object")
    from lxml import etree
    ole = etree.Element("{urn:schemas-microsoft-com:office:office}OLEObject")
    ole.set("ProgID", "Equation.DSMT4")
    obj.append(ole)
    p.add_run()._r.append(obj)
    results = check(a, b)
    assert {"090", "091", "092", "096", "097"} <= {f.rule_id for f in results}
    assert all(f.status == Status.UNKNOWN for f in results)
    assert all(UNREADABLE in f.message for f in results)
    assert not {"089", "094", "095", "098", "099"} & {f.rule_id for f in results}


def test_read_only_xml_even_font_proxy_and_table_reads():
    a, b = documents()
    equation(a)
    equation(b, size=20)
    before = [doc.element.xml for doc in (a, b)]
    before_styles = [doc.styles.element.xml for doc in (a, b)]
    before_settings = [doc.settings.element.xml for doc in (a, b)]
    check(a, b)
    assert before == [doc.element.xml for doc in (a, b)]
    assert before_styles == [doc.styles.element.xml for doc in (a, b)]
    assert before_settings == [doc.settings.element.xml for doc in (a, b)]


def test_incomplete_field_and_char_unit_spacing_return_read_limits():
    a, b = documents()
    equation(a)
    bp, _ = equation(b)
    run = bp.add_run()._r
    run.append(node("w:fldChar", fldCharType="begin"))
    instr = node("w:instrText")
    instr.text = "SEQ Equation"
    run.append(instr)
    bp._p.get_or_add_pPr().find(qn("w:spacing")).set(qn("w:afterLines"), "100")
    results = check(a, b)
    assert select(results, "091", Status.UNKNOWN)
    assert select(results, "097", Status.UNKNOWN)
    assert all(UNREADABLE in f.message for f in select(results, "091", Status.UNKNOWN))


def test_conflicting_baseline_does_not_pick_arbitrary_formula():
    a, b = documents()
    equation(a, "(1)", jc="left")
    equation(a, "式（2）", jc="right")
    equation(b, "(3)")
    results = check(a, b)
    assert not select(results, "090")
    assert not select(results, "091")
    assert select(results, "090", Status.UNKNOWN)
    assert select(results, "091", Status.UNKNOWN)


def test_partial_template_opaque_does_not_reduce_baseline_to_readable_sample():
    a, b = documents()
    equation(a, jc="left")
    p = a.add_paragraph("(2)")
    p.add_run()._r.append(node("w:object"))
    equation(b, jc="right", size=19)
    results = check(a, b)
    assert not select(results, "090")
    assert not select(results, "096")
    assert select(results, "096", Status.UNKNOWN)


def test_greek_math_symbols_and_control_properties_are_checked():
    a, b = documents()
    for doc, size in ((a, 13), (b, 19)):
        _, formula = equation(doc)
        formula.append(math_run("αβ", size=size))
        nary = node("m:nary")
        props = node("m:naryPr")
        props.append(node("m:chr", val="∑"))
        control = node("m:ctrlPr")
        word_props = node("w:rPr")
        word_props.append(node("w:sz", val=2 * size))
        control.append(word_props)
        props.append(control)
        nary.append(props)
        formula.append(nary)
    failures = select(check(a, b), "096")
    assert any("αβ" in f.message for f in failures)
    assert any("∑" in f.message and "control" in f.message for f in failures)


def test_revision_and_malformed_number_do_not_make_false_missing_reference():
    a, b = documents()
    equation(a)
    a.add_paragraph("由本文式(1)可得结果。")
    p, _ = equation(b)
    p.add_run()._r.append(node("w:fldChar", fldCharType="begin"))
    b.add_paragraph("由本文式(999)可得结果。")
    results = check(a, b)
    assert not select(results, "093")
    assert select(results, "093", Status.UNKNOWN)
    p._p.get_or_add_pPr().append(node("w:pPrChange", id=3))
    results = check(a, b)
    assert not select(results, "093")
    assert select(results, "096", Status.UNKNOWN)


def test_unreadable_object_poison_reference_inventory(monkeypatch):
    import thesis_format_checker.detectors.equations as module
    a, b = documents()
    equation(a)
    a.add_paragraph("本文式(1)")
    equation(b)
    b.add_paragraph("本文式(999)")
    baseline = extract_equation_baseline(DocxContext(a))
    original = module._read_records
    monkeypatch.setattr(module, "_read_records", lambda ctx: (original(ctx)[0], [(ctx.paragraphs[0], "读取失败")]))
    results = check_equation_errors(baseline, DocxContext(b))
    assert not select(results, "093")
    assert select(results, "093", Status.UNKNOWN)


def test_range_formula_reference_is_not_partially_declared_good():
    a, b = documents()
    equation(a)
    a.add_paragraph("本文式(1)")
    equation(b)
    b.add_paragraph("本文式(1)—(9)")
    results = check(a, b)
    assert not select(results, "093")
    assert any("范围" in f.message for f in select(results, "093", Status.UNKNOWN))
