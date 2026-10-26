"""模板体检工具（tools/check_template.py）的单元测试。"""

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def _tool():
    spec = importlib.util.spec_from_file_location("check_template", ROOT / "tools" / "check_template.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # 数据类需要能按模块名回查，先登记再执行
    spec.loader.exec_module(module)
    return module


def _items(module, path):
    return {item.name: item for item in module.inspect(path)}


def test_missing_heading_samples_are_reported_as_required(runtime_dir):
    module = _tool()
    document = Document()
    document.add_paragraph("正文示例内容。")
    path = runtime_dir / "template.docx"
    document.save(path)

    items = _items(module, path)

    assert items["文件格式"].status == module.PASS
    assert items["1 级标题样例"].status == module.FAIL
    assert items["2 级标题样例"].status == module.FAIL
    assert "必须补齐" in module.render(list(items.values()))


def test_heading_samples_make_the_level_pass(runtime_dir):
    module = _tool()
    document = Document()
    for index in range(2):
        document.add_paragraph(f"第{index + 1}章 绪论", style="Heading 1")
        document.add_paragraph(f"{index + 1} 节标题", style="Heading 2")
        document.add_paragraph(f"{index + 1}.1 小节标题", style="Heading 3")
        for paragraph_index in range(2):
            document.add_paragraph(f"正文示例内容 {paragraph_index + 1}，用于比对字体字号与段落格式。")
    path = runtime_dir / "template.docx"
    document.save(path)

    items = _items(module, path)

    assert items["1 级标题样例"].status == module.PASS
    assert items["2 级标题样例"].status == module.PASS
    assert items["3 级标题样例"].status == module.PASS
    assert items["正文样例"].status == module.PASS


def test_mixed_heading_styles_are_reported(runtime_dir):
    module = _tool()
    document = Document()
    for index in range(2):
        document.add_paragraph(f"第{index + 1}章 绪论", style="Heading 1")
    # 手工设置大纲级别的普通段落：同一层级出现两种样式，应提示统一。
    paragraph = document.add_paragraph("第一章 绪论（手工大纲级别）")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    paragraph._p.get_or_add_pPr().append(outline)
    path = runtime_dir / "template.docx"
    document.save(path)

    items = _items(module, path)

    assert items["1 级标题样例"].status == module.WARN
    assert "多种样式" in items["1 级标题样例"].detail
