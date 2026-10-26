"""一致性、附录与特殊版式规则的真实 DOCX 夹具。"""
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


EXPECTED = {"134", "136", "137", "139", "155", "156", "157", "158", "159", "164"}


def fixture(path, *, template):
    from test_detectors_toc import fixture as toc_fixture
    from test_pipeline_headers_toc import field
    from test_detectors_equations import equation, math_defaults
    doc, _, _, headings = toc_fixture(labels=("第一章 绪论", "第二章 方法"))
    math_defaults(doc)
    for style_name in ("Normal", "Heading 1", "Caption"):
        style = doc.styles[style_name]
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.attrib.clear()
        for slot, value in (("eastAsia", "宋体"), ("ascii", "Times New Roman"), ("hAnsi", "Times New Roman")):
            fonts.set(qn("w:" + slot), value)
        style.font.size = Pt(12 if style_name != "Heading 1" else 16)
        style.font.bold = style_name == "Heading 1"
    first = doc.paragraphs[0]
    first.insert_paragraph_before("论文题目：模板示例题名" if template else "论文题目：论文实际题名")
    first.insert_paragraph_before("摘要")
    first.insert_paragraph_before("论文题目：模板示例题名" if template else "论文题目：另一实际题名")
    if not template:
        # Edit only the bookmarked title text, preserving the target bookmark.
        headings[1].runs[0].text = "第二章 新方法"
        headings[1].runs[0].font.bold = False
    for number in range(20):
        p = doc.add_paragraph(f"这是第{number + 1}段论文正文，用于检查 English research 123 和原生批注。")
        p.paragraph_format.first_line_indent = Pt(24)
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")
        p.paragraph_format.line_spacing = 1.5
    doc.add_paragraph("本文第一章给出方法。" if template else "本文第九章给出方法。")
    ref = doc.add_paragraph("本文章节引用：")
    field(ref, "REF _Toc0", "第一章 绪论" if template else "第九章 旧题名")
    caption = doc.add_paragraph("图 1 正文示意", style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation(doc, "(1)", size=12)
    doc.add_paragraph("附录中的图表、公式格式与正文一致。")
    doc.add_paragraph("本文附录 A 给出资料。")
    doc.add_paragraph("本文附录 B 给出补充资料。" if template else "本文附录 Z 给出补充资料。")
    doc.add_heading("附录 A 原始资料", 1)
    appendix_caption = doc.add_paragraph("图 A.1 附录示意" if template else "图 1 附录示意", style="Caption")
    appendix_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not template:
        appendix_caption.runs[0].font.size = Pt(22)
    equation(doc, "(A.1)", size=12)
    appendix_title = doc.add_heading("附录 B 其他资料" if template else "附录 C 其他资料", 1)
    if not template:
        appendix_title.runs[0].font.size = Pt(22)
    doc.add_paragraph("附录可读内容。")
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if template else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Mm(297), Mm(210)) if template else (Mm(210), Mm(297))
    doc.add_paragraph("特殊版式：横向页面")
    special = doc.add_paragraph("本页使用模板明确的特殊节设置。")
    if not template:
        doc.add_comment(special.runs, "保留的人工批注", author="人工校阅者")
    doc.save(path)
