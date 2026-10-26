from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
import pytest

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.appendices import extract_appendix_baseline, check_appendix_errors, collect_appendix_inventory
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def node(tag, **attrs):
    result = OxmlElement(tag)
    for key, value in attrs.items():
        result.set(qn(tag.split(":")[0] + ":" + key), str(value))
    return result


def setup():
    doc = Document()
    for name, size in (("Normal", 11), ("Heading 1", 16), ("Heading 2", 14), ("Caption", 10)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn("w:" + key), "宋体" if key == "eastAsia" else "Arial")
            theme = qn("w:" + key + "Theme")
            if theme in fonts.attrib:
                del fonts.attrib[theme]
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def title(doc, label="C", text="数据"):
    return doc.add_paragraph("附录 " + label + " " + text, "Heading 1")


def caption(doc, number="C.1", kind="图"):
    return doc.add_paragraph(kind + " " + number + " 示例", "Caption")


def formula(doc, number="C.1", size=11, alignment="center", opaque=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if opaque:
        obj = node("w:object")
        ole = etree.Element("{urn:schemas-microsoft-com:office:office}OLEObject", ProgID="Equation.3")
        obj.append(ole)
        p.add_run()._r.append(obj)
    else:
        display = node("m:oMathPara")
        props = node("m:oMathParaPr")
        props.append(node("m:jc", val=alignment))
        display.append(props)
        math = node("m:oMath")
        run = node("m:r")
        rpr = node("w:rPr")
        rpr.append(node("w:rFonts", ascii="Cambria Math", hAnsi="Cambria Math", eastAsia="宋体"))
        rpr.append(node("w:sz", val=size * 2))
        run.append(rpr)
        text = node("m:t")
        text.text = "x=1"
        run.append(text)
        math.append(run)
        display.append(math)
        p._p.append(display)
    if number is not None:
        p.add_run("(" + number + ")")
    return p


def field(p, instr="REF item", cache="附录 C"):
    f = node("w:fldSimple", instr=instr)
    r, t = node("w:r"), node("w:t")
    t.text = cache
    r.append(t)
    f.append(r)
    p._p.append(f)
    return f


def check(template, actual):
    return check_appendix_errors(extract_appendix_baseline(DocxContext(template)), DocxContext(actual))


def select(findings, rule, status=Status.FAIL):
    # 模板未涉及的结论现在是“不适用”（不写批注），测试中与“无法判定”一并接受。
    statuses = (Status.UNKNOWN, Status.NOT_APPLICABLE) if status == Status.UNKNOWN else (status,)
    return [f for f in findings if family_of(f.rule_id) == rule and f.status in statuses]


def template_with_titles():
    template = setup()
    title(template, "C")
    title(template, "D")
    return template


def test_155_missing_and_uncited_appendices_are_internal_comparisons():
    template, actual = setup(), setup()
    ref = actual.add_paragraph("详见本文附录 Z。")
    end = title(actual, "C")
    findings = check(template, actual)
    failures = select(findings, "155")
    assert len(failures) == 2
    assert any(f.anchor._p is ref._p and "Z" in f.actual for f in failures)
    assert any(f.anchor._p is end._p and "C" in f.actual for f in failures)


def test_155_valid_explicit_references_external_and_toc_exclusion():
    template, actual = setup(), setup()
    actual.add_paragraph("正文数据详见本文附录 C。")
    actual.add_paragraph("外部文献附录 Z 介绍原始资料。")
    toc = actual.add_paragraph()
    field(toc, "TOC \\o \"1-3\"", "附录 Q 目录缓存")
    title(actual)
    assert not select(check(template, actual), "155")
    inventory = collect_appendix_inventory(DocxContext(actual))
    assert [r["token"] for r in inventory["titles"]] == ["C"]


def test_155_opaque_later_content_and_dynamic_reference_do_not_false_fail():
    for opaque in ("drawing", "field", "revision"):
        template, actual = setup(), setup()
        actual.add_paragraph("正文详见本文附录 Z。")
        p = actual.add_paragraph()
        if opaque == "drawing":
            p.add_run()._r.append(node("w:drawing"))
        elif opaque == "field":
            field(p, cache="附录 C")
        else:
            p._p.append(node("w:ins", id="1"))
        title(actual)
        findings = check(template, actual)
        assert select(findings, "155", Status.UNKNOWN)
        assert not any("未发现正文引用" in f.actual for f in select(findings, "155"))
        if opaque != "field":
            assert not select(findings, "155")


def test_155_optional_template_allows_unreferenced_appendix():
    template, actual = setup(), setup()
    template.add_paragraph("附录可以不在正文引用。")
    title(actual)
    assert not select(check(template, actual), "155")


def test_156_dynamic_non_a_start_sequence_and_shape():
    template, actual = template_with_titles(), setup()
    title(actual, "C")
    second = title(actual, "E")
    assert select(check(template, actual), "156")
    second.text = "附录 D 数据"
    assert not select(check(template, actual), "156")
    second.text = "附录D 数据"
    assert select(check(template, actual), "156")


def test_156_decimal_numbering_and_duplicate():
    template, actual = setup(), setup()
    title(template, "3")
    title(template, "4")
    title(actual, "3")
    second = title(actual, "3")
    assert select(check(template, actual), "156")
    second.text = "附录 4 数据"
    assert not select(check(template, actual), "156")


def test_156_single_template_sample_does_not_invent_continuity():
    template, actual = setup(), setup()
    title(template)
    title(actual, "C")
    title(actual, "H")
    findings = check(template, actual)
    assert not select(findings, "156")
    assert select(findings, "156", Status.UNKNOWN)


def test_157_effective_inherited_font_and_direct_override():
    template, actual = template_with_titles(), setup()
    child = actual.styles.add_style("Appendix Heading Child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = actual.styles["Heading 1"]
    p = actual.add_paragraph("附录 C 数据", child)
    title(actual, "D")
    assert not select(check(template, actual), "157")
    p.runs[0].font.size = Pt(18)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    findings = select(check(template, actual), "157")
    assert findings and any("18" in f.actual for f in findings)


def test_157_total_and_subheading_roles_are_separate():
    template, actual = setup(), setup()
    for doc in (template, actual):
        doc.add_paragraph("附录", "Heading 1")
        title(doc)
        doc.add_paragraph("子标题", "Heading 2")
    actual.paragraphs[-1].runs[0].font.size = Pt(20)
    findings = select(check(template, actual), "157")
    assert len(findings) == 1
    assert findings[0].anchor._p is actual.paragraphs[-1]._p


def test_157_revised_or_automatic_titles_unknown():
    template, actual = template_with_titles(), setup()
    p = title(actual)
    rpr = p.runs[0]._r.get_or_add_rPr()
    rpr.append(node("w:rPrChange", id="2"))
    findings = check(template, actual)
    assert select(findings, "157", Status.UNKNOWN)
    assert not select(findings, "157")
    assert select(findings, "156", Status.UNKNOWN)


def test_158_letter_captions_and_formula_supported_and_source_immutable():
    template, actual = setup(), setup()
    for doc in (template, actual):
        title(doc)
        caption(doc, "C.1")
        caption(doc, "C.2")
        caption(doc, "C.1", "表")
        formula(doc)
    before = actual.element.xml
    assert not select(check(template, actual), "158")
    assert actual.element.xml == before
    actual.paragraphs[1].text = "图 2.1 示例"
    actual.paragraphs[-1].runs[-1].text = "(B.1)"
    findings = select(check(template, actual), "158")
    assert any("2.1" in f.actual for f in findings)
    assert any("B.1" in f.actual for f in findings)


def test_158_non_default_separator_and_continuity():
    template, actual = setup(), setup()
    for doc in (template, actual):
        title(doc, "7")
        caption(doc, "7-3")
        caption(doc, "7-4")
    assert not select(check(template, actual), "158")
    actual.paragraphs[-1].text = "图 7-6 示例"
    assert select(check(template, actual), "158")


def test_158_reset_only_when_multiple_appendices_demonstrate_it():
    template, actual = setup(), setup()
    for doc in (template, actual):
        title(doc, "C")
        caption(doc, "C.3")
        title(doc, "D")
        caption(doc, "D.3")
    actual.paragraphs[-1].text = "图 D.4 示例"
    failures = select(check(template, actual), "158")
    assert any("重新从序号3" in f.expected for f in failures)


def test_158_does_not_force_independent_numbering_if_template_continues_body():
    template, actual = setup(), setup()
    for doc in (template, actual):
        title(doc)
        caption(doc, "8")
        formula(doc, "9")
    assert not select(check(template, actual), "158")


def test_158_fields_ole_and_missing_baseline_unknown():
    template, actual = setup(), setup()
    title(template)
    formula(template)
    title(actual)
    formula(actual, opaque=True)
    findings = check(template, actual)
    assert select(findings, "158", Status.UNKNOWN)
    assert not select(findings, "158")
    assert all(UNREADABLE in f.message for f in select(findings, "158", Status.UNKNOWN))
    p = caption(actual)
    field(p, "SEQ Figure", "1")
    assert select(check(template, actual), "158", Status.UNKNOWN)


def test_158_same_row_formula_number_cell():
    template, actual = setup(), setup()
    for doc in (template, actual):
        title(doc)
        table = doc.add_table(rows=1, cols=2)
        generated = formula(doc, None)
        table.cell(0, 0)._tc.append(generated._p)
        table.cell(0, 1).text = "(C.1)"
    assert not select(check(template, actual), "158")
    actual.tables[0].cell(0, 1).text = "(D.1)"
    assert select(check(template, actual), "158")


def inherited_objects(kind="caption"):
    template, actual = setup(), setup()
    for doc in (template, actual):
        doc.add_paragraph("附录中的图表和公式格式与正文一致。")
        if kind == "caption":
            caption(doc, "1.1")
        elif kind == "table":
            caption(doc, "1.1", "表")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "示例1"
        else:
            formula(doc, "1.1")
        title(doc)
        if kind == "caption":
            caption(doc)
        elif kind == "table":
            caption(doc, "C.1", "表")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "示例1"
        else:
            formula(doc)
    return template, actual


def test_159_explicit_relation_compares_template_not_wrong_document_body():
    template, actual = inherited_objects()
    assert not select(check(template, actual), "159")
    actual.paragraphs[1].runs[0].font.size = Pt(20)
    assert not select(check(template, actual), "159")
    actual.paragraphs[-1].runs[0].font.size = Pt(12)
    assert select(check(template, actual), "159")


def test_159_missing_relation_and_conflicting_template_examples_unknown():
    template, actual = inherited_objects()
    template.paragraphs[0].text = "附录可能采用不同的图表格式。"
    actual.paragraphs[-1].runs[0].font.size = Pt(12)
    assert not select(check(template, actual), "159")
    assert select(check(template, actual), "159", Status.UNKNOWN)
    template.paragraphs[0].text = "附录中的图表和公式格式与正文一致。"
    template.paragraphs[-1].runs[0].font.size = Pt(9)
    assert not select(check(template, actual), "159")
    assert any("冲突" in f.actual for f in select(check(template, actual), "159", Status.UNKNOWN))


def test_159_readable_table_cell_and_formula_formats():
    template, actual = inherited_objects("table")
    actual.tables[-1].cell(0, 0).paragraphs[0].runs[0].font.size = Pt(15)
    assert select(check(template, actual), "159")
    template, actual = inherited_objects("formula")
    actual.paragraphs[-1].paragraph_format.space_after = Pt(17)
    actual.paragraphs[-1]._p.xpath(".//m:r/w:rPr/w:sz")[0].set(qn("w:val"), "32")
    findings = select(check(template, actual), "159")
    assert any("17" in f.actual for f in findings)
    assert any("16" in f.actual for f in findings)


def test_unreadable_textbox_exception_isolated_and_no_semantic_ids():
    template, actual = template_with_titles(), setup()
    title(actual)
    p = actual.add_paragraph()
    box = node("w:txbxContent")
    box.append(deepcopy(title(setup(), "Z")._p))
    p.add_run()._r.append(box)
    before_a, before_t = actual.element.xml, template.element.xml
    findings = check(template, actual)
    assert {f.rule_id for f in findings} <= {"155", "156", "157", "158", "159"}
    assert any(UNREADABLE in f.message for f in findings)
    assert before_a == actual.element.xml and before_t == template.element.xml


def test_empty_document_and_bibliography_source_text_not_detected():
    assert check(setup(), setup()) == []
    template, actual = setup(), setup()
    title(actual)
    actual.add_paragraph("参考资料：文献[999]和未知数据来源。")
    findings = check(template, actual)
    assert not select(findings, "159")


def test_155_bibliography_is_not_body_reference_evidence():
    template, actual = setup(), setup()
    actual.add_paragraph("参考文献", "Heading 1")
    actual.add_paragraph("详见本文附录 C。")
    title(actual)
    failures = select(check(template, actual), "155")
    assert len(failures) == 1 and "未发现正文引用" in failures[0].actual


def test_159_unreadable_template_sample_does_not_choose_other_examples():
    template, actual = inherited_objects()
    template.paragraphs[-1].runs[0]._r.get_or_add_rPr().append(node("w:rPrChange", id="1"))
    actual.paragraphs[-1].runs[0].font.size = Pt(30)
    findings = check(template, actual)
    assert not select(findings, "159")
    assert any("模板附录示例含不可读格式" in f.actual for f in select(findings, "159", Status.UNKNOWN))


def test_malformed_numbering_is_caught_instead_of_crashing():
    template, actual = template_with_titles(), setup()
    p = title(actual)
    numbering = node("w:numPr")
    numbering.append(node("w:numId", val="broken"))
    p._p.get_or_add_pPr().append(numbering)
    findings = check(template, actual)
    assert findings and all(f.status == Status.UNKNOWN for f in findings)
    assert all(UNREADABLE in f.message for f in findings)


@pytest.mark.parametrize("declaration", [
    "附录图表格式不必与正文一致。", "附录图表格式无需与正文一致。",
    "附录图表格式不要求与正文一致。", "不要求附录图表格式与正文一致。",
    "附录图表格式与正文一致。附录图表格式不必与正文一致。",
    "附录图表格式无需与正文一致。附录图表格式与正文一致。",
])
def test_159_negated_or_conflicting_inheritance_is_not_a_mandatory_standard(declaration):
    template, actual = inherited_objects()
    template.paragraphs[0].text = declaration
    actual.paragraphs[-1].runs[0].font.size = Pt(30)
    findings = check(template, actual)
    assert not select(findings, "159")
    assert select(findings, "159", Status.UNKNOWN)


@pytest.mark.parametrize("declaration", [
    "附录编号不要求连续编号。", "附录无需连续编号。", "附录不必顺序编号。",
    "附录必须连续编号。附录编号不要求连续编号。",
    "附录允许跳号。附录应连续编号。",
])
def test_156_negated_or_conflicting_sequence_not_overridden_by_samples(declaration):
    template, actual = setup(), setup()
    template.add_paragraph(declaration)
    title(template, "C")
    title(template, "D")
    title(actual, "C")
    title(actual, "E")
    findings = check(template, actual)
    assert not select(findings, "156")
    assert any("否定或冲突" in f.actual for f in select(findings, "156", Status.UNKNOWN))


@pytest.mark.parametrize("text", ["附录C数据", "附录 C数据", "附录一 原始数据"])
def test_malformed_styled_appendix_heading_is_not_proof_of_missing_target(text):
    template, actual = template_with_titles(), setup()
    actual.add_paragraph("详细数据见本文附录 C。")
    p = actual.add_paragraph(text, "Heading 1")
    findings = check(template, actual)
    assert not select(findings, "155")
    assert select(findings, "155", Status.UNKNOWN)
    assert select(findings, "156", Status.UNKNOWN)
    assert any(f.anchor._p is p._p for f in select(findings, "157", Status.UNKNOWN))
