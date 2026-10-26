"""Public command round trip; opted in only when desktop Word is available."""
from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from thesis_format_checker.core.point_comments import content_tree as semantic_tree
from conftest import confirmed_heads, family_titles


@pytest.mark.integration
def test_public_cmd_and_native_word_reopen(runtime_dir):
    if os.environ.get("RUN_WORD_INTEGRATION") != "1":
        pytest.skip("设置 RUN_WORD_INTEGRATION=1 后运行真实Word测试")
    if os.name != "nt":
        pytest.skip("需要Windows桌面Word")
    from test_pipeline_headers_toc import fixture
    root = Path(__file__).resolve().parents[1]
    template, paper, output = [runtime_dir / name for name in ("模板.docx", "论文.docx", "批注.docx")]
    fixture(template, template=True)
    fixture(paper, template=False)
    before = [p.read_bytes() for p in (template, paper)]
    for phase, args in (
        ("compare", [str(root / "thesis-checker.cmd"), "compare", "--template", str(template), "--document", str(paper), "--output", str(output)]),
        ("reopen", [sys.executable, str(Path(__file__).resolve()), str(output)]),
    ):
        stdout, stderr = runtime_dir / (phase + "_stdout.log"), runtime_dir / (phase + "_stderr.log")
        with stdout.open("wb") as out, stderr.open("wb") as err:
            try:
                process = subprocess.run(args, cwd=root, env=os.environ.copy(), stdout=out, stderr=err, timeout=120)
            except subprocess.TimeoutExpired:
                pytest.fail("本地Word测试超时，已停止等待，不自动重试，不终止用户Word实例；请核对本次子进程状态")
        assert process.returncode == 0, stderr.read_text(encoding="utf-8", errors="replace")[-2500:]
    assert before == [p.read_bytes() for p in (template, paper)]
    comments = [c.text for c in Document(output).comments]
    heads = confirmed_heads(Document(output).comments)
    for rule in ("036", "039", "056", "064"):
        # 批注不再写规则编号：按该族的标题核对确定问题批注。
        assert heads & family_titles(rule), rule
    assert not any("既有检测暂未完成" in text for text in comments)
    assert "保留的人工批注" in comments
    with ZipFile(paper) as source, ZipFile(output) as annotated:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in source.namelist():
            if name not in allowed:
                assert source.read(name) == annotated.read(name), name
        assert semantic_tree(etree.fromstring(source.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))


if __name__ == "__main__":
    # A separate process bounds even a COM open/recovery-dialog stall. Owns only
    # the Word instance created by WordSession; never closes user's Word window.
    from thesis_format_checker.core.word_session import WordSession
    path = Path(sys.argv[1]).resolve()
    expected = len(Document(path).comments)
    with WordSession() as session:
        document = session.open_readonly(path)
        assert document.Comments.Count == expected
        assert document.Comments.Count > 1
        for index in range(1, document.Comments.Count + 1):
            assert document.Comments(index).Scope.Start <= document.Comments(index).Scope.End
    print("原生批注通过Word只读重开校验")
