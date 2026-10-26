from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.detectors.fields import scan_fields


def marker(paragraph, kind):
    r = paragraph.add_run()
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    r._r.append(node)


def instruction(paragraph, value):
    node = OxmlElement("w:instrText")
    node.text = value
    paragraph.add_run()._r.append(node)


def test_nested_toc_preserves_paragraphs_fields_and_cache_without_mutation():
    doc = Document()
    p = doc.add_paragraph()
    marker(p, "begin")
    instruction(p, 'TOC \\o "1-3" \\h')
    marker(p, "separate")
    entry = doc.add_paragraph("绪论\t")
    marker(entry, "begin")
    instruction(entry, 'PAGEREF _Toc100 \\h')
    marker(entry, "separate")
    entry.add_run("99")
    marker(entry, "end")
    blank = doc.add_paragraph()
    marker(blank, "end")
    before = doc.element.xml
    fields = scan_fields(doc.element.body)
    assert not fields.errors
    assert [f["kind"] for f in fields] == ["TOC", "PAGEREF"]
    assert fields[0]["has_nested"] and fields[1]["nested"]
    assert fields[0]["result_text"] == "绪论\t99"
    assert fields[1]["result_text"] == "99"
    assert len(fields[0]["result_paragraph_elements"]) == 3
    assert fields[1]["paragraph_element"] is entry._p
    assert before == doc.element.xml


def test_simple_and_unclosed_fields_are_read_without_evaluation():
    doc = Document()
    p = doc.add_paragraph()
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), "PAGE \\* roman")
    simple.set(qn("w:dirty"), "true")
    r, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "iv"
    r.append(t)
    simple.append(r)
    p._p.append(simple)
    marker(p, "begin")
    instruction(p, "STYLEREF 1")
    before = doc.element.xml
    fields = scan_fields(doc.element.body)
    assert fields[0]["kind"] == "PAGE"
    assert fields[0]["result_text"] == "iv" and fields[0]["dirty"]
    assert fields[1]["malformed"] and fields.errors
    assert before == doc.element.xml


def test_field_instructions_do_not_leak_into_outer_cached_text():
    doc = Document()
    p = doc.add_paragraph()
    marker(p, "begin")
    instruction(p, "IF ")
    marker(p, "begin")
    instruction(p, "PAGE")
    marker(p, "separate")
    p.add_run("2")
    marker(p, "end")
    instruction(p, ' = 2 "second" "other"')
    marker(p, "separate")
    p.add_run("second")
    marker(p, "end")
    fields = scan_fields(doc.element.body)
    assert fields[0]["result_text"] == "second"
    assert fields[1]["result_text"] == "2"
    assert not fields.errors


def test_fields_in_deleted_revision_are_excluded_and_orphans_reported():
    doc = Document()
    p = doc.add_paragraph()
    deleted = OxmlElement("w:del")
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), "PAGE")
    deleted.append(simple)
    p._p.append(deleted)
    marker(p, "end")
    fields = scan_fields(doc.element.body)
    assert not fields and fields.errors
