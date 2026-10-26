from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from thesis_format_checker.engine.pipeline import run_comparison, semantic_tree
from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.core.models import Finding, RunSummary, Status
from thesis_format_checker.checks.catalog import LEGACY_RULE_IDS
from thesis_format_checker.engine.pipeline import SECTION_RULES, check_sections, extract_baseline


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def field(paragraph, instruction, cache):
    for kind in ("begin", "instruction", "separate", "cache", "end"):
        run = paragraph.add_run()
        if kind == "instruction":
            node = OxmlElement("w:instrText")
            node.text = instruction
        elif kind == "cache":
            run.text = cache
            continue
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)


def fixture(path, *, template):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for key in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, key, Mm(25))
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Normal"]._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    head = section.header.paragraphs[0]
    head.text = "学校固定页眉"
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT if template else WD_ALIGN_PARAGRAPH.LEFT
    head.runs[0].font.name = "Arial"
    head.runs[0].font.size = Pt(11 if template else 16)
    foot = section.footer.paragraphs[0]
    foot.text = "学校固定页脚" if template else "多余内容"
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.runs[0].font.name = "Arial"
    foot.runs[0].font.size = Pt(10)
    title = doc.add_paragraph("目录")
    title.paragraph_format.line_spacing = 1.0
    title.runs[0].font.name = "Arial"
    title.runs[0].font.size = Pt(19 if template else 11)
    toc = doc.add_paragraph()
    if template:
        field(toc, 'TOC \\o "1-2" \\h', "第一章 绪论\t1")
    else:
        toc.add_run("第一章 绪论\t1")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_heading("1.1 方法", level=2)
    for index in range(20):
        p = doc.add_paragraph(f"这是第{index + 1}段用于格式对比和批注验证的正常正文，包含 English scientific research 123。")
        p.paragraph_format.first_line_indent = Pt(24)
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
    if not template:
        doc.add_comment(doc.paragraphs[-1].runs, "保留的人工批注", author="测试者")
    doc.save(path)


def test_exact_rule_coverage_and_baseline_namespace():
    """大类四、五、六（034–066）的规则全部有结论：docx 检测 + 旧版 Word 流程。"""
    covered = set()
    for section in ("headers", "headings", "toc"):
        covered |= set(SECTION_RULES[section])
    legacy_here = {"035", "045", "046", "047"}
    assert legacy_here <= set(LEGACY_RULE_IDS)
    assert covered | legacy_here == {f"{i:03}" for i in range(34, 67)}
    ctx = DocxContext(Document())
    old = {"previous": {"keep": 1}}
    data = extract_baseline(ctx, old)
    assert data["previous"] is old["previous"]
    assert set(old) == {"previous"}
    assert {"headers", "headings", "toc"} <= set(data)
    with pytest.raises(ValueError, match="未覆盖"):
        extract_baseline(ctx, data)


def test_detector_failure_isolated_and_commented(monkeypatch):
    """任一取证模块读取失败，只会把该组规则降级为“无法判定”。"""
    import thesis_format_checker.detectors.headers as headers
    import thesis_format_checker.detectors.headings as headings
    import thesis_format_checker.detectors.toc as toc

    def broken(*args, **kwargs):
        raise ValueError("object cannot be read")

    monkeypatch.setattr(headers, "check_header_footer_errors", broken)
    monkeypatch.setattr(headings, "check_heading_numbering_errors", broken)
    monkeypatch.setattr(toc, "check_toc_errors", broken)
    ctx = DocxContext(Document())
    sections = ("headers", "headings", "toc")
    result = check_sections(extract_baseline(ctx, sections=sections), ctx, sections=sections)
    assert len(result) == 29
    assert all(f.status == Status.UNKNOWN and UNREADABLE in f.message for f in result)


def test_full_pipeline_reuses_legacy_and_preserves_source_package(runtime_dir):
    a, b, out = [runtime_dir / name for name in ("template.docx", "paper.docx", "out.docx")]
    fixture(a, template=True)
    fixture(b, template=False)
    hashes = [sha256(p.read_bytes()).hexdigest() for p in (a, b)]
    calls = []
    def legacy(*args, **kwargs):
        calls.append(args)
        return RunSummary(findings=[Finding("045", "一级标题", Status.FAIL, "已有证据", expected="模板值", actual="实际值", paragraph_index=4)])
    summary = run_comparison(legacy, a, b, out)
    assert len(calls) == 1
    # 细粒度拆分后按规则族统计（036.3、039.1 …）
    families = {rule_id.split(".")[0] for rule_id in summary.counts_by_rule()}
    assert {"036", "039", "056", "064", "045"} <= families
    assert hashes == [sha256(p.read_bytes()).hexdigest() for p in (a, b)]
    with ZipFile(b) as before, ZipFile(out) as after:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in before.namelist():
            if name not in allowed:
                assert before.read(name) == after.read(name), name
        # 页眉页脚批注要锚在正文的某个字符上，允许“无损拆分 run”：用 content_tree
        # 校验（忽略批注、规范化相邻同格式 run），文字与格式仍必须一致。
        from thesis_format_checker.core.point_comments import content_tree

        assert content_tree(etree.fromstring(before.read("word/document.xml"))) == content_tree(
            etree.fromstring(after.read("word/document.xml")))
    loaded = Document(out)
    comments = list(loaded.comments)
    assert comments[0].text == "保留的人工批注"
    for comment in comments[1:]:
        assert "应该：" in comment.text and "实际：" in comment.text
        # 批注已锚定在出问题的文字上，正文里不再重复写“位置：”一行。
        assert "位置：" not in comment.text


def test_comments_on_sdt_toc_hyperlink_keep_body_and_bookmarks(runtime_dir):
    from test_detectors_toc import fixture as toc_fixture
    from thesis_format_checker.core.comments import add_word_comment
    from thesis_format_checker.engine.pipeline import save_comments_only
    a, b, out = [runtime_dir / name for name in ("template.docx", "paper.docx", "out.docx")]
    template, _, _, _ = toc_fixture(sdt=True)
    paper, _, entries, headings = toc_fixture(sdt=True)
    headings[0].runs[0].text = "修改后的绪论"
    paper.add_comment(entries[1].runs, "目录已有人工批注", author="读者")
    template.save(a)
    paper.save(b)
    loaded = Document(b)
    ctx = DocxContext(loaded)
    result = check_sections(extract_baseline(DocxContext(Document(a))), ctx)
    toc_findings = [f for f in result if family_of(f.rule_id) in {"058", "066"} and f.status == Status.FAIL]
    assert {f.rule_id for f in toc_findings} == {"058", "066"}
    for finding in toc_findings:
        add_word_comment(loaded, finding, finding.anchor)
    save_comments_only(b, loaded, out)
    assert len(Document(out).comments) == 1 + len(toc_findings)
    with ZipFile(b) as original, ZipFile(out) as annotated:
        assert semantic_tree(etree.fromstring(original.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))


def test_broken_header_reference_produces_limits_without_crashing(runtime_dir):
    a, b, out = [runtime_dir / name for name in ("template.docx", "paper.docx", "out.docx")]
    fixture(a, template=True)
    fixture(b, template=False)
    paper = Document(b)
    ref = paper.sections[0]._sectPr.find(qn("w:headerReference"))
    ref.set(qn("r:id"), "missing-header-part")
    paper.save(b)
    summary = run_comparison(lambda *a, **k: RunSummary(), a, b, out)
    assert out.exists()
    assert any(family_of(f.rule_id) == "034" and f.status == Status.UNKNOWN and UNREADABLE in f.message for f in summary.findings)


def test_035_adapter_failure_preserves_prior_evidence_and_other_rules(monkeypatch):
    from thesis_format_checker.engine.pipeline import adapt_legacy_findings
    import thesis_format_checker.detectors.headers as headers
    def broken(*args, **kwargs):
        raise ValueError("broken header object")
    monkeypatch.setattr(headers, "adapt_legacy_header_findings", broken)
    header = Finding("035", "页眉文字", Status.FAIL, "原比对证据：学校A与学校B\n已有值")
    heading = Finding("045", "一级标题", Status.FAIL, "字号差异")
    adapted = adapt_legacy_findings([header, heading], {}, DocxContext(Document()))
    assert adapted[0].status == Status.UNKNOWN and UNREADABLE in adapted[0].message
    assert "学校A与学校B" in adapted[0].message
    # 统一批注格式：抬头 + 应该 + 实际
    assert len(adapted[0].message.splitlines()) == 3
    assert header.status == Status.FAIL
    assert adapted[1] is heading
