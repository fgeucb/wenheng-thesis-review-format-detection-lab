from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Pt
from lxml import etree

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.notes import extract_notes_baseline, check_note_reference_errors
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def node(tag, **attrs):
    out = OxmlElement(tag)
    for key, value in attrs.items():
        out.set(qn("w:" + key), str(value))
    return out


def note_part(doc, kind="footnote", size=11, note_id=7, separator=True, text="示例注释 note 1", marker=True):
    root = node("w:" + kind + "s")
    if separator:
        part_note = node("w:" + kind, id=-1, type="separator")
        p = node("w:p")
        r = node("w:r")
        r.append(node("w:separator"))
        p.append(r)
        part_note.append(p)
        root.append(part_note)
    entry = node("w:" + kind, id=note_id)
    style = "Footnote Text" if kind == "footnote" else "Endnote Text"
    if style not in doc.styles:
        doc.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
    doc.styles[style].font.name = "Arial"
    doc.styles[style].font.size = Pt(size)
    fonts = doc.styles[style].element.get_or_add_rPr().get_or_add_rFonts()
    for key, value in (("eastAsia", "宋体"), ("ascii", "Arial"), ("hAnsi", "Arial")):
        fonts.set(qn("w:" + key), value)
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.0
    p.add_run(text)
    entry.append(deepcopy(p._p))
    p._p.getparent().remove(p._p)
    root.append(entry)
    part = Part(PackURI("/word/" + kind + "s.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml." + kind + "s+xml",
                etree.tostring(root), doc.part.package)
    doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + kind + "s")
    anchor = doc.add_paragraph("正文注释标记")
    if marker:
        anchor.add_run()._r.append(node("w:" + kind + "Reference", id=note_id))
    return anchor, part, root


def note_settings(doc, restart="continuous", fmt="decimal", start=1, section=None):
    parent = doc.settings.element if section is None else doc.sections[section]._sectPr
    props = parent.find(qn("w:footnotePr"))
    if props is not None:
        parent.remove(props)
    props = node("w:footnotePr")
    for tag, value in (("numFmt", fmt), ("numStart", start), ("numRestart", restart)):
        props.append(node("w:" + tag, val=value))
    parent.append(props)


def field(p, instruction="REF Fig1 \\h", cache="图1", simple=True):
    if simple:
        fld = node("w:fldSimple", instr=instruction)
        run = node("w:r")
        leaf = node("w:t")
        leaf.text = cache
        run.append(leaf)
        fld.append(run)
        p._p.append(fld)
        return fld
    p.add_run()._r.append(node("w:fldChar", fldCharType="begin"))
    instruction_node = node("w:instrText")
    instruction_node.text = instruction
    p.add_run()._r.append(instruction_node)
    p.add_run()._r.append(node("w:fldChar", fldCharType="separate"))
    p.add_run(cache)
    p.add_run()._r.append(node("w:fldChar", fldCharType="end"))
    return p._p


def bookmark(doc, text="图1", name="Fig1"):
    p = doc.add_paragraph()
    p._p.append(node("w:bookmarkStart", id=8, name=name))
    p.add_run(text)
    p._p.append(node("w:bookmarkEnd", id=8))
    return p


def check(template, document):
    return check_note_reference_errors(extract_notes_baseline(DocxContext(template)), DocxContext(document))


def select(items, rule, status=Status.FAIL):
    statuses = (Status.UNKNOWN, Status.NOT_APPLICABLE) if status == Status.UNKNOWN else (status,)
    return [f for f in items if family_of(f.rule_id) == rule and f.status in statuses]


def test_105_dynamic_note_types_and_body_anchor():
    a, b = Document(), Document()
    note_part(a)
    anchor, _, _ = note_part(b, kind="endnote")
    result = select(check(a, b), "105")
    assert len(result) == 1 and result[0].anchor._p is anchor._p
    assert "footnote" in result[0].expected
    note_part(a, kind="endnote")
    assert not select(check(a, b), "105")


def test_105_template_absence_is_unknown_not_guess_school():
    a, b = Document(), Document()
    note_part(b)
    results = check(a, b)
    assert not select(results, "105")
    assert select(results, "105", Status.UNKNOWN)


def test_106_global_and_section_overrides_are_read():
    a, b = Document(), Document()
    note_part(a)
    note_part(b)
    note_settings(a, restart="eachSect", fmt="upperRoman", start=4)
    note_settings(b, restart="continuous")
    assert select(check(a, b), "106")
    note_settings(b, restart="eachSect", fmt="upperRoman", start=4)
    assert not select(check(a, b), "106")
    note_settings(b, restart="continuous", section=0)
    assert select(check(a, b), "106")


def test_106_ids_are_not_display_numbers_and_page_restart_is_unknown():
    a, b = Document(), Document()
    note_part(a, note_id=2)
    note_part(b, note_id=900)
    assert not select(check(a, b), "106")
    note_settings(a, restart="eachPage")
    note_settings(b, restart="eachPage")
    results = check(a, b)
    assert not select(results, "106")
    assert select(results, "106", Status.UNKNOWN)


def test_106_custom_mark_and_malformed_number_setting_unknown():
    a, b = Document(), Document()
    note_part(a)
    anchor, _, _ = note_part(b)
    anchor._p.xpath(".//w:footnoteReference")[0].set(qn("w:customMarkFollows"), "1")
    assert select(check(a, b), "106", Status.UNKNOWN)
    note_settings(b, start="invalid")
    assert select(check(a, b), "106", Status.UNKNOWN)


def test_107_inherited_and_direct_note_format_and_preservation():
    a, b = Document(), Document()
    note_part(a, size=11)
    anchor, part, root = note_part(b, size=14)
    before = part.blob
    original = b.element.xml
    result = select(check(a, b), "107")
    assert result and result[0].anchor._p is anchor._p
    assert part.blob == before and b.element.xml == original
    b.styles["Footnote Text"].font.size = Pt(11)
    assert not select(check(a, b), "107")
    body = next(n for n in root if n.get(qn("w:id")) == "7")
    run = body.find("./" + qn("w:p") + "/" + qn("w:r"))
    props = node("w:rPr")
    props.append(node("w:sz", val=30))
    run.insert(0, props)
    part._blob = etree.tostring(root)
    assert select(check(a, b), "107")


def test_107_corrupt_part_does_not_crash_or_suppress_fields():
    a, b = Document(), Document()
    note_part(a)
    _, part, _ = note_part(b)
    part._blob = b"<broken"
    field(b.add_paragraph(), cache="错误！未找到引用源。")
    result = check(a, b)
    assert select(result, "107", Status.UNKNOWN)
    assert select(result, "110")
    assert UNREADABLE in select(result, "107", Status.UNKNOWN)[0].message


def test_107_separator_marker_and_paragraph_format():
    a, b = Document(), Document()
    note_part(a)
    _, part, root = note_part(b)
    sep = next(n for n in root if n.get(qn("w:type")) == "separator").find(qn("w:p"))
    props = node("w:pPr")
    props.append(node("w:ind", left=120))
    sep.insert(0, props)
    part._blob = etree.tostring(root)
    assert select(check(a, b), "107")


def test_108_plain_nonbibliographic_ref_cache_is_compared_to_bookmark():
    a, b = Document(), Document()
    bookmark(a, "图1")
    field(a.add_paragraph("参见"))
    bookmark(b, "图2")
    anchor = b.add_paragraph("参见")
    field(anchor, cache="图1")
    failure = select(check(a, b), "108")
    assert len(failure) == 1 and failure[0].anchor._p is anchor._p
    assert "已保存REF结果" in failure[0].actual
    assert "尚未刷新域" in failure[0].actual


def test_108_manual_reference_requires_template_mechanism_and_explicit_cue():
    a, b = Document(), Document()
    b.add_paragraph("参见图3-1")
    b.add_paragraph("图3-1 本图标题")
    assert not select(check(a, b), "108")
    field(a.add_paragraph("参见"), cache="图1-1")
    assert not select(check(a, b), "108")
    assert select(check(a, b), "108", Status.UNKNOWN)
    a.add_paragraph("应使用自动交叉引用")
    assert len(select(check(a, b), "108")) == 1


def test_108_number_only_field_with_external_label_detects_manual_reference():
    a, b = Document(), Document()
    field(a.add_paragraph("参见图"), cache="1")
    a.add_paragraph("应使用自动交叉引用")
    b.add_paragraph("参见图2")
    assert select(check(a, b), "108")


def test_108_field_result_is_not_mistaken_for_manual_reference():
    a, b = Document(), Document()
    bookmark(a)
    bookmark(b)
    field(a.add_paragraph("参见"))
    field(b.add_paragraph("参见"), simple=False)
    assert not select(check(a, b), "108")


def test_108_numbering_switch_missing_bookmark_and_revisions_are_unknown():
    a, b = Document(), Document()
    bookmark(a)
    field(a.add_paragraph("参见"))
    bookmark(b, "图3 title")
    field(b.add_paragraph("参见"), instruction="REF Fig1 \\n", cache="图1")
    result = check(a, b)
    assert not select(result, "108")
    assert select(result, "108", Status.UNKNOWN)


def test_108_bibliographic_refs_are_deferred():
    a, b = Document(), Document()
    bookmark(b, "[2] Reference", name="Bib")
    field(b.add_paragraph(), instruction="REF Bib", cache="[1]")
    results = check(a, b)
    assert not select(results, "108")
    # 103/104（引用序号格式与顺序）已实现：正文里的 “[2]” 会按新规则给出结论；
    # 仍未实现的是 100/101/102/109。
    assert not any(family_of(f.rule_id) in {"100", "101", "102", "109"} for f in results)


def test_110_only_domain_field_caches_trigger_not_plain_explanations():
    a, b = Document(), Document()
    b.add_paragraph("Word有时显示错误！未找到引用源。")
    assert not select(check(a, b), "110")
    anchor = b.add_paragraph()
    field(anchor, cache="Error! Reference source not found.")
    assert len(select(check(a, b), "110")) == 1
    assert select(check(a, b), "110")[0].anchor._p is anchor._p


def test_110_malformed_and_nested_field_are_unknown():
    a, b = Document(), Document()
    p = b.add_paragraph()
    p.add_run()._r.append(node("w:fldChar", fldCharType="begin"))
    instr = node("w:instrText")
    instr.text = "REF Fig1"
    p.add_run()._r.append(instr)
    assert select(check(a, b), "110", Status.UNKNOWN)


def test_source_xml_and_note_blobs_unchanged_for_all_checks():
    a, b = Document(), Document()
    _, ap, _ = note_part(a)
    _, bp, _ = note_part(b)
    bookmark(a)
    bookmark(b)
    field(a.add_paragraph("参见"))
    field(b.add_paragraph("参见"))
    before = (a.element.xml, b.element.xml, a.styles.element.xml, b.styles.element.xml, ap.blob, bp.blob)
    results = check(a, b)
    after = (a.element.xml, b.element.xml, a.styles.element.xml, b.styles.element.xml, ap.blob, bp.blob)
    assert before == after
    assert not select(results, "105") and not select(results, "106") and not select(results, "107")


def test_template_corrupt_notes_never_supplies_partial_type_or_format_baseline():
    a, b = Document(), Document()
    _, part, _ = note_part(a)
    part._blob = b"<broken"
    note_part(b, kind="endnote")
    result = check(a, b)
    assert not select(result, "105")
    assert select(result, "105", Status.UNKNOWN)


def test_template_revised_note_marker_cannot_define_allowed_type():
    a, b = Document(), Document()
    anchor, _, _ = note_part(a)
    wrapper = node("w:ins", id=1, author="tester")
    run = anchor._p[-1]
    anchor._p.remove(run)
    wrapper.append(run)
    anchor._p.append(wrapper)
    note_part(b, kind="endnote")
    result = check(a, b)
    assert not select(result, "105")
    assert select(result, "105", Status.UNKNOWN)


def test_auto_numbered_bookmark_target_is_unknown_even_numpr_before_bookmark():
    a, b = Document(), Document()
    field(a.add_paragraph("参见"), cache="图1")
    target = bookmark(b, "图2")
    num = node("w:numPr")
    num.append(node("w:numId", val=1))
    target._p.get_or_add_pPr().append(num)
    field(b.add_paragraph("参见"), cache="图1")
    results = check(a, b)
    assert not select(results, "108")
    assert select(results, "108", Status.UNKNOWN)


def test_template_section_variation_is_not_matched_by_same_index():
    a, b = Document(), Document()
    note_part(a)
    note_part(b)
    a.add_section(WD_SECTION_START.NEW_PAGE)
    b.add_section(WD_SECTION_START.NEW_PAGE)
    note_settings(a, restart="eachSect", section=1)
    results = check(a, b)
    assert not select(results, "106")
    assert select(results, "106", Status.UNKNOWN)


def test_manual_reference_excludes_bibliography_explanations_and_captions():
    a, b = Document(), Document()
    a.add_paragraph("应使用自动交叉引用")
    field(a.add_paragraph("参见"), cache="图1")
    b.add_paragraph("例如参见图2")
    b.add_paragraph("如图3", style="Caption")
    b.add_paragraph("参考文献")
    b.add_paragraph("参见图4")
    assert not select(check(a, b), "108")


def test_110_note_story_error_anchors_at_body_reference_and_preserves_part():
    a, b = Document(), Document()
    note_part(a)
    anchor, part, root = note_part(b)
    entry = next(n for n in root if n.get(qn("w:id")) == "7")
    p = entry.find(qn("w:p"))
    fld = node("w:fldSimple", instr="REF Missing")
    run = node("w:r")
    leaf = node("w:t")
    leaf.text = "错误！未找到引用源。"
    run.append(leaf)
    fld.append(run)
    p.append(fld)
    part._blob = etree.tostring(root)
    before = part.blob
    result = select(check(a, b), "110")
    assert result and result[0].anchor._p is anchor._p
    assert part.blob == before


def test_110_header_story_error_anchors_in_body_without_modifying_header():
    a, b = Document(), Document()
    body = b.add_paragraph("正文")
    header = b.sections[0].header.paragraphs[0]
    field(header, cache="错误！未找到引用源。")
    before = header._p.xml
    result = select(check(a, b), "110")
    assert result and result[0].anchor._p is body._p
    assert header._p.xml == before


def test_107_automatic_list_indent_is_unknown():
    a, b = Document(), Document()
    note_part(a)
    _, part, root = note_part(b)
    entry = next(n for n in root if n.get(qn("w:id")) == "7")
    p = entry.find(qn("w:p"))
    num = node("w:numPr")
    num.append(node("w:numId", val=1))
    p.get_or_add_pPr().append(num)
    part._blob = etree.tostring(root)
    results = check(a, b)
    assert not select(results, "107")
    assert select(results, "107", Status.UNKNOWN)


def test_107_template_missing_font_group_does_not_infer_from_partial_examples():
    a, b = Document(), Document()
    _, apart, root = note_part(a)
    entry = next(n for n in root if n.get(qn("w:id")) == "7")
    second = deepcopy(entry)
    second.set(qn("w:id"), "8")
    for leaf in second.iter(qn("w:t")):
        leaf.text = "中文注释"
    root.append(second)
    apart._blob = etree.tostring(root)
    a.add_paragraph("第二处标记").add_run()._r.append(node("w:footnoteReference", id=8))
    note_part(b, text="note 1", size=14)
    results = check(a, b)
    assert select(results, "107", Status.UNKNOWN)


def test_field_checks_do_not_mutate_unknown_or_header_relationships():
    a, b = Document(), Document()
    b.add_paragraph("正文")
    b.sections[0].header.paragraphs[0].add_run("页眉")
    before = [(rid, relation.reltype, relation.target_part.blob) for rid, relation in b.part.rels.items() if not relation.is_external]
    check(a, b)
    after = [(rid, relation.reltype, relation.target_part.blob) for rid, relation in b.part.rels.items() if not relation.is_external]
    assert before == after


def test_manual_reference_keeps_bibliography_subheadings_protected():
    a, b = Document(), Document()
    a.add_paragraph("应使用自动交叉引用")
    field(a.add_paragraph("参见"), cache="图1")
    b.add_paragraph("第十章 参考文献", style="Heading 1")
    b.add_paragraph("中文文献", style="Heading 2")
    b.add_paragraph("参见图2")
    b.add_paragraph("外文文献", style="Heading 2")
    b.add_paragraph("参见图3")
    b.add_paragraph("后续章节", style="Heading 1")
    anchor = b.add_paragraph("参见图4")
    failures = select(check(a, b), "108")
    assert len(failures) == 1 and failures[0].anchor._p is anchor._p


def test_manual_reference_unknown_bibliography_level_requires_explicit_end():
    a, b = Document(), Document()
    a.add_paragraph("应使用自动交叉引用")
    field(a.add_paragraph("参见"), cache="图1")
    b.add_paragraph("第10章 参考文献")
    b.add_paragraph("期刊文献", style="Heading 2")
    b.add_paragraph("参见图2")
    b.add_paragraph("致谢", style="Heading 1")
    anchor = b.add_paragraph("参见图3")
    failures = select(check(a, b), "108")
    assert len(failures) == 1 and failures[0].anchor._p is anchor._p


def test_template_revised_ref_cannot_be_discarded_to_define_manual_failure():
    a, b = Document(), Document()
    a.add_paragraph("应使用自动交叉引用")
    field(a.add_paragraph("参见"), cache="图1")
    p = a.add_paragraph("参见")
    bad = field(p, cache="图2")
    p._p.remove(bad)
    revision = node("w:del", id=5, author="tester")
    revision.append(bad)
    p._p.append(revision)
    b.add_paragraph("参见图3")
    results = check(a, b)
    assert not select(results, "108")
    assert select(results, "108", Status.UNKNOWN)


def test_template_broken_ref_cannot_be_discarded_to_define_manual_failure():
    a, b = Document(), Document()
    a.add_paragraph("应使用自动交叉引用")
    field(a.add_paragraph("参见"), cache="图1")
    p = a.add_paragraph()
    p.add_run()._r.append(node("w:fldChar", fldCharType="begin"))
    instruction = node("w:instrText")
    instruction.text = "REF broken"
    p.add_run()._r.append(instruction)
    b.add_paragraph("参见图3")
    results = check(a, b)
    assert not select(results, "108")
    assert select(results, "108", Status.UNKNOWN)


def _citation_paragraph(document, prefix, marker, *, superscript=True):
    paragraph = document.add_paragraph(prefix)
    run = paragraph.add_run(marker)
    run.font.superscript = superscript
    return paragraph


def test_citation_ranges_expand_and_numeric_intervals_are_not_references():
    template, document = Document(), Document()
    _citation_paragraph(document, "引用", "[1]")
    document.add_paragraph("概率值映射到 [0, 1] 区间。")
    _citation_paragraph(document, "连续引用", "[2-4]")
    gap = _citation_paragraph(document, "下一条引用", "[6]")
    findings = check(template, document)
    assert not select(findings, "103")
    missing = select(findings, "104")
    assert len(missing) == 1
    assert missing[0].anchor._p is gap._p
    assert missing[0].excerpt == "[6]"
    assert gap.text[missing[0].start:missing[0].end] == "[6]"
    assert "缺少序号：5" in missing[0].actual


def test_citation_format_difference_anchors_exact_outlier_marker():
    template, document = Document(), Document()
    _citation_paragraph(document, "引用", "[1]")
    _citation_paragraph(document, "引用", "[2]")
    outlier = _citation_paragraph(document, "引用", "[3]", superscript=False)
    findings = check(template, document)
    mixed = [finding for finding in select(findings, "103")
             if finding.rule_id == "103.2"]
    assert len(mixed) == 1
    assert mixed[0].anchor._p is outlier._p
    assert outlier.text[mixed[0].start:mixed[0].end] == "[3]"
