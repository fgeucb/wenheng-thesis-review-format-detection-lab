"""错误样本批处理工具的完整性验证回归测试。"""

import importlib.util
import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def _tool():
    spec = importlib.util.spec_from_file_location(
        "run_error_samples", ROOT / "tools" / "run_error_samples.py"
    )
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _finalize_tool():
    spec = importlib.util.spec_from_file_location(
        "finalize_error_sample_results",
        ROOT / "tools" / "finalize_error_sample_results.py",
    )
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _source_and_commented_copy(runtime_dir, text="【错误】测试问题\n应该：正确\n实际：错误"):
    source = runtime_dir / "source.docx"
    output = runtime_dir / "output.docx"
    document = Document()
    paragraph = document.add_paragraph("测试正文")
    document.save(source)

    revised = Document(source)
    revised.add_comment(
        revised.paragraphs[0].runs,
        text=text,
        author="论文格式检测",
        initials="TC",
    )
    revised.save(output)
    return source, output


def test_integrity_verifier_accepts_comment_only_change(runtime_dir):
    tool = _tool()
    source, output = _source_and_commented_copy(runtime_dir)

    result = tool.verify_output_integrity(
        source, output, source_sha256=tool._digest(source)
    )

    assert result["passed"]
    assert result["package"]["body_semantic_equal"]
    assert result["package"]["style_parts_identical"]
    assert result["package"]["media_identical"]
    assert result["package"]["sections_equal"]
    assert result["comments"]["anchors_complete"]


def test_integrity_verifier_rejects_report_only_text_in_comment(runtime_dir):
    tool = _tool()
    source, output = _source_and_commented_copy(
        runtime_dir,
        text="无法判定：只应保留在 JSON 报告中",
    )

    result = tool.verify_output_integrity(source, output)

    assert not result["passed"]
    assert result["comments"]["unknown_system_comments"] == 1


def test_integrity_verifier_rejects_body_change(runtime_dir):
    tool = _tool()
    source, output = _source_and_commented_copy(runtime_dir)
    revised = Document(output)
    revised.paragraphs[0].runs[0].text = "正文被改动"
    revised.save(output)

    result = tool.verify_output_integrity(source, output)

    assert not result["passed"]
    assert not result["package"]["body_semantic_equal"]


def test_result_finalizer_removes_one_comment_and_renames_another(runtime_dir):
    tool = _tool()
    finalizer = _finalize_tool()
    source = runtime_dir / "source.docx"
    output = runtime_dir / "output.docx"
    document = Document()
    document.add_paragraph("目    录")
    document.add_paragraph("正文")
    document.save(source)

    revised = Document(source)
    remove_text = (
        "【错误】标题与后续内容不在同一页\n"
        "应该：标题与其后的正文在同一页\n"
        "实际：标题在第 9 页末，后续内容在第 11 页"
    )
    revised.add_comment(
        revised.paragraphs[0].runs, text=remove_text,
        author="论文格式检测", initials="TC",
    )
    revised.add_comment(
        revised.paragraphs[1].runs,
        text="【警告】前置部分出现页眉\n应该：无内容\n实际：有文字",
        author="论文格式检测", initials="TC",
    )
    revised.save(output)

    removed, renamed = finalizer._patch_comment_package(output, {remove_text})
    texts = [item.text for item in Document(output).comments]

    assert (removed, renamed) == (1, 1)
    assert len(texts) == 1
    assert "前置部分页眉与模板不一致" in texts[0]
    assert tool.verify_output_integrity(source, output)["passed"]

