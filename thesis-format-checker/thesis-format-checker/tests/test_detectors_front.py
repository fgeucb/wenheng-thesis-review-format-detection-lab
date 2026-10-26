import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from thesis_format_checker.core.docx_context import DocxContext, effective_paragraph, effective_runs
from thesis_format_checker.detectors.front import (
    _declaration_lines, check_front_errors, extract_front_baseline,
)
from thesis_format_checker.core.models import Status


def formatted(paragraph, cn="楷体", latin="Arial", size=13, bold=False):
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = latin
        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cn)
        run.font.size = Pt(size)
        run.bold = bold
    return paragraph


def paragraph(doc, text, **kwargs):
    return formatted(doc.add_paragraph(text), **kwargs)


def inspect(template, document):
    template_before = template.element.xml
    document_before = document.element.xml
    baseline = extract_front_baseline(DocxContext(template))
    findings = check_front_errors(baseline, DocxContext(document))
    assert template.element.xml == template_before
    assert document.element.xml == document_before
    for finding in findings:
        lines = finding.message.splitlines()
        # 批注抬头只写等级与问题标题，不写规则编号。
        assert lines[0].startswith("【") and finding.title in lines[0]
        assert f"｜{finding.rule_id}" not in lines[0]
        assert lines[1].startswith("应该：")
        assert lines[2].startswith("实际：")
    return findings


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（013.3 → 013）。"""

    return rule_id.split(".")[0]


def failed(findings, rule):
    """按规则族查询失败项：族编号或其细粒度规则都算命中。"""

    return [f for f in findings if family_of(f.rule_id) == rule and f.status == Status.FAIL]


def with_text(findings, rule, text):
    """按规则族 + 文案片段查询结论。"""

    return [f for f in findings if family_of(f.rule_id) == rule and text in f.message]


def test_title_compares_internal_values_and_dynamic_font():
    template, document = Document(), Document()
    paragraph(template, "论文题目：示例题目A", cn="仿宋", latin="Courier New", size=19, bold=True)
    cover = paragraph(document, "论文题目：本人的题目B", cn="仿宋", latin="Courier New", size=19, bold=True)
    paragraph(document, "题名页")
    paragraph(document, "论文题目：本人的题目B", cn="仿宋", latin="Courier New", size=19, bold=True)
    assert not failed(inspect(template, document), "012")
    assert not failed(inspect(template, document), "013")
    cover.runs[0].font.size = Pt(11)
    document.paragraphs[-1].runs[0].text = "论文题目：另一个题目C"
    result = inspect(template, document)
    assert failed(result, "012")
    assert failed(result, "013")
    assert "19" in str(failed(result, "013")[0].expected)


def test_unlabelled_cover_title_placeholder_detects_missing_title():
    template, document = Document(), Document()
    paragraph(template, "硕 士 学 位 论 文")
    paragraph(template, "模板示例题目")
    paragraph(template, "（题名和副题名）")
    paragraph(template, "硕士学位论文")
    paragraph(template, "模板示例题目")
    paragraph(template, "作    者：×××")

    paragraph(document, "硕 士 学 位 论 文")
    marker = paragraph(document, "（题名和副题名）")
    paragraph(document, "硕士学位论文")
    paragraph(document, "论文实际题目")
    paragraph(document, "作    者：张三")

    missing = failed(inspect(template, document), "012")
    assert len(missing) == 1
    assert missing[0].rule_id == "012.1"
    assert missing[0].anchor._p is marker._p


@pytest.mark.parametrize("label,rule", [("论文题目", "012"), ("指导教师", "015"), ("提交日期", "016"), ("密级", "019")])
def test_template_placeholder_makes_field_required(label, rule):
    template, document = Document(), Document()
    paragraph(template, label + "：________")
    paragraph(document, "这是待检测文档")
    assert failed(inspect(template, document), rule)


def test_guide_style_field_without_colon_is_recognized_as_required():
    """格式指南用大段空格排版字段时，不能因没有冒号而漏掉必填项。"""

    template, document = Document(), Document()
    paragraph(template, "学科名称            ______")
    paragraph(document, "论文题目：测试论文")

    findings = inspect(template, document)

    assert failed(findings, "014")
    assert any(item.rule_id == "014.4" for item in failed(findings, "014"))


def test_sample_value_does_not_make_missing_field_mandatory():
    template, document = Document(), Document()
    paragraph(template, "指导教师：张教授")
    paragraph(document, "论文题目：课题")
    results = inspect(template, document)
    assert not failed(results, "015")
    assert with_text(results, "015", "未明确必填")


def test_optional_and_conflicting_requirement_are_not_failed():
    for directive in ("密级可选", "密级必填，但密级可选"):
        template, document = Document(), Document()
        paragraph(template, "密级：________")
        paragraph(template, directive)
        paragraph(document, "论文题目：课题")
        assert not failed(inspect(template, document), "019")


def test_adjacent_table_cells_recognized_and_location_change():
    template, document = Document(), Document()
    for doc in (template, document):
        row = doc.add_table(rows=1, cols=2).rows[0]
        formatted(row.cells[0].paragraphs[0])
        row.cells[0].paragraphs[0].add_run("分类号")
        row.cells[1].paragraphs[0].add_run("TP391")
        formatted(row.cells[1].paragraphs[0])
    assert not failed(inspect(template, document), "019")
    document.tables[0].cell(0, 0).text = "分类号：TP391"
    document.tables[0].cell(0, 1).text = ""
    result = inspect(template, document)
    assert failed(result, "019")
    assert failed(result, "018")


def test_advisor_fields_use_role_font_not_template_name():
    template, document = Document(), Document()
    paragraph(template, "指导教师：张教授", cn="宋体", size=15)
    advisor = paragraph(document, "指导教师：李教授", cn="宋体", size=15)
    assert not failed(inspect(template, document), "015")
    advisor.runs[0].bold = True
    assert failed(inspect(template, document), "015")


@pytest.mark.parametrize("pattern,good,bad", [
    ("YYYY年MM月DD日", "2027年03月09日", "2027/03/09"),
    ("2026-09-07", "2028-05-11", "2028-5-11"),
    ("2026年9月", "2029年11月", "2029年13月"),
])
def test_date_format_is_dynamic_and_checks_calendar(pattern, good, bad):
    template, document = Document(), Document()
    paragraph(template, "提交日期：" + pattern)
    target = paragraph(document, "提交日期：" + good)
    assert not failed(inspect(template, document), "016")
    target.runs[0].text = "提交日期：" + bad
    assert failed(inspect(template, document), "016")


def test_cover_alignment_dynamic_and_unavailable_coordinates_reported():
    template, document = Document(), Document()
    p = paragraph(template, "作者姓名：张三")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(21)
    paragraph(document, "作者姓名：李四")
    result = inspect(template, document)
    assert failed(result, "018")
    assert any(family_of(f.rule_id) == "018" and f.status == Status.UNKNOWN and "实际页面坐标" in f.message
               for f in result)


def test_abstract_languages_compared_independently():
    template, document = Document(), Document()
    for doc in (template, document):
        paragraph(doc, "摘要", cn="黑体", size=19)
        paragraph(doc, "研究摘要。", cn="宋体", size=12)
        paragraph(doc, "关键词：检测；模板", cn="楷体", size=11)
        paragraph(doc, "Abstract", latin="Arial", size=18)
        paragraph(doc, "This is the abstract.", latin="Times New Roman", size=13)
        paragraph(doc, "Key words: checking; template", latin="Courier New", size=10)
        doc.add_heading("1 正文", level=1)
    assert not failed(inspect(template, document), "021")
    document.paragraphs[4].runs[0].font.size = Pt(15)
    result = failed(inspect(template, document), "021")
    assert result and any("英文摘要正文" in f.message for f in result)


def test_textbox_presence_prevents_false_missing_and_has_limit_comment():
    template, document = Document(), Document()
    paragraph(template, "指导教师：________")
    p = paragraph(document, "封面")
    textbox = OxmlElement("w:txbxContent")
    box_p = OxmlElement("w:p")
    box_r = OxmlElement("w:r")
    box_t = OxmlElement("w:t")
    box_t.text = "指导教师：王老师"
    box_r.append(box_t)
    box_p.append(box_r)
    textbox.append(box_p)
    p._p.append(textbox)
    results = inspect(template, document)
    assert not failed(results, "015")
    assert with_text(results, "015", "该对象docx库无法读取，本条暂无法执行检测")


def test_date_field_is_unknown_instead_of_false_format_failure():
    template, document = Document(), Document()
    paragraph(template, "提交日期：YYYY年MM月DD日")
    target = paragraph(document, "提交日期：2026年09月27日")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    target._p.append(field)
    results = inspect(template, document)
    assert not failed(results, "016")
    assert with_text(results, "016", "该对象docx库无法读取，本条暂无法执行检测")


def test_no_semantic_or_reference_rule_is_emitted():
    results = inspect(Document(), Document())
    # 014/020 已改为可确定性检测（封面基本信息、声明签名/日期）。
    assert set(family_of(f.rule_id) for f in results) <= {"012", "013", "014", "015", "016", "018", "019", "020", "021"}


def test_corrupt_property_is_unknown_and_no_crash():
    template, document = Document(), Document()
    paragraph(template, "论文题目：题目A")
    target = paragraph(document, "论文题目：题目A")
    target._p.get_or_add_pPr().get_or_add_jc().set(qn("w:val"), "bogus")
    results = inspect(template, document)
    assert any(f.status == Status.UNKNOWN and "该对象docx库无法读取，本条暂无法执行检测" in f.message for f in results)


def test_following_requirement_note_is_not_consumed_as_field_value():
    template, document = Document(), Document()
    paragraph(template, "指导教师：")
    paragraph(template, "指导教师必填")
    paragraph(document, "论文题目：课题")
    baseline = extract_front_baseline(DocxContext(template))
    assert baseline["fields"][0]["value"] == ""
    assert failed(inspect(template, document), "015")


def test_multiple_template_advisors_are_not_forced_to_first_format():
    template, document = Document(), Document()
    paragraph(template, "指导教师：李老师", cn="楷体")
    paragraph(template, "指导教师：王老师", cn="宋体")
    paragraph(document, "指导教师：王老师", cn="宋体")
    findings = inspect(template, document)
    assert not failed(findings, "015")
    assert with_text(findings, "015", "多个")


def test_populated_brackets_are_not_mistaken_for_placeholders():
    template, document = Document(), Document()
    paragraph(template, "论文题目：〈论文题目〉")
    paragraph(template, "密级：〈密级〉")
    paragraph(document, "论文题目：（实验研究）")
    paragraph(document, "密级：（公开）")
    results = inspect(template, document)
    assert not failed(results, "012")
    assert not failed(results, "019")


def test_manual_chinese_title_line_break_does_not_cause_content_mismatch():
    template, document = Document(), Document()
    paragraph(template, "论文题目：〈论文题目〉")
    paragraph(document, "论文题目：格式检测\n研究")
    paragraph(document, "题名页")
    paragraph(document, "论文题目：格式检测研究")
    assert not failed(inspect(template, document), "012")


def test_bold_style_toggles_and_direct_run_override():
    doc = Document()
    parent = doc.styles.add_style("FrontBoldBase", WD_STYLE_TYPE.PARAGRAPH)
    parent.font.bold = True
    child = doc.styles.add_style("FrontBoldChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    child.font.bold = True
    p = doc.add_paragraph("论文题目：研究", style=child)
    assert effective_runs(p)[0][1]["bold"] is False
    char = doc.styles.add_style("FrontBoldChar", WD_STYLE_TYPE.CHARACTER)
    char.font.bold = True
    p.runs[0].style = char
    assert effective_runs(p)[0][1]["bold"] is True
    p.runs[0].bold = False
    assert effective_runs(p)[0][1]["bold"] is False
    p.runs[0].bold = True
    assert effective_runs(p)[0][1]["bold"] is True


def test_false_bold_in_style_does_not_invert_inherited_bold():
    doc = Document()
    parent = doc.styles.add_style("FrontOn", WD_STYLE_TYPE.PARAGRAPH)
    parent.font.bold = True
    child = doc.styles.add_style("FrontNoToggle", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    child.font.bold = False
    p = doc.add_paragraph("题目", style=child)
    assert effective_runs(p)[0][1]["bold"] is True


@pytest.mark.parametrize("xml_attr,key", [
    ("firstLineChars", "first_line_indent"), ("hangingChars", "first_line_indent"),
    ("leftChars", "left_indent"), ("startChars", "left_indent"),
    ("rightChars", "right_indent"), ("endChars", "right_indent"),
])
def test_character_unit_indents_convert_to_points(xml_attr, key):
    """字符单位缩进按“1 字符 = 当前字号”换算成磅值（不再判成不可读）。"""

    doc = Document()
    p = doc.add_paragraph("正文")
    p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:" + xml_attr), "200")
    # 默认字号 11 磅：2 个字符 = 22 磅；hangingChars 为负值。
    size = effective_runs(p)[0][1]["size_pt"]
    expected = -2 * size if xml_attr == "hangingChars" else 2 * size
    assert abs(effective_paragraph(p)[key] - expected) < 0.01


@pytest.mark.parametrize("xml_attr,key", [("start", "left_indent"), ("end", "right_indent")])
def test_twips_only_indents_are_unknown(xml_attr, key):
    doc = Document()
    p = doc.add_paragraph("正文")
    p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:" + xml_attr), "200")
    assert effective_paragraph(p)[key] is None


@pytest.mark.parametrize("xml_attr,key", [
    ("beforeLines", "space_before"), ("afterLines", "space_after"),
    ("beforeAutospacing", "space_before"), ("afterAutospacing", "space_after"),
])
def test_unparsed_paragraph_line_spacing_units_are_unknown(xml_attr, key):
    doc = Document()
    p = doc.add_paragraph("正文")
    p._p.get_or_add_pPr().get_or_add_spacing().set(qn("w:" + xml_attr), "100" if "Lines" in xml_attr else "1")
    assert effective_paragraph(p)[key] is None


def test_front_reports_unknown_for_unparsed_indent_without_false_failure():
    template, document = Document(), Document()
    paragraph(template, "指导教师：张老师").paragraph_format.first_line_indent = Pt(20)
    p = paragraph(document, "指导教师：李老师")
    p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")
    results = inspect(template, document)
    # 模板首行缩进 20 磅、论文为 2 字符（不同字号换算出不同磅值）：给出确定结论，
    # 不再因为“字符单位”而报“无法判定”。
    assert not any(family_of(f.rule_id) == "015" and f.status == Status.UNKNOWN for f in results)
    assert any(family_of(f.rule_id) in {"015", "018"} for f in results)


def test_declaration_region_ends_at_authorization_page():
    """原创性声明缺签名时，不能把授权声明页的签名当成它的签名。"""

    document = Document()
    document.add_paragraph("论文原创性声明")
    document.add_paragraph("本人郑重声明：所呈交的论文是本人在导师指导下独立完成的研究成果。")
    document.add_paragraph("使用授权声明")
    document.add_paragraph("论文作者签名：张三")
    document.add_heading("摘  要", 1)
    lines = _declaration_lines(DocxContext(document))
    assert "author_signature" not in lines


def test_declaration_region_ends_at_front_matter_boundary():
    """摘要（前置区结束标志）之后的“签名：”不属于声明页区域。"""

    document = Document()
    document.add_paragraph("论文原创性声明")
    document.add_paragraph("本人郑重声明：所呈交的论文是本人在导师指导下独立完成的研究成果。")
    document.add_paragraph("摘  要")  # 未套标题样式，只能靠前置区边界收敛区域
    document.add_paragraph("论文作者签名：张三")
    lines = _declaration_lines(DocxContext(document))
    assert "author_signature" not in lines


def _with_signature_image(paragraph):
    drawing, inline, extent = (OxmlElement(tag) for tag in ("w:drawing", "wp:inline", "wp:extent"))
    extent.set("cx", "1000000")
    extent.set("cy", "400000")
    inline.append(extent)
    drawing.append(inline)
    paragraph.add_run()._r.append(drawing)
    return paragraph


def test_signature_image_counts_as_filled_and_advisor_lines_are_checked():
    """作者签名是图片时不算缺失；授权声明的导师签名/日期为空必须报出来。"""

    template, document = Document(), Document()
    for doc in (template, document):
        doc.add_paragraph("论文原创性声明")
        doc.add_paragraph("本人郑重声明：论文系本人独立完成。")
        doc.add_paragraph("论文作者签名：________    日期：____年__月__日")
        doc.add_paragraph("使用授权声明")
        doc.add_paragraph("论文作者签名：__________    导师签名：__________")
        doc.add_paragraph("日期：____年__月__日        日期：____年__月__日")
    document.paragraphs[2].text = "论文作者签名：            日期：2026年5月10日"
    _with_signature_image(document.paragraphs[2])   # 作者手写签名扫描件
    document.paragraphs[4].text = "论文作者签名：            导师签名："
    _with_signature_image(document.paragraphs[4])   # 授权页作者签名也是图片
    document.paragraphs[5].text = "日期：2026年5月10日        日期：     年    月    日"

    findings = check_front_errors(extract_front_baseline(DocxContext(template)),
                                  DocxContext(document))
    codes = [item.rule_id for item in findings if item.status == Status.FAIL]
    messages = " ".join(item.actual for item in findings if item.status == Status.FAIL)
    assert "020.1" in codes and "导师签名" in messages      # 导师签名没填
    assert "020.2" in codes and "导师日期" in messages      # 导师日期没填
    assert "作者签名未填写" not in messages                 # 图片签名不算缺失


def test_statement_heading_and_unlabelled_dates_keep_page_ownership():
    """“声 明”与无“日期：”年月日行仍应区分原创性声明和授权声明。"""

    template, document = Document(), Document()
    for target in (template, document):
        target.add_paragraph("声 明")
        target.add_paragraph("本人声明论文系本人独立完成。")
        target.add_paragraph("研究生签名：________")
        target.add_paragraph("年    月    日")
        target.add_paragraph("学位论文使用授权声明")
        target.add_paragraph("研究生签名：________    导师签名：________")
        target.add_paragraph("年    月    日            年    月    日")

    document.paragraphs[2].text = "研究生签名：张三"
    document.paragraphs[3].text = "2026年5月10日"
    document.paragraphs[5].text = "研究生签名：张三    导师签名："
    document.paragraphs[6].text = "2026年5月10日        年    月    日"

    findings = check_front_errors(
        extract_front_baseline(DocxContext(template)), DocxContext(document)
    )
    failures = [item for item in findings if item.status == Status.FAIL]

    assert {item.title for item in failures} == {
        "使用授权声明缺少签名", "使用授权声明缺少日期"
    }
    assert all("原创性声明" not in item.title for item in failures)


def test_authorization_missing_fields_have_correct_title_and_field_range():
    from docx import Document
    from thesis_format_checker.core.docx_context import DocxContext
    from thesis_format_checker.detectors.front import _declaration_checks
    doc = Document()
    doc.add_paragraph("使用授权声明")
    doc.add_paragraph("作者签名：张三    导师签名：      ")
    doc.add_paragraph("日期：2026年5月1日    日期：年  月  日")
    findings = _declaration_checks({"declaration": {
        "authorization_advisor_signature": True, "authorization_advisor_date": True,
    }}, DocxContext(doc))
    assert {f.title for f in findings} == {"使用授权声明缺少签名", "使用授权声明缺少日期"}
    assert all("原创性" not in f.message for f in findings)
    assert findings[0].anchor.text[findings[0].start:findings[0].end].startswith("导师签名")
    assert "2026" not in findings[1].anchor.text[findings[1].start:findings[1].end]

