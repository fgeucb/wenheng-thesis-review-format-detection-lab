from copy import deepcopy
import os

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches

from thesis_format_checker.core.comments import add_comment, add_word_comment
from thesis_format_checker.core.docx_context import DocxContext, make_finding, unknown
from thesis_format_checker.core.models import Finding, RunSummary, Status
from thesis_format_checker.core.point_comments import add_boundary_comment, add_point_comment, content_tree, marker_offset
from thesis_format_checker.core.word_layout import BodyPoint, Page, Pagination, read_pagination
from thesis_format_checker.detectors.structure import _check_order
from thesis_format_checker.engine.pipeline import dedupe_findings, run_comparison, save_comments_only, write_page_comments


@pytest.mark.parametrize("status,message", [
    (Status.UNKNOWN, "读取受限"), (Status.FAIL, "无法判定"),
    (Status.SKIPPED, "待实现"), (Status.NOT_APPLICABLE, "不适用"),
])
def test_all_finding_writers_reject_undetermined(status, message):
    d = Document()
    p = d.add_paragraph("锚点")
    f = Finding("010.1", "顺序", status, message, paragraph_index=1)
    add_word_comment(d, f, p)
    # 门禁必须在访问 COM 前生效；object 不提供任何 Word 属性。
    add_comment(object(), f)
    assert len(d.comments) == 0
    assert not d.element.body.xpath(".//w:commentRangeStart")


@pytest.mark.parametrize("writer", [add_point_comment, add_boundary_comment])
def test_low_level_text_writers_reject_unknown_without_splitting(writer):
    d = Document()
    p = d.add_paragraph("无损原文")
    before = p._p.xml
    assert writer(d, p, 1, "【无法判定】测试") is None
    assert before == p._p.xml and len(d.comments) == 0


def test_missing_anchor_never_becomes_document_start():
    d = Document()
    d.add_paragraph("封面题目")
    add_word_comment(d, make_finding("010.1", "顺序", "甲乙", "乙甲"))
    class COM:
        class Comments:
            @staticmethod
            def Add(**kwargs):
                pytest.fail("不得写到文首")
    add_comment(COM(), make_finding("010.1", "顺序", "甲乙", "乙甲"))
    assert len(d.comments) == 0


@pytest.mark.parametrize("pieces", [("前文AUC与F1后文",), ("前文A", "U", "C与F1后文")])
def test_text_range_comments_select_only_target_across_runs(pieces):
    d = Document()
    p = d.add_paragraph()
    for i, piece in enumerate(pieces):
        p.add_run(piece).bold = bool(i % 2)
    before = content_tree(p._p)
    for token in ("AUC", "F1"):
        f = make_finding("149.1", "全称", "给出全称", token, p)
        f.start = p.text.index(token)
        f.end = f.start + len(token)
        comment = add_word_comment(d, f, p)
        start = next(n for n in p._p.xpath("./w:commentRangeStart")
                     if n.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") == str(comment.comment_id))
        end = next(n for n in p._p.xpath("./w:commentRangeEnd")
                   if n.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") == str(comment.comment_id))
        assert p.text[marker_offset(p._p, start):marker_offset(p._p, end)] == token
    assert content_tree(p._p) == before


def test_visible_page_problem_annotates_each_matching_page_once(tmp_path):
    d, ctx, pages = context_with_pages()
    pages[3].section = 0
    f = page_finding(page_variant="even", page_scope="variant")
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "input", d, [f, deepcopy(f)])
    assert not unplaced
    assert [x["page"] for x in written] == [3, 4]
    assert len(d.comments) == 2


def test_distance_transition_and_direct_rule_share_one_comment_per_page(tmp_path):
    d, ctx, pages = context_with_pages()
    first = page_finding(page_scope="section", page_issue="header_distance:10:20")
    first.rule_id = "027.1"
    transition = deepcopy(first)
    transition.rule_id = "029"
    transition.expected, transition.actual = "与前节一致", "与前节距离不同"
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "input", d,
                                           [first, transition])
    assert not unplaced
    assert [x["page"] for x in written] == [1, 2, 3]
    assert all(x["rules"] == ["027.1", "029"] for x in written)


def test_tab_stop_properties_do_not_count_as_text_offsets():
    from docx.shared import Inches
    d = Document()
    p = d.add_paragraph("AUC")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1))
    add_boundary_comment(d, p, 0, "页眉")
    assert marker_offset(p._p, p._p.xpath("./w:commentRangeStart")[0]) == 0


def test_balanced_field_outer_edges_do_not_touch_field_code():
    from thesis_format_checker.core.xmlutil import qn
    d = Document()
    p = d.add_paragraph()
    for kind in ("begin", "separate"):
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        p.add_run()._r.append(node)
    p.add_run("目录条目")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)
    before = content_tree(p._p)
    add_boundary_comment(d, p, 0, "页眉", edge="start")
    add_boundary_comment(d, p, len(p.text), "页脚", edge="end")
    assert before == content_tree(p._p)
    assert len(d.comments) == 2


def test_exact_text_beside_signature_image_preserves_image():
    from thesis_format_checker.core.xmlutil import qn
    d = Document()
    p = d.add_paragraph("作者签名：")
    p.add_run()._r.append(OxmlElement("w:drawing"))
    p.add_run("  导师签名：  ")
    before = content_tree(p._p)
    f = make_finding("020.1", "使用授权声明缺少签名", "填写", "空白", p)
    f.start, f.end = p.text.index("导师"), len(p.text)
    add_word_comment(d, f, p)
    start = p._p.find(qn("w:commentRangeStart"))
    assert marker_offset(p._p, start) == f.start
    assert content_tree(p._p) == before


def test_order_unknown_has_no_anchor_and_confirmed_inversion_has_real_component():
    d = Document()
    d.add_paragraph("")  # 无论首段是否有校徽，都不能成为顺序限制说明的锚点
    toc = d.add_paragraph("目录")
    d.add_paragraph("摘要")
    ctx = DocxContext(d)
    data = {"order_conflict": False, "order": [], "samples": {"toc": 0, "abstract_zh": 1}}
    positions = {"toc": 1, "abstract_zh": 2}
    f = _check_order(data, ctx, positions)[0]
    assert f.status == Status.UNKNOWN and f.anchor is None
    data["order_conflict"] = True
    assert _check_order(data, ctx, positions)[0].anchor is None
    data.update(order_conflict=False, order=[["abstract_zh", "toc"]])
    f = _check_order(data, ctx, positions)[0]
    assert f.status == Status.FAIL and f.anchor._p is toc._p


def context_with_pages():
    d = Document()
    for n in range(4):
        d.add_paragraph(f"页{n}的顶部和底部文字")
    ctx = DocxContext(d)
    ctx.render_facts = {"total_pages": 4}
    # 文件页序与印刷页码奇偶相反，第一张印刷偶数页又恰为首页变体。
    pages = [
        Page(1, 0, "first", "first", BodyPoint(0, 0, 0), BodyPoint(0, 10, 10), adjusted_number=2),
        Page(2, 0, "default", "default", BodyPoint(1, 0, 100), BodyPoint(1, 10, 110), adjusted_number=3),
        Page(3, 0, "even", "even", BodyPoint(2, 0, 200), BodyPoint(2, 10, 210), adjusted_number=4),
        Page(4, 1, "even", "even", BodyPoint(3, 0, 300), BodyPoint(3, 10, 310), adjusted_number=6),
    ]
    ctx.render_layout = {"pagination": Pagination(pages=pages, boundaries_verified=True)}
    return d, ctx, pages


def page_finding(story="header", **kwargs):
    return Finding("037.2", "奇偶页设置", Status.FAIL, "【警告】奇偶页设置",
                   expected="相同", actual="不同", section_index=1, page_story=story, **kwargs)


def test_even_setting_uses_printed_parity_excludes_first_and_deduplicates(tmp_path):
    d, ctx, _ = context_with_pages()
    a = page_finding(page_scope="even_pages", page_skip_first=True, page_config="odd_even:0:1")
    b = deepcopy(a)
    b.section_index = 2
    b.page_skip_first = False
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "in.docx", d, [a, b])
    assert not unplaced and len(written) == len(d.comments) == 1
    assert written[0]["page"] == 3
    assert written[0]["affected_pages"] == [3, 4]
    assert a.affected_pages == [3] and b.affected_pages == [4]
    assert list(d.comments)[0].text.splitlines() == ["【警告】奇偶页设置", "应该：相同", "实际：不同"]


@pytest.mark.parametrize("variant,page", [("default", 2), ("first", 1), ("even", 3)])
def test_exact_header_variant_is_not_section_first(tmp_path, variant, page):
    d, ctx, _ = context_with_pages()
    written, _ = write_page_comments(RunSummary(), ctx, tmp_path / "in", d,
                                    [page_finding(page_scope="variant", page_variant=variant)])
    assert written[0]["page"] == page


def test_unused_header_variant_is_not_an_error_or_unplaced_comment(tmp_path):
    d, ctx, pages = context_with_pages()
    pages[3].header_variant = "default"
    finding = page_finding(page_scope="variant", page_variant="even")
    finding.section_index = 2
    summary = RunSummary()

    written, unplaced = write_page_comments(
        summary, ctx, tmp_path / "in", d, [finding])

    assert not written and not unplaced and not d.comments
    assert finding.status == Status.NOT_APPLICABLE
    assert summary.not_applicable


def test_header_content_difference_is_written_on_every_affected_page(tmp_path):
    d, ctx, pages = context_with_pages()
    pages.append(Page(
        5, 0, "default", "default",
        BodyPoint(3, 0, 400), BodyPoint(3, 10, 410), adjusted_number=7))
    finding = page_finding(page_scope="variant", page_variant="default")
    written, unplaced = write_page_comments(
        RunSummary(), ctx, tmp_path / "in", d, [finding])
    assert not unplaced
    assert [item["page"] for item in written] == [2, 5]
    assert len(d.comments) == 2


def test_footer_is_zero_length_at_bottom_and_never_falls_back_to_top(tmp_path):
    d, ctx, pages = context_with_pages()
    before = content_tree(d.element)
    f = page_finding("footer", page_variant="even")
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "in", d, [f])
    assert not unplaced
    assert written[0]["offset"] == 10 and written[0]["length"] == 0
    start = d.element.body.xpath(".//w:commentRangeStart")[0]
    end = d.element.body.xpath(".//w:commentRangeEnd")[0]
    assert marker_offset(start.getparent(), start) == marker_offset(end.getparent(), end) == 10
    assert before == content_tree(d.element)
    pages[2].bottom = None
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "in", d, [f])
    assert not written and unplaced == [f]


def test_variants_survive_deduplication():
    _, ctx, _ = context_with_pages()
    a = page_finding(page_variant="first")
    b = page_finding(page_variant="even")
    result, dropped = dedupe_findings([a, b, deepcopy(b)], ctx)
    assert result == [a, b] and dropped == 1


def test_same_underlying_configuration_across_rules_writes_once(tmp_path):
    d, ctx, _ = context_with_pages()
    a = page_finding(page_scope="even_pages", page_config="odd_even:0:1")
    b = deepcopy(a)
    b.rule_id = "031.2"
    b.expected, b.actual = "奇偶页不同应为关闭", "奇偶页不同当前为开启"
    written, unplaced = write_page_comments(RunSummary(), ctx, tmp_path / "in", d, [a, b])
    assert not unplaced and len(written) == len(d.comments) == 1
    assert written[0]["rules"] == ["031.2", "037.2"]


def test_different_first_configuration_targets_section_first_page(tmp_path):
    d, ctx, _ = context_with_pages()
    f = page_finding(page_scope="first_page")
    written, _ = write_page_comments(RunSummary(), ctx, tmp_path / "in", d, [f])
    assert written[0]["page"] == 1 and f.affected_pages == [1]


def test_explicit_page_break_is_preserved_in_boundary_offsets():
    from docx.enum.text import WD_BREAK
    from thesis_format_checker.core.point_comments import boundary_text
    d = Document()
    p = d.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("下一页题目😀")
    before = content_tree(p._p)
    assert boundary_text(p) == "\f下一页题目😀"
    with pytest.raises(Exception, match="段前距"):
        add_boundary_comment(d, p, 0, "【警告】页脚")
    assert content_tree(p._p) == before and len(d.comments) == 0
    add_boundary_comment(d, p, 1, "【警告】页眉")
    start = p._p.xpath("./w:commentRangeStart")[0]
    end = p._p.xpath("./w:commentRangeEnd")[0]
    assert marker_offset(p._p, start) == marker_offset(p._p, end) == 1
    assert content_tree(p._p) == before


def test_display_equation_is_report_only_without_changing_math(tmp_path, monkeypatch):
    d = Document()
    p = d.add_paragraph()
    math = OxmlElement("m:oMathPara")
    equation, run, text = (OxmlElement(tag) for tag in ("m:oMath", "m:r", "m:t"))
    text.text = "x=1"
    run.append(text)
    equation.append(run)
    math.append(equation)
    p._p.append(math)
    original = p._p.xml
    f = make_finding("139", "公式格式", "居中", "左对齐", paragraph=p)
    add_word_comment(d, f, p)
    assert p._p.xml == original and len(d.comments) == 0
    source, output = tmp_path / "math.docx", tmp_path / "out.docx"
    d.save(source)
    f.anchor = None
    f.paragraph_index = 1
    monkeypatch.setattr("thesis_format_checker.engine.pipeline.check_sections", lambda *a, **k: [f])
    summary = run_comparison(lambda *a, **k: RunSummary(), source, source, output)
    assert any(x.rule_id == "139" and x.status == Status.FAIL for x in summary.findings)
    assert any("显示公式" in x for x in summary.report_only)
    assert len(Document(output).comments) == 0


def test_pipeline_does_not_fallback_when_pagination_is_unavailable(tmp_path, monkeypatch):
    d = Document()
    d.add_paragraph("无关的封面标题")
    source = tmp_path / "paper.docx"
    d.save(source)
    f = page_finding()
    monkeypatch.setattr("thesis_format_checker.engine.pipeline.check_sections", lambda *a, **k: [f])
    summary = run_comparison(lambda *a, **k: RunSummary(), source, source, tmp_path / "out.docx")
    assert len(Document(tmp_path / "out.docx").comments) == 0
    assert summary.findings and summary.report_only


@pytest.mark.integration
def test_word_boundary_comments_roundtrip_across_long_paragraph(tmp_path):
    if os.environ.get("RUN_WORD_INTEGRATION") != "1":
        pytest.skip("设置 RUN_WORD_INTEGRATION=1 后运行真实 Word 验证")
    from thesis_format_checker.core.word_session import WordSession
    d = Document()
    d.sections[0].page_height = Inches(5)
    d.sections[0].header.paragraphs[0].text = "页眉"
    d.sections[0].footer.paragraphs[0].text = "页脚"
    d.add_paragraph("这是一段跨页正文，用于检查页末不会误指到下一页。" * 150)
    d.add_paragraph("最后一段")
    from docx.enum.text import WD_BREAK
    p = d.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("分页符后的页面😀")
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    d.save(source)
    ctx = DocxContext(d)
    pagination = read_pagination(source, ctx)
    assert len(pagination.pages) >= 3
    assert pagination.boundaries_verified
    # Word attributes the leading page break to the preceding page; its safe
    # footer endpoint is the previous paragraph mark, never before that break.
    assert pagination.pages[-2].bottom.paragraph == 1
    assert pagination.pages[-2].bottom.offset == len("最后一段")
    points = []
    for page in pagination.pages:
        assert page.top is not None and page.bottom is not None
        for story, point in (("header", page.top), ("footer", page.bottom)):
            add_boundary_comment(d, ctx.paragraphs[point.paragraph], point.offset,
                                 f"{story} page {page.number}")
            points.append((page.number, point.position))
    save_comments_only(source, d, output, allow_run_splits=True)
    with WordSession() as session:
        wd = session.open_readonly(output.resolve())
        wd.Repaginate()
        assert wd.ComputeStatistics(2) == len(pagination.pages)
        assert wd.Comments.Count == len(points)
        for i, (page, position) in enumerate(points, 1):
            scope = wd.Comments(i).Scope
            assert scope.Start == scope.End == position
            assert scope.Information(3) == page
        for story in (wd.Sections(1).Headers(1), wd.Sections(1).Footers(1)):
            with pytest.raises(Exception, match="主体部分"):
                wd.Comments.Add(Range=story.Range, Text="禁止直接写入")
