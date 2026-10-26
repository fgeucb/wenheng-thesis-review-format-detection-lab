from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.models import Status
from thesis_format_checker.detectors.outline import (
    check_outline_errors,
    extract_outline_baseline,
)


def _failures(findings, family):
    return [
        item
        for item in findings
        if item.status is Status.FAIL and item.rule_id.split(".", 1)[0] == family
    ]


def test_stale_toc_entries_do_not_hide_missing_abstracts_or_references():
    template = Document()
    template.add_heading("摘 要", level=1)
    template.add_heading("Abstract", level=1)
    template.add_heading("参考文献", level=1)

    actual = Document()
    actual.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    actual.add_paragraph("摘 要 I", style="TOC 1")
    actual.add_paragraph("Abstract II", style="TOC 1")
    actual.add_paragraph("参考文献 50", style="TOC 1")

    findings = check_outline_errors(
        extract_outline_baseline(DocxContext(template)),
        DocxContext(actual),
    )

    assert _failures(findings, "003")
    assert _failures(findings, "004")
    assert _failures(findings, "008")


def test_cover_and_title_page_are_recognized_independently():
    template = Document()
    template.add_paragraph("学位论文封面格式")
    template.add_paragraph("封二格式")
    baseline = extract_outline_baseline(DocxContext(template))

    title_page_only = Document()
    title_page_only.add_paragraph("作者：张三")
    title_page_only.add_paragraph("指导教师：李四")
    findings = check_outline_errors(baseline, DocxContext(title_page_only))
    assert [item.rule_id for item in _failures(findings, "001")] == ["001.1"]

    cover_only = Document()
    cover_only.add_paragraph("分类号：TP391 密级：公开 UDC：004 学位论文")
    findings = check_outline_errors(baseline, DocxContext(cover_only))
    assert [item.rule_id for item in _failures(findings, "001")] == ["001.2"]


def test_standalone_statement_heading_is_not_confused_with_authorization():
    template = Document()
    template.add_paragraph("声  明")
    template.add_paragraph("学位论文使用授权声明")
    baseline = extract_outline_baseline(DocxContext(template))

    findings = check_outline_errors(baseline, DocxContext(Document()))
    codes = {item.rule_id for item in _failures(findings, "002")}
    assert {"002.1", "002.2"} <= codes


def test_toc_chapter_missing_from_body_is_reported_at_the_toc_entry():
    template = Document()
    template.add_heading("摘 要", level=1)
    template.add_heading("Abstract", level=1)
    template.add_heading("参考文献", level=1)

    actual = Document()
    actual.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    toc_entry = actual.add_paragraph("1\t绪论\t9", style="TOC 1")
    actual.add_heading("相关理论", level=1)
    actual.add_paragraph("正文内容。")

    findings = check_outline_errors(
        extract_outline_baseline(DocxContext(template)),
        DocxContext(actual),
    )
    missing = _failures(findings, "007")
    assert len(missing) == 1
    assert missing[0].anchor._p is toc_entry._p
    assert "绪论" in missing[0].expected
