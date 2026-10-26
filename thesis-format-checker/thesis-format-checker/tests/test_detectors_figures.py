from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.figures import (
    collect_caption_records, extract_figure_table_baseline, check_figure_table_errors,
)
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def node(name, **attrs):
    element = OxmlElement(name)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    return element


def clone(document):
    from io import BytesIO
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


def format_text(paragraph, size=11):
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        fonts = run._r.get_or_add_rPr().rFonts
        fonts.set(qn("w:eastAsia"), "仿宋")


def picture(document, floating=False):
    p = document.add_paragraph()
    run = p.add_run()
    drawing = node("w:drawing")
    drawing.append(node("wp:anchor" if floating else "wp:inline"))
    run._r.append(drawing)
    return p


def caption(document, text):
    p = document.add_paragraph(text, style="Caption")
    format_text(p)
    return p


def figure(document, number, floating=False):
    p = picture(document, floating)
    c = caption(document, "图 " + number + " 测试示意")
    return p, c


def table(document, number="1", header=True):
    c = caption(document, "表 " + number + " 测试数据")
    t = document.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.text = "表头" if ri == 0 else "数据"
            format_text(cell.paragraphs[0])
    if header:
        t.rows[0]._tr.get_or_add_trPr().append(node("w:tblHeader"))
    return c, t


def check(template, actual):
    return check_figure_table_errors(extract_figure_table_baseline(DocxContext(template)), DocxContext(actual))


def findings(result, rule, status=Status.FAIL):
    statuses = (Status.UNKNOWN, Status.NOT_APPLICABLE) if status == Status.UNKNOWN else (status,)
    return [item for item in result if family_of(item.rule_id) == rule and item.status in statuses]


def test_caption_detection_excludes_prose_mentions_and_toc():
    d = Document()
    d.add_paragraph("图 1 显示了测试结果，我们继续讨论。")
    p, c = figure(d, "1")
    assert [r["paragraph"]._p for r in collect_caption_records(DocxContext(d))] == [c._p]


def test_english_caption_abbreviations_are_recognised_by_kind():
    """中英双语题注常用的 Fig./Tab. 缩写要分别按图、表识别。"""

    document = Document()
    picture(document)
    document.add_paragraph("Fig.1-1  System architecture")
    document.add_paragraph("Tab.1-1  Technical specifications")
    document.add_table(rows=2, cols=2)

    records = {(record["text"], record["kind"]) for record in collect_caption_records(DocxContext(document))}

    assert ("Fig.1-1  System architecture", "figure") in records
    assert ("Tab.1-1  Technical specifications", "table") in records


def test_078_uses_template_position_in_both_directions():
    a = Document()
    p, c = figure(a, "1")
    b = clone(a)
    b.element.body.insert(0, b.paragraphs[1]._p)
    result = check(a, b)
    assert findings(result, "078")
    assert not findings(check(b, b), "078")
    assert findings(check(b, a), "078")


def test_078_table_below_caption_and_floating_unknown():
    a = Document()
    c, t = table(a)
    b = clone(a)
    b.element.body.insert(0, b.tables[0]._tbl)
    assert findings(check(a, b), "078")
    d = Document()
    figure(d, "1", floating=True)
    actual = check(d, d)
    assert findings(actual, "078", Status.UNKNOWN)
    assert not findings(actual, "078")


def test_079_local_run_overlay_and_dynamic_baseline():
    a = Document()
    figure(a, "1")
    b = clone(a)
    b.paragraphs[1].runs[0].font.size = Pt(15)
    assert findings(check(a, b), "079")
    a.paragraphs[1].runs[0].font.size = Pt(15)
    assert not findings(check(a, b), "079")


def test_079_inherited_style_and_local_override():
    a = Document()
    style = a.styles.add_style("图题专用", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = a.styles["Caption"]
    style.font.name = "Courier New"
    style.font.size = Pt(12)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体")
    picture(a)
    a.add_paragraph("图 1 示例", style=style)
    b = clone(a)
    b.paragraphs[1].runs[0].font.name = "Calibri"
    assert findings(check(a, b), "079")
    assert not findings(check(a, a), "079")


def test_080_skip_repeat_and_template_dynamic_step():
    a = Document()
    for number in ("5", "7", "9"):
        figure(a, number)
        a.add_paragraph("正文隔开")
    b = clone(a)
    b.paragraphs[4].runs[0].text = "图 9 测试示意"
    assert findings(check(a, b), "080")
    assert not findings(check(a, a), "080")


def test_080_chapter_prefix_uses_heading_records():
    a = Document()
    for chapter in (1, 2):
        a.add_heading(f"{chapter} 章节", 1)
        for number in (1, 2):
            figure(a, f"{chapter}-{number}")
            a.add_paragraph("正文")
    b = clone(a)
    for p in b.paragraphs:
        if p.text.startswith("图 2-"):
            p.runs[0].text = p.text.replace("图 2-", "图 3-")
    assert any("章节前缀" in f.actual for f in findings(check(a, b), "080"))


def test_081_literal_spacing_separator_and_digits():
    a = Document()
    figure(a, "1-1")
    b = clone(a)
    b.paragraphs[1].runs[0].text = "图1.1 测试示意"
    assert findings(check(a, b), "081")


def test_fields_cached_values_never_fail_sequence():
    a = Document()
    picture(a)
    p = caption(a, "图 ")
    field = node("w:fldSimple", instr="SEQ 图 \\* ARABIC")
    run = node("w:r")
    value = node("w:t")
    value.text = "999"
    run.append(value)
    field.append(run)
    p._p.append(field)
    p.add_run(" 示例")
    result = check(a, a)
    for rule in ("080", "081"):
        assert findings(result, rule, Status.UNKNOWN)
        assert not findings(result, rule)


def test_083_style_inherited_borders_and_cell_overlay():
    a = Document()
    table(a)
    b = clone(a)
    borders = node("w:tcBorders")
    borders.append(node("w:left", val="double", sz="8", color="FF0000"))
    b.tables[0].cell(1, 0)._tc.get_or_add_tcPr().append(borders)
    assert findings(check(a, b), "083")
    assert not findings(check(a, a), "083")


def test_083_no_school_three_line_default():
    a = Document()
    table(a)
    assert not findings(check(a, a), "083")
    b = clone(a)
    b.tables[0].style = "Normal Table"
    assert findings(check(a, b), "083")


def test_083_conditional_style_and_merged_cells_unknown():
    a = Document()
    table(a)
    b = clone(a)
    b.tables[0].cell(1, 0).merge(b.tables[0].cell(1, 1))
    result = check(a, b)
    assert findings(result, "083", Status.UNKNOWN)
    assert not findings(result, "083")


def test_084_header_body_height_shading_and_local_font():
    a = Document()
    table(a)
    b = clone(a)
    t = b.tables[0]
    t.rows[0].height = Pt(25)
    t.cell(0, 0)._tc.get_or_add_tcPr().append(node("w:shd", val="clear", fill="FF0000"))
    t.cell(1, 0).paragraphs[0].runs[0].font.size = Pt(18)
    result = findings(check(a, b), "084")
    assert any("行高" in f.expected for f in result)
    assert any("底纹" in f.expected for f in result)
    assert any("字号" in f.expected for f in result)


def test_084_explicit_total_label_matches_template():
    a = Document()
    _, t = table(a)
    t.cell(2, 0).paragraphs[0].runs[0].text = "合计"
    b = clone(a)
    b.tables[0].cell(2, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    assert any("对齐" in f.expected for f in findings(check(a, b), "084"))


def test_unknown_has_required_message_and_no_semantic_rules():
    a = Document()
    figure(a, "1", floating=True)
    result = check(a, a)
    assert all(UNREADABLE in item.message for item in result if item.status == Status.UNKNOWN)
    assert not {"082", "085", "086", "087", "088"}.intersection(f.rule_id for f in result)


def test_reading_does_not_modify_any_part():
    d = Document()
    figure(d, "1")
    d.add_paragraph("正文")
    table(d)
    before = {part.partname: part.blob for part in d.part.package.parts}
    check(d, d)
    after = {part.partname: part.blob for part in d.part.package.parts}
    assert before == after


def test_ambiguous_table_baseline_unknown_no_arbitrary_choice():
    a = Document()
    table(a, "1")
    a.add_paragraph("正文")
    _, t = table(a, "2")
    t.cell(1, 0).paragraphs[0].runs[0].font.size = Pt(20)
    b = Document()
    table(b)
    assert findings(check(a, b), "084", Status.UNKNOWN)


def test_shared_cell_border_conflicts_unknown():
    a = Document()
    table(a)
    b = clone(a)
    borders = node("w:tcBorders")
    borders.append(node("w:right", val="double", sz="8", color="FF0000"))
    b.tables[0].cell(1, 0)._tc.get_or_add_tcPr().append(borders)
    result = check(a, b)
    assert findings(result, "083", Status.UNKNOWN)
    assert not findings(result, "083")


def test_unassociated_table_not_silently_ignored_and_formula_table_excluded():
    a = Document()
    table(a)
    b = Document()
    t = b.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "数据"
    result = check(a, b)
    assert findings(result, "083", Status.UNKNOWN)
    assert findings(result, "084", Status.UNKNOWN)
    t.cell(0, 0).paragraphs[0]._p.append(node("m:oMath"))
    result = check(a, b)
    assert not findings(result, "083", Status.UNKNOWN)


def test_caption_revisions_unknown_without_sequence_fail():
    a = Document()
    figure(a, "1")
    a.add_paragraph("正文")
    figure(a, "2")
    b = clone(a)
    p = b.paragraphs[1]
    p.runs[0].text = "图 88 测试示意"
    p._p.get_or_add_pPr().append(node("w:pPrChange"))
    result = check(a, b)
    assert findings(result, "080", Status.UNKNOWN)
    assert not findings(result, "080")


def test_caption_number_and_title_can_have_different_sizes():
    a = Document()
    picture(a)
    p = a.add_paragraph(style="Caption")
    p.add_run("图 1 ")
    p.add_run("名称")
    format_text(p)
    p.runs[0].font.size = Pt(9)
    p.runs[1].font.size = Pt(12)
    assert not findings(check(a, a), "079")
    b = clone(a)
    b.paragraphs[1].runs[1].font.size = Pt(13)
    assert findings(check(a, b), "079")
