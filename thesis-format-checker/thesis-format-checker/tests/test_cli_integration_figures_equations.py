"""可选的真实 Word 端到端测试：公共命令 + 原生批注。"""
import os
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from thesis_format_checker.core.point_comments import content_tree as semantic_tree
from conftest import assert_reported_comments


@pytest.mark.integration
def test_public_cmd_and_native_word_reopen(runtime_dir):
    if os.environ.get("RUN_WORD_INTEGRATION") != "1" or os.name != "nt":
        pytest.skip("需设置 RUN_WORD_INTEGRATION=1 并使用Windows桌面Word")
    from test_pipeline_figures_equations import fixture
    root = Path(__file__).resolve().parents[1]
    template, paper, output = [runtime_dir / name for name in ("模板.docx", "论文.docx", "批注.docx")]
    fixture(template, template=True)
    fixture(paper, template=False)
    before = [p.read_bytes() for p in (template, paper)]
    report = runtime_dir / "report.json"
    for phase, args in (
        ("compare", [str(root / "thesis-checker.cmd"), "compare", "--template", str(template), "--document", str(paper), "--output", str(output), "--report-json", str(report)]),
        ("reopen", [sys.executable, str(Path(__file__).with_name("test_cli_integration_headers_toc.py")), str(output)]),
    ):
        stdout, stderr = runtime_dir / (phase + "_stdout.log"), runtime_dir / (phase + "_stderr.log")
        with stdout.open("wb") as out, stderr.open("wb") as err:
            try:
                process = subprocess.run(args, cwd=root, env=os.environ.copy(), stdout=out, stderr=err, timeout=120)
            except subprocess.TimeoutExpired:
                pytest.fail("本地Word测试超时；已停止等待，不自动重试，不终止用户Word实例")
        assert process.returncode == 0, stderr.read_text(encoding="utf-8", errors="replace")[-2500:]
    assert before == [p.read_bytes() for p in (template, paper)]
    comments = [c.text for c in Document(output).comments]
    expected = ("070", "071", "073", "076", "077", "078", "079", "080", "081", "083", "084", "093")
    assert_reported_comments(report, Document(output).comments, expected)
    # Display-math findings must survive as confirmed report entries; inserting
    # native comments on these anchors changes Word's equation layout.
    import json
    payload = json.loads(report.read_text(encoding="utf-8"))
    for family in ("090", "091", "092", "096", "097"):
        assert any(f["rule_id"].split(".")[0] == family and f["status"] == "fail"
                   for f in payload["findings"])
        assert any(note.startswith("【" + family) and "显示公式" in note
                   for note in payload["report_only"])
    assert not Document(output).element.body.xpath(".//m:oMathPara/../w:commentReference")
    assert "保留的人工批注" in comments
    assert not any("既有检测暂未完成" in c for c in comments)
    with ZipFile(paper) as source, ZipFile(output) as annotated:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in source.namelist():
            if name not in allowed:
                assert source.read(name) == annotated.read(name), name
        assert semantic_tree(etree.fromstring(source.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))
