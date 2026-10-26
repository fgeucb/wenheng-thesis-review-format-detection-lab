from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.detectors.structure import check_structure_errors, extract_structure_baseline
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def context(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return DocxContext(document)


def check(template, actual):
    return check_structure_errors(extract_structure_baseline(template), actual)


def matches(item, rule):
    """按规则族查询：族编号或其细粒度规则（001.1、005.2 …）都算命中。"""

    return family_of(item.rule_id) == rule or item.rule_id.startswith(rule + ".")


def failures(findings, rule):
    return [item for item in findings if matches(item, rule) and item.status == Status.FAIL]


def uncertain(findings, rule):
    return [item for item in findings if matches(item, rule) and item.status == Status.UNKNOWN]


def append_field(ctx, instruction, close=True):
    paragraph = ctx.document.add_paragraph()
    for kind in ("begin",):
        run = paragraph.add_run()
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), kind)
        run._r.append(field)
    run = paragraph.add_run()
    node = OxmlElement("w:instrText")
    node.text = instruction
    run._r.append(node)
    if close:
        run = paragraph.add_run()
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "end")
        run._r.append(field)
    return DocxContext(ctx.document)


def test_mandatory_components_and_no_reference_detection():
    template = context("必须包含：封面、原创性声明、目录、致谢、附录、参考文献。", "封面", "原创性声明", "声明正文", "目录", "致谢", "附录")
    actual = context("摘要", "内容")
    findings = check(template, actual)
    assert failures(findings, "001")
    assert failures(findings, "002")
    assert failures(findings, "006")
    assert len(failures(findings, "009")) == 2
    assert not any(family_of(item.rule_id) == "008" for item in findings)


def test_sample_does_not_establish_required_component():
    findings = check(context("封面", "致谢", "附录"), context("摘要", "内容"))
    assert not failures(findings, "001")
    assert not failures(findings, "009")
    # 模板只给了示例、没写“必须包含”时按“模板未涉及”处理：不写批注、不计入无法判定。
    assert not uncertain(findings, "001")
    assert not uncertain(findings, "009")
    assert any(item.rule_id.startswith("009") and item.status == Status.NOT_APPLICABLE
               for item in findings)


def test_optional_and_conditional_items_are_not_missing_failures():
    findings = check(context("附录为可选部分；如有成果时需包含攻读学位期间成果。", "附录", "攻读学位期间成果"), context("摘要"))
    assert not failures(findings, "009")
    assert uncertain(findings, "009")


def test_missing_component_textbox_is_unknown_not_false_failure():
    actual = context("摘要")
    actual.has_textboxes = True
    findings = check(context("必须包含：封面。", "封面"), actual)
    assert not failures(findings, "001")
    assert "该对象docx库无法读取，本条暂无法执行检测" in uncertain(findings, "001")[0].message


def test_declared_order_uses_available_components_and_ignores_references():
    template = context("组成部分顺序：封面、目录、摘要、参考文献、致谢。", "封面", "目录", "摘要", "致谢")
    actual = context("封面", "摘要", "内容", "目录", "参考文献", "条目", "致谢", "内容")
    findings = check(template, actual)
    ordering = failures(findings, "010")
    assert len(ordering) == 1
    assert "目录" in ordering[0].message and "摘要" in ordering[0].message
    assert "参考文献" not in ordering[0].message


def test_template_sample_order_is_the_default_order_requirement():
    findings = check(context("封面", "目录", "摘要"), context("封面", "摘要", "目录"))
    ordering = failures(findings, "010")
    assert len(ordering) == 1
    assert "目录" in ordering[0].message and "摘要" in ordering[0].message


def test_toc_back_matter_order_is_used_and_toc_entries_are_not_actual_parts():
    template_document = Document()
    template_document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    template_document.add_paragraph("封面")
    template_document.add_paragraph("摘要")
    template_document.add_paragraph("目录")
    template_document.add_paragraph("致谢\t1", style="TOC 1")
    template_document.add_paragraph("参考文献\t2", style="TOC 1")
    template_document.add_paragraph("附录\t3", style="TOC 1")

    actual_document = Document()
    actual_document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    actual_document.add_paragraph("封面")
    actual_document.add_paragraph("摘要")
    actual_document.add_paragraph("目录")
    actual_document.add_paragraph("致谢\t1", style="TOC 1")
    actual_document.add_paragraph("参考文献\t2", style="TOC 1")
    actual_document.add_paragraph("附录\t3", style="TOC 1")
    actual_document.add_paragraph("致谢")
    actual_document.add_paragraph("附录")
    actual_document.add_paragraph("参考文献")

    findings = check(DocxContext(template_document), DocxContext(actual_document))
    ordering = failures(findings, "010")
    assert len(ordering) == 1
    assert "参考文献" in ordering[0].message and "附录" in ordering[0].message


def test_keyword_count_is_dynamic_and_english_phrase_spaces_are_not_separators():
    template = context("中文关键词数量：2—3个。英文关键词数量：2—3个。", "关键词：示例甲；示例乙", "Keywords: sample one; sample two")
    actual = context("关键词：词一；词二；词三；词四", "Keywords: machine learning; semantic search")
    findings = check(template, actual)
    items = failures(findings, "005")
    assert len(items) == 1
    assert "中文关键词" in items[0].message
    assert "4" in items[0].message


def test_sample_keyword_count_is_not_a_school_requirement():
    findings = check(context("关键词：一；二；三"), context("关键词：甲；乙；丙；丁；戊"))
    assert not failures(findings, "005")


def test_template_keyword_section_makes_the_corresponding_language_required():
    findings = check(
        context("关键词：示例甲；示例乙", "Key words: sample one; sample two"),
        context("关键词：实际甲；实际乙"),
    )
    items = failures(findings, "005")
    assert len(items) == 1
    assert items[0].rule_id == "005.3"
    assert "英文关键词" in items[0].message


def test_conflicting_keyword_count_is_unknown():
    findings = check(context("中文关键词数量：2—3个。", "中文关键词数量：4—6个。"), context("关键词：甲；乙"))
    assert uncertain(findings, "005")
    assert not failures(findings, "005")


def test_automatic_toc_template_detects_plain_text_toc():
    template = append_field(context("目录"), ' TOC \\o "1-3" \\h ')
    findings = check(template, context("目录", "第一章……1"))
    assert failures(findings, "006")


def test_automatic_toc_accepts_custom_style_source():
    template = append_field(context("目录"), ' TOC \\o "1-3" ')
    actual = append_field(context("目录"), ' TOC \\t "研究章节,1" ')
    assert not failures(check(template, actual), "006")


def test_malformed_toc_degrades_to_boundary_comment():
    template = append_field(context("目录"), ' TOC \\o "1-3" ')
    actual = append_field(context("目录"), ' TOC \\o "1-3" ', close=False)
    findings = check(template, actual)
    assert uncertain(findings, "006")
    assert "该对象docx库无法读取，本条暂无法执行检测" in uncertain(findings, "006")[0].message


def test_empty_declaration_is_detected_but_image_content_is_not_empty():
    template = context("必须包含：原创性声明、目录。", "原创性声明", "模板正文", "目录")
    empty = check(template, context("原创性声明", "目录"))
    assert failures(empty, "002")
    actual = context("原创性声明", "", "目录")
    actual.paragraphs[1].add_run()._r.append(OxmlElement("w:drawing"))
    assert not failures(check(template, actual), "002")


def test_standalone_statement_heading_is_originality_declaration():
    template = context("声  明", "声明正文", "学位论文使用授权声明", "授权正文")
    samples = extract_structure_baseline(template)["structure"]["samples"]
    assert "originality" in samples
    assert "authorization" in samples


def test_plain_cover_labels_inside_table_are_recognized():
    template = context("必须包含：封面。", "封面")
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "论文题目：研究主题"
    table.cell(1, 0).text = "作者姓名：张三"
    document.add_paragraph("摘要")
    assert not failures(check(template, DocxContext(document)), "001")


def test_input_body_is_never_changed_by_detection():
    template = context("必须包含：封面、致谢。", "封面", "致谢")
    actual = context("摘要", "正文")
    before = actual.document.element.xml
    check(template, actual)
    assert actual.document.element.xml == before


def test_cyclic_order_requirements_are_unknown_not_document_failure():
    template = context("组成部分顺序：封面、目录。", "组成部分顺序：目录、摘要。", "组成部分顺序：摘要、封面。")
    actual = context("封面", "目录", "摘要")
    findings = check(template, actual)
    assert uncertain(findings, "010")
    assert not failures(findings, "010")


def test_keyword_label_in_separate_table_cell_does_not_false_report_empty():
    template = context("中文关键词数量：2—3个。")
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "关键词："
    table.cell(0, 1).text = "机器学习；信息检索"
    findings = check(template, DocxContext(document))
    assert uncertain(findings, "005")
    assert not failures(findings, "005")


def test_statement_heading_followed_by_real_chapter_is_empty():
    template = context("必须包含：原创性声明。", "原创性声明", "声明正文")
    document = Document()
    document.add_paragraph("原创性声明")
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("绪论正文不属于声明。")
    assert failures(check(template, DocxContext(document)), "002")


def test_title_page_fields_are_not_mistaken_for_a_separate_cover():
    template = context("必须包含：封面、题名页。", "封面", "题名页")
    actual = context("题名页", "论文题目：研究主题", "作者姓名：张三", "摘要")
    findings = check(template, actual)
    assert len(failures(findings, "001")) == 1
    assert "封面" in failures(findings, "001")[0].actual


def test_common_subheading_names_are_allowed_under_different_chapters():
    document = Document()
    for chapter in ("第一章 方法", "第二章 实验"):
        document.add_heading(chapter, level=1)
        document.add_heading("引言", level=2)
        document.add_paragraph("本节正文。")
        document.add_heading("本章小结", level=2)
        document.add_paragraph("小结正文。")
    findings = check(context("模板说明"), DocxContext(document))
    assert not failures(findings, "011")


def test_same_subheading_name_is_duplicate_within_one_parent_chapter():
    document = Document()
    document.add_heading("第一章 方法", level=1)
    document.add_heading("实验设置", level=2)
    document.add_paragraph("第一处正文。")
    document.add_heading("实验设置", level=2)
    document.add_paragraph("第二处正文。")
    findings = check(context("模板说明"), DocxContext(document))
    items = [item for item in failures(findings, "011") if item.rule_id == "011.2"]
    assert len(items) == 1
