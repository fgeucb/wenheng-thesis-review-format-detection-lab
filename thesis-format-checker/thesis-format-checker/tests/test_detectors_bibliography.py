from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.bibliography import (
    check_bibliography_errors, extract_bibliography_baseline,
)
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



FIRST = "[1] 张三，李四. 示例甲[J]. 示例期刊, 2024, 12(3): 4-9. DOI: 10.1234/abc https://example.org/a [2024-01-02]"
SECOND = "[2] 王五，赵六. 示例乙[J]. 示例期刊, 2023, 13(4): 8-12. DOI: 10.1234/def https://example.org/b [2024-02-03]"
REQUIRED = "[J]必填字段：作者、题名、期刊名、年份、卷、期、页码、DOI、URL、引用日期"
AUTHORS = "中文作者超过3人，列前3人，后加等"
ORDER = "参考文献按文末编号升序排列"


def doc(entries=(FIRST, SECOND), instructions=(REQUIRED, AUTHORS, ORDER), tail=None):
    document = Document()
    document.add_heading("正文", level=1)
    document.add_paragraph("本文不进行引用判定。")
    document.add_heading("参考文献", level=1)
    for value in instructions:
        document.add_paragraph(value)
    for value in entries:
        document.add_paragraph(value)
    if tail:
        document.add_heading(tail, level=1)
        document.add_paragraph("这是一段非参考文献正文。")
    return document


def run(document, template=None):
    baseline = extract_bibliography_baseline(DocxContext(template or doc()))
    return check_bibliography_errors(baseline, DocxContext(document))


def fail_ids(findings):
    """按规则族返回失败项编号（114.1 → 114），用于族级断言。"""

    return {item.rule_id.split(".")[0] for item in findings if item.status == Status.FAIL}


def unknowns(findings, rule):
    # 模板未涉及的结论现在是“不适用”（不写批注），测试中与“无法判定”一并接受。
    return [item for item in findings if family_of(item.rule_id) == rule
            and item.status in (Status.UNKNOWN, Status.NOT_APPLICABLE)]


def test_valid_manual_entries_do_not_fail_and_links_are_not_claimed_verified():
    findings = run(doc())
    assert not fail_ids(findings)
    assert len(unknowns(findings, "117")) == 2
    assert all("未联网" in item.message for item in unknowns(findings, "117"))


def test_all_seven_api_rules_can_report_concrete_template_based_errors():
    changed = FIRST.replace("[1]", "（3）").replace("张三，李四", "张三，李四，王五，赵六").replace("示例甲", "")
    changed = changed.replace("[J].", "[J]。").replace("(3):", "(3)：").replace("DOI:", "DOI：")
    findings = run(doc((changed, SECOND)))
    assert fail_ids(findings) == {"111", "113", "114", "115", "117", "118", "121"}
    assert all(f.anchor is not None for f in findings)
    # 统一批注格式：抬头 + 应该 + 实际（定位由调度层补充）
    assert all(f.message.count("\n") == 2 for f in findings)


def test_template_punctuation_and_numbering_are_dynamic():
    changed = tuple(s.replace("[1]", "（1）").replace("[2]", "（2）").replace("[J].", "[J]。").replace("(3):", "(3)：").replace("(4):", "(4)：") for s in (FIRST, SECOND))
    template = doc(changed)
    assert not fail_ids(run(doc(changed), template))
    assert {"111", "115", "118"} <= fail_ids(run(doc(), template))


def test_author_threshold_only_from_explicit_template_statement():
    changed = FIRST.replace("张三，李四", "张三，李四，王五，赵六")
    assert "113" in fail_ids(run(doc((changed,))))
    template = doc(instructions=(REQUIRED, ORDER))
    assert "113" not in fail_ids(run(doc((changed,)), template))
    template = doc(instructions=(REQUIRED, "中文作者超过5人，列前4人，后加等", ORDER))
    assert "113" not in fail_ids(run(doc((changed,)), template))


def test_author_elision_checks_visible_count_and_marker_without_inventing_full_list():
    good = FIRST.replace("张三，李四", "张三，李四，王五，等")
    bad_count = FIRST.replace("张三，李四", "张三，李四，等")
    wrong_marker = FIRST.replace("张三，李四", "张三，李四，王五，et al")
    assert "113" not in fail_ids(run(doc((good,))))
    assert "113" in fail_ids(run(doc((bad_count,))))
    assert "113" in fail_ids(run(doc((wrong_marker,))))


def test_author_separator_comes_from_template_examples():
    changed = FIRST.replace("张三，李四", "张三、李四")
    assert "113" in fail_ids(run(doc((changed,))))
    template = doc(tuple(s.replace("，", "、") for s in (FIRST, SECOND)))
    assert "113" not in fail_ids(run(doc((changed,)), template))


def test_explicit_empty_title_and_issue_are_missing_not_semantically_inferred():
    changed = FIRST.replace("示例甲", "").replace("(3)", "()")
    findings = run(doc((changed,)))
    assert "114" in fail_ids(findings)
    # 一个字段一条结论：题名（114.1）与期号（114.7）分别成条，合起来看
    texts = "；".join(f.actual for f in findings
                     if family_of(f.rule_id) == "114" and f.status == Status.FAIL)
    assert "题名" in texts and "期" in texts


def test_unparsed_year_is_unknown_not_guessed_missing_fields():
    changed = FIRST.replace("2024, 12", "缺年, 12")
    findings = run(doc((changed,)))
    assert "114" not in fail_ids(findings)
    assert unknowns(findings, "114") and unknowns(findings, "115")


def test_placeholder_is_missing_only_for_declared_required_field():
    changed = FIRST.replace("示例甲", "题名")
    assert "114" in fail_ids(run(doc((changed,))))
    findings = run(doc((changed,)), doc(instructions=(ORDER,)))
    assert "114" not in fail_ids(findings)
    assert unknowns(findings, "114")


def test_a_standard_name_does_not_supply_hardcoded_rules():
    template = doc((), instructions=("参考文献著录要求：按GB/T 7714—2015。",))
    findings = run(doc(), template)
    assert not fail_ids(findings)
    assert {"111", "113", "114", "115", "117", "118"} <= {f.rule_id for f in findings}


def test_digital_example_presence_does_not_mean_required():
    changed = FIRST.split(" DOI:")[0]
    assert "117" in fail_ids(run(doc((changed,))))
    template = doc(instructions=("[J]必填字段：题名、年份", ORDER))
    assert "117" not in fail_ids(run(doc((changed,)), template))


def test_doi_syntax_and_access_date_shape_are_template_based():
    malformed = FIRST.replace("10.1234/abc", "10.1234abc")
    slash_date = FIRST.replace("2024-01-02", "2024/01/02")
    assert "117" in fail_ids(run(doc((malformed,))))
    assert "117" in fail_ids(run(doc((slash_date,))))
    template = doc(tuple(s.replace("-01-02", "/01/02").replace("-02-03", "/02/03") for s in (FIRST, SECOND)))
    assert "117" not in fail_ids(run(doc((slash_date,)), template))


def test_year_volume_order_comparison_does_not_guess_reversed_numbers():
    changed = FIRST.replace("2024, 12(3)", "12(3), 2024")
    findings = run(doc((changed,)))
    assert "115" in fail_ids(findings)
    assert "114" not in fail_ids(findings)
    assert unknowns(findings, "114")


def test_title_internal_punctuation_does_not_trigger_punctuation_rule():
    changed = FIRST.replace("示例甲", "研究：数据，算法（第一部分）")
    assert "118" not in fail_ids(run(doc((changed,))))


def test_digits_in_journal_title_are_not_misclassified_as_volume():
    changed = FIRST.replace("示例期刊", "数据2期刊")
    assert "115" not in fail_ids(run(doc((changed,))))


def test_unmatched_type_is_unknown_and_does_not_infer_true_source_kind():
    changed = FIRST.replace("[J]", "[M]")
    findings = run(doc((changed,)))
    # 模板只出现过 [J] 之类的标识：用 [M] 时按“标识使用不当”（112.2）报出，
    # 但不会把它推断成别的真实文献类型。
    assert fail_ids(findings) == {"112"}
    assert unknowns(findings, "111")
    # 112 已实现（标识使用不当）；116/119/120 属于已删除/待确认族，不应出现。
    assert not ({"116", "119", "120"} & {f.rule_id for f in findings})


def test_stop_at_acknowledgements_appendix_and_arbitrary_next_peer_heading():
    for heading in ("致谢", "附录A", "研究成果"):
        findings = run(doc(tail=heading))
        assert not fail_ids(findings)
        assert len(findings) == 2


def test_reference_like_body_paragraphs_without_bibliography_heading_are_ignored():
    document = Document()
    document.add_paragraph(FIRST)
    assert run(document) == []


def test_paper_title_mentioning_standard_is_not_mistaken_for_instruction():
    changed = FIRST.replace("示例甲", "GB/T格式要求研究")
    findings = run(doc((changed,)))
    assert len(unknowns(findings, "117")) == 1


def test_plain_latin_author_order_requires_explicit_instruction():
    a = "[1] ADAMS A, BAKER B. First[J]. Journal, 2024, 12(3): 4-9."
    b = "[2] ZHANG Z, ZHOU Y. Second[J]. Journal, 2023, 13(4): 8-12."
    template = doc((a, b), instructions=("参考文献按作者字母顺序排列", "[J]必填字段：题名、年份"))
    assert "121" in fail_ids(run(doc((b, a)), template))
    template = doc((a, b), instructions=("[J]必填字段：题名、年份",))
    findings = run(doc((b, a)), template)
    assert "121" not in fail_ids(findings)
    assert unknowns(findings, "121")


def test_chinese_author_collation_is_not_guessed_and_citation_order_is_deferred():
    template = doc(instructions=(REQUIRED, "参考文献按作者字母顺序排列"))
    findings = run(doc(), template)
    assert unknowns(findings, "121")
    template = doc(instructions=(REQUIRED, "参考文献按首次引用顺序排列"))
    assert "121" not in {f.rule_id for f in run(doc(), template)}


def test_conflicting_numeric_examples_do_not_choose_one_as_standard():
    template = doc((FIRST, SECOND.replace("(4):", "(4)：")))
    findings = run(doc((FIRST,)), template)
    assert "115" not in fail_ids(findings)
    assert unknowns(findings, "115")


def test_ambiguous_author_initials_are_unknown_not_split_as_title():
    text = "[1] SMITH J. A. Title[J]. Journal, 2024, 12(3): 4-9."
    findings = run(doc((text,)))
    assert not fail_ids(findings)
    assert unknowns(findings, "113")


def test_bibliography_field_cache_is_not_treated_as_a_fresh_manual_entry():
    document = doc(())
    p = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "BIBLIOGRAPHY")
    run_element = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = FIRST
    run_element.append(value)
    field.append(run_element)
    p._p.append(field)
    findings = run(document)
    assert not fail_ids(findings)
    assert {"111", "113", "114", "115", "117", "118", "121"} <= {f.rule_id for f in findings}
    assert all(UNREADABLE in f.message for f in findings)


def test_source_document_and_template_xml_remain_unchanged():
    template, document = doc(), doc((FIRST.replace("[J].", "[J]。"),))
    before_template = template.element.xml
    before_document = document.element.xml
    before_styles = document.styles.element.xml
    assert fail_ids(run(document, template))
    assert template.element.xml == before_template
    assert document.element.xml == before_document
    assert document.styles.element.xml == before_styles


def test_input_baseline_is_not_overwritten():
    baseline = {"old_key": {"important": 123}}
    addition = extract_bibliography_baseline(DocxContext(doc()))
    baseline.update(addition)
    assert baseline["old_key"] == {"important": 123}
    assert set(addition) == {"bibliography"}


def test_unreadable_template_example_does_not_leave_a_false_unanimous_baseline():
    template = doc((FIRST, SECOND.replace("示例乙", "含句点.题名")))
    changed = FIRST.replace("[J].", "[J]。")
    findings = run(doc((changed,)), template)
    assert not fail_ids(findings)
    assert unknowns(findings, "111")


def test_revised_template_example_blocks_its_reference_group():
    template = doc()
    paragraph = template.paragraphs[-1]
    revision = OxmlElement("w:rPrChange")
    paragraph.runs[0]._r.get_or_add_rPr().append(revision)
    changed = FIRST.replace("[J].", "[J]。")
    findings = run(doc((changed,)), template)
    assert not fail_ids(findings)
    assert unknowns(findings, "111")


def test_dynamic_template_entry_does_not_make_other_manual_examples_unanimous():
    template = doc((FIRST,))
    p = template.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "BIBLIOGRAPHY")
    value = OxmlElement("w:t")
    value.text = SECOND
    r = OxmlElement("w:r")
    r.append(value)
    field.append(r)
    p._p.append(field)
    findings = run(doc((FIRST.replace("[J].", "[J]。"),)), template)
    assert not fail_ids(findings)
    assert unknowns(findings, "111")


def test_present_unparsed_doi_url_or_date_is_unknown_instead_of_missing():
    for before, after in (("DOI: 10.1234/abc", "DOI："),
                          ("https://example.org/a", "URL：待确认"),
                          ("[2024-01-02]", "引用日期：2024年1月2日")):
        findings = run(doc((FIRST.replace(before, after),)))
        assert "117" not in fail_ids(findings)
        assert any("不能认定缺失" in f.message for f in unknowns(findings, "117"))


def test_one_padded_date_example_does_not_establish_padding_requirement():
    template = doc((FIRST,))
    actual = FIRST.replace("2024-01-02", "2024-1-2")
    assert "117" not in fail_ids(run(doc((actual,)), template))
    assert "117" in fail_ids(run(doc((actual,)), doc()))


def test_repeating_identical_date_does_not_count_as_independent_padding_evidence():
    template = doc((FIRST, SECOND.replace("2024-02-03", "2024-01-02")))
    actual = FIRST.replace("2024-01-02", "2024-1-2")
    assert "117" not in fail_ids(run(doc((actual,)), template))


def test_order_resets_at_unreadable_actual_entry():
    document = doc((FIRST.replace("[1]", "[3]"), SECOND, FIRST))
    paragraph = document.paragraphs[-2]
    revision = OxmlElement("w:rPrChange")
    paragraph.runs[0]._r.get_or_add_rPr().append(revision)
    findings = run(document)
    assert "121" not in fail_ids(findings)
    assert unknowns(findings, "121")


def test_order_resets_at_actual_entry_without_readable_number():
    document = doc((FIRST.replace("[1]", "[3]"), SECOND.replace("[2] ", ""), FIRST))
    findings = run(document)
    assert "121" not in fail_ids(findings)
    assert unknowns(findings, "121")


def test_hidden_hyperlink_target_is_not_reported_missing():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    document = doc((FIRST.replace(" https://example.org/a", ""),))
    paragraph = document.paragraphs[-1]
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to("https://example.org/a", RT.HYPERLINK, is_external=True))
    r, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "访问地址"
    r.append(t)
    link.append(r)
    paragraph._p.append(link)
    findings = run(document)
    assert "117" not in fail_ids(findings)
    assert unknowns(findings, "117")


def test_numbered_explicit_instructions_do_not_poison_bibliography_examples():
    instructions = ("1. " + REQUIRED, "2. " + AUTHORS, "3. " + ORDER)
    template = doc(instructions=instructions)
    assert not fail_ids(run(doc(), template))
    assert "113" in fail_ids(run(doc((FIRST.replace("张三，李四", "张三，李四，王五，赵六"),)), template))
