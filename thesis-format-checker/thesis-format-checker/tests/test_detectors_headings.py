from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.headings import (
    _heading_level, collect_heading_records, extract_heading_numbering_baseline,
    check_heading_numbering_errors,
)
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def element(name, **attrs):
    node = OxmlElement("w:" + name)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    return node


def define_list(doc, num_id=41, abstract_id=31, start=1, suffix="tab", fmt="decimal", pattern=None,
                restart=None, override_start=None, explicit_levels=3):
    root = doc.part.numbering_part.element
    abstract = element("abstractNum", abstractNumId=abstract_id)
    abstract.append(element("multiLevelType", val="multilevel"))
    for index in range(explicit_levels):
        level = element("lvl", ilvl=index)
        level.append(element("start", val=start))
        level.append(element("numFmt", val=fmt))
        level.append(element("lvlText", val=pattern if pattern is not None else ".".join("%" + str(i + 1) for i in range(index + 1))))
        level.append(element("suff", val=suffix))
        level.append(element("pStyle", val="Heading" + str(index + 1)))
        if restart is not None:
            level.append(element("lvlRestart", val=restart))
        abstract.append(level)
    root.append(abstract)
    num = element("num", numId=num_id)
    num.append(element("abstractNumId", val=abstract_id))
    if override_start is not None:
        override = element("lvlOverride", ilvl=0)
        override.append(element("startOverride", val=override_start))
        num.append(override)
    root.append(num)
    return abstract, num


def heading(doc, text, level=1, num_id=None):
    paragraph = doc.add_heading(text, level=level)
    if num_id is not None:
        num_pr = element("numPr")
        num_pr.append(element("ilvl", val=level - 1))
        num_pr.append(element("numId", val=num_id))
        paragraph._p.get_or_add_pPr().append(num_pr)
    return paragraph


def check(template, actual):
    baseline = extract_heading_numbering_baseline(DocxContext(template))
    return check_heading_numbering_errors(baseline, DocxContext(actual))


def failed(findings, rule):
    return [item for item in findings if family_of(item.rule_id) == rule and item.status == Status.FAIL]


def auto_template():
    template = Document()
    define_list(template)
    for level, text in ((1, "绪论"), (2, "研究背景"), (3, "技术现状"), (2, "研究目标"), (1, "方法")):
        heading(template, text, level, 41)
    return template


def test_collect_list_definitions_ids_local_and_actual_numbers():
    template = auto_template()
    actual = Document()
    define_list(actual, 77, 62)
    for level, text in ((1, "第一篇"), (2, "背景"), (3, "应用"), (2, "目标"), (1, "第二篇")):
        heading(actual, text, level, 77)
    records = collect_heading_records(DocxContext(actual))
    assert [record["number"] for record in records] == ["1", "1.1", "1.1.1", "1.2", "2"]
    assert records[1]["display_text"] == "1.1\t背景"
    findings = check(template, actual)
    assert not [finding for finding in findings if finding.status == Status.FAIL]
    assert any(family_of(item.rule_id) == "054" and item.status == Status.UNKNOWN for item in findings)


def test_template_defined_non_decimal_numbering_system_is_supported():
    template, actual = Document(), Document()
    define_list(template, fmt="upperRoman", pattern="编%1)", explicit_levels=1)
    define_list(actual, num_id=77, abstract_id=62, fmt="upperRoman",
                pattern="编%1)", explicit_levels=1)
    for doc, num_id in ((template, 41), (actual, 77)):
        heading(doc, "概述", 1, num_id)
        heading(doc, "方法", 1, num_id)

    assert not failed(check(template, actual), "051")
    assert not failed(check(template, actual), "052")

    different = Document()
    define_list(different, num_id=88, abstract_id=63, fmt="upperLetter",
                pattern="%1.", explicit_levels=1)
    heading(different, "概述", 1, 88)
    heading(different, "方法", 1, 88)
    assert failed(check(template, different), "051")


def test_jump_down_reported_return_up_allowed_and_template_exemptions():
    template, actual = auto_template(), Document()
    heading(actual, "1\t首章", 1)
    heading(actual, "1.1.1\t三级", 3)
    heading(actual, "2\t第二章", 1)
    assert len(failed(check(template, actual), "048")) == 1
    template = Document()
    heading(template, "1 章", 1)
    heading(template, "1.1.1 特殊层级", 3)
    assert not failed(check(template, actual), "048")


def test_first_heading_missing_template_parent():
    actual = Document()
    heading(actual, "1.1 小节", 2)
    assert failed(check(auto_template(), actual), "048")


def test_049_only_inconsistent_outside_template_not_uniform_wrong():
    template = Document()
    for text in ("1 甲", "2 乙"):
        heading(template, text).runs[0].font.size = Pt(18)
    actual = Document()
    heading(actual, "1 丙").runs[0].font.size = Pt(18)
    wrong = heading(actual, "2 丁")
    wrong.runs[0].font.size = Pt(13)
    findings = failed(check(template, actual), "049")
    assert len(findings) == 1 and findings[0].anchor._p is wrong._p
    actual.paragraphs[0].runs[0].font.size = Pt(13)
    assert not failed(check(template, actual), "049")


def test_049_template_allowed_variants_are_legal():
    template, actual = Document(), Document()
    for doc in (template, actual):
        heading(doc, "1 甲").runs[0].font.size = Pt(18)
        heading(doc, "2 乙").runs[0].font.size = Pt(13)
    assert not failed(check(template, actual), "049")


def test_050_manual_sequence_start_gap_duplicate_and_parent_values():
    actual = Document()
    for level, text in ((1, "1\t开头"), (2, "1.1\t第一节"), (2, "1.3\t跳号"),
                        (2, "1.3\t重号"), (1, "2\t另一章"), (2, "1.1\t错误父号")):
        heading(actual, text, level)
    matches = failed(check(auto_template(), actual), "050")
    assert len(matches) == 3
    assert any("父编号" in item.actual for item in matches)


def test_050_valid_nested_manual_restarts_and_055_manual():
    actual = Document()
    for level, text in ((1, "1\t开头"), (2, "1.1\t第一节"), (2, "1.2\t第二节"),
                        (1, "2\t另一章"), (2, "2.1\t第一节")):
        heading(actual, text, level)
    findings = check(auto_template(), actual)
    assert not failed(findings, "050")
    assert len(failed(findings, "055")) == 5


def test_050_sequences_restart_per_chapter_when_template_has_no_level1_basis():
    """论文每章的子级编号从 1 重新开始：序列必须按“章”分组，不能按“级”串线。

    模板只提供二级标题基准（没有一级标题基准）时，一级标题不会进入比较，
    旧实现按“级”共享状态，于是第二章的“1”被当成第一章的延续而误报跳号。
    """

    template = Document()
    heading(template, "第一章\t章标题", 1)
    heading(template, "1 研究背景", 2)
    heading(template, "2 研究目标", 2)

    def build(extra_gap=False):
        actual = Document()
        for chapter in ("第一章", "第二章"):
            heading(actual, chapter + "\t章标题", 1)
            heading(actual, "1\t第一节", 2)
            heading(actual, "2\t第二节", 2)
        if extra_gap:
            heading(actual, "4\t跳号小节", 2)
        return actual

    assert not failed(check(template, build()), "050")
    # 同一章内真正的跳号仍要报出来。
    assert len(failed(check(template, build(extra_gap=True)), "050")) == 1


def test_050_heading_without_unique_template_basis_is_recorded_not_judged():
    """模板对某一条标题没有唯一序列基准时只记账不判定，下一条不能误报“应为 1”。"""

    template = Document()
    define_list(template)
    heading(template, "自动编号节", 3, 41)
    heading(template, "2.1 实验目的", 3)  # 与论文标题同名 → 候选被缩成这一条，无法确认步长
    actual = Document()
    heading(actual, "第一章\t章标题", 1)
    heading(actual, "2\t实验", 2)
    heading(actual, "2.1 实验目的", 3)
    heading(actual, "2.2 实验方法", 3)
    assert not failed(check(template, actual), "050")


def test_front_matter_before_abstract_is_not_a_heading():
    """摘要之前的封面/声明只按模板同位置比对，不进入标题层级。"""

    actual = Document()
    declaration = actual.add_heading("原创性声明", 1)
    actual.add_heading("基于深度学习的××研究", 1)
    heading(actual, "摘  要", 1)
    heading(actual, "第一章 绪论", 1)
    ctx = DocxContext(actual)
    assert _heading_level(ctx.paragraphs[0]) is None
    assert _heading_level(ctx.paragraphs[1]) is None
    assert _heading_level(ctx.paragraphs[2]) == 1
    assert _heading_level(ctx.paragraphs[3]) == 1
    assert not any(record["paragraph"]._p is declaration._p
                   for record in collect_heading_records(ctx))


def test_050_start_override_mismatch_preserves_definitions():
    actual = Document()
    define_list(actual, override_start=7)
    heading(actual, "错误开始", num_id=41)
    result = collect_heading_records(DocxContext(actual))
    assert result[0]["number"] == "7"
    assert failed(check(auto_template(), actual), "050")


def test_050_two_independent_lists_not_merged_as_duplicate():
    template = auto_template()
    actual = Document()
    define_list(actual, 81, 91)
    define_list(actual, 82, 92)
    heading(actual, "列表甲", num_id=81)
    heading(actual, "列表乙", num_id=82)
    findings = check(template, actual)
    assert not failed(findings, "050")
    assert any(family_of(item.rule_id) == "050" and "多个实例" in item.message for item in findings)


def test_050_manual_template_requires_repeat_evidence_and_dynamic_start_step():
    template, actual = Document(), Document()
    heading(template, "5 第一个")
    heading(template, "7 第二个")
    heading(actual, "5 对应")
    heading(actual, "6 错误步长")
    assert failed(check(template, actual), "050")
    template = Document()
    heading(template, "5 孤立示例")
    findings = check(template, actual)
    assert not failed(findings, "050")
    # 模板只有一个孤立编号，无法确认步长：现在是“不适用”（不写批注），仍不猜结论。
    assert any(family_of(item.rule_id) == "050"
               and item.status in (Status.UNKNOWN, Status.NOT_APPLICABLE) for item in findings)


def test_051_chapter_and_dotted_patterns_and_052_separator():
    actual = Document()
    heading(actual, "第1章\t开头")
    heading(actual, "1.1 研究", 2)
    findings = check(auto_template(), actual)
    assert len(failed(findings, "051")) == 1
    assert len(failed(findings, "052")) == 1


def test_052_auto_suffix_and_extra_literal_tab():
    actual = Document()
    define_list(actual, suffix="space")
    heading(actual, "\t前置多余Tab", num_id=41)
    assert failed(check(auto_template(), actual), "052")


def test_053_numbered_and_unnumbered_roles_not_confused():
    template, actual = auto_template(), Document()
    heading(template, "摘要", 1)
    heading(actual, "缺少章号", 1)
    heading(actual, "1 摘要", 1)
    findings = failed(check(template, actual), "053")
    assert len(findings) == 2


def test_number_in_title_year_not_assumed_chapter_number():
    template, actual = Document(), Document()
    heading(template, "2026年研究进展")
    heading(actual, "2027年研究进展")
    assert not failed(check(template, actual), "053")
    assert collect_heading_records(DocxContext(actual))[0]["number"] is None


def test_054_never_fabricates_page_placement_even_keep_with_next_on():
    template, actual = Document(), Document()
    heading(template, "标题").paragraph_format.keep_with_next = True
    heading(actual, "标题").paragraph_format.keep_with_next = True
    findings = check(template, actual)
    assert not failed(findings, "054")
    assert any(family_of(item.rule_id) == "054" and UNREADABLE in item.message for item in findings)


def test_055_manual_template_does_not_require_automatic_lists():
    template, actual = Document(), Document()
    for doc in (template, actual):
        heading(doc, "第1章 甲")
        heading(doc, "第2章 乙")
    assert not failed(check(template, actual), "055")


def test_style_inherited_numpr_and_local_pstyle_level():
    actual = Document()
    define_list(actual)
    for level in (1, 2):
        num_pr = element("numPr")
        num_pr.append(element("numId", val=41))
        actual.styles["Heading " + str(level)].element.get_or_add_pPr().append(num_pr)
    heading(actual, "父", 1)
    heading(actual, "子", 2)
    records = collect_heading_records(DocxContext(actual))
    assert [item["number"] for item in records] == ["1", "1.1"]
    paragraph = actual.paragraphs[1]
    pnum = element("numPr")
    pnum.append(element("numId", val=0))
    paragraph._p.get_or_add_pPr().append(pnum)
    assert collect_heading_records(DocxContext(actual))[1]["automatic"] is False


def test_custom_heading_style_outline_inherited_and_no_text_bold_guessing():
    doc = Document()
    custom = doc.styles.add_style("School chapter", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles["Heading 1"]
    p = doc.add_paragraph("1 绪论", custom)
    doc.add_paragraph("2 加粗正文").runs[0].bold = True
    records = collect_heading_records(DocxContext(doc))
    assert len(records) == 1 and records[0]["paragraph"]._p is p._p


def test_automatic_nonheading_list_items_increment_real_sequence():
    doc = Document()
    define_list(doc)
    heading(doc, "第一章", num_id=41)
    p = doc.add_paragraph("共享列表的非标题")
    num_pr = element("numPr")
    num_pr.append(element("numId", val=41))
    p._p.get_or_add_pPr().append(num_pr)
    heading(doc, "第三章", num_id=41)
    assert [item["number"] for item in collect_heading_records(DocxContext(doc))] == ["1", "3"]
    assert failed(check(auto_template(), doc), "050")


def test_broken_list_or_fields_reports_unknown_not_missing_number():
    actual = Document()
    paragraph = heading(actual, "坏列表", num_id=9876)
    findings = check(auto_template(), actual)
    assert not failed(findings, "053")
    assert any(family_of(item.rule_id) == "051" and UNREADABLE in item.message for item in findings)
    actual = Document()
    paragraph = heading(actual, "动态标题")
    field = element("fldSimple", instr=" SEQ Chapter ")
    paragraph._p.append(field)
    findings = check(auto_template(), actual)
    assert not failed(findings, "055")
    assert any(family_of(item.rule_id) == "055" and item.status == Status.UNKNOWN for item in findings)


def test_unsupported_chinese_auto_display_still_checks_readable_definition():
    template, actual = Document(), Document()
    define_list(template, fmt="chineseCounting", pattern="第%1章")
    define_list(actual, fmt="chineseCounting", pattern="第%1节")
    heading(template, "甲", num_id=41)
    heading(actual, "乙", num_id=41)
    records = collect_heading_records(DocxContext(actual))
    assert records[0]["number_known"] is False and records[0]["display_text"] is None
    findings = check(template, actual)
    assert failed(findings, "051")
    assert not failed(findings, "053")


def test_lvl_override_replacement_and_restart_none_read():
    doc = Document()
    abstract, num = define_list(doc)
    override = element("lvlOverride", ilvl=1)
    replacement = deepcopy(abstract.findall(qn("w:lvl"))[1])
    replacement.find(qn("w:lvlText")).set(qn("w:val"), "%2)")
    replacement.append(element("lvlRestart", val=0))
    override.append(replacement)
    num.append(override)
    for level, text in ((1, "甲"), (2, "乙"), (1, "丙"), (2, "丁")):
        heading(doc, text, level, 41)
    assert [item["number"] for item in collect_heading_records(DocxContext(doc))] == ["1", "1)", "2", "2)"]


def test_toc_field_cache_not_repeated_as_body_headings():
    doc = Document()
    heading(doc, "1 正文")
    begin = doc.add_paragraph()
    run = element("r")
    run.append(element("fldChar", fldCharType="begin"))
    instruction = element("instrText")
    instruction.text = ' TOC \\o "1-3" '
    run.append(instruction)
    begin._p.append(run)
    heading(doc, "1 正文缓存")
    end = doc.add_paragraph()
    run = element("r")
    run.append(element("fldChar", fldCharType="end"))
    end._p.append(run)
    assert len(collect_heading_records(DocxContext(doc))) == 1


def test_template_and_document_parts_never_mutated_by_extraction_or_checks():
    template, actual = auto_template(), auto_template()
    parts = [template.element, template.styles.element, template.part.numbering_part.element,
             actual.element, actual.styles.element, actual.part.numbering_part.element]
    before = [item.xml for item in parts]
    check(template, actual)
    assert [item.xml for item in parts] == before


def test_numbering_exception_local_to_rules_not_crash():
    actual = Document()
    abstract, _ = define_list(actual)
    abstract.find(qn("w:lvl")).find(qn("w:start")).set(qn("w:val"), "broken")
    heading(actual, "坏数字", num_id=41)
    findings = check(auto_template(), actual)
    assert any(item.status == Status.UNKNOWN and UNREADABLE in item.message for item in findings)


def test_matching_template_literal_extra_separator_is_not_hardcoded_failure():
    template, actual = Document(), Document()
    for doc in (template, actual):
        define_list(doc)
        heading(doc, "\t额外前导Tab由模板示例明确给出", num_id=41)
    assert not failed(check(template, actual), "052")


def test_deleted_heading_does_not_consume_number_and_live_revision_unknown():
    template, actual = auto_template(), Document()
    define_list(actual)
    deleted = heading(actual, "已删除标题", num_id=41)
    wrapper = element("del")
    index = actual.element.body.index(deleted._p)
    actual.element.body.remove(deleted._p)
    wrapper.append(deleted._p)
    actual.element.body.insert(index, wrapper)
    heading(actual, "当前标题", num_id=41)
    records = collect_heading_records(DocxContext(actual))
    assert len(records) == 1 and records[0]["number"] == "1"
    revised = actual.paragraphs[0]
    revised._p.get_or_add_pPr().append(element("pPrChange", id=7, author="测试"))
    findings = check(template, actual)
    assert any(family_of(item.rule_id) == "050" and "修订" in item.message for item in findings)
    assert not failed(findings, "053")


def test_list_indent_not_misread_as_zero_or_implicitly_verified():
    template, actual = auto_template(), auto_template()
    abstract = next(node for node in actual.part.numbering_part.element.findall(qn("w:abstractNum"))
                    if node.get(qn("w:abstractNumId")) == "31")
    props = element("pPr")
    props.append(element("ind", left=720, hanging=360))
    abstract.find(qn("w:lvl")).append(props)
    findings = check(template, actual)
    assert any(family_of(item.rule_id) == "049" and "列表缩进" in item.message and item.status == Status.UNKNOWN for item in findings)
    assert any(family_of(item.rule_id) == "052" and "列表缩进" in item.message and item.status == Status.UNKNOWN for item in findings)


def test_manual_punctuation_without_space_is_numbered_and_compares_shape():
    template, actual = Document(), Document()
    heading(template, "1、标题")
    heading(template, "2、另一标题")
    heading(actual, "1.标题")
    assert failed(check(template, actual), "051")
    assert not failed(check(template, actual), "053")


def test_broken_outline_level_prompts_instead_of_dropping_heading():
    actual = Document()
    paragraph = heading(actual, "1 标题")
    paragraph._p.get_or_add_pPr().append(element("outlineLvl", val="broken"))
    findings = check(auto_template(), actual)
    assert any(family_of(item.rule_id) == "048" and UNREADABLE in item.message for item in findings)
    assert not failed(findings, "048")
