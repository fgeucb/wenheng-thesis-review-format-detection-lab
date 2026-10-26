from __future__ import annotations

from copy import deepcopy

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.shared import Pt, RGBColor

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.paragraphs import (
    _body_records, check_paragraph_errors, extract_paragraph_baseline,
)
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]


def test_body_records_exclude_cover_and_declaration():
    """摘要之前的封面/声明不属于正文范围，即使显式配置了正文字体样式也不采集。"""

    document = Document()
    document.add_paragraph("论文原创性声明及使用授权说明")
    document.add_paragraph("本人郑重声明：所呈交的论文是本人在导师指导下独立完成的研究成果。")
    document.add_heading("摘  要", 1)
    document.add_paragraph("摘要正文。")
    mapping = {"document": {"body_styles": ["Normal"]}}
    texts = [record["paragraph"].text for record in _body_records(DocxContext(document), mapping)]
    assert texts == ["摘要正文。"]



def _document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.font.bold = False
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = False
    normal.paragraph_format.keep_with_next = False
    normal.paragraph_format.left_indent = Pt(0)
    normal.paragraph_format.right_indent = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(24)
    doc.add_paragraph("正文中文 text 123。")
    doc.add_paragraph("正文中文 words 456。")
    return doc


def _check(template, document):
    return check_paragraph_errors(extract_paragraph_baseline(DocxContext(template)), DocxContext(document))


def _fails(findings, rule):
    return [f for f in findings if family_of(f.rule_id) == rule and f.status == Status.FAIL]


def _unknown(findings, rule):
    # 模板未涉及的结论现在是“不适用”（不写批注），测试中与“无法判定”一并接受。
    return [f for f in findings if family_of(f.rule_id) == rule
            and f.status in (Status.UNKNOWN, Status.NOT_APPLICABLE)]


def test_matching_template_is_read_only_and_has_no_failures():
    template, document = _document(), _document()
    before = (template.element.xml, template.styles.element.xml, document.element.xml, document.styles.element.xml)
    findings = _check(template, document)
    assert not [f for f in findings if f.status == Status.FAIL]
    assert before == (template.element.xml, template.styles.element.xml, document.element.xml, document.styles.element.xml)


@pytest.mark.parametrize("property,value,rule", [
    ("alignment", WD_ALIGN_PARAGRAPH.CENTER, "071"),
    ("left_indent", Pt(19), "076"),
    ("right_indent", Pt(9), "076"),
    ("first_line_indent", Pt(-10), "076"),
    ("widow_control", False, "077"),
    ("keep_with_next", True, "077"),
    ("keep_together", True, "077"),
])
def test_each_new_paragraph_setting(property, value, rule):
    template, document = _document(), _document()
    setattr(document.paragraphs[0].paragraph_format, property, value)
    findings = _check(template, document)
    assert _fails(findings, rule)
    assert _fails(findings, rule)[0].anchor._p is document.paragraphs[0]._p


@pytest.mark.parametrize("property,value", [
    ("name", "Arial"), ("size", Pt(15)), ("bold", True), ("color", RGBColor(220, 20, 30)),
])
def test_run_overlay_inconsistent_body_format(property, value):
    template, document = _document(), _document()
    font = document.paragraphs[1].runs[0].font
    if property == "color":
        font.color.rgb = value
    else:
        setattr(font, property, value)
    findings = _check(template, document)
    assert _fails(findings, "070")
    assert all(f.anchor._p is document.paragraphs[1]._p for f in _fails(findings, "070"))


def test_uniform_wrong_font_not_duplicated_under_consistency_rule():
    template, document = _document(), _document()
    document.styles["Normal"].font.name = "Arial"
    assert not _fails(_check(template, document), "070")


def test_character_style_emphasis_compared_as_separate_use():
    template, document = _document(), _document()
    for doc in (template, document):
        style = doc.styles.add_style("EmphasisLocal", WD_STYLE_TYPE.CHARACTER)
        style.font.bold = True
        doc.paragraphs[0].add_run("重要 term").style = style
    assert not _fails(_check(template, document), "070")


def test_custom_body_inherited_settings_and_direct_override():
    template, document = _document(), _document()
    for doc in (template, document):
        style = doc.styles.add_style("正文段", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        for p in doc.paragraphs:
            p.style = style
    document.paragraphs[1].runs[0].font.size = Pt(17)
    document.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    findings = _check(template, document)
    assert _fails(findings, "070") and _fails(findings, "071")


def test_template_conflict_compares_with_same_position_sample():
    """模板有两种样例时按“同一层级 + 同一位置”取基准，不再全篇无法判定。"""

    template, document = _document(), _document()
    template.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    findings = _check(template, document)
    fails = _fails(findings, "071")
    assert len(fails) == 1
    assert fails[0].anchor._p is document.paragraphs[0]._p
    assert not _unknown(findings, "071")


def test_template_conflict_does_not_poison_other_paragraphs():
    """只有与模板同位置样例不一致的那一段报错，其他段落不受牵连。"""

    template, document = _document(), _document()
    for doc in (template, document):
        doc.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.paragraphs[1].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    findings = _check(template, document)
    fails = _fails(findings, "071")
    assert len(fails) == 1
    assert fails[0].anchor._p is document.paragraphs[1]._p
    assert not _unknown(findings, "071")


def test_body_style_number_suffix_shares_normal_body_alignment_basis():
    template, document = Document(), Document()
    body_style = template.styles.add_style("正文1", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = template.styles["Normal"]
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    title = template.add_paragraph("TITLE TITLE TITLE")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ("模板正文一。", "模板正文二。", "模板正文三。"):
        template.add_paragraph(text, body_style)
    document.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    actual = document.add_paragraph("实际正文。")
    findings = _check(template, document)
    assert not _fails(findings, "071")
    assert _body_records(DocxContext(document))[0]["role"] == "正文"


def test_front_toc_back_table_and_formula_not_treated_as_body():
    template, document = _document(), _document()
    for doc in (template, document):
        doc.paragraphs[0].insert_paragraph_before("封面")
        p = doc.paragraphs[1].insert_paragraph_before("论文题目：普通文本")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[2].insert_paragraph_before("第一章 绪论", "Heading 1")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格文字"
        p = doc.add_paragraph("公式段")
        p._p.append(OxmlElement("m:oMath"))
        doc.add_paragraph("参考文献")
        doc.add_paragraph("文献分类", "Heading 2")
        doc.add_paragraph("文献格式不执行")
    document.paragraphs[1].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.paragraphs[-1].paragraph_format.keep_with_next = True
    assert not _fails(_check(template, document), "071")
    assert not _fails(_check(template, document), "077")


def test_no_first_indent_allowed_for_heading_and_caption():
    template, document = _document(), _document()
    for doc in (template, document):
        for text, style in (("第一章 绪论", "Heading 1"), ("图1 样例", "Caption"), ("表1 样例", "Caption")):
            p = doc.add_paragraph(text, style)
            p.paragraph_format.first_line_indent = Pt(0)
    for p in document.paragraphs[-3:]:
        p.paragraph_format.first_line_indent = Pt(14)
    findings = _check(template, document)
    assert len(_fails(findings, "073")) == 3


def test_template_allows_heading_indent_not_hardcoded_zero():
    template, document = _document(), _document()
    for doc in (template, document):
        p = doc.add_paragraph("第一章 绪论", "Heading 1")
        p.paragraph_format.first_line_indent = Pt(12)
    assert not _fails(_check(template, document), "073")


def test_character_first_indent_is_not_converted_to_points():
    template, document = _document(), _document()
    for doc in (template, document):
        p = doc.add_paragraph("第一章 绪论", "Heading 1")
        p.paragraph_format.first_line_indent = Pt(0)
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0" if doc is template else "100")
    findings = _check(template, document)
    assert _fails(findings, "073")
    assert "1字符" in _fails(findings, "073")[0].message


def test_numbered_heading_without_extra_indent_is_comparable():
    template, document = _document(), _document()
    for doc in (template, document):
        p = doc.add_paragraph("第一章 绪论", "Heading 1")
        p.paragraph_format.first_line_indent = Pt(0)
        numpr = OxmlElement("w:numPr")
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), "1")
        numpr.append(numid)
        p._p.get_or_add_pPr().append(numpr)
    findings = _check(template, document)
    assert not _fails(findings, "073")
    assert not _unknown(findings, "073")


def test_numbered_heading_with_explicit_first_line_indent_is_reported():
    template, document = _document(), _document()
    for doc in (template, document):
        p = doc.add_paragraph("第一章 绪论", "Heading 1")
        numpr = OxmlElement("w:numPr")
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), "1")
        numpr.append(numid)
        p._p.get_or_add_pPr().append(numpr)
    document.paragraphs[-1].paragraph_format.first_line_indent = Pt(24)
    findings = _check(template, document)
    assert _fails(findings, "073")


def test_revision_is_unknown_and_char_indent_converts_to_points():
    template, document = _document(), _document()
    document.paragraphs[0]._p.get_or_add_pPr().get_or_add_ind().set(qn("w:leftChars"), "200")
    document.paragraphs[1]._p.get_or_add_pPr().append(OxmlElement("w:pPrChange"))
    findings = _check(template, document)
    # 修订段落无法判断；字符单位缩进按字号换算成磅值后照常比较（模板为 0 → 报左缩进差异）。
    assert _unknown(findings, "071")
    assert _fails(findings, "076")


def _add_note(document, size=10, kind="footnote", omit_relationship=False):
    reference = OxmlElement("w:" + kind + "Reference")
    reference.set(qn("w:id"), "1")
    document.paragraphs[0].add_run()._r.append(reference)
    if omit_relationship:
        return
    root = parse_xml(
        f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:{kind} w:id="1"><w:p><w:r><w:rPr><w:sz w:val="{size*2}"/></w:rPr>'
        f'<w:t>注释 note 123</w:t></w:r></w:p></w:{kind}></w:{kind}s>'
    )
    part = Part(PackURI(f"/word/{kind}s.xml"),
                f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
                root.xml.encode() if hasattr(root, "xml") else __import__("lxml.etree", fromlist=["tostring"]).tostring(root),
                document.part.package)
    document.part.relate_to(part, f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s")


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_note_font_size_compares_by_purpose_and_anchors_body_reference(kind):
    template, document = _document(), _document()
    _add_note(template, 10, kind)
    _add_note(document, 13, kind)
    before = document.element.xml
    findings = _check(template, document)
    assert _fails(findings, "069")
    assert _fails(findings, "069")[0].anchor._p is document.paragraphs[0]._p
    assert "13磅" in _fails(findings, "069")[0].message
    assert document.element.xml == before


def test_missing_note_relationship_and_missing_baseline_do_not_pass():
    template, document = _document(), _document()
    _add_note(document, omit_relationship=True)
    findings = _check(template, document)
    assert _unknown(findings, "069")
    assert UNREADABLE in _unknown(findings, "069")[0].message
    document = _document()
    _add_note(document)
    assert _unknown(_check(template, document), "069")


def test_reused_rules_are_not_reimplemented_in_additions():
    template, document = _document(), _document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(16)
    document.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
    document.paragraphs[0].paragraph_format.line_spacing = 2
    document.paragraphs[0].paragraph_format.space_after = Pt(20)
    findings = _check(template, document)
    assert not {f.rule_id for f in findings}.intersection({"067", "068", "069", "072", "074", "075"})


def test_mapping_explicit_body_style_and_range():
    template, document = _document(), _document()
    for doc in (template, document):
        style = doc.styles.add_style("SpecialBody", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        doc.paragraphs[0].style = style
    document.paragraphs[1].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mapping = {side: {"body_styles": ["SpecialBody"]} for side in ("template", "document")}
    baseline = extract_paragraph_baseline(DocxContext(template), mapping)
    assert not _fails(check_paragraph_errors(baseline, DocxContext(document), mapping), "071")


def test_nondefault_template_alignment_and_pagination_are_respected():
    template, document = _document(), _document()
    for doc in (template, document):
        doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.styles["Normal"].paragraph_format.keep_together = True
        doc.styles["Normal"].paragraph_format.widow_control = False
    assert not _fails(_check(template, document), "071")
    assert not _fails(_check(template, document), "077")


@pytest.mark.parametrize("kind", ["heading", "caption", "footnote"])
def test_consistency_covers_other_explicit_roles(kind):
    template, document = _document(), _document()
    if kind == "footnote":
        for doc in (template, document):
            _add_note(doc)
            rel = next(r for r in doc.part.rels.values() if r.reltype.endswith("/footnotes"))
            root = parse_xml(rel.target_part.blob)
            p = next(root.iter(qn("w:p")))
            second = deepcopy(p[0])
            second.find(qn("w:t")).text = "第二段注释 words 456"
            if doc is document:
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "FF0000")
                second.find(qn("w:rPr")).append(color)
            p.append(second)
            rel.target_part._blob = __import__("lxml.etree", fromlist=["tostring"]).tostring(root)
    else:
        for doc in (template, document):
            style = "Heading 1" if kind == "heading" else "Caption"
            texts = ("第一章 绪论", "第二章 方法") if kind == "heading" else ("图1 示例", "图2 示例")
            for text in texts:
                p = doc.add_paragraph(text, style)
                p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        document.paragraphs[-1].runs[0].font.color.rgb = RGBColor(255, 0, 0)
    findings = _check(template, document)
    assert _fails(findings, "070")
    if kind == "footnote":
        assert _fails(findings, "070")[0].anchor._p is document.paragraphs[0]._p


def test_template_revision_poisoned_role_not_partial_majority():
    template, document = _document(), _document()
    template.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement("w:pPrChange"))
    document.paragraphs[1].runs[0].font.size = Pt(20)
    findings = _check(template, document)
    assert not _fails(findings, "070")
    assert _unknown(findings, "070")


def test_automatic_vs_rgb_color_is_not_claimed_visually_different():
    template, document = _document(), _document()
    for doc in (template, document):
        doc.styles["Normal"].font.color.rgb = None
    document.paragraphs[1].runs[0].font.color.rgb = RGBColor(0, 0, 0)
    findings = _check(template, document)
    assert not _fails(findings, "070")
    assert _unknown(findings, "070")
