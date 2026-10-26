"""结果缓存（性能优化第 5 条）与静态先筛（第 3 条）的回归测试。"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from thesis_format_checker.core.cache import (
    ResultCache,
    code_fingerprint,
    file_fingerprint,
    summary_from_payload,
    summary_payload,
)
from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.models import Finding, RunSummary, Severity, Status
from thesis_format_checker.detectors.rendered import layout_requirements


# ---------------------------------------------------------------------------
# 序列化与指纹
# ---------------------------------------------------------------------------

def test_file_fingerprint_changes_with_content(runtime_dir):
    first, second = runtime_dir / "a.bin", runtime_dir / "b.bin"
    first.write_bytes(b"same")
    second.write_bytes(b"same")
    assert file_fingerprint(first) == file_fingerprint(second)
    second.write_bytes(b"different")
    assert file_fingerprint(first) != file_fingerprint(second)
    assert file_fingerprint(runtime_dir / "missing.bin") == ""


def test_summary_roundtrip_keeps_findings_and_render_facts():
    summary = RunSummary(
        sections_checked=3,
        body_paragraphs_checked=812,
        findings=[Finding(rule_id="069", title="正文字号", status=Status.FAIL,
                          message="正文应 12 磅", severity=Severity.ERROR,
                          expected=12.0, actual=10.5, paragraph_index=870,
                          excerpt="本科生毕业论文（设计）", start=0, end=3,
                          merged_locations=[12, 15], merged_count=2)],
        unknowns=["【035】页眉动态内容"],
        render_facts={"total_pages": 67, "pages": {1: 1, 2: 1, 3: 2},
                      "section_of": {1: 1, 3: 2},
                      "sections": [{"index": 1, "first_page": 1, "last_page": 2}]},
    )
    restored = summary_from_payload(summary_payload(summary))
    assert restored.sections_checked == 3
    assert restored.body_paragraphs_checked == 812
    assert restored.unknowns == ["【035】页眉动态内容"]
    # 整数键（Word 段号 → 页码）必须原样还原，否则分页规则会全部读不到。
    assert restored.render_facts["pages"][3] == 2
    assert restored.render_facts["section_of"][3] == 2
    finding = restored.findings[0]
    assert finding.status is Status.FAIL and finding.severity is Severity.ERROR
    assert finding.expected == 12.0 and finding.actual == 10.5
    assert finding.merged_locations == [12, 15] and finding.merged_count == 2


def test_layout_roundtrip_restores_dataclasses(runtime_dir):
    """排版快照里的 BodyPoint/Pagination 必须还原成对象，页眉批注才落得下去。"""

    from thesis_format_checker.core.cache import decode, encode
    from thesis_format_checker.core.word_layout import BodyPoint, Page, Pagination

    layout = {
        "paragraphs": {5: {"box": {"page": 2, "y": 100.0}, "printed_number": 3}},
        "pagination": Pagination(pages=[Page(2, 0, top=BodyPoint(5, 0, 42, 100.0),
                                            bottom=BodyPoint(7, 3, 90, 700.0))],
                                 abstract_page=3),
    }
    restored = decode(json.loads(json.dumps(encode(layout))))
    pagination = restored["pagination"]
    assert isinstance(pagination, Pagination)
    assert pagination.abstract_page == 3
    page = pagination.pages[0]
    assert isinstance(page.top, BodyPoint) and page.top.paragraph == 5
    assert page.bottom.y == 700.0
    assert pagination.page_of_section(0) == [page]
    assert restored["paragraphs"][5]["printed_number"] == 3


# ---------------------------------------------------------------------------
# 缓存读写
# ---------------------------------------------------------------------------

def test_cache_store_and_load_roundtrip(runtime_dir):
    cache = ResultCache(directory=runtime_dir)
    key = cache.key("layout", document=None, params={"a": 1})
    assert cache.load("layout", key, "排版几何快照") is None
    # 段落号是整数键，缓存层要能原样还原（否则几何规则会按“字符串段落号”查不到）。
    assert cache.store("layout", key, {"paragraphs": {1: {"page": 2}}, "pages": (1, 2)})
    assert cache.load("layout", key, "排版几何快照") == {"paragraphs": {1: {"page": 2}},
                                                         "pages": (1, 2)}
    assert "排版几何快照" in cache.hits


def test_cache_key_changes_with_input_and_source(runtime_dir):
    document = runtime_dir / "d.bin"
    document.write_bytes(b"first")
    cache = ResultCache(directory=runtime_dir)
    first = cache.key("layout", document=document)
    document2 = runtime_dir / "d2.bin"
    document2.write_bytes(b"second")
    assert cache.key("layout", document=document2) != first
    assert cache.key("layout", document=document) == first        # 内容相同 → 键相同
    assert code_fingerprint() and len(code_fingerprint()) == 16


def test_closed_cache_never_reads_or_writes(runtime_dir):
    cache = ResultCache(directory=runtime_dir, enabled=False)
    key = cache.key("layout", params={"x": 1})
    assert cache.store("layout", key, {"a": 1}) is False
    assert cache.load("layout", key) is None
    assert cache.stats()["enabled"] is False


def test_broken_cache_file_is_ignored_and_removed(runtime_dir):
    cache = ResultCache(directory=runtime_dir)
    key = cache.key("layout", params={"x": 2})
    path = runtime_dir / "layout" / f"{key}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("{ not json", encoding="utf-8")
    assert cache.load("layout", key) is None
    assert not path.exists()


# ---------------------------------------------------------------------------
# 静态先筛：不需要排版几何时不读 Word
# ---------------------------------------------------------------------------

def _context_with_figure():
    document = Document()
    caption = document.add_paragraph("图 2-1 系统结构")
    # 题注必须挂“题注样式”或紧邻图表对象才算数（与检测器口径一致）。
    caption.style = document.styles["Caption"]
    document.add_paragraph("正文内容。")
    return DocxContext(document)


def test_layout_requirements_none_when_no_geometry_rule_needs_it():
    context = _context_with_figure()
    findings = [Finding(rule_id="032.1", title="空白页", status=Status.UNKNOWN,
                        message="需要分页"), Finding(rule_id="069", title="正文字号",
                                                status=Status.FAIL, message="字号不符")]
    assert layout_requirements(context, findings) is None


def test_layout_requirements_selects_only_needed_paragraphs():
    context = _context_with_figure()
    findings = [Finding(rule_id="078.1", title="题注位置", status=Status.UNKNOWN,
                        message="需要坐标")]
    request = layout_requirements(context, findings)
    assert request is not None
    assert 0 in request["indices"]          # 题注所在段落
    assert request["want_figures"] and request["want_tables"]


def test_layout_requirements_toc_uses_heading_paragraphs():
    document = Document()
    document.add_paragraph("1 绪论", style="Heading 1")
    context = DocxContext(document)
    findings = [Finding(rule_id="059", title="目录页码", status=Status.UNKNOWN,
                        message="需要实际页码")]
    request = layout_requirements(context, findings)
    assert request is not None and 0 in request["indices"]
    assert not request["want_figures"] and not request["want_tables"]


def test_layout_requirements_asks_for_tables_when_document_has_one():
    """087（跨页表格）没有静态来源：只要文档里真有表格就必须量一次分页。"""

    document = Document()
    document.add_paragraph("表 2-1 数据")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "标题行"
    request = layout_requirements(DocxContext(document), [])
    assert request is not None and request["want_tables"]


def test_layout_requirements_asks_for_equations_geometry():
    """098（公式越界）同样没有静态来源：文档里有公式就要量公式段落坐标。"""

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(
        f'<m:oMathPara {nsdecls("m", "w")}>'
        f'<m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath></m:oMathPara>'))
    request = layout_requirements(DocxContext(document), [])
    assert request is not None
    assert 0 in request["indices"]


def test_layout_requirements_none_for_plain_text_document():
    """纯文字文档（无表格/公式/题注，也没有几何类无法判定）完全不读 Word 排版。"""

    document = Document()
    document.add_paragraph("这是一段普通正文。")
    assert layout_requirements(DocxContext(document), []) is None


# ---------------------------------------------------------------------------
# 端到端：命中缓存后不再调用 Word 流程
# ---------------------------------------------------------------------------

def test_repeated_run_reuses_cached_word_flow(runtime_dir, monkeypatch):
    from thesis_format_checker.engine import pipeline

    template, document = runtime_dir / "t.docx", runtime_dir / "p.docx"
    Document().save(template)
    paragraph = Document()
    paragraph.add_paragraph("正文内容。")
    paragraph.save(document)
    calls: list[int] = []

    def legacy(template_path, document_path, output_path, **kwargs):
        calls.append(1)
        return RunSummary(sections_checked=1, body_paragraphs_checked=2,
                          findings=[Finding(rule_id="069", title="正文字号",
                                            status=Status.FAIL, message="字号不符",
                                            paragraph_index=1)])

    outputs = []
    for _ in range(2):
        output = runtime_dir / f"out{len(outputs)}.docx"
        summary = pipeline.run_comparison(legacy, template, document, output,
                                          overwrite=True, use_cache=True,
                                          cache_dir=runtime_dir / "cache")
        outputs.append(summary)
    # Word 流程只跑一次：第二次直接复用缓存里的段落属性与结论。
    assert len(calls) == 1
    assert "Word 段落属性与分页事实" in outputs[1].cache_stats["hits"]
    assert outputs[1].sections_checked == 1


def test_cache_disabled_runs_word_flow_every_time(runtime_dir):
    from thesis_format_checker.engine import pipeline

    template, document = runtime_dir / "t.docx", runtime_dir / "p.docx"
    Document().save(template)
    Document().save(document)
    calls: list[int] = []

    def legacy(template_path, document_path, output_path, **kwargs):
        calls.append(1)
        return RunSummary()

    for index in range(2):
        pipeline.run_comparison(legacy, template, document, runtime_dir / f"o{index}.docx",
                                overwrite=True, use_cache=False)
    assert len(calls) == 2
