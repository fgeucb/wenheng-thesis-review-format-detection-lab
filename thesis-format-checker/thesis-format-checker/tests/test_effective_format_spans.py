"""``read_span_formats`` 的 COM 兼容与性能回归。

背景：旧版 Word“有效格式”流程传入的是 ``win32com`` 段落对象（没有 ``w:p``），
早期实现只按 OOXML run 边界细化，一访问 ``paragraph._p`` 就抛 ``AttributeError``，
整条 Word 检测链路（023/025/035/045–047/067–075/122/123）因此全部退化成“无法判定”。
本文件锁死两点：

1. COM 段落不会再抛异常，且能按字符给出正确字体/字号；
2. 混合格式段落**不是**逐字符读，而是二分收敛（调用次数远小于字符数）。
"""

from __future__ import annotations

from docx import Document

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.effective_format import read_span_formats

WD_UNDEFINED = 9999999


class FakeFont:
    def __init__(self, name: str, size: float) -> None:
        self.NameFarEast = name
        self.NameAscii = name
        self.Size = size
        self.ThemeFont = 0


class FakeRange:
    def __init__(self, text: str, font: FakeFont) -> None:
        self.Text = text
        self.Font = font
        self.Start = 0


class FakeComParagraph:
    """最小可用的 COM 段落替身：只有 Range.Text / Range.Start / Range.Font。"""

    def __init__(self, text: str, fonts: dict[int, str], size: float = 12.0) -> None:
        self._text = text
        self._fonts = fonts
        self._size = size
        self.calls: list[tuple[int, int]] = []
        self.Range = FakeRange(text, FakeFont("宋体", size))

    def font_for(self, start: int, end: int) -> FakeFont:
        names = {self._fonts.get(index, "宋体") for index in range(start, end)}
        if len(names) > 1:
            # Word 对“混合直接格式”的返回值：wdUndefined。
            return FakeFont(str(WD_UNDEFINED), WD_UNDEFINED)
        return FakeFont(names.pop() if names else "宋体", self._size)


class FakeDocument:
    def __init__(self, paragraph: FakeComParagraph) -> None:
        self._paragraph = paragraph

    def Range(self, start: int, end: int) -> FakeRange:
        self._paragraph.calls.append((start, end))
        return FakeRange("", self._paragraph.font_for(start, end))


def test_com_paragraph_does_not_raise_and_splits_mixed_format():
    text = "中文" * 6 + "ABC" * 4
    fonts = {index: "黑体" for index in range(0, 12)}
    paragraph = FakeComParagraph(text, fonts)
    spans = read_span_formats(FakeDocument(paragraph), paragraph)
    chinese = [span for span in spans if span.script == "chinese"]
    latin = [span for span in spans if span.script == "latin"]
    assert chinese and all(span.font_name == "黑体" for span in chinese)
    assert latin and all(span.font_name == "宋体" for span in latin)
    # 覆盖完整、无空洞
    assert spans[0].start == 0 and spans[-1].end == len(text)
    assert all(left.end == right.start for left, right in zip(spans, spans[1:]))


def test_mixed_format_uses_bisect_not_per_character():
    text = "A" * 40 + "B" * 40
    fonts = {index: "Times New Roman" for index in range(40)}
    paragraph = FakeComParagraph(text, fonts)
    read_span_formats(FakeDocument(paragraph), paragraph)
    # 逐字符读会是 80 次以上；二分收敛应显著更少。
    assert len(paragraph.calls) < len(text) // 2, paragraph.calls


def test_docx_paragraph_still_uses_run_boundaries():
    """python-docx 段落继续走 run 边界优化（不需要 COM 也能算出分段）。"""

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("中文标题")
    paragraph.add_run("English 123")
    context = DocxContext(document)
    # 直接调用内部边界函数：run 边界应把 4 个字与英文数字切开。
    from thesis_format_checker.core.effective_format import _run_boundaries

    text = paragraph.text
    boundaries = _run_boundaries(paragraph, text, 0, len(text))
    assert boundaries is not None
    assert (0, 4) in boundaries and (4, len(text)) in boundaries
    assert context.index(paragraph) == 0
