from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from thesis_format_checker.engine.pipeline import extract_baseline, save_comments_only, run_comparison
from thesis_format_checker.core.docx_context import DocxContext, effective_paragraph, effective_runs, make_finding, unknown
from thesis_format_checker.core.comments import add_word_comment
from thesis_format_checker.core.models import Finding, RunSummary, Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def test_native_comments_preserve_all_noncomment_parts_and_existing_comments(runtime_dir):
    source, output = runtime_dir / "input.docx", runtime_dir / "out.docx"
    doc = Document()
    p = doc.add_paragraph("论文题目：保留文字、格式和修订")
    p.runs[0].bold = True
    p.paragraph_format.space_after = Pt(9)
    doc.add_comment(p.runs, "已有用户批注", author="用户")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格信息"
    doc.sections[0].header.paragraphs[0].text = "保留页眉"
    doc.save(source)
    loaded = Document(source)
    add_word_comment(loaded, make_finding("012", "题目", "一致", "不同"), loaded.paragraphs[0])
    save_comments_only(source, loaded, output)
    result = Document(output)
    assert result.paragraphs[0].text == p.text
    assert result.paragraphs[0].runs[0].bold is True
    assert len(result.comments) == 2
    assert list(result.comments)[0].text == "已有用户批注"
    with ZipFile(source) as before, ZipFile(output) as after:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        for name in before.namelist():
            if name not in allowed:
                assert before.read(name) == after.read(name), name


def test_comment_only_write_rejects_any_body_change(runtime_dir):
    source, output = runtime_dir / "source.docx", runtime_dir / "out.docx"
    doc = Document()
    doc.add_paragraph("原始文字")
    doc.save(source)
    doc.paragraphs[0].text = "被改写"
    with pytest.raises(Exception, match="改变了非批注"):
        save_comments_only(source, doc, output)
    assert not output.exists()


def test_empty_paragraph_comment_no_visible_text_added(runtime_dir):
    source, output = runtime_dir / "empty.docx", runtime_dir / "out.docx"
    doc = Document()
    doc.add_paragraph()
    doc.save(source)
    add_word_comment(doc, make_finding("001", "缺少封面", "模板要求封面", "未找到"), doc.paragraphs[0])
    save_comments_only(source, doc, output)
    assert Document(output).paragraphs[0].text == ""
    assert len(Document(output).comments) == 1


def test_common_reads_inherited_format_without_mutation():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(13)
    style.paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph("English 123")
    before = doc.element.xml
    assert effective_runs(p)[0][1]["size_pt"] == 13
    assert effective_paragraph(p)["space_after"] == 7
    assert doc.element.xml == before


def test_pipeline_retains_legacy_findings_and_returns_native_comments(runtime_dir):
    template, paper, output = [runtime_dir / name for name in ("a.docx", "b.docx", "out.docx")]
    doc = Document()
    doc.add_paragraph("论文题目：基准题目")
    doc.add_paragraph("第一章 绪论", "Heading 1")
    doc.add_paragraph("本文是用于流程验证的正文。")
    doc.save(template)
    doc.save(paper)
    calls = []
    def legacy(a, b, out, **kwargs):
        calls.append((a, b, out))
        return RunSummary(findings=[Finding("025", "页边距", Status.FAIL, "原规则证据", expected="动态模板", actual="当前数值", section_index=1)])
    summary = run_comparison(legacy, template, paper, output)
    assert len(calls) == 1 and calls[0][2] != output
    assert "025" in summary.counts_by_rule()
    # 现在已无“整族待实现”的规则；摘要里仍会列出已删除/待确认的族。
    assert any("已删除" in entry or "028" in entry for entry in summary.skipped)
    assert not any(family_of(f.rule_id) == "008" for f in summary.findings)
    comments = list(Document(output).comments)
    assert any("动态模板" in c.text for c in comments)      # 025 的“应该”值
    for comment in comments:
        assert "应该：" in comment.text and "实际：" in comment.text
        # 批注已锚定在出问题的文字上，正文里不再重复写“位置：”一行。
        assert "位置：" not in comment.text


def test_pipeline_legacy_unavailable_does_not_claim_successful_check(runtime_dir):
    paths = [runtime_dir / name for name in ("a.docx", "b.docx", "out.docx")]
    doc = Document()
    doc.add_paragraph("普通正文。")
    for path in paths[:2]:
        doc.save(path)
    def broken(*args, **kwargs):
        raise RuntimeError("COM unavailable")
    summary = run_comparison(broken, *paths)
    assert paths[2].exists()
    assert any(family_of(f.rule_id) == "025" and f.status == Status.UNKNOWN for f in summary.findings)


def test_baseline_extends_existing_without_overwriting():
    baseline = extract_baseline(DocxContext(Document()), {"old_key": "keep"})
    assert baseline["old_key"] == "keep"
    assert {"structure", "front", "pages"} <= set(baseline)
    assert "fields" in baseline["front"]
    with pytest.raises(ValueError, match="未覆盖"):
        extract_baseline(DocxContext(Document()), baseline)


def test_existing_nonstandard_comment_part_is_preserved_and_extended(runtime_dir):
    from lxml import etree
    original, renamed, output = [runtime_dir / name for name in ("original.docx", "renamed.docx", "out.docx")]
    doc = Document()
    p = doc.add_paragraph("已有批注的正文")
    doc.add_comment(p.runs, "原批注", author="作者")
    doc.save(original)
    with ZipFile(original) as source, ZipFile(renamed, "w") as target:
        for info in source.infolist():
            blob, name = source.read(info.filename), info.filename
            if name == "word/comments.xml":
                name = "word/comments2.xml"
            elif name == "word/_rels/document.xml.rels":
                root = etree.fromstring(blob)
                for rel in root:
                    if rel.get("Type", "").endswith("/comments"):
                        rel.set("Target", "comments2.xml")
                blob = etree.tostring(root)
            elif name == "[Content_Types].xml":
                blob = blob.replace(b"/word/comments.xml", b"/word/comments2.xml")
            target.writestr(name, blob)
    loaded = Document(renamed)
    add_word_comment(loaded, make_finding("012", "题目", "一致", "不同"), loaded.paragraphs[0])
    save_comments_only(renamed, loaded, output)
    assert [c.text for c in Document(output).comments][0] == "原批注"
    assert len(Document(output).comments) == 2
    with ZipFile(output) as result:
        assert "word/comments2.xml" in result.namelist()
        assert "word/comments.xml" not in result.namelist()


def test_output_cannot_overwrite_input_or_template(runtime_dir):
    a, b = runtime_dir / "a.docx", runtime_dir / "b.docx"
    doc = Document()
    doc.add_paragraph("待检测正文")
    doc.save(a)
    doc.save(b)
    before = (a.read_bytes(), b.read_bytes())
    for path in (a, b):
        with pytest.raises(Exception, match="不能覆盖"):
            run_comparison(lambda *args, **kwargs: RunSummary(), a, b, path, overwrite=True)
    assert before == (a.read_bytes(), b.read_bytes())


def test_original_unknown_findings_are_report_only(runtime_dir):
    a, b, out = [runtime_dir / name for name in ("a.docx", "b.docx", "out.docx")]
    doc = Document()
    doc.add_paragraph("普通正文")
    doc.save(a)
    doc.save(b)
    def legacy(*args, **kwargs):
        return RunSummary(unknowns=["【025】文档 B 第 1 节无法匹配模板特殊节，未检测边距"])
    summary = run_comparison(legacy, a, b, out)
    assert any(family_of(f.rule_id) == "025" and f.status == Status.UNKNOWN for f in summary.findings)
    assert not any("无法匹配模板特殊节" in c.text for c in Document(out).comments)
    assert any("无法匹配模板特殊节" in text for text in summary.report_only)
