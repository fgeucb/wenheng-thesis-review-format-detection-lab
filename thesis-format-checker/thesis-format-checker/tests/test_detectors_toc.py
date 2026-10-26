"""TOC fixtures retain their cached field results and all source XML."""
from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.toc import extract_toc_baseline, check_toc_errors
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def marker(paragraph, kind):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    paragraph.add_run()._r.append(node)


def instruction(paragraph, text):
    node = OxmlElement("w:instrText")
    node.text = text
    paragraph.add_run()._r.append(node)


def bookmark(paragraph, name, number):
    node = OxmlElement("w:bookmarkStart")
    node.set(qn("w:name"), name)
    node.set(qn("w:id"), str(number))
    paragraph._p.insert(0, node)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.append(end)


def fixture(labels=("绪论", "方法"), levels=(1, 1), field='TOC \\o "1-3" \\h', *, sdt=False, nested=True):
    doc = Document()
    for level in (1, 2, 3):
        style = doc.styles.add_style("TOC " + str(level), WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(12 + level)
        style.font.bold = False
        style.paragraph_format.left_indent = Inches((level - 1) * .2)
        style.paragraph_format.tab_stops.add_tab_stop(Inches(5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    title = doc.add_paragraph("目录")
    title.runs[0].font.size = Pt(18)
    start = doc.add_paragraph()
    marker(start, "begin")
    instruction(start, field)
    marker(start, "separate")
    entries = []
    for index, (label, level) in enumerate(zip(labels, levels)):
        paragraph = doc.add_paragraph(style="TOC " + str(level))
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "_Toc" + str(index))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        paragraph.add_run("\t")
        if nested:
            marker(paragraph, "begin")
            instruction(paragraph, "PAGEREF _Toc" + str(index) + " \\h")
            marker(paragraph, "separate")
        paragraph.add_run(str(index + 1))
        if nested:
            marker(paragraph, "end")
        entries.append(paragraph)
    finish = doc.add_paragraph()
    marker(finish, "end")
    headings = []
    for index, (label, level) in enumerate(zip(labels, levels)):
        paragraph = doc.add_heading(label, level=level)
        bookmark(paragraph, "_Toc" + str(index), index)
        headings.append(paragraph)
        doc.add_paragraph("正文内容。")
    if sdt:
        content = OxmlElement("w:sdtContent")
        container = OxmlElement("w:sdt")
        container.append(content)
        start._p.addprevious(container)
        for paragraph in (start, *entries, finish):
            content.append(paragraph._p)
    return doc, title, entries, headings


def detect(template, document):
    before_template, before_document = template.element.xml, document.element.xml
    result = check_toc_errors(extract_toc_baseline(DocxContext(template)), DocxContext(document))
    assert template.element.xml == before_template
    assert document.element.xml == before_document
    return result


def failed(result, rule):
    return [f for f in result if family_of(f.rule_id) == rule and f.status == Status.FAIL]


def uncertain(result, rule):
    # 模板未涉及的结论现在是“不适用”（不写批注），测试中与“无法判定”一并接受。
    return [f for f in result if family_of(f.rule_id) == rule
            and f.status in (Status.UNKNOWN, Status.NOT_APPLICABLE)]


def test_identical_nested_hyperlinked_sdt_directory_has_no_definite_errors():
    template = fixture(sdt=True)[0]
    document = fixture(sdt=True)[0]
    result = detect(template, document)
    assert not [f for f in result if f.status == Status.FAIL]
    assert uncertain(result, "059") and uncertain(result, "066")
    assert UNREADABLE in uncertain(result, "059")[0].message


def test_056_plain_directory_lacks_template_automatic_structure():
    template = fixture()[0]
    document = Document()
    document.add_paragraph("目录")
    document.add_paragraph("绪论\t1")
    assert failed(detect(template, document), "056")


def test_manual_toc_still_has_independent_064_title_format_check():
    template = fixture()[0]
    document = Document()
    title = document.add_paragraph("目录")
    title.runs[0].font.size = Pt(8)
    document.add_paragraph("绪论\t1")
    result = detect(template, document)
    assert failed(result, "056") and failed(result, "064")


def test_multiple_tabs_are_not_guessed_as_page_number_tab():
    template, _, template_entries, _ = fixture()
    document = fixture()[0]
    for entry in template_entries:
        entry.paragraph_format.tab_stops.add_tab_stop(Inches(.2))
    result = detect(template, document)
    assert uncertain(result, "062") and uncertain(result, "063")
    assert not failed(result, "062") and not failed(result, "063")


def test_057_missing_heading_uses_source_range_and_heading_anchor():
    template = fixture()[0]
    document, _, entries, headings = fixture()
    entries[1]._p.getparent().remove(entries[1]._p)
    result = detect(template, document)
    missing = failed(result, "057")
    assert len(missing) == 1
    assert missing[0].anchor._p is headings[1]._p
    assert not failed(result, "058")


def test_058_066_renamed_same_target_is_stale_not_missing_plus_extra():
    template = fixture()[0]
    document, _, entries, headings = fixture()
    headings[1].runs[0].text = "新方法"
    result = detect(template, document)
    assert failed(result, "058") and failed(result, "066")
    assert not failed(result, "057")
    assert failed(result, "058")[0].anchor._p is entries[1]._p


def test_059_never_treats_saved_page_cache_as_actual_page():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    entries[0].runs[-2].text = "999"
    result = detect(template, document)
    assert uncertain(result, "059")
    assert not failed(result, "059")


def test_060_compares_explicit_dynamic_range():
    template = fixture(field='TOC \\o "1-2" \\h')[0]
    document = fixture(field='TOC \\o "1-3" \\h')[0]
    assert failed(detect(template, document), "060")


def test_061_level_format_and_hyperlink_run_override():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    run = entries[0]._p.find(qn("w:hyperlink")).find(qn("w:r"))
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "44")
    props.append(size)
    run.insert(0, props)
    assert failed(detect(template, document), "061")


def test_toc_entry_paragraph_alignment_does_not_override_page_tab_evidence():
    template, _, template_entries, _ = fixture()
    document, _, entries, _ = fixture()
    for entry in template_entries:
        entry.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for entry in entries:
        entry.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    assert not failed(detect(template, document), "061")


def test_empty_toc_field_owner_with_toc_style_is_not_an_entry():
    template = fixture()[0]
    document = fixture()[0]
    template.paragraphs[1].style = "TOC 1"
    document.paragraphs[1].style = "TOC 1"

    result = detect(template, document)

    assert not failed(result, "065")
    assert not failed(result, "061")


def test_repeated_shared_toc_style_issue_is_reported_once():
    template = fixture(labels=("甲", "乙", "丙"), levels=(1, 1, 1))[0]
    document = fixture(labels=("甲", "乙", "丙"), levels=(1, 1, 1))[0]
    stops = document.styles["TOC 1"].paragraph_format.tab_stops
    stops.clear_all()
    stops.add_tab_stop(Inches(5.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    faults = failed(detect(template, document), "063")

    assert len(faults) == 1


def test_062_tab_leader_follows_template_not_fixed_dots():
    template, _, template_entries, _ = fixture()
    document, _, entries, _ = fixture()
    for entry in template_entries:
        entry.paragraph_format.tab_stops.add_tab_stop(Inches(5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES)
    assert failed(detect(template, document), "062")
    for entry in entries:
        entry.paragraph_format.tab_stops.add_tab_stop(Inches(5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES)
    assert not failed(detect(template, document), "062")


def test_063_alignment_and_tab_position():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    entries[0].paragraph_format.tab_stops.add_tab_stop(Inches(5), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
    assert failed(detect(template, document), "063")


def test_064_title_format_not_title_physical_page():
    template = fixture()[0]
    document, title, _, _ = fixture()
    title.runs[0].font.size = Pt(9)
    title.paragraph_format.space_after = Pt(19)
    assert failed(detect(template, document), "064")


def test_065_duplicate_target_not_duplicate_label():
    template = fixture(labels=("相同标题", "相同标题"))[0]
    document, _, entries, _ = fixture(labels=("相同标题", "相同标题"))
    assert not failed(detect(template, document), "065")
    entries[1]._p.addnext(deepcopy(entries[1]._p))
    assert len(failed(detect(template, document), "065")) == 1


def test_065_empty_title_with_target_does_not_flag_marker_paragraph():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    entries[0]._p.find(qn("w:hyperlink")).find(qn("w:r")).find(qn("w:t")).text = ""
    result = detect(template, document)
    assert len(failed(result, "065")) == 1


def test_unlinked_toc_does_not_guess_same_named_heading_membership():
    template = fixture()[0]
    document, _, entries, _ = fixture(nested=False)
    for entry in entries:
        entry._p.find(qn("w:hyperlink")).attrib.clear()
    result = detect(template, document)
    assert uncertain(result, "057")
    assert not failed(result, "057")


def test_broken_bookmark_does_not_prove_heading_nonexistence():
    template = fixture()[0]
    document, _, _, headings = fixture()
    bookmark_node = headings[0]._p.find(qn("w:bookmarkStart"))
    headings[0]._p.remove(bookmark_node)
    result = detect(template, document)
    assert uncertain(result, "057")
    assert not failed(result, "057")


def test_custom_scope_and_tc_sources_are_unknown_not_guessed():
    for instruction_text in ('TOC \\o "1-3" \\b ChapterScope', 'TOC \\f C'):
        template = fixture(field=instruction_text)[0]
        document = fixture(field=instruction_text)[0]
        result = detect(template, document)
        assert uncertain(result, "057") and uncertain(result, "060")
        assert not failed(result, "057")


def test_suppressed_page_numbers_do_not_require_leaders_or_pagination():
    template, _, template_entries, _ = fixture(field='TOC \\o "1-3" \\h \\n', nested=False)
    document, _, entries, _ = fixture(field='TOC \\o "1-3" \\h \\n', nested=False)
    for paragraph in (*template_entries, *entries):
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
    result = detect(template, document)
    assert not any(family_of(f.rule_id) in {"059", "062", "063"} for f in result)


def test_malformed_toc_is_unknown_for_all_affected_rules():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    finish = entries[-1]._p.getnext()
    finish.getparent().remove(finish)
    result = detect(template, document)
    assert all(f.status == Status.UNKNOWN for f in result)
    assert {f.rule_id for f in result} >= {str(n).zfill(3) for n in range(56, 67)} - {"064"}


def test_custom_toc_style_mapping_can_collect_non_outline_source():
    template, _, _, _ = fixture(field='TOC \\t "Research topic,1" \\h')
    document, _, _, headings = fixture(field='TOC \\t "Research topic,1" \\h')
    for doc in (template, document):
        doc.styles.add_style("Research topic", WD_STYLE_TYPE.PARAGRAPH)
    for heading in headings:
        heading.style = "Research topic"
    result = detect(template, document)
    assert not failed(result, "057") and not failed(result, "058")


def test_heading_beyond_requested_toc_level_is_not_missing():
    template = fixture(field='TOC \\o "1-1" \\h')[0]
    document = fixture(field='TOC \\o "1-1" \\h')[0]
    document.add_heading("不收录的小节", level=2)
    assert not failed(detect(template, document), "057")


def test_template_mixed_level_format_does_not_pick_majority():
    template, _, entries, _ = fixture()
    entries[1].paragraph_format.left_indent = Inches(.7)
    document = fixture()[0]
    result = detect(template, document)
    assert uncertain(result, "061")
    assert not failed(result, "061")


def test_revisions_prevent_title_membership_and_cached_text_guessing():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    entry = entries[0]
    revision = OxmlElement("w:ins")
    revision.append(deepcopy(entry.runs[-2]._r))
    entry._p.append(revision)
    result = detect(template, document)
    assert uncertain(result, "057") and uncertain(result, "058")
    assert not failed(result, "057") and not failed(result, "058")


def test_empty_unupdated_toc_cache_is_unknown_not_all_titles_missing():
    template = fixture()[0]
    document, _, entries, _ = fixture()
    for entry in entries:
        entry._p.getparent().remove(entry._p)
    result = detect(template, document)
    assert uncertain(result, "057")
    assert not failed(result, "057")


def test_toc_o_without_u_does_not_assume_custom_based_on_heading_is_included():
    template = fixture()[0]
    document = fixture()[0]
    style = document.styles.add_style("Special local heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Heading 1"]
    document.add_paragraph("不属于显式来源的样式标题", style=style)
    assert not failed(detect(template, document), "057")


def test_heading_dynamic_field_cache_is_not_compared_as_plain_title():
    template = fixture()[0]
    document, _, _, headings = fixture()
    heading = headings[0]
    marker(heading, "begin")
    instruction(heading, 'DOCPROPERTY "ChapterName"')
    marker(heading, "separate")
    heading.add_run("动态标题缓存")
    marker(heading, "end")
    result = detect(template, document)
    assert uncertain(result, "058")
    assert not failed(result, "058")


def test_template_empty_cache_marks_unavailable_level_formats_and_tabs():
    template, _, entries, _ = fixture()
    for entry in entries:
        entry._p.getparent().remove(entry._p)
    document = fixture()[0]
    result = detect(template, document)
    assert all(uncertain(result, rule) for rule in ("061", "062", "063"))
    assert not any(failed(result, rule) for rule in ("061", "062", "063"))
