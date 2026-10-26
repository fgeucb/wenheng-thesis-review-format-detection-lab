from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from thesis_format_checker.detectors.text_runs import effective_text_runs, resolve_color


def fonts(props, **attributes):
    node = props.get_or_add_rFonts()
    for key, value in attributes.items():
        node.set(qn("w:" + key), value)


def test_effective_fonts_include_inheritance_character_styles_and_local_overrides():
    document = Document()
    base = document.styles.add_style("模板正文", WD_STYLE_TYPE.PARAGRAPH)
    base.font.name = "Template Latin"
    base.font.size = Pt(13)
    fonts(base.element.get_or_add_rPr(), eastAsia="模板中文")
    derived = document.styles.add_style("继承正文", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    character = document.styles.add_style("局部用途", WD_STYLE_TYPE.CHARACTER)
    character.font.size = Pt(15)
    character.font.bold = True
    p = document.add_paragraph(style=derived)
    a = p.add_run("中文 Latin123")
    b = p.add_run("局部456")
    b.style = character
    b.font.bold = False
    b.font.color.rgb = RGBColor(18, 52, 86)
    fonts(b._r.get_or_add_rPr(), ascii="Local Latin")
    before = etree.tostring(document.element)
    runs = effective_text_runs(p)
    assert runs[0][1]["east_asia_font"] == "模板中文"
    assert runs[0][1]["ascii_font"] == "Template Latin"
    assert runs[0][1]["size_pt"] == 13
    assert runs[1][1]["ascii_font"] == "Local Latin"
    assert runs[1][1]["hansi_font"] == "Template Latin"
    assert runs[1][1]["east_asia_font"] == "模板中文"
    assert runs[1][1]["size_pt"] == 15
    assert runs[1][1]["bold"] is False
    assert runs[1][1]["color"] == "123456"
    assert etree.tostring(document.element) == before


def test_theme_font_uses_language_and_does_not_assume_simplified_chinese():
    document = Document()
    p = document.add_paragraph("模板主题字体")
    prop = p.runs[0]._r.get_or_add_rPr()
    fonts(prop, eastAsiaTheme="minorEastAsia", asciiTheme="minorAscii")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:eastAsia"), "zh-TW")
    prop.append(lang)
    relation = next(r for r in document.part.rels.values() if r.reltype.endswith("/theme"))
    theme = etree.fromstring(relation.target_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    expected = theme.find(".//a:minorFont/a:font[@script='Hant']", ns).get("typeface")
    assert effective_text_runs(p)[0][1]["east_asia_font"] == expected
    lang.set(qn("w:eastAsia"), "und")
    assert effective_text_runs(p)[0][1]["east_asia_font"] is None


def test_theme_color_compares_resolved_color_and_tints_are_unknown():
    document = Document()
    assert resolve_color(document, {"themeColor": "text1", "val": "FF0000"}) == "000000"
    assert resolve_color(document, {"val": "123abc"}) == "123ABC"
    assert resolve_color(document, {"themeColor": "accent1", "themeTint": "80"}) is None
    assert resolve_color(document, {"val": "auto"}) == "auto"


def test_hyperlink_simple_field_revision_and_deleted_text_are_read_without_mutation():
    document = Document()
    p = document.add_paragraph("正文")
    original = deepcopy(p.runs[0]._r)
    for tag, text in (("w:hyperlink", "链接"), ("w:fldSimple", "域缓存"),
                      ("w:ins", "修订插入"), ("w:del", "已删除")):
        container = OxmlElement(tag)
        node = deepcopy(original)
        node.find(qn("w:t")).text = text
        container.append(node)
        p._p.append(container)
    before = etree.tostring(p._p)
    values = effective_text_runs(p)
    assert [text for text, props in values] == ["正文", "链接", "域缓存", "修订插入"]
    assert [props["field"] for _, props in values] == [False, False, True, False]
    assert all(props["revision"] for _, props in values)
    assert before == etree.tostring(p._p)


def test_table_style_does_not_lose_unresolved_east_asia_when_only_ascii_overridden():
    document = Document()
    style = document.styles.add_style("表格中文规则", WD_STYLE_TYPE.TABLE)
    fonts(style.element.get_or_add_rPr(), eastAsia="表格字体")
    table = document.add_table(rows=1, cols=1)
    table.style = style
    p = table.cell(0, 0).paragraphs[0]
    run = p.add_run("中文 abc")
    fonts(run._r.get_or_add_rPr(), ascii="Explicit Latin")
    props = effective_text_runs(p)[0][1]
    assert props["east_asia_font"] is None
    assert props["ascii_font"] == "Explicit Latin"
    assert props["limitations"]
    fonts(run._r.rPr, eastAsia="直接中文字体")
    assert effective_text_runs(p)[0][1]["east_asia_font"] == "直接中文字体"
