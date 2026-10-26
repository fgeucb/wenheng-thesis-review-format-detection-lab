"""标题层级树：编号 / 目录缓存 / 样式三层证据，并把误用样式的正文剔除。"""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.detectors.headings import _heading_level


def paragraph(document, text, style=None, outline=None):
    item = document.add_paragraph(text)
    if style is not None:
        item.style = document.styles[style]
    if outline is not None:
        node = OxmlElement("w:outlineLvl")
        node.set(qn("w:val"), str(outline))
        item._p.get_or_add_pPr().append(node)
    return item


def toc_paragraph(document, text, level):
    name = f"toc {level}"
    if name not in [style.name for style in document.styles]:
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    return paragraph(document, text, name)


def level_of(document, text):
    for item in DocxContext(document).paragraphs:
        if item.text.strip() == text:
            return _heading_level(item)
    raise AssertionError(f"未找到段落：{text}")


def test_long_body_sentence_in_heading_style_is_not_a_heading():
    document = Document()
    paragraph(document, "自1990年以来，中国资本市场规模稳步扩张，沪深两市上市公司数量已逾5500家，"
                       "总市值超过80万亿元，市场结构日趋成熟。", "Heading 2")
    assert level_of(document, document.paragraphs[0].text) is None


def test_numbered_ladder_under_chapter_convention():
    document = Document()
    for text, style in (("第一章 绪论", "Heading 1"), ("1 研究背景与意义", "Heading 2"),
                        ("1.1 研究背景", "Heading 3"), ("1.1.1 数据来源", None)):
        paragraph(document, text, style)
    assert level_of(document, "第一章 绪论") == 1
    assert level_of(document, "1 研究背景与意义") == 2
    assert level_of(document, "1.1 研究背景") == 3
    assert level_of(document, "1.1.1 数据来源") == 4


def test_single_number_needs_style_evidence():
    document = Document()
    paragraph(document, "1 绪论", "Heading 1")
    paragraph(document, "2 附录")
    assert level_of(document, "1 绪论") == 1
    assert level_of(document, "2 附录") is None


def test_numbers_inside_body_text_are_not_headings():
    document = Document()
    paragraph(document, "1990年上海证券交易所成立以来，市场快速发展。", "Heading 2")
    paragraph(document, "0.5320")
    paragraph(document, "31.5% 的样本来自制造业。", "Heading 2")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "1.1 表格里的编号"
    assert level_of(document, "1990年上海证券交易所成立以来，市场快速发展。") is None
    assert level_of(document, "0.5320") is None
    assert level_of(document, "31.5% 的样本来自制造业。") is None
    assert level_of(document, "1.1 表格里的编号") is None


def test_toc_cache_decides_the_level_without_any_style():
    document = Document()
    toc_paragraph(document, "1 研究背景\t3", 2)
    paragraph(document, "1 研究背景")
    assert level_of(document, "1 研究背景") == 2


def test_styled_heading_without_numbering_is_kept():
    document = Document()
    paragraph(document, "参考文献", "Heading 1")
    assert level_of(document, "参考文献") == 1


def test_parent_chain_guard_rejects_orphan_numbering():
    document = Document()
    paragraph(document, "1.1 没有父级的小节")
    assert level_of(document, "1.1 没有父级的小节") is None


def test_broken_outline_level_still_raises_for_caller():
    document = Document()
    item = paragraph(document, "1 绪论", "Heading 1")
    broken = OxmlElement("w:outlineLvl")
    broken.set(qn("w:val"), "broken")
    item._p.get_or_add_pPr().append(broken)
    with pytest.raises(ValueError):
        level_of(document, "1 绪论")
