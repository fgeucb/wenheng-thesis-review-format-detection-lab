"""011（空/重复章节）、014（封面基本信息）、020（声明签名/日期）、022.3（前置页码）测试。"""

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.detectors.front import check_front_errors, extract_front_baseline
from thesis_format_checker.detectors.pages import check_page_errors, extract_page_baseline
from thesis_format_checker.detectors.structure import _check_chapters


def codes(findings, family):
    return [finding.rule_id for finding in findings if finding.rule_id.split(".")[0] == family]


def test_empty_chapter_is_reported_but_non_empty_is_not():
    document = Document()
    document.add_heading("第一章 绪论", level=1)          # 没有任何正文
    document.add_heading("第二章 方法", level=1)
    document.add_paragraph("本章给出具体方法。")

    findings = [finding for finding in _check_chapters(DocxContext(document))
                if finding.rule_id == "011.1"]

    assert [finding.anchor.text for finding in findings] == ["第一章 绪论"]


def test_chapter_with_only_subsection_content_is_not_empty():
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_heading("1.1 研究背景", level=2)
    document.add_paragraph("背景内容。")

    assert not [finding for finding in _check_chapters(DocxContext(document))
                if finding.rule_id == "011.1"]


def test_figure_caption_in_heading_style_is_not_a_chapter():
    """论文常把题注误设为“标题 2”样式：题注不是章节，也不该被判成空章节。"""

    document = Document()
    document.add_heading("第二章 方法", level=1)
    document.add_paragraph("本章内容。")
    document.add_paragraph("图2-1  系统总体架构图", style="Heading 2")

    assert not [finding for finding in _check_chapters(DocxContext(document))
                if finding.rule_id in {"011.1", "011.2"}]


def test_duplicate_chapter_title_is_reported():
    document = Document()
    document.add_heading("第二章 方法", level=1)
    document.add_paragraph("内容一。")
    document.add_heading("第二章 方法", level=1)
    document.add_paragraph("内容二。")

    findings = [finding for finding in _check_chapters(DocxContext(document))
                if finding.rule_id == "011.2"]

    assert len(findings) == 1
    assert findings[0].anchor.text == "第二章 方法"


def test_missing_cover_fields_follow_the_template():
    template, document = Document(), Document()
    for doc in (template, document):
        doc.add_paragraph("姓    名：张三")
        doc.add_paragraph("学    号：2023001")
        doc.add_paragraph("学    院：农学院")
        doc.add_paragraph("专    业：农学")
    document.paragraphs[1].text = "学    号："       # 论文把学号清空

    findings = check_front_errors(extract_front_baseline(DocxContext(template)), DocxContext(document))

    assert "014.2" in codes(findings, "014")          # 学号缺失
    assert "014.5" not in codes(findings, "014")      # 模板封面没有班级字段 → 不检查
    assert not [finding for finding in findings
                if finding.rule_id in {"014.1", "014.3", "014.4"}]


def test_declaration_signature_and_date_are_checked():
    template, document = Document(), Document()
    for doc in (template, document):
        doc.add_heading("原创性声明", level=1)
        doc.add_paragraph("本人郑重声明：论文系本人独立完成。")
        doc.add_paragraph("论文作者签名：________    日期：____年__月__日")
    document.paragraphs[2].text = "论文作者签名：            日期：2026年5月10日"

    findings = check_front_errors(extract_front_baseline(DocxContext(template)), DocxContext(document))

    assert "020.1" in codes(findings, "020")          # 签名没填
    assert "020.2" not in codes(findings, "020")      # 日期已填


def _set_page_number_format(section, fmt):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    node.set(qn("w:start"), "1")
    if fmt:
        node.set(qn("w:fmt"), fmt)
    elif qn("w:fmt") in node.attrib:
        del node.attrib[qn("w:fmt")]


def test_front_matter_arabic_page_numbers_are_reported():
    template, document = Document(), Document()
    for doc in (template, document):
        doc.add_paragraph("摘要")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading("第一章 绪论", level=1)
    _set_page_number_format(template.sections[0], "upperRoman")
    _set_page_number_format(document.sections[0], None)      # 默认阿拉伯数字
    for doc in (template, document):
        _set_page_number_format(doc.sections[1], None)

    result = check_page_errors(extract_page_baseline(DocxContext(template)), DocxContext(document))
    assert "022.3" in codes(result, "022")


def test_roman_front_matter_page_numbers_pass():
    template, document = Document(), Document()
    for doc in (template, document):
        doc.add_paragraph("摘要")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading("第一章 绪论", level=1)
    for doc in (template, document):
        _set_page_number_format(doc.sections[0], "upperRoman")
        _set_page_number_format(doc.sections[1], None)

    result = check_page_errors(extract_page_baseline(DocxContext(template)), DocxContext(document))
    assert "022.3" not in codes(result, "022")
