from copy import deepcopy

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt
from lxml import etree

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.formats import extract_format_consistency_baseline, check_format_consistency_errors
from thesis_format_checker.core.models import Status


def node(tag, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    return element


def style(doc, name):
    if name not in doc.styles:
        out = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        out.base_style = doc.styles["Heading 1"]
    else:
        out = doc.styles[name]
    out.font.name = "Arial"
    out.font.size = Pt(13)
    out.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    out.paragraph_format.line_spacing = 1.25
    return out


def heading(doc, text, bold=True, style_name="Heading 1"):
    p = doc.add_paragraph(text, style(doc, style_name))
    p.runs[0].font.bold = bold
    return p


def compare(template, document, mapping=None):
    base = extract_format_consistency_baseline(DocxContext(template), mapping)
    return check_format_consistency_errors(base, DocxContext(document), mapping)


def fail_ids(findings):
    """按规则族返回失败项编号（139.2 → 139）。"""

    return {f.rule_id.split(".")[0] for f in findings if f.status == Status.FAIL}


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（139.2 → 139）。"""

    return rule_id.split(".")[0]


def special(doc, label="特殊版式：横向页面", landscape=True, new=False):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE) if new else doc.sections[0]
    p = doc.add_paragraph(label)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    return section, p


def math(doc, size=13, subscript_size=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    maths = OxmlElement("m:oMathPara")
    props = OxmlElement("m:oMathParaPr")
    align = OxmlElement("m:jc")
    align.set(qn("m:val"), "center")
    props.append(align)
    maths.append(props)
    formula = OxmlElement("m:oMath")
    def run(text, points):
        r = OxmlElement("m:r")
        pr = node("w:rPr")
        pr.append(node("w:sz", val=points * 2))
        r.append(pr)
        t = OxmlElement("m:t")
        t.text = text
        r.append(t)
        return r
    formula.append(run("x=1", size))
    if subscript_size is not None:
        sub = OxmlElement("m:sub")
        sub.append(run("2", subscript_size))
        formula.append(sub)
    maths.append(formula)
    p._p.append(maths)
    settings = doc.settings.element.find(qn("m:mathPr"))
    if settings is None:
        settings = OxmlElement("m:mathPr")
        doc.settings.element.append(settings)
    font = settings.find(qn("m:mathFont"))
    if font is None:
        font = OxmlElement("m:mathFont")
        settings.append(font)
    font.set(qn("m:val"), "Cambria Math")
    return p


def inline_math(doc, alignment, indent):
    p = doc.add_paragraph("公式前说明文字 ")
    p.alignment = alignment
    p.paragraph_format.first_line_indent = Pt(indent)
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    props = node("w:rPr")
    props.append(node("w:sz", val=24))
    run.append(props)
    text = OxmlElement("m:t")
    text.text = "x=1"
    run.append(text)
    formula.append(run)
    p._p.append(formula)
    return p


def notes(doc, kind, sizes):
    name = "Footnote Text" if kind == "footnote" else "Endnote Text"
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style(doc, name)
    root = node("w:" + kind + "s")
    anchors = []
    for identity, points in enumerate(sizes, 1):
        entry = node("w:" + kind, id=identity)
        p = doc.add_paragraph("解释内容 note", name)
        p.runs[0].font.size = Pt(points)
        entry.append(deepcopy(p._p))
        p._p.getparent().remove(p._p)
        root.append(entry)
        anchor = doc.add_paragraph("正文注释位置")
        anchor.add_run()._r.append(node("w:" + kind + "Reference", id=identity))
        anchors.append(anchor)
    part = Part(PackURI("/word/" + kind + "s.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml." + kind + "s+xml", etree.tostring(root), doc.part.package)
    doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + kind + "s")
    return anchors, part


def test_heading_inconsistency_uses_pair_evidence_not_majority_or_template_rewrite():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "第一节")
        heading(document, "第二节")
    actual.paragraphs[1].runs[0].font.bold = False
    findings = compare(template, actual)
    faults = [f for f in findings if f.rule_id == "139" and f.status == Status.FAIL]
    assert faults and faults[0].anchor._p is actual.paragraphs[1]._p
    assert "哪处正确" in faults[0].message


def test_internal_consistency_does_not_require_template_majority():
    template, actual = Document(), Document()
    heading(actual, "绪论")
    heading(actual, "方法", False)
    assert "139" in fail_ids(compare(template, actual))


def test_same_class_same_style_template_conflicts_are_unknown():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "第一节", True)
        heading(document, "第二节", False)
    findings = compare(template, actual)
    assert "139" not in fail_ids(findings)
    assert any(f.status == Status.UNKNOWN and "模板同用途样例冲突" in f.message for f in findings)


def test_explicit_template_style_variants_are_legal_separate_groups():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "主组甲", True, "主组标题")
        heading(document, "主组乙", True, "主组标题")
        heading(document, "副组甲", False, "副组标题")
        heading(document, "副组乙", False, "副组标题")
    assert "139" not in fail_ids(compare(template, actual))


def test_appendix_body_and_english_chinese_headings_not_forced_into_one_role():
    template, actual = Document(), Document()
    heading(actual, "正文甲", True)
    heading(actual, "正文乙", True)
    heading(actual, "English heading", False)
    heading(actual, "附录 A", False)
    heading(actual, "附录 B", False)
    assert "139" not in fail_ids(compare(template, actual))


def test_caption_actual_direct_run_format_difference_detected():
    template, actual = Document(), Document()
    for document in (template, actual):
        style(document, "Caption")
        document.add_paragraph("图1 示例甲", "Caption")
        document.add_paragraph("图2 示例乙", "Caption")
    actual.paragraphs[1].runs[0].font.size = Pt(19)
    assert "139" in fail_ids(compare(template, actual))


def test_table_same_role_font_inconsistency_detected():
    template, actual = Document(), Document()
    for document in (template, actual):
        for number in (1, 2):
            document.add_paragraph("相关正文说明")
            document.add_paragraph(f"表{number} 表格示例", "Caption")
            table = document.add_table(rows=1, cols=1)
            p = table.cell(0, 0).paragraphs[0]
            p.add_run("单元格文字").font.size = Pt(12 if document is template or number == 1 else 18)
    findings = compare(template, actual)
    # 同类图表格式不一致 → 139.2
    assert any(family_of(f.rule_id) == "139" and f.status == Status.FAIL and "单元格" in f.message
               for f in findings)


def test_omml_structure_roles_separate_subscripts_and_compare_base_size():
    template, actual = Document(), Document()
    math(template, 13, 9)
    math(template, 13, 9)
    math(actual, 13, 9)
    math(actual, 17, 9)
    assert "139" in fail_ids(compare(template, actual))
    assert "139" not in fail_ids(compare(template, template))


def test_inline_formula_does_not_compare_surrounding_prose_layout():
    template, actual = Document(), Document()
    inline_math(actual, WD_ALIGN_PARAGRAPH.JUSTIFY, 0)
    inline_math(actual, WD_ALIGN_PARAGRAPH.LEFT, 24)

    assert not [f for f in compare(template, actual)
                if f.rule_id == "139.3" and f.status == Status.FAIL]


def test_formula_consistency_uses_compact_object_locators():
    template, actual = Document(), Document()
    first = math(actual)
    second = math(actual)
    first.paragraph_format.first_line_indent = Pt(0)
    second.paragraph_format.first_line_indent = Pt(24)

    faults = [f for f in compare(template, actual)
              if f.rule_id == "139.3" and f.status == Status.FAIL]

    assert faults
    assert "独立公式（第" in faults[0].expected
    assert "独立公式（第" in faults[0].actual
    assert "公式：" not in faults[0].message


def test_native_notes_compare_inside_each_kind_and_anchor_body_references():
    template, actual = Document(), Document()
    notes(template, "footnote", [11, 11])
    notes(template, "endnote", [15, 15])
    anchors, _ = notes(actual, "footnote", [11, 17])
    notes(actual, "endnote", [15, 15])
    findings = compare(template, actual)
    faults = [f for f in findings if family_of(f.rule_id) == "139" and f.status == Status.FAIL]
    assert faults and faults[0].anchor._p is anchors[1]._p


def test_labelled_special_page_changed_to_portrait_still_detected():
    template, actual = Document(), Document()
    special(template, landscape=True)
    _, anchor = special(actual, landscape=False)
    findings = compare(template, actual)
    faults = [f for f in findings if f.rule_id == "164" and f.status == Status.FAIL]
    assert faults and faults[0].anchor._p is anchor._p
    assert "orientation" in faults[0].message


def test_unlabelled_landscape_is_not_arbitrarily_invalid():
    template, actual = Document(), Document()
    special(template)
    special(actual, "普通内容")
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN for f in findings)


def test_conflicting_special_page_templates_do_not_select_first():
    template, actual = Document(), Document()
    special(template, landscape=True)
    special(template, landscape=False, new=True)
    special(actual, landscape=True)
    actual.sections[0].page_width = Cm(31)
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and "冲突" in f.message for f in findings)


def test_explicit_section_mapping_narrows_same_purpose_template_variants():
    template, actual = Document(), Document()
    special(template, landscape=True)
    special(template, landscape=False, new=True)
    special(actual, landscape=False)
    actual.sections[0].page_width = Cm(31)
    assert "164" in fail_ids(compare(template, actual, {"section_map": {"1": 2}}))


def test_fold_and_print_color_are_unknown_even_if_page_settings_match():
    for label in ("特殊版式：折页", "特殊版式：彩色页面"):
        template, actual = Document(), Document()
        special(template, label)
        special(actual, label)
        findings = compare(template, actual)
        assert "164" not in fail_ids(findings)
        assert any(f.rule_id == "164" and UNREADABLE in f.message for f in findings)


def test_oversized_table_fixed_grid_difference_and_render_limit_are_both_reported():
    template, actual = Document(), Document()
    for document in (template, actual):
        special(document, "特殊版式：超大表格：数据矩阵")
        table = document.add_table(rows=2, cols=2)
        table.autofit = False
        table.columns[0].width = Cm(4 if document is template else 6)
        for row in table.rows:
            row.cells[0].width = table.columns[0].width
    findings = compare(template, actual)
    assert "164" in fail_ids(findings)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN and "真实跨页" in f.message for f in findings)


def test_template_and_original_package_parts_stay_unmodified():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "标题甲")
        heading(document, "标题乙", False)
        notes(document, "footnote", [11, 15])
        special(document, new=True)
        math(document)
    before = [{str(p.partname): p.blob for p in doc.part.package.parts} for doc in (template, actual)]
    compare(template, actual)
    after = [{str(p.partname): p.blob for p in doc.part.package.parts} for doc in (template, actual)]
    assert before == after


def test_local_corrupt_math_does_not_block_special_page_checks(monkeypatch):
    template, actual = Document(), Document()
    special(template)
    special(actual, landscape=False)
    import thesis_format_checker.detectors.formats as module
    monkeypatch.setattr(module, "_font_samples", lambda *_: (_ for _ in ()).throw(ValueError("bad math")))
    math(actual)
    findings = compare(template, actual)
    assert "164" in fail_ids(findings)
    assert any(f.rule_id == "139" and UNREADABLE in f.message for f in findings)


def test_oversized_table_cell_grid_conflict_needs_renderer():
    template, actual = Document(), Document()
    for document in (template, actual):
        special(document, "超大表格")
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
    actual.tables[0].columns[0].width = Cm(9)
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and "首选宽度与网格不一致" in f.message for f in findings)


def test_style_revision_is_unknown_not_effective_format_failure():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "标题甲")
        heading(document, "标题乙", False)
    revised = node("w:rPrChange", id=1)
    revised.append(node("w:rPr"))
    actual.styles["Heading 1"].element.get_or_add_rPr().append(revised)
    findings = compare(template, actual)
    assert "139" not in fail_ids(findings)
    assert any(f.rule_id == "139" and "修订" in f.message for f in findings)


def test_bad_template_note_part_does_not_suppress_readable_heading_pairs():
    template, actual = Document(), Document()
    for document in (template, actual):
        heading(document, "标题甲")
        heading(document, "标题乙")
    actual.paragraphs[1].runs[0].font.bold = False
    template.add_paragraph("注释标记").add_run()._r.append(node("w:footnoteReference", id=7))
    findings = compare(template, actual)
    assert "139" in fail_ids(findings)


def test_single_ole_math_object_reports_exact_unreadable_notice():
    template, actual = Document(), Document()
    p = actual.add_paragraph("(1)")
    obj = node("w:object")
    obj.set("ProgID", "Equation.3")
    p.add_run()._r.append(obj)
    findings = compare(template, actual)
    assert any(f.rule_id == "139" and f.status == Status.UNKNOWN and UNREADABLE in f.message for f in findings)


def test_unlabelled_bad_orientation_produces_unknown():
    template, actual = Document(), Document()
    actual.add_paragraph("普通页面")
    actual.sections[0]._sectPr.find(qn("w:pgSz")).set(qn("w:orient"), "invalid")
    findings = compare(template, actual)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN and "方向读取失败" in f.message for f in findings)


def test_special_label_field_cache_is_unknown_not_page_failure():
    template, actual = Document(), Document()
    special(template)
    _, anchor = special(actual, landscape=False)
    first = anchor.runs[0]._r
    first.addprevious(node("w:r"))
    first.getprevious().append(node("w:fldChar", fldCharType="begin"))
    instr = node("w:r")
    instruction = node("w:instrText")
    instruction.text = "REF SpecialPurpose"
    instr.append(instruction)
    first.addprevious(instr)
    sep = node("w:r")
    sep.append(node("w:fldChar", fldCharType="separate"))
    first.addprevious(sep)
    end = node("w:r")
    end.append(node("w:fldChar", fldCharType="end"))
    first.addnext(end)
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN and "动态域" in f.message for f in findings)


def test_special_label_with_inline_picture_is_unknown():
    template, actual = Document(), Document()
    special(template)
    _, anchor = special(actual, landscape=False)
    anchor.add_run()._r.append(node("w:drawing"))
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN and "图片" in f.message for f in findings)


def test_simple_field_only_special_label_is_unknown_even_if_portrait():
    template, actual = Document(), Document()
    special(template)
    _, anchor = special(actual, landscape=False)
    run = anchor.runs[0]._r
    field = node("w:fldSimple", instr="REF SpecialPurpose")
    anchor._p.append(field)
    field.append(run)
    findings = compare(template, actual)
    assert "164" not in fail_ids(findings)
    assert any(f.rule_id == "164" and f.status == Status.UNKNOWN and "动态域" in f.message for f in findings)
