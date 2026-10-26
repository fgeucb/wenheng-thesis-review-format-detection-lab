"""需要排版几何的规则：078 坐标判位、087 跨页表格、033 越界。"""

from __future__ import annotations

from docx import Document

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.models import Status
from thesis_format_checker.detectors.rendered import _overflow, _table_pagination, _toc_pages


def _context(*texts):
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    return DocxContext(document)


def test_cross_page_table_without_repeated_header_is_reported():
    context = _context("表 1 示例", "表格内容")
    layout = {"tables": {0: {"pages": (3, 4), "headers": [True, False, False],
                             "floating": False,
                             "box": {"page": 3, "end_page": 4, "x": 10, "y": 20, "bottom": 700}}}}
    findings = _table_pagination(context, layout)
    assert [finding.rule_id for finding in findings] == ["087.1"]
    assert "重复标题行" in findings[0].expected
    assert "跨第 3–4 页" in findings[0].actual


def test_cross_page_table_with_repeated_header_passes():
    context = _context("表 1 示例")
    layout = {"tables": {0: {"pages": (1, 2), "headers": [True, True, True], "floating": False}}}
    assert _table_pagination(context, layout) == []


def test_floating_cross_page_table_is_unknown_not_failure():
    context = _context("表 1 示例")
    layout = {"tables": {0: {"pages": (1, 2), "headers": [True, False], "floating": True}}}
    findings = _table_pagination(context, layout)
    assert findings and findings[0].status == Status.UNKNOWN


def test_single_page_table_is_not_reported():
    context = _context("表 1 示例")
    layout = {"tables": {0: {"pages": (2, 2), "headers": [True, False], "floating": False}}}
    assert _table_pagination(context, layout) == []


def test_figure_wider_than_printable_area_is_reported():
    document = Document()
    document.add_paragraph("图 1 示例")
    context = DocxContext(document)
    section = document.sections[0]
    right_edge = (float(section.page_width) - float(section.right_margin)) / 12700
    layout = {"figures": {0: {"boxes": [
        {"page": 1, "end_page": 1, "x": right_edge - 5, "y": 100,
         "bottom": 300, "width": 40.0}]}}}
    findings = _overflow(context, layout)
    assert [finding.rule_id for finding in findings] == ["033.1"]
    assert "超出可打印区" in findings[0].actual


def test_figure_inside_printable_area_passes():
    document = Document()
    document.add_paragraph("图 1 示例")
    context = DocxContext(document)
    layout = {"figures": {0: {"boxes": [
        {"page": 1, "end_page": 1, "x": 60.0, "y": 100, "bottom": 300, "width": 200.0}]}}}
    assert _overflow(context, layout) == []


def _toc_context():
    document = Document()
    document.add_paragraph("1 绪论", style="Heading 1")
    document.add_paragraph("正文内容。")
    return DocxContext(document)


def test_toc_cached_page_mismatch_is_reported(monkeypatch):
    context = _toc_context()
    monkeypatch.setattr("thesis_format_checker.detectors.toc._snapshot",
                        lambda ctx: {"groups": [{"entries": [
                            {"label": "1绪论", "page": "3", "level": 1}]}]})
    layout = {"paragraphs": {0: {"printed_number": 5}}}
    findings = _toc_pages(context, layout)
    assert [finding.rule_id for finding in findings] == ["059"]
    assert "应为 3" in findings[0].expected and "第 5 页" in findings[0].actual


def test_toc_matching_page_passes(monkeypatch):
    context = _toc_context()
    monkeypatch.setattr("thesis_format_checker.detectors.toc._snapshot",
                        lambda ctx: {"groups": [{"entries": [
                            {"label": "1绪论", "page": "5", "level": 1}]}]})
    assert _toc_pages(context, {"paragraphs": {0: {"printed_number": 5}}}) == []
