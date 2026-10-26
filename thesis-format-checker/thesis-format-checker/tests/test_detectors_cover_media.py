"""封面图片/艺术字检测（规则 017）的单元测试。

覆盖四种情况：内容一致 → 不报；论文缺图 → 报缺失；尺寸变了 → 报尺寸不符；
图片被挪到别处 → 报位置不符。
"""

from __future__ import annotations

import base64
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from thesis_format_checker.core.docx_context import DocxContext
from thesis_format_checker.core.models import Status
from thesis_format_checker.detectors.cover_media import (
    check_cover_media_errors,
    extract_cover_media_baseline,
)

#: 1×1 透明 PNG，用来冒充校徽/校名图片
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
)


def build_document(*, with_logo: bool = True, logo_index: int = 0, logo_width_mm: float = 26.7,
                   filler: int = 3) -> Document:
    """构造一份“封面 + 若干正文段”的文档；logo_index 指定图片放在第几段。"""

    document = Document()
    for position in range(filler):
        paragraph = document.add_paragraph("封面占位文字")
        if with_logo and position == logo_index:
            paragraph.add_run().add_picture(BytesIO(PNG), width=Mm(logo_width_mm))
    if with_logo and logo_index >= filler:
        for position in range(filler, logo_index):
            document.add_paragraph(f"第 {position + 1} 段")
        document.add_paragraph().add_run().add_picture(BytesIO(PNG), width=Mm(logo_width_mm))
    return document


def check(template: Document, paper: Document):
    baseline = extract_cover_media_baseline(DocxContext(template))
    return check_cover_media_errors(baseline, DocxContext(paper))


def codes(findings) -> set[str]:
    return {item.rule_id for item in findings if item.status == Status.FAIL}


def test_same_logo_and_position_pass():
    """图片数据一致、位置一致 → 不报任何问题。"""

    assert check(build_document(), build_document()) == []


def test_missing_logo_is_reported():
    """模板有校徽、论文没有 → 报 017.1 缺少校徽。"""

    findings = check(build_document(), build_document(with_logo=False))
    assert codes(findings) == {"017.1"}
    assert "校徽" in findings[0].message


def test_resized_logo_is_reported():
    """同一张图但尺寸放大一倍 → 报 017.3 尺寸不符。"""

    findings = check(build_document(), build_document(logo_width_mm=53.4))
    assert codes(findings) == {"017.3"}


def test_moved_logo_is_reported():
    """图片内容一致但被挪到第 20 段 → 报 017.3 位置不符。"""

    findings = check(build_document(), build_document(logo_index=20, filler=20))
    assert "017.3" in codes(findings)
    assert "位置" in findings[0].message or "第 20" in findings[0].message
