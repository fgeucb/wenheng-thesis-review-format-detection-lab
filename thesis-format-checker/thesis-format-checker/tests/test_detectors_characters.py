from __future__ import annotations

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.detectors.characters import extract_character_baseline, check_character_errors
from thesis_format_checker.core.models import Status


def family_of(rule_id):
    """细粒度规则编号 → 所属规则族（036.1 → 036）。"""

    return rule_id.split(".")[0]



def _doc(*texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    return doc


def _check(template, document):
    return check_character_errors(extract_character_baseline(DocxContext(template)), DocxContext(document))


def _matches(findings, rule, status=Status.FAIL):
    statuses = (Status.UNKNOWN, Status.NOT_APPLICABLE) if status == Status.UNKNOWN else (status,)
    return [f for f in findings if family_of(f.rule_id) == rule and f.status in statuses]


def test_equal_body_raw_characters_no_failure_and_no_mutation():
    texts = ["中文 ABC 与数字 123，正常：格式。", "其他 DEF 与数据 456，正文：字符。", "内容 ghi 789。"]
    template, document = _doc(*texts), _doc(*texts)
    before = template.element.xml, template.styles.element.xml, document.element.xml, document.styles.element.xml
    assert not [f for f in _check(template, document) if f.status == Status.FAIL]
    assert before == (template.element.xml, template.styles.element.xml, document.element.xml, document.styles.element.xml)


@pytest.mark.parametrize("expected,actual", [("ABC123", "ＡＢＣ１２３"), ("ＡＢＣ１２３", "ABC123")])
def test_width_standard_is_template_dynamic(expected, actual):
    template = _doc("正文" + expected + "。", "文字" + expected + "。")
    document = _doc("正文" + actual + "。")
    failures = _matches(_check(template, document), "124")
    assert failures and failures[0].anchor._p is document.paragraphs[0]._p
    assert "应该：" in failures[0].message and "实际：" in failures[0].message


def test_width_preserves_legitimate_chinese_and_english_punctuation_contexts():
    template = _doc("中文，词语；语言。 alpha,beta;gamma", "正文，内容；段落。 delta,epsilon;theta")
    document = _doc("中文，词语；语言。 alpha,beta;gamma")
    assert not _matches(_check(template, document), "124")
    document.paragraphs[0].text = "中文,词语；语言。 alpha,beta;gamma"
    findings = _matches(_check(template, document), "124")
    assert findings and "中文语境逗号" in findings[0].message


def test_width_conflicts_are_unknown_not_majority_baseline():
    template = _doc("ABC 123", "ABC 456", "ＡＢＣ 789")
    findings = _check(template, _doc("ＡＢＣ 123"))
    assert not _matches(findings, "124")
    assert _matches(findings, "124", Status.UNKNOWN)


def test_one_body_example_is_insufficient_even_with_many_characters():
    template = _doc("多个英文 ABC DEF GHI 与数字 123 456 789。")
    findings = _check(template, _doc("中文ＡＢＣ１２３。"))
    assert not _matches(findings, "124")
    assert _matches(findings, "124", Status.UNKNOWN)


@pytest.mark.parametrize("expected,actual", [("", " "), (" ", ""), ("\u00a0", " "), ("\u3000", " ")])
def test_boundary_spacing_is_dynamic_and_codepoint_exact(expected, actual):
    template = _doc("中文" + expected + "ABC" + expected + "内容" + expected + "123" + expected + "数据。",
                    "正文" + expected + "DEF" + expected + "文字" + expected + "456" + expected + "信息。")
    document = _doc("中文" + actual + "ABC" + actual + "内容" + actual + "123" + actual + "数据。")
    failures = _matches(_check(template, document), "125")
    assert failures and "自动间距" in failures[0].actual
    assert _check(template, document)


def test_boundary_rules_keep_directions_and_letter_digit_uses_separate():
    template = _doc("中文ABC 结束123内容。", "文字DEF 结束456正文。")
    document = _doc("中文ABC 结束123内容。")
    assert not _matches(_check(template, document), "125")


def test_url_doi_hyperlink_code_formula_and_reference_regions_are_excluded():
    template = _doc("中文 ABC 内容 123 数据。", "正文 DEF 文字 456 信息。")
    document = _doc("链接https://example.com/a123中文", "DOI:10.1000/Abc中文", "邮箱a@b.com中文")
    document.styles.add_style("Code Listing", WD_STYLE_TYPE.PARAGRAPH).base_style = document.styles["Normal"]
    document.add_paragraph("代码ＡＢＣ１２３中文", "Code Listing")
    p = document.add_paragraph("正文ＡＢＣ１２３中文")
    math = OxmlElement("m:oMath")
    p._p.append(math)
    p = document.add_paragraph("正文ＡＢＣ１２３中文")
    p._p.append(OxmlElement("w:hyperlink"))
    document.add_heading("参考文献", level=1)
    document.add_paragraph("[1]ＡＢＣ１２３中文")
    findings = _check(template, document)
    assert not _matches(findings, "124")
    assert not _matches(findings, "125")


def test_split_runs_preserve_boundary_and_character_evidence():
    template = _doc("中文 ABC 内容。", "正文 DEF 文字。")
    document = _doc()
    p = document.add_paragraph()
    p.add_run("中文")
    p.add_run("A")
    p.add_run("BC")
    p.add_run(" 内容。")
    assert _matches(_check(template, document), "125")


def test_auto_spacing_does_not_count_as_real_space():
    template = _doc("中文 ABC 内容。", "正文 DEF 文字。")
    document = _doc("中文ABC 内容。")
    ppr = document.paragraphs[0]._p.get_or_add_pPr()
    node = OxmlElement("w:autoSpaceDE")
    node.set(qn("w:val"), "1")
    ppr.append(node)
    assert _matches(_check(template, document), "125")


def test_consecutive_empty_paragraphs_compare_boundary_not_fixed_threshold():
    template = _doc("第一段正文。", "第二段正文。", "第三段正文。")
    document = _doc("第一段正文。", "", "", "第二段正文。", "第三段正文。")
    findings = _matches(_check(template, document), "128")
    assert len(findings) == 1
    assert findings[0].anchor._p is document.paragraphs[3]._p
    assert "连续空段数=2" in findings[0].actual


def test_template_two_blank_paragraphs_are_allowed_and_single_blank_left_to_075():
    template = _doc("第一段。", "", "", "第二段。", "", "", "第三段。")
    document = _doc("第一段。", "", "", "第二段。", "", "", "第三段。")
    assert not _matches(_check(template, document), "128")
    other = _doc("第一段。", "", "第二段。", "第三段。")
    assert not _matches(_check(template, other), "128")


def test_conflicting_template_blank_boundaries_are_unknown():
    template = _doc("第一段。", "", "第二段。", "第三段。")
    document = _doc("第一段。", "", "", "第二段。")
    findings = _check(template, document)
    assert not _matches(findings, "128")
    assert _matches(findings, "128", Status.UNKNOWN)


def test_template_revision_poisoned_gap_and_manual_line_baselines():
    template = _doc("第一段。", "第二段。", "第三段。")
    change = OxmlElement("w:pPrChange")
    template.paragraphs[1]._p.get_or_add_pPr().append(change)
    document = _doc("第一段。", "", "", "第二段。", "末段\n\n下行。")
    findings = _check(template, document)
    assert not _matches(findings, "128")
    assert _matches(findings, "128", Status.UNKNOWN)


def test_template_revision_in_empty_paragraph_poisoned_gap_baseline():
    template = _doc("第一段。", "", "第二段。", "", "第三段。")
    template.paragraphs[1]._p.get_or_add_pPr().append(OxmlElement("w:pPrChange"))
    document = _doc("第一段。", "", "", "第二段。")
    findings = _check(template, document)
    assert not _matches(findings, "128")
    assert _matches(findings, "128", Status.UNKNOWN)


def test_pagebreak_section_bookmark_and_table_end_are_not_empty_filler():
    template = _doc("第一段。", "第二段。", "第三段。")
    document = _doc("第一段。")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("第二段。")
    document.add_paragraph()._p.append(OxmlElement("w:bookmarkStart"))
    document.add_paragraph()._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    document.add_paragraph("第三段。")
    table = document.add_table(rows=3, cols=1)
    for cell in table.column_cells(0):
        cell.add_paragraph("")
    assert not _matches(_check(template, document), "128")


def test_consecutive_manual_breaks_and_explicit_page_break_not_conflated():
    template = _doc("第一段。", "第二段。", "第三段。")
    document = _doc("正文\n\n\n后文。")
    assert _matches(_check(template, document), "128")
    document = _doc()
    p = document.add_paragraph("正文")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("后文")
    assert not _matches(_check(template, document), "128")


@pytest.mark.parametrize("samples,actual", [
    (("提交日期：2026年05月02日", "提交日期：2027年06月03日"), "提交日期：2028/07/04"),
    (("提交日期：2026/5/2", "提交日期：2027/6/3"), "提交日期：2028年7月4日"),
    (("开始时间：08:30", "开始时间：09:45"), "开始时间：10时30分"),
    (("适用年份：2026年", "适用年份：2027年"), "适用年份：2028"),
    (("项目编号：AB-001", "项目编号：CD-002"), "项目编号：EF/003"),
])
def test_labelled_date_time_year_number_format_dynamically_compares(samples, actual):
    template, document = _doc(*samples), _doc(actual)
    failures = _matches(_check(template, document), "129")
    assert failures and failures[0].anchor._p is document.paragraphs[0]._p


def test_numeric_padding_requires_repeated_demonstrated_zero_padding():
    padded = _doc("提交日期：2026-05-02", "提交日期：2027-06-03")
    assert _matches(_check(padded, _doc("提交日期：2028-7-4")), "129")
    unproven = _doc("提交日期：2026-11-12", "提交日期：2027-12-13")
    assert not _matches(_check(unproven, _doc("提交日期：2028-7-4")), "129")


def test_same_pattern_values_not_compared_as_identity_or_factual_truth():
    template = _doc("项目编号：AB-001", "项目编号：CD-002", "提交日期：2026年5月2日", "提交日期：2027年6月3日")
    document = _doc("项目编号：XY-999", "提交日期：9999年99月99日")
    assert not _matches(_check(template, document), "129")


def test_free_text_decimals_versions_citations_are_not_dates_or_number_fields():
    template = _doc("提交日期：2026.5", "提交日期：2027.6")
    document = _doc("实验结果2028.7，版本2026.5，文献[123]。", "日期意义是统计区间而非具体日子。")
    assert not [f for f in _check(template, document) if f.rule_id == "129"]


def test_date_roles_are_not_pooled_or_inferred_from_one_example():
    template = _doc("提交日期：2026年5月", "审核日期：2027年6月")
    findings = _check(template, _doc("提交日期：2028/7"))
    assert not _matches(findings, "129")
    assert _matches(findings, "129", Status.UNKNOWN)


def test_bibliography_subheading_does_not_end_reference_exclusion():
    template = _doc("项目编号：AB-001", "项目编号：CD-002")
    document = _doc()
    document.add_heading("参考文献", level=1)
    document.add_heading("期刊论文", level=2)
    document.add_paragraph("项目编号：EF/003")
    assert not [f for f in _check(template, document) if f.rule_id == "129"]


def test_conflicting_date_template_is_unknown_not_majority_vote():
    template = _doc("提交日期：2026年5月", "提交日期：2027年6月", "提交日期：2028/7")
    findings = _check(template, _doc("提交日期：2029/8"))
    assert not _matches(findings, "129")
    assert _matches(findings, "129", Status.UNKNOWN)


def test_date_field_does_not_trust_stale_result_and_has_library_limit_notice():
    template = _doc("提交日期：2026年5月", "提交日期：2027年6月")
    document = _doc("提交日期：")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'DATE \\@ "yyyy/M"')
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "2028/7"
    run.append(text)
    field.append(run)
    document.paragraphs[0]._p.append(field)
    findings = _check(template, document)
    assert not _matches(findings, "129")
    assert UNREADABLE in _matches(findings, "129", Status.UNKNOWN)[0].actual


def test_tracked_revision_character_check_is_unknown_not_final_text_error():
    template = _doc("正文 ABC 内容。", "文字 DEF 数据。")
    document = _doc("正文ＡＢＣ内容。")
    change = OxmlElement("w:rPrChange")
    document.paragraphs[0].runs[0]._r.get_or_add_rPr().append(change)
    findings = _check(template, document)
    assert not _matches(findings, "124")
    assert _matches(findings, "124", Status.UNKNOWN)


def test_only_requested_api_rules_are_emitted_even_on_unreadable_baseline():
    findings = check_character_errors({"characters": {"unreadable": "XML错误"}}, DocxContext(_doc("正文。")))
    # 126/127/130/131/132 同样是本模块负责的细粒度规则族。
    assert {f.rule_id for f in findings} == {"124", "125", "126", "127", "128", "129",
                                             "130", "131", "132"}
    assert all(f.status == Status.UNKNOWN and UNREADABLE in f.actual for f in findings)


def test_unexpected_per_paragraph_read_error_is_isolated(monkeypatch):
    import thesis_format_checker.detectors.characters as module
    template, document = _doc("正文 ABC 内容。", "文字 DEF 数据。"), _doc("正文 ABC 内容。", "文字 DEF 数据。")
    baseline = extract_character_baseline(DocxContext(template))
    original = module._check_text

    def broken_first(data, record):
        if record["paragraph"].text.startswith("正文"):
            raise ValueError("模拟读取异常")
        return original(data, record)

    monkeypatch.setattr(module, "_check_text", broken_first)
    findings = module.check_character_errors(baseline, DocxContext(document))
    assert _matches(findings, "124", Status.UNKNOWN)
    assert _matches(findings, "125", Status.UNKNOWN)
