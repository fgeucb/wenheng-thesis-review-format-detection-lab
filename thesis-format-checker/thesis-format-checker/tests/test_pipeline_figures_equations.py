from base64 import b64decode
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

from thesis_format_checker.engine.pipeline import run_comparison, semantic_tree
from thesis_format_checker.core.docx_context import DocxContext, UNREADABLE
from thesis_format_checker.core.models import Finding, RunSummary, Status
from thesis_format_checker.checks.catalog import (DEFERRED_RULE_IDS, LEGACY_RULE_IDS,
                                                  PENDING_RULE_IDS, REMOVED_RULE_IDS,
                                                  RENDERED_FAMILIES)
from thesis_format_checker.engine.pipeline import SECTION_RULES, check_sections, extract_baseline
from conftest import comment_heads, family_titles


def fixture(path, *, template):
    from test_pipeline_headers_toc import fixture as prior_fixture
    from test_detectors_figures import caption, table, node
    from test_detectors_equations import equation, math_defaults
    prior_fixture(path, template=True)
    doc = Document(path)
    math_defaults(doc)
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.styles["Normal"].paragraph_format.keep_with_next = False
    doc.styles["Normal"].paragraph_format.keep_together = False
    paragraphs = [p for p in doc.paragraphs if p.text.startswith("这是第")]
    if not template:
        paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraphs[2].paragraph_format.left_indent = Pt(20)
        paragraphs[3].paragraph_format.keep_with_next = True
        paragraphs[4].add_run("这是局部格式异常文字 ABC123。").font.bold = True
        next(p for p in doc.paragraphs if p.text.startswith("1.1 ")).paragraph_format.first_line_indent = Pt(12)
    png = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
    for number in (1, 2):
        p = doc.add_paragraph()
        p.add_run().add_picture(BytesIO(png), width=Inches(1))
        label = "图 1 测试示意" if number == 1 else "图 2 测试示意" if template else "图4 测试示意"
        c = caption(doc, label)
        if not template and number == 1:
            c.runs[0].font.size = Pt(19)
            p._p.addprevious(c._p)
        doc.add_paragraph("下面继续讨论实验结果，示意图采用嵌入型图片。")
    _, tbl = table(doc)
    if not template:
        borders = node("w:tcBorders")
        borders.append(node("w:left", val="double", sz="8", color="FF0000"))
        tbl.cell(1, 0)._tc.get_or_add_tcPr().append(borders)
        tbl.cell(1, 1).paragraphs[0].runs[0].font.size = Pt(18)
        tbl.rows[0].height = Pt(25)
    doc.add_paragraph("公式全文连续编号，公式编号从1开始。")
    equation(doc, "(1)")
    p, _ = equation(doc, "(2)" if template else "（4）", jc="center" if template else "left", size=None if template else 19)
    if not template:
        p.paragraph_format.space_after = Pt(20)
    doc.add_paragraph("由本文式(1)可得结果。" if template else "由本文式(9)可得结果。")
    if not template:
        doc.add_comment(paragraphs[0].runs, "保留的人工批注", author="人工校阅者")
    doc.save(path)


def test_exact_coverage_reuse_extension_and_namespaced_baseline():
    """大类七、八、九（067–099）的规则全部有结论：docx 检测 + 旧版 Word 流程。"""
    sections = ("paragraphs", "figures", "equations")
    new = [rule for section in sections for rule in SECTION_RULES[section]]
    assert len(new) == len(set(new)) == 18
    legacy_here = {rule for rule in LEGACY_RULE_IDS if 67 <= int(rule) <= 99}
    assert set(new) & legacy_here == {"069"}  # 069 由 docx 注释字号与旧版正文字号共同提供
    api = set(new) | legacy_here
    assert len(api) == 23
    removed_here = {rule for rule in REMOVED_RULE_IDS if 67 <= int(rule) <= 99}
    pending_here = {rule for rule in PENDING_RULE_IDS if 67 <= int(rule) <= 99}
    deferred_here = {rule for rule in DEFERRED_RULE_IDS if 67 <= int(rule) <= 99}
    rendered_here = {rule for rule in RENDERED_FAMILIES if 67 <= int(rule) <= 99}
    assert removed_here == {"086", "099"}
    assert pending_here == set()
    # 087/098 已由排版几何阶段实现（跨页表格、公式越界），不再是待确认；
    # 其余这些族明确标注为“待确认/依赖语义”，不再标成可确定性待实现。
    assert rendered_here == {"087", "098"}
    assert deferred_here == {"082", "085", "088", "089", "094", "095"}
    assert not api & (removed_here | pending_here | deferred_here)
    assert api | rendered_here | removed_here | pending_here | deferred_here == {
        f"{number:03}" for number in range(67, 100)}
    old = {"previous": {"untouched": True}}
    ctx = DocxContext(Document())
    baseline = extract_baseline(ctx, old)
    assert baseline["previous"] is old["previous"]
    assert set(old) == {"previous"}
    assert {"paragraphs", "figures", "equations"} <= set(baseline)
    with pytest.raises(ValueError, match="未覆盖"):
        extract_baseline(ctx, baseline)


def test_detector_exceptions_produce_read_limits_and_do_not_abort_others(monkeypatch):
    """段落/图表/公式任一模块抛异常时，只影响该组的 18 条规则。"""
    import thesis_format_checker.detectors.equations as equations
    import thesis_format_checker.detectors.figures as figures
    import thesis_format_checker.detectors.paragraphs as paragraphs

    def broken(*args, **kwargs):
        raise ValueError("invalid object")

    monkeypatch.setattr(paragraphs, "check_paragraph_errors", broken)
    monkeypatch.setattr(figures, "check_figure_table_errors", broken)
    monkeypatch.setattr(equations, "check_equation_errors", broken)
    ctx = DocxContext(Document())
    sections = ("paragraphs", "figures", "equations")
    baseline = extract_baseline(ctx, sections=sections)
    result = check_sections(baseline, ctx, sections=sections)
    assert len(result) == 18
    assert all(f.status == Status.UNKNOWN and UNREADABLE in f.message for f in result)
    assert not (set(REMOVED_RULE_IDS) | set(PENDING_RULE_IDS)) & {f.rule_id for f in result}


def test_pipeline_preserves_legacy_findings_and_noncomment_package_parts(runtime_dir):
    a, b, out = [runtime_dir / name for name in ("template.docx", "paper.docx", "out.docx")]
    fixture(a, template=True)
    fixture(b, template=False)
    original = [p.read_bytes() for p in (a, b)]
    calls = []
    def legacy(*args, **kwargs):
        calls.append(args)
        return RunSummary(findings=[Finding("067", "原有中文字体", Status.FAIL, "既有检测证据", expected="模板字体", actual="不同字体", paragraph_index=5)])
    summary = run_comparison(legacy, a, b, out)
    assert len(calls) == 1
    failed = {f.rule_id for f in summary.findings if f.status == Status.FAIL}
    # 细粒度拆分后按规则族统计（073.1、076.2 …）
    families = {rule_id.split(".")[0] for rule_id in failed}
    assert {"067", "070", "071", "073", "076", "077", "078", "079", "080", "081", "083", "084",
            "090", "091", "092", "093", "096", "097"} <= families
    assert not (set(REMOVED_RULE_IDS) | set(PENDING_RULE_IDS)) & {f.rule_id for f in summary.findings}
    assert any("已删除" in note for note in summary.skipped)
    assert original == [p.read_bytes() for p in (a, b)]
    with ZipFile(b) as source, ZipFile(out) as annotated:
        allowed = {"word/document.xml", "word/comments.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        assert set(source.namelist()) == set(annotated.namelist())
        for name in source.namelist():
            if name not in allowed:
                assert source.read(name) == annotated.read(name), name
        assert semantic_tree(etree.fromstring(source.read("word/document.xml"))) == semantic_tree(etree.fromstring(annotated.read("word/document.xml")))
    comments = list(Document(out).comments)
    assert comments[0].text == "保留的人工批注"
    # Word 会把有批注引用 run 的显示公式改按行内排版；结论保留但不写入。
    assert not (comment_heads(comments) & family_titles("096"))
    assert any(note.startswith("【096") and "显示公式" in note for note in summary.report_only)
    for comment in comments[1:]:
        assert "应该：" in comment.text and "实际：" in comment.text
        # 批注已锚定在出问题的文字上，正文里不再重复写“位置：”一行。
        assert "位置：" not in comment.text
