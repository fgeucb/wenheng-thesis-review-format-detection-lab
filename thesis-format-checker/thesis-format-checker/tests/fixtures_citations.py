"""引用、注释与字符类规则的真实 DOCX 夹具。"""
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.shared import Pt
from lxml import etree


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    return node


def add_notes(doc, paragraph, kind="footnote", size=10):
    ref = paragraph.add_run()
    ref.font.superscript = True
    ref._r.append(element("w:" + kind + "Reference", id="1"))
    root = parse_xml(f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    for identity, tag in ((-1, "separator"), (0, "continuationSeparator")):
        note = element("w:" + kind, id=identity, type=tag)
        p = element("w:p")
        run = element("w:r")
        run.append(element("w:" + tag))
        p.append(run)
        note.append(p)
        root.append(note)
    note = element("w:" + kind, id="1")
    p = element("w:p")
    prop = element("w:pPr")
    prop.append(element("w:spacing", before="0", after="0", line="240", lineRule="auto"))
    prop.append(element("w:ind", left="0", right="0", firstLine="0"))
    p.append(prop)
    marker = element("w:r")
    marker.append(element("w:" + kind + "Ref"))
    p.append(marker)
    run = element("w:r")
    rpr = element("w:rPr")
    rpr.append(element("w:rFonts", ascii="Arial", hAnsi="Arial", eastAsia="仿宋"))
    rpr.append(element("w:sz", val=str(size * 2)))
    run.append(rpr)
    text = element("w:t")
    text.text = "这是注释说明 note 123。"
    run.append(text)
    p.append(run)
    note.append(p)
    root.append(note)
    part = Part(PackURI(f"/word/{kind}s.xml"),
                f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
                etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), doc.part.package)
    doc.part.relate_to(part, f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s")
    settings = doc.settings.element.find(qn("w:" + kind + "Pr"))
    if settings is None:
        settings = element("w:" + kind + "Pr")
        doc.settings.element.append(settings)
    for name, value in (("numFmt", "decimal"), ("numStart", "1"), ("numRestart", "continuous")):
        settings.append(element("w:" + name, val=value))
    for identity in (-1, 0):
        settings.append(element("w:" + kind, id=identity))


def field(paragraph, instruction, cache):
    node = element("w:fldSimple", instr=instruction)
    run = element("w:r")
    text = element("w:t")
    text.text = cache
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


def bibliography_entry(number, title="测试题名"):
    return f"[{number}] 张三，李四. {title}[J]. 示例期刊, 2024, 12(3): 4-9. DOI: 10.1234/abc https://example.org/a [2024-01-02]"


def fixture(path, *, template):
    from test_pipeline_headers_toc import fixture as prior_fixture
    prior_fixture(path, template=True)
    doc = Document(path)
    body = [p for p in doc.paragraphs if p.text.startswith("这是第")]
    for i, p in enumerate(body):
        p.runs[0].text = f"这是第 {i + 1} 段普通正文，测试中文 English 混排及数字 123 的格式。"
    if not template:
        body[2].add_run("正文 ABC 123 示例中的ＡＢＣ１２３字符。")
        body[3].add_run("这是English混排以及数字123边界。")
        body[4].insert_paragraph_before("")
        body[4].insert_paragraph_before("")
        body[5].add_run("中文,标点与（未闭合括号。")
    add_notes(doc, body[0], size=10 if template else 15)
    if not template:
        add_notes(doc, body[1], kind="endnote")
        doc.settings.element.find(qn("w:footnotePr")).find(qn("w:numRestart")).set(qn("w:val"), "eachSect")
    doc.add_paragraph("提交日期：2026年05月02日" if template else "提交日期：2028/07/04")
    doc.add_paragraph("提交日期：2027年06月03日")
    doc.add_paragraph("交叉引用必须使用REF域。", style="Title")
    caption = doc.add_paragraph("图1", style="Caption")
    caption._p.insert(0, element("w:bookmarkStart", id="40", name="figure_one"))
    caption._p.append(element("w:bookmarkEnd", id="40"))
    reference = doc.add_paragraph("参见")
    if template:
        field(reference, "REF figure_one \\h", "图1")
    else:
        reference.add_run("图1")
    reference.add_run("的结果。")
    stale = doc.add_paragraph("另见")
    field(stale, "REF figure_one \\h", "图1" if template else "图9")
    broken = doc.add_paragraph("交叉引用显示：")
    field(broken, "REF figure_one \\h" if template else "REF missing_target \\h", "图1" if template else "错误！未找到引用源。")
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[J]必填字段：作者、题名、期刊名、年份、卷、期、页码、DOI、URL、引用日期")
    doc.add_paragraph("中文作者超过3人，列前3人，后加等")
    doc.add_paragraph("参考文献按文末编号升序排列")
    if template:
        for i in (1, 2):
            doc.add_paragraph(bibliography_entry(i, f"测试题名{i}"))
    else:
        texts = [bibliography_entry(i, f"测试题名{i}") for i in range(1, 8)]
        texts[0] = texts[0].replace("[1]", "（1）", 1)
        texts[1] = texts[1].replace("张三，李四", "张三，李四，王五，赵六", 1)
        texts[2] = texts[2].replace("测试题名3", "")
        texts[3] = texts[3].replace("12(3):", "12(3)：")
        texts[4] = texts[4].replace("DOI:", "DOI：")
        texts[5] = texts[5].replace("[J].", "[J]。")
        texts[5], texts[6] = texts[6], texts[5]
        for text in texts:
            doc.add_paragraph(text)
        doc.add_comment(body[0].runs, "保留的人工批注", author="人工校阅")
    doc.save(path)
