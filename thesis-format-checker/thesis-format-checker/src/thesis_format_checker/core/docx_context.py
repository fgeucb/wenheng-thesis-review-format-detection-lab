"""只读的 docx 段落与格式上下文。

本模块封装 python-docx，提供段落/文本枚举、有效段落格式与有效文字格式读取，
供结构、封面、页面等大类共用；所有读取都不会修改文档。"""
from __future__ import annotations

from docx.text.paragraph import Paragraph

from thesis_format_checker.core.models import Finding, Severity, Status, build_comment
from thesis_format_checker.core.xmlutil import qn


def memo(ctx, key: str, produce):
    """在**同一份文档**上缓存"纯读取"结果（同一份文档只算一次）。

    检测器里有不少全篇扫描（标题记录、题注记录、域扫描、编号树…）会被多个大类
    反复调用：实测标题记录一次运行被算了 40 次、题注 10 次、域扫描 6 次。
    这些扫描只读文档、结果只取决于文档内容，因此把结果挂在上下文对象上复用；
    缓存随上下文一起释放，不会跨文档串味。

    约定：被缓存的返回值视为**只读**（调用方不得就地修改记录）。
    """

    store = getattr(ctx, "_memo", None)
    if store is None:
        store = {}
        setattr(ctx, "_memo", store)
    if key not in store:
        store[key] = produce()
    return store[key]


def install_style_cache(document) -> None:
    """给 python-docx 的样式解析加一层缓存（同一份文档只解析一次）。

    ``paragraph.style`` / ``run.style`` 在检测里被调用十几万次，python-docx 每次都要
    用 ``Styles.get_by_id`` 扫一遍 styles.xml，未命中时还会 ``default_for`` 全表扫描
    （实测累计约 46 秒，是 docx 侧最大的单项开销）。样式部件在检测期间只读，
    因此按 ``(style_id, style_type)`` 缓存解析结果即可，语义完全不变。

    注意：``document.styles`` 每次访问都会**新建**一个 ``Styles`` 代理对象，
    所以缓存必须挂在 ``DocumentPart`` 上（``paragraph.style`` / ``run.style``
    最终都走 ``part.get_style``）。
    """

    part = getattr(document, "part", None)
    if part is None or getattr(part, "_tfc_style_cache", None) is not None:
        return
    original = part.get_style
    cache: dict = {}

    def cached(style_id, style_type):
        key = (style_id, style_type)
        if key not in cache:
            cache[key] = original(style_id, style_type)
        return cache[key]

    try:
        part.get_style = cached
        part._tfc_style_cache = cache
    except Exception:      # 某些 python-docx 版本的对象不可写属性：忽略即可
        pass

UNREADABLE = "该对象docx库无法读取，本条暂无法执行检测"


class DocxContext:
    def __init__(self, document):
        install_style_cache(document)
        self.document = document
        self.paragraphs = [Paragraph(p, document._body) for p in document.element.body.iter(qn("w:p"))
                           if not any(a.tag == qn("w:txbxContent") for a in p.iterancestors())]
        self.texts = [p.text for p in self.paragraphs]
        self.text = "\n".join(self.texts)
        self.has_textboxes = bool(document.element.xpath(".//w:txbxContent"))
        self.issues = []
        self._sections = []
        self._front_matter_end: int | None = None
        section = 0
        for p in self.paragraphs:
            self._sections.append(section)
            if p._p.xpath("./w:pPr/w:sectPr"):
                section += 1

    @property
    def front_matter_end(self) -> int:
        """封面/题名页/声明等前置区的结束下标；0 表示未识别到摘要或目录标志。

        前置区（摘要页之前）只由封面模块按模板同位置比对，不参与标题层级、
        正文段落、摘要等通用判定；本属性是这一口径的唯一来源，结果按上下文缓存。
        """

        if self._front_matter_end is None:
            from thesis_format_checker.core.scope import front_matter_end

            self._front_matter_end = front_matter_end(self.texts)
        return self._front_matter_end

    def is_front_matter(self, paragraph) -> bool:
        """段落是否属于封面/声明等前置区（摘要、目录及其后一律不算）。"""

        index = self.index(paragraph)
        return 0 <= index < self.front_matter_end

    def index(self, paragraph):
        return next((i for i, p in enumerate(self.paragraphs) if p._p is paragraph._p), -1)

    def section_index(self, paragraph):
        index = self.index(paragraph)
        return self._sections[index] if index >= 0 else 0


def make_finding(rule, title, expected, actual, paragraph=None, status=Status.FAIL,
                 severity: Severity = Severity.ERROR) -> Finding:
    """构造一条结论；批注文案统一由 :func:`core.models.build_comment` 渲染。

    等级在检测调度层会按规则目录统一覆盖（见 ``engine.pipeline.normalize_finding``），
    所以各取证模块不必各自维护等级，只需给出“应该是什么、实际是什么”。
    """

    rule_id = str(rule).zfill(3)
    finding = Finding(rule_id, title, status, "",
                      severity=severity, expected=expected, actual=actual)
    finding.message = build_comment(severity, rule_id, title, expected, actual, status=status)
    finding.anchor = paragraph
    return finding


#: 说明“模板里没有对应对象/样例”的措辞。命中这些措辞时，规则对该论文并不适用：
#: 既不写批注，也不计入“无法判定”；模板本身的完备性由模板体检工具负责提示。
TEMPLATE_MISSING_MARKERS = (
    "模板没有", "模板缺少", "模板未提供", "模板未明确", "模板未形成", "模板中不存在",
    "模板中没有", "模板不存在", "模板相关示例含不可读", "模板同角色基准不唯一",
    # 页眉页脚类：模板节基准不唯一/生效选项不一致时，本条对这份论文没有可比基准。
    "模板页眉页脚基准不唯一", "首页/奇偶页的生效选项不一致", "首页/奇偶页选项不一致",
    # Word 路径：模板没有同级标题样例 / 页眉基准无法匹配时，不写“无法判定”。
    "模板 A 中没有", "无法匹配的模板基准",
    "模板同角色文字格式不唯一", "模板同角色题注存在不同格式",
    "模板对应编号形式缺失", "模板对应独立/行内/表格公式段落基准缺失",
    "模板对应角色格式缺失或不唯一", "模板对应公式布局缺失",
    "模板含超出内容区域", "模板排序方式存在冲突", "模板未给出",
    "模板同类年卷期页数字结构不唯一", "缺少至少两处", "同用途段内换行样例不足",
    "模板对应用途部分文字读取失败",
    # 节对应关系：论文与模板分节数不同、模板同角色节基准不一致时，
    # 结论只进报告，不要求使用者手写 section_map。
    "请配置 section_map", "显式指定对应关系", "模板存在不同节基准",
    # 模板只给了示例、没写明“必须包含”的组成部分：不能据此判定论文缺失。
    "未明确其必需性",
)


def is_template_missing(reason) -> bool:
    """原因是否属于“模板里没有对应内容”，而不是论文本身的问题。"""

    text = str(reason or "")
    return any(marker in text for marker in TEMPLATE_MISSING_MARKERS)


def unknown(rule, title, reason, paragraph=None):
    """构造“无法判定”结论。

    若原因说明的是模板里没有对应对象/样例（见 :data:`TEMPLATE_MISSING_MARKERS`），
    说明这条规则对这份论文并不适用，直接跳过：不写批注、不计入“无法判定”。
    """

    if is_template_missing(reason):
        return not_applicable(rule, title, reason, paragraph)
    return make_finding(rule, title, "仅按模板中可确认的对应要求检测",
                        "暂无法判定：" + str(reason), paragraph, Status.UNKNOWN)


def not_applicable(rule, title, reason, paragraph=None):
    """模板里没有对应对象/样例：本条对这份论文不适用。

    这类结论**既不写批注也不计入“无法判定”**，只在报告里作为“已跳过”统计，
    避免“模板没有密级、分类号”之类的规则在论文里刷屏。
    """

    return make_finding(rule, title, "模板未涉及此项",
                        "不适用：" + str(reason), paragraph, Status.NOT_APPLICABLE)


def _styles(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def _default_prop(paragraph, name):
    roots = paragraph.part.document.styles.element.xpath(f"./w:docDefaults/w:{name}Default/w:{name}")
    return roots[0] if roots else None


def _effective_size_pt(paragraph):
    """段落的有效字号（磅）：取字符数最多的字号；读不到时退回文档默认字号。

    字符单位缩进（``firstLineChars``/``leftChars``）要按字号换算成磅值，
    因此这里必须给出一个确定值，否则该属性只能报“无法完整确认”。
    """

    weights: dict[float, int] = {}
    try:
        for text, props in effective_runs(paragraph):
            size = props.get("size_pt")
            if size:
                weights[size] = weights.get(size, 0) + len(text or "")
    except Exception:
        pass
    if weights:
        return max(weights.items(), key=lambda item: item[1])[0]
    try:
        default = _default_prop(paragraph, "rPr")
        node = default.find(qn("w:sz")) if default is not None else None
        if node is not None and node.get(qn("w:val")):
            return float(node.get(qn("w:val"))) / 2
    except Exception:
        pass
    return None


def effective_paragraph(paragraph):
    keys = ("alignment", "first_line_indent", "left_indent", "right_indent", "space_before", "space_after",
            "line_spacing", "line_spacing_rule", "keep_with_next", "keep_together", "widow_control", "page_break_before")
    result = {}
    formats = [paragraph.paragraph_format] + [s.paragraph_format for s in _styles(paragraph.style)]
    default = _default_prop(paragraph, "pPr")
    if default is not None:
        from docx.oxml import OxmlElement
        shell = OxmlElement("w:p")
        from copy import deepcopy
        shell.append(deepcopy(default))
        formats.append(Paragraph(shell, paragraph._parent).paragraph_format)
    for key in keys:
        value = next((getattr(fmt, key) for fmt in formats if getattr(fmt, key) is not None), None)
        if hasattr(value, "pt"):
            value = round(value.pt, 5)
        elif value is not None and hasattr(value, "numerator") and not isinstance(value, bool):
            value = float(value) if isinstance(value, float) else int(value)
        result[key] = value
    # These are WordprocessingML absence semantics, not school formatting standards.
    for key in ("alignment", "first_line_indent", "left_indent", "right_indent", "space_before", "space_after"):
        if result[key] is None:
            result[key] = 0
    for key in ("keep_with_next", "keep_together", "page_break_before"):
        if result[key] is None:
            result[key] = False
    # python-docx exposes point indents/spacing but cannot resolve character-
    # or line-unit attributes. These override the corresponding point values
    # in Word; treating an absent point value as zero would produce false FAILs.
    props = [paragraph._p.pPr] + [style.element.pPr for style in _styles(paragraph.style)]
    props.append(default)

    def raw_attr(element, attr):
        for prop in props:
            if prop is not None:
                node = prop.find(qn("w:" + element))
                if node is not None and node.get(qn("w:" + attr)) is not None:
                    return node.get(qn("w:" + attr))
        return None

    # 字符单位缩进（w:ind/@w:firstLineChars 等）：Word 按“1 字符 = 当前字号”换算，
    # 这里用段落有效字号换算成磅值，而不是把有效值判成“不可读”（封面常用字符单位）。
    char_size = _effective_size_pt(paragraph)
    for key, attrs in {
        "first_line_indent": ("firstLineChars", "hangingChars"),
        "left_indent": ("leftChars", "startChars", "start"),
        "right_indent": ("rightChars", "endChars", "end"),
    }.items():
        values = {attr: raw_attr("ind", attr) for attr in attrs}
        values = {attr: value for attr, value in values.items() if value is not None}
        if not values:
            continue
        chars = None
        if key == "first_line_indent":
            if "hangingChars" in values:
                chars = -float(values["hangingChars"]) / 100
            elif "firstLineChars" in values:
                chars = float(values["firstLineChars"]) / 100
        elif key == "left_indent":
            chars = next((float(values[name]) / 100 for name in ("leftChars", "startChars")
                          if name in values), None)
        else:
            chars = next((float(values[name]) / 100 for name in ("rightChars", "endChars")
                          if name in values), None)
        if chars is None or char_size is None:
            result[key] = None
        else:
            result[key] = round(chars * char_size, 4)
    for key, prefix in (("space_before", "before"), ("space_after", "after")):
        auto = raw_attr("spacing", prefix + "Autospacing")
        if raw_attr("spacing", prefix + "Lines") is not None or auto not in (None, "0", "false", "off"):
            result[key] = None
    for prop in props:
        if prop is None:
            continue
        contextual = prop.find(qn("w:contextualSpacing"))
        if contextual is not None:
            if contextual.get(qn("w:val"), "1") not in ("0", "false", "off"):
                result["space_before"] = result["space_after"] = None
            break
    return result


def _theme_font(paragraph, token, east_asia):
    from lxml import etree
    try:
        relation = next(r for r in paragraph.part.rels.values() if r.reltype.endswith("/theme"))
        root = etree.fromstring(relation.target_part.blob)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        group = "majorFont" if token.startswith("major") else "minorFont"
        if east_asia:
            node = root.find(f".//a:{group}/a:ea", ns)
            if node is not None and node.get("typeface"):
                return node.get("typeface")
            node = root.find(f".//a:{group}/a:font[@script='Hans']", ns)
        else:
            node = root.find(f".//a:{group}/a:latin", ns)
        return node.get("typeface") if node is not None else None
    except Exception:
        return None


def effective_runs(paragraph):
    result = []
    # iter_inner_content includes hyperlinks without changing the document.
    runs = []
    for item in paragraph.iter_inner_content():
        if hasattr(item, "runs"):
            runs.extend(item.runs)
        else:
            runs.append(item)
    for run in runs:
        if not run.text:
            continue
        props = []
        if run._r.rPr is not None:
            props.append(run._r.rPr)
        for style in [*_styles(run.style), *_styles(paragraph.style)]:
            if style.element.rPr is not None:
                props.append(style.element.rPr)
        default = _default_prop(paragraph, "rPr")
        if default is not None:
            props.append(default)

        def attr(element, name="val"):
            for prop in props:
                node = prop.find(qn("w:" + element))
                if node is not None and node.get(qn("w:" + name)) is not None:
                    return node.get(qn("w:" + name))
            return None

        def font(name, theme):
            for prop in props:
                node = prop.find(qn("w:rFonts"))
                if node is None:
                    continue
                if node.get(qn("w:" + theme)):
                    return _theme_font(paragraph, node.get(qn("w:" + theme)), name == "eastAsia")
                if node.get(qn("w:" + name)):
                    return node.get(qn("w:" + name))
            return None

        size = attr("sz")
        # Bold is a toggle in style definitions, unlike its absolute on/off
        # meaning in direct run formatting (ISO 29500 toggle properties).
        # Apply each style base-to-derived in paragraph then character order.
        default_bold = default.find(qn("w:b")) if default is not None else None
        bold = (default_bold is not None and
                default_bold.get(qn("w:val"), "1") not in ("0", "false", "off"))
        for style in [*reversed(list(_styles(paragraph.style))), *reversed(list(_styles(run.style)))]:
            style_prop = style.element.rPr
            node = style_prop.find(qn("w:b")) if style_prop is not None else None
            if node is not None and node.get(qn("w:val"), "1") not in ("0", "false", "off"):
                bold = not bold
        direct = run._r.rPr
        direct_bold = direct.find(qn("w:b")) if direct is not None else None
        if direct_bold is not None:
            bold = direct_bold.get(qn("w:val"), "1") not in ("0", "false", "off")
        result.append((run.text, {"east_asia_font": font("eastAsia", "eastAsiaTheme"),
                                 "ascii_font": font("ascii", "asciiTheme"),
                                 "size_pt": float(size) / 2 if size else None,
                                 "bold": bold,
                                 "color": attr("color")}))
    return result


def describe_value(key, value):
    if value is None:
        return "无法读取有效值"
    if isinstance(value, bool):
        return "启用" if value else "未启用"
    if key == "alignment":
        return {0: "左对齐", 1: "居中", 2: "右对齐", 3: "两端对齐", 4: "分散对齐"}.get(value, str(value))
    if key == "line_spacing_rule":
        return {0: "单倍", 1: "1.5倍", 2: "双倍", 3: "最小值", 4: "固定值", 5: "多倍"}.get(value, str(value))
    if key in {"first_line_indent", "left_indent", "right_indent", "space_before", "space_after", "size_pt"}:
        return f"{value:g}磅"
    return str(value)
