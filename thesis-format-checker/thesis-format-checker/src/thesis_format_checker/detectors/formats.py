"""Read-only internal object consistency (139) and labelled special layout (164).

No majority value becomes a formatting standard. Figure pixels, printed colour,
folds and physical pagination remain outside python-docx's readable evidence.
140、160、161、162、163、165：本条依赖大模型语义，本阶段暂不实现。
参考文献引用判定：待开会讨论暂不实现。
"""
from __future__ import annotations

import json

from collections import defaultdict
from copy import deepcopy
import re

from thesis_format_checker.core.xmlutil import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from thesis_format_checker.checks.catalog import FINE_RULES
from thesis_format_checker.core.docx_context import UNREADABLE, effective_paragraph, make_finding, unknown, memo
from thesis_format_checker.detectors.pages import _read_section, _section_paragraphs
from thesis_format_checker.detectors.headings import collect_heading_records
from thesis_format_checker.detectors.text_runs import effective_text_runs
from thesis_format_checker.detectors.equations import _read_records, _font_samples, _layout
from thesis_format_checker.detectors.figures import collect_caption_records, _table_records
from thesis_format_checker.detectors.notes import _note_data, _has_revision

TITLES = {"139": "同类对象内部格式不一致", "164": "特殊版式设置不符合对应模板"}

#: 139 内部格式键 → 普通用户看得懂的中文名（不暴露 alignment / line_spacing 这类字段名）
_PROFILE_NAMES = {
    "alignment": "对齐方式", "line_spacing": "行距数值", "line_spacing_rule": "行距类型",
    "first_line_indent": "首行缩进", "left_indent": "左缩进", "right_indent": "右缩进",
    "space_before": "段前间距", "space_after": "段后间距",
    "keep_with_next": "与下段同页", "keep_together": "段中不分页",
    "paragraph_alignment": "段落对齐", "math_alignment": "公式对象对齐", "number_alignment": "编号对齐",
    "east_asia_font": "中文字体", "ascii_font": "英文字体", "size_pt": "字号",
    "bold": "加粗", "color": "颜色", "italic": "倾斜",
}
_ALIGNMENT_NAMES = {0: "左对齐", 1: "居中", 2: "右对齐", 3: "两端对齐", 4: "分散对齐"}
_ALIGNMENT_NAMES.update({"left": "左对齐", "center": "居中", "right": "右对齐",
                         "justify": "两端对齐", "distribute": "分散对齐"})
_LINE_RULE_NAMES = {0: "单倍行距", 1: "1.5 倍行距", 2: "双倍行距", 3: "最小值",
                    4: "固定值", 5: "多倍行距"}
_POINT_KEYS = {"first_line_indent", "left_indent", "right_indent", "space_before", "space_after", "size_pt"}


def _display_key(key: str) -> str:
    """把 ``段落/alignment``、``布局/number_alignment``、``数学/中文/bold`` 等键写成中文。"""

    parts = key.split("/")
    name = _PROFILE_NAMES.get(parts[-1], parts[-1])
    if key.startswith("布局/"):
        return name if name.startswith("公式") else "公式" + name
    if key.startswith("数学/"):
        role = "、".join(part for part in parts[1:-1] if part)
        return ("公式内文字（" + role + "）" + name) if role else ("公式内文字" + name)
    if key.startswith("单元格/"):
        role = "、".join(part for part in parts[1:-1] if part)
        return ("表格单元格（" + role + "）" + name) if role else ("表格单元格" + name)
    return name


def _display_profile_value(key: str, value) -> str:
    """把内部取值写成中文描述（避免 0/1、True/False 这类原始值）。"""

    name = key.split("/")[-1]
    if name in {"alignment", "paragraph_alignment", "math_alignment", "number_alignment"}:
        return _ALIGNMENT_NAMES.get(value, str(value))
    if name == "line_spacing_rule":
        return _LINE_RULE_NAMES.get(value, str(value))
    if isinstance(value, bool):
        return "是" if value else "否"
    if name in _POINT_KEYS and isinstance(value, (int, float)):
        return f"{value:g} 磅"
    return str(value)


def _formula_label(number, index: int, kind: str) -> str:
    """Use a stable object locator instead of truncating surrounding prose."""

    prefix = "行内公式" if kind == "inline" else "独立公式"
    value = (number or {}).get("value")
    if value:
        return f"{prefix}（编号 {value}，第 {index} 段）"
    return f"{prefix}（第 {index} 段）"
_PKEYS = ("alignment", "first_line_indent", "left_indent", "right_indent", "space_before", "space_after",
          "line_spacing", "line_spacing_rule", "keep_with_next", "keep_together")
_PAGE_KEYS = ("orientation", "page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin")
_LABEL = re.compile(r"^(?:特殊版式\s*[：:]\s*)?(超大表格|横向页面|折页|彩色页面)(?:\s*[：:]\s*(.{1,60}))?$")


def _notice(rule, reason, anchor=None):
    return unknown(rule, _title(rule), UNREADABLE + "；" + reason, anchor)


def _title(code: str) -> str:
    """细粒度规则标题取自规则目录；传入族编号时用本模块的族标题。"""

    item = FINE_RULES.get(code)
    return item.title if item else TITLES[code]


def _same_group_code(label) -> str:
    """同类对象 → 139.1 标题 / 139.2 图表 / 139.3 公式 / 139.4 注释。"""

    value = str(label)
    folded = value.casefold()
    if "标题" in value or "heading" in folded:
        return "139.1"
    if any(token in value for token in ("图", "表", "题注")) or "caption" in folded:
        return "139.2"
    if "公式" in value or "equation" in folded:
        return "139.3"
    if any(token in value for token in ("注释", "脚注", "尾注")) or "note" in folded:
        return "139.4"
    return "139"


def _single(values):
    values = list(values)
    return values[0] if values and values[0] is not None and all(v == values[0] for v in values) else None


def _language(text):
    return "zh" if re.search(r"[\u3400-\u9fff]", text) else "en" if re.search(r"[A-Za-z]", text) else "other"


def _revised_style(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        if _has_revision(style.element):
            return True
        style = style.base_style
    return False


def _text_profile(paragraph, start=0, end=None):
    groups = defaultdict(list)
    offset = 0
    for text, props in effective_text_runs(paragraph):
        stop = offset + len(text)
        fragment = text[max(0, start-offset):None if end is None else max(0, end-offset)] if stop > start else ""
        offset = stop
        if not fragment.strip():
            continue
        if props.get("revision") or props.get("field") or _revised_style(props["run"].style):
            raise ValueError("含修订或动态域，文字角色及最终显示格式未确认")
        for script, pattern, font in (("中文", r"[\u3400-\u9fff]", "east_asia_font"),
                                      ("英文数字", r"[A-Za-z0-9]", "ascii_font"),
                                      ("其他字符", r"[^\x00-\x7f\u3400-\u9fff\u3000-\u303f\uff00-\uffef\s]", "hansi_font")):
            if not re.search(pattern, fragment):
                continue
            role = (props["character_style"], props["vertical_align"], script)
            for name in (font, "size_pt", "bold", "color"):
                groups[role + (name,)].append(props.get(name))
    return {"文字/" + "/".join(role): _single(values) for role, values in groups.items()}


def _profile(paragraph, start=0, end=None):
    if _has_revision(paragraph._p) or _revised_style(paragraph.style):
        raise ValueError("段落存在未接受修订")
    profile = {"段落/" + k: v for k, v in effective_paragraph(paragraph).items() if k in _PKEYS}
    # A list level can override indents; do not compare a guessed paragraph-only value.
    styles, style, seen = [], paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        styles.append(style.element)
        style = style.base_style
    if any(n.xpath("./w:pPr/w:numPr") for n in [paragraph._p] + styles):
        for key in ("first_line_indent", "left_indent", "right_indent"):
            profile["段落/" + key] = None
    profile.update(_text_profile(paragraph, start, end))
    return profile


def _record(group, paragraph, profile, anchor=None, label=None, error=None):
    return {"group": group, "variant": paragraph.style.name if paragraph.style else "",
            "profile": profile, "anchor": anchor if anchor is not None else paragraph,
            "label": label or paragraph.text[:70], "error": error}


def __objects_impl(ctx, mapping=None):
    records, errors = [], []
    fallback = ctx.paragraphs[0] if ctx.paragraphs else None
    regions = {}
    try:
        headings = collect_heading_records(ctx, mapping)
        by_node = {r["paragraph"]._p: r for r in headings}
        region = "body"
        for paragraph in ctx.paragraphs:
            heading = by_node.get(paragraph._p)
            if heading:
                if re.match(r"^\s*(?:附录(?=\s|[A-Za-z０-９0-9一二三四五六七八九十：:]|$)|Appendix\b)", heading["text"], re.I):
                    region = "appendix"
                elif heading["role"] not in ("body", "appendix"):
                    region = heading["role"]
                elif heading["level"] == 1 and region != "appendix":
                    region = "body"
            regions[paragraph._p] = region
        for item in headings:
            p = item["paragraph"]
            group = ("标题", item["level"], item["role"], regions.get(p._p, "body"), _language(item["title"]))
            try:
                start = p.text.find(item["title"]) if item["title"] else 0
                records.append(_record(group, p, _profile(p, max(0, start))))
            except Exception as exc:
                records.append(_record(group, p, {}, error=str(exc)))
    except Exception as exc:
        errors.append((fallback, "标题对象读取失败：" + str(exc)))
    captions = []
    try:
        captions = collect_caption_records(ctx)
        for item in captions:
            p = item["paragraph"]
            group = ("题注", item["kind"], item["role"], regions.get(p._p, "body"), item["language"])
            try:
                profile = _profile(p, item["title_start"])
                profile.update({"编号/" + key: val for key, val in _text_profile(p, 0, item["number_end"]).items()})
                records.append(_record(group, p, profile))
            except Exception as exc:
                records.append(_record(group, p, {}, error=str(exc)))
            if item["kind"] == "table" and item["object"] is None and item["adjacent"]:
                records.append(_record(("表格", item["role"], regions.get(p._p, "body"), item["language"]),
                                       p, {}, error="题注前后存在多个候选表格，无法唯一关联同类表格"))
        for table in _table_records(ctx, captions):
            caption = table["caption"]
            p = caption["paragraph"]
            profile, samples = {}, defaultdict(list)
            error = table.get("error")
            if not error and (table["format_issue"] or table["border_issue"]):
                error = "表格合并、条件样式或边框冲突需要Word布局解析"
            if not error:
                # Do not demand equal row/column counts or content-driven heights.
                # Only internally uniform cell roles are candidates for comparison.
                for row in table["rows"]:
                    for cell in row["cells"]:
                        samples["单元格/" + row["role"] + "/shading"].append(cell["shading"])
                        samples["单元格/" + row["role"] + "/vertical"].append(cell["vertical"])
                        for paragraph, fmt in cell["paragraphs"]:
                            if not paragraph.text.strip():
                                continue
                            style_name = paragraph.style.name if paragraph.style else ""
                            for key, value in fmt.items():
                                value = _single(value) if isinstance(value, set) else value
                                samples["单元格/" + row["role"] + "/" + style_name + "/" + key].append(value)
                profile = {key: _single(values) for key, values in samples.items()}
            rec = _record(("表格", caption["role"], regions.get(p._p, "body"), caption["language"]), p, profile, error=error)
            rec["variant"] = table.get("style_name", "")
            records.append(rec)
    except Exception as exc:
        errors.append((fallback, "图表题注或表格读取失败：" + str(exc)))
    try:
        equations, issues = _read_records(ctx)
        errors.extend((p, "公式对象读取失败：" + reason) for p, reason in issues)
        for item in equations:
            p = item["paragraph"]
            group = ("公式", regions.get(p._p, item["region"]), item["kind"])
            try:
                if item["opaque"] or item["revision"] or item["dynamic"]:
                    raise ValueError("公式为OLE/图片、含修订或动态域，不能确认有效内部格式")
                profile = {}
                if item["kind"] != "inline":
                    # Inline equations inherit the surrounding prose layout;
                    # paragraph alignment/indent/spacing are not properties of
                    # the formula and must not be compared across prose roles.
                    profile.update({"段落/" + k: v for k, v in effective_paragraph(p).items()
                                    if k in _PKEYS and k != "alignment"})
                    profile.update({"布局/" + k: v for k, v in _layout(item, ctx).items()
                                    if k in ("math_alignment", "paragraph_alignment", "number_alignment")})
                samples, issues = _font_samples(item, ctx)
                groups = defaultdict(list)
                for role, text, attrs in samples:
                    for key, value in attrs.items():
                        groups["数学/" + "/".join(map(str, role)) + "/" + key].append(value)
                profile.update({key: _single(values) for key, values in groups.items()})
                records.append(_record(group, p, profile,
                                       label=_formula_label(item.get("number"), ctx.index(p) + 1,
                                                            item["kind"])))
                if issues:
                    errors.append((p, "数学自动符号/缩放：" + "；".join(dict.fromkeys(issues))))
            except Exception as exc:
                records.append(_record(group, p, {}, error=str(exc)))
    except Exception as exc:
        errors.append((fallback, "公式对象读取失败：" + str(exc)))
    try:
        notes = _note_data(ctx)
        errors.extend((fallback, kind + "：" + reason) for kind, reason in notes["errors"])
        for note in notes["records"]:
            if note["type"] != "normal":
                continue
            for index, element in enumerate(note["element"].iter(qn("w:p"))):
                p = Paragraph(deepcopy(element), ctx.document._body)
                group = ("注释", note["kind"], "首段" if index == 0 else "续段", _language(p.text))
                try:
                    if any(a.tag == qn("w:tbl") for a in element.iterancestors()) or element.xpath(".//w:object|.//w:drawing|.//w:txbxContent"):
                        raise ValueError("注释含表格、文本框或嵌入对象")
                    records.append(_record(group, p, _profile(p), note["anchor"], f'{note["kind"]} {note["id"]}：{p.text[:40]}'))
                except Exception as exc:
                    records.append(_record(group, p, {}, note["anchor"], error=str(exc)))
    except Exception as exc:
        errors.append((fallback, "脚注/尾注读取失败：" + str(exc)))
    return records, errors


def _special_label(paragraph):
    # Only a standalone purpose label is accepted, never a prose mention.
    # Read visible cached text as evidence of a *candidate* only; the caller
    # rejects fields/textboxes/objects before it may become a layout baseline.
    composite = paragraph._p.xpath(".//w:fldChar|.//w:fldSimple|.//w:instrText|.//w:object|.//w:drawing|.//w:pict|.//w:txbxContent")
    text = "".join(t.text or "" for t in paragraph._p.iter(qn("w:t"))) if composite else paragraph.text
    match = _LABEL.fullmatch(text.strip())
    return (match[1], (match[2] or "").strip()) if match else None


def _table_geometry(paragraph, ctx):
    node = paragraph._p.getnext()
    while node is not None and node.tag == qn("w:p") and not "".join(node.itertext()).strip():
        node = node.getnext()
    if node is None or node.tag != qn("w:tbl"):
        raise ValueError("超大表格标签后未紧邻唯一原生Word表格")
    table = Table(node, ctx.document._body)
    if _has_revision(node) or node.xpath(".//w:tblpPr|.//w:gridSpan|.//w:vMerge|.//w:tbl//w:tbl"):
        raise ValueError("浮动/合并/嵌套/修订表格的实际几何布局无法确认")
    grid = node.find(qn("w:tblGrid"))
    if grid is None or not len(grid):
        raise ValueError("表格缺少网格宽度")
    widths = tuple(int(n.get(qn("w:w"))) for n in grid if n.tag == qn("w:gridCol"))
    if not widths or any(v <= 0 for v in widths):
        raise ValueError("表格网格宽度无效")
    if table.autofit:
        raise ValueError("自动适应表格的实际列宽必须由Word重新排版")
    for row in table.rows:
        if len(row.cells) != len(widths):
            raise ValueError("表格行单元格数与网格列数不一致")
        for cell, grid_width in zip(row.cells, widths):
            width = cell._tc.tcPr.find(qn("w:tcW")) if cell._tc.tcPr is not None else None
            if width is not None and (width.get(qn("w:type")) != "dxa" or int(width.get(qn("w:w"))) != grid_width):
                raise ValueError("单元格首选宽度与网格不一致，需要Word布局解析")
    # Widths are the declared grid, not proof of physical rendering/pagination.
    return {"grid_twips": widths, "fixed_layout": True}


def _special_sections(ctx):
    records, errors = [], []
    groups = _section_paragraphs(ctx)
    try:
        sections = list(ctx.document.sections)
    except Exception as exc:
        return [], [(ctx.paragraphs[0] if ctx.paragraphs else None, "特殊节无法读取：" + str(exc))]
    for index, section in enumerate(sections):
        paragraphs = groups.get(index, [])
        labels = [(p, _special_label(p)) for p in paragraphs if _special_label(p)]
        anchor = next((p for p in paragraphs if p.text.strip()), paragraphs[0] if paragraphs else None)
        if not labels:
            try:
                landscape = int(section.orientation) == 1
            except Exception as exc:
                records.append({"purpose": None, "anchor": anchor, "index": index, "profile": {},
                                "error": "无明确用途标签的节方向读取失败：" + type(exc).__name__})
                continue
            if landscape:
                records.append({"purpose": None, "anchor": anchor, "index": index, "profile": {}, "error": "横向节未提供可与模板明确关联的特殊版式标签"})
            continue
        for p, purpose in labels:
            record = {"purpose": purpose, "anchor": p, "index": index, "profile": {}, "error": None, "geometry_error": None}
            try:
                if len({label for _, label in labels}) != 1 or _has_revision(section._sectPr) or _has_revision(p._p):
                    raise ValueError("同节存在多个特殊用途或未接受的设置修订")
                if p._p.xpath(".//w:fldChar|.//w:fldSimple|.//w:instrText|.//w:object|.//w:drawing|.//w:pict|.//w:txbxContent"):
                    raise ValueError("特殊用途标签含动态域、图片、文本框或嵌入对象，不能以缓存确认用途")
                read = _read_section(ctx, sections, index, paragraphs)
                for key in _PAGE_KEYS:
                    value = read.get(key) if key in read else getattr(section, key)
                    record["profile"][key] = int(value) if value is not None else None
                if purpose[0] == "超大表格":
                    try:
                        record["geometry"] = _table_geometry(p, ctx)
                    except Exception as exc:
                        record["geometry_error"] = str(exc)
            except Exception as exc:
                record["error"] = str(exc)
            records.append(record)
    return records, errors


def extract_format_consistency_baseline(ctx, mapping=None):
    """Append-only namespace; template evidence is never majority-selected."""
    records, errors = _objects(ctx, mapping)
    special, special_errors = _special_sections(ctx)
    return {"formats": {"objects": records, "errors": errors,
                                "special": special, "special_errors": special_errors}}


def _check_consistency(base, ctx, mapping=None):
    records, errors = _objects(ctx, mapping)
    findings = [_notice("139", reason, p) for p, reason in errors]
    if base.get("errors"):
        findings.append(_notice("139", "模板对象集合不完整：" + "；".join(reason for _, reason in base["errors"]),
                                ctx.paragraphs[0] if ctx.paragraphs else None))
    actual_groups, template_groups = defaultdict(list), defaultdict(list)
    for record in records:
        actual_groups[record["group"]].append(record)
    for record in base.get("objects", []):
        template_groups[record["group"]].append(record)
    for group, objects in actual_groups.items():
        if len(objects) < 2:
            # A single OLE formula is still not declared checked successfully.
            if objects[0]["error"]:
                findings.append(_notice("139", objects[0]["error"], objects[0]["anchor"]))
            continue
        templates = template_groups[group]
        error_words = {"标题": ("标题",), "题注": ("图表",), "表格": ("图表",),
                       "公式": ("公式", "数学"), "注释": ("注释", "footnote", "endnote", "脚注", "尾注")}
        template_incomplete = any(any(word in reason for word in error_words.get(group[0], ()))
                                  for _, reason in base.get("errors", []))
        variants = {r["variant"] for r in templates}
        subgroups = defaultdict(list)
        for obj in objects:
            subgroups[obj["variant"] if len(variants) > 1 else None].append(obj)
        for variant, same_objects in subgroups.items():
            if variant is not None and variant not in variants:
                findings.append(_notice("139", "对象样式无法对应模板允许的同用途变体", same_objects[0]["anchor"]))
                continue
            for obj in same_objects:
                if obj["error"]:
                    findings.append(_notice("139", obj["error"], obj["anchor"]))
            if len(same_objects) < 2:
                continue
            examples = [r for r in templates if variant is None or r["variant"] == variant]
            good = [r for r in same_objects if not r["error"]]
            if len(good) < 2:
                continue
            keys = set.intersection(*(set(r["profile"]) for r in good))
            uncertain, differences = [], defaultdict(list)
            for key in sorted(keys):
                observed = [r["profile"][key] for r in good]
                if any(value is None for value in observed):
                    uncertain.append(key)
                    continue
                if all(value == observed[0] for value in observed):
                    continue
                expected = [r["profile"].get(key) for r in examples]
                if template_incomplete or any(r["error"] for r in examples) or expected and (_single(expected) is None):
                    uncertain.append(key + "（模板同用途样例冲突或有效值不明）")
                    continue
                for index, obj in enumerate(good[1:], 1):
                    if observed[index] != observed[0]:
                        differences[index].append((key, observed[0], observed[index]))
            for index, changes in differences.items():
                original, obj = good[0], good[index]
                expected_values = list(dict.fromkeys(
                    f"{_display_key(k)}为{_display_profile_value(k, a)}" for k, a, _ in changes))
                actual_values = list(dict.fromkeys(
                    f"{_display_key(k)}为{_display_profile_value(k, b)}" for k, _, b in changes))
                code = _same_group_code(obj.get("label", ""))
                if code == "139.3":
                    expected = ("规范模板未提供可对应的公式样例；同类公式格式应保持一致。"
                                f"请与{original['label']}核对（本工具不判断哪一处正确）；"
                                + "；".join(expected_values))
                    actual = obj["label"] + "；" + "；".join(actual_values)
                else:
                    expected = ("文档内部同类对象格式应一致；与“" + original["label"] + "”相互核对，"
                                "不以多数或先后位置判定哪处正确；" + "；".join(expected_values))
                    actual = obj["label"] + "；" + "；".join(actual_values)
                findings.append(make_finding(code, _title(code), expected, actual, obj["anchor"]))
            if uncertain:
                findings.append(_notice("139", "以下同角色属性不能可靠互比：" + "、".join(uncertain), good[-1]["anchor"]))
    return findings


def _check_special(base, ctx, mapping=None):
    records, errors = _special_sections(ctx)
    findings = [_notice("164", reason, p) for p, reason in errors]
    if base.get("special_errors"):
        findings.append(_notice("164", "模板特殊节读取不完整：" + "；".join(reason for _, reason in base["special_errors"]),
                                ctx.paragraphs[0] if ctx.paragraphs else None))
    for obj in records:
        anchor, purpose = obj["anchor"], obj["purpose"]
        if obj["error"]:
            findings.append(_notice("164", obj["error"], anchor))
            continue
        candidates = [r for r in base.get("special", []) if r["purpose"] == purpose]
        section_map = (mapping or {}).get("section_map", {})
        mapped = section_map.get(str(obj["index"] + 1), section_map.get(obj["index"] + 1)) if isinstance(section_map, dict) else None
        if mapped is not None:
            try:
                candidates = [r for r in candidates if r["index"] == int(mapped) - 1]
            except (TypeError, ValueError):
                candidates = []
        if not candidates:
            findings.append(_notice("164", "模板中未找到同名同用途特殊版式样例：" + str(purpose), anchor))
            continue
        changes, uncertain = [], []
        for key, value in obj["profile"].items():
            values = [r["profile"].get(key) if not r["error"] else None for r in candidates]
            expected = _single(values) if not base.get("special_errors") else None
            if expected is None or value is None:
                uncertain.append(key + "（缺失/冲突）")
            elif expected != value:
                changes.append((key, expected, value))
        if purpose[0] == "超大表格":
            expected = _single([r.get("geometry") if not r.get("geometry_error") and not r["error"] else None for r in candidates])
            actual = obj.get("geometry")
            if expected is None or actual is None:
                uncertain.append("表格网格/固定布局：" + (obj.get("geometry_error") or "模板样例无法确定或相互冲突"))
            elif len(expected["grid_twips"]) != len(actual["grid_twips"]):
                uncertain.append("表格列数不同，未提供逐列用途对应，不能强制列宽相同")
            elif expected != actual:
                changes.append(("同用途固定表格网格", expected, actual))
            uncertain.append("超大表格真实跨页、缩放及打印裁切需Word排版或人工核对")
        if purpose[0] in ("折页", "彩色页面"):
            uncertain.append("实体折页/印刷颜色及打印机输出不能从docx页面设置确认")
        if changes:
            findings.append(make_finding("164", TITLES["164"], str(purpose) + "；" + "；".join(f"{k}={a}" for k, a, _ in changes),
                                         "；".join(f"{k}={b}" for k, _, b in changes), anchor))
        if uncertain:
            findings.append(_notice("164", "；".join(uncertain), anchor))
    return findings


def check_format_consistency_errors(baseline, ctx, mapping=None):
    base = baseline.get("formats", {})
    findings = []
    for rule, check in (("139", _check_consistency), ("164", _check_special)):
        try:
            findings.extend(check(base, ctx, mapping))
        except Exception as exc:
            findings.append(_notice(rule, "读取局部对象失败：" + type(exc).__name__, ctx.paragraphs[0] if ctx.paragraphs else None))
    return findings

def _objects(ctx, mapping=None):
    """带缓存的入口：同一份文档只扫描一次（键包含映射配置）。"""

    fingerprint = json.dumps(mapping or {}, sort_keys=True, ensure_ascii=False, default=str)
    return memo(ctx, "format_objects:" + fingerprint, lambda: __objects_impl(ctx, mapping))
