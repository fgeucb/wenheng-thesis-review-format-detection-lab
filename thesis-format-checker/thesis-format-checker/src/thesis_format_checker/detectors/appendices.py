"""Read-only appendix checks, 155--159; no displayed field is refreshed.

160、161、162、163、165：本条依赖大模型语义，本阶段暂不实现。
159中的参考资料/数据来源语义识别及参考文献引用判定：
【待开会讨论，本版本暂不开发】。不以正文实际错误替代模板基准。
"""
from __future__ import annotations

from collections import defaultdict
import re
from types import SimpleNamespace

from thesis_format_checker.core.xmlutil import qn
from docx.table import Table

from thesis_format_checker.checks.catalog import FINE_RULES
from thesis_format_checker.core.docx_context import UNREADABLE, effective_paragraph, make_finding, unknown, _styles
from thesis_format_checker.detectors.fields import scan_fields
from thesis_format_checker.detectors.headings import _heading_level, _toc_paragraph_ids, _num_pr
from thesis_format_checker.detectors.figures import (collect_caption_records, _visible, _revised, _caption_format,
                             _text_signature, _table_snapshot, _table_checks)
from thesis_format_checker.detectors.equations import (_read_records, _layout, _font_samples, _check_fonts,
                              _check_layout, _check_spacing)
from thesis_format_checker.core.models import Status

TITLES = {"155": "附录与正文引用不对应", "156": "附录编号或连续性不符合模板",
          "157": "附录标题格式不符合模板", "158": "附录图表公式编号不符合模板",
          "159": "附录对象格式与模板正文共同规则冲突"}
_TITLE = re.compile(r"^(?P<label>附录|Appendix)(?P<gap>[ \t\u3000]*)(?P<token>[A-Za-z]+|[0-9]+)(?P<tail>(?:[ \t\u3000:：.．、—\-]+.*)?)$", re.I)
_LABEL = re.compile(r"(?:附录|Appendix)[ \t\u3000]*(?P<token>[A-Za-z]+|[0-9]+)(?![A-Za-z0-9])", re.I)
_CAPTION = re.compile(r"^(?P<label>图|表|Figure|Table)(?P<gap>[ \t\u3000]*)(?P<number>(?:[A-Za-z]+|[0-9]+)(?:[.．\-－][0-9]+)+|[0-9]+)(?P<end>[ \t\u3000]*)(?P<title>.*)$", re.I)
_FORMULA = re.compile(r"(?P<label>(?:式|Equation|Eq\.?)[ \t]*)?(?P<open>[（(])(?P<left>[ \t]*)(?P<number>(?:[A-Za-z]+|[0-9]+)(?:[.．\-－][0-9]+)+|[0-9]+)(?P<right>[ \t]*)(?P<close>[）)])", re.I)
_PARA = ("alignment", "space_before", "space_after")


def _notice(rule, reason, paragraph=None):
    return unknown(rule, _title(rule), UNREADABLE + "；" + reason, paragraph)


def _emit(rule, expected, actual, paragraph):
    return make_finding(rule, _title(rule), expected, actual, paragraph)


def _title(code: str) -> str:
    """细粒度规则标题取自规则目录；传入族编号时用本模块的族标题。"""

    item = FINE_RULES.get(code)
    return item.title if item else TITLES[code]


def _token(token):
    if token.isascii() and token.isdigit():
        return ("decimal", int(token), "padded" if len(token) > 1 and token.startswith("0") else "plain")
    if token.isascii() and token.isalpha() and (token.isupper() or token.islower()):
        value = 0
        for letter in token.upper():
            value = value * 26 + ord(letter) - ord("A") + 1
        return ("upper" if token.isupper() else "lower", value, "plain")
    return None


def _revised_style(p):
    return _revised(p._p) or any(_revised(style.element) for style in _styles(p.style))


def _unsafe(p):
    return (_revised_style(p) or bool(p._p.xpath(".//w:fldSimple|.//w:fldChar|.//w:instrText|.//w:drawing|.//w:pict|.//w:object"))
            or _num_pr(p) is not None)


def _title_info(p, role):
    text = _visible(p._p).strip()
    match = _TITLE.fullmatch(text)
    number = _token(match["token"]) if match else None
    # Label, spaces, numbering alphabet and delimiter are template evidence.
    delimiter = re.match(r"[ \t\u3000:：.．、—\-]*", match["tail"])[0] if match else None
    return {"paragraph": p, "text": text, "role": role,
            "token": match["token"] if match else None, "number": number,
            "shape": (match["label"], match["gap"], number[0], number[2], delimiter) if number else None,
            "unsafe": _unsafe(p)}


def collect_appendix_inventory(ctx, mapping=None):
    """Identify explicit appendix headings, excluding TOC and deleted text.

    Numbered appendix titles must have a heading/appendix style or consist
    solely of their label. A normal sentence beginning with 附录 is not a title.
    """
    # Existing heading mapping is reserved by the shared heading collector;
    # do not invent a new mapping syntax or change its public contract here.
    del mapping
    toc = _toc_paragraph_ids(ctx)
    titles, owners, body, problems = [], {}, [], []
    current, appendix_seen, excluded_region = None, False, False
    for p in ctx.paragraphs:
        if id(p._p) in toc or p._p in toc:
            continue
        if any(a.tag in {qn("w:del"), qn("w:moveFrom")} for a in p._p.iterancestors()):
            continue
        text = _visible(p._p).strip()
        level = _heading_level(p)
        if level == 1:
            excluded_region = re.sub(r"\s+", "", text).casefold() in {
                "参考文献", "references", "bibliography", "致谢", "acknowledgements", "acknowledgments",
                "摘要", "abstract", "学位论文原创性声明", "学位论文版权使用授权书"}
        appendix_style = any(re.search(r"appendix|附录", s.name or "", re.I) for s in _styles(p.style))
        match = _TITLE.fullmatch(text)
        short = bool(match and not match["tail"].strip())
        total = text.casefold() in {"附录", "appendix", "appendices"}
        # A styled heading with an appendix prefix remains an appendix
        # candidate even if its number/title boundary is malformed. Dropping
        # it would incorrectly prove a cited appendix absent (e.g. 附录C数据).
        prefix_candidate = bool(level and re.match(r"^(?:附录|Appendix\b)", text, re.I))
        is_appendix = bool((match and (level or appendix_style or short)) or total or appendix_style or prefix_candidate)
        if is_appendix:
            role = "total" if total else "section"
            item = _title_info(p, role)
            if role == "section" and not item["number"]:
                item["unsafe"] = True
            if any(a.tag == qn("w:tbl") for a in p._p.iterancestors()):
                item["unsafe"] = True
                problems.append((p, "附录标题位于表格中，附录边界不能可靠确定"))
            titles.append(item)
            appendix_seen = True
            if role == "section":
                current = item
            else:
                current = None
            owners[p._p] = current
            continue
        if current is not None and level == 1:
            # A subsequent unrelated top-level part terminates the appendix.
            current = None
        if current is not None and level and level > 1:
            item = _title_info(p, "sub" + str(level))
            titles.append(item)
        owners[p._p] = current
        if current is None and not appendix_seen and not excluded_region:
            body.append(p)
    if ctx.has_textboxes:
        problems.append((titles[0]["paragraph"] if titles else (ctx.paragraphs[0] if ctx.paragraphs else None),
                         "文本框不属于可完整枚举的正文/附录清单"))
    if ctx.document.element.body.xpath(".//w:altChunk"):
        problems.append((ctx.paragraphs[0] if ctx.paragraphs else None, "外部导入内容无法完整枚举"))
    if any(t["unsafe"] or (t["role"] == "section" and not t["number"]) for t in titles):
        problems.append((titles[0]["paragraph"], "附录标题包含自动编号、域、修订、对象或未支持的编号体系"))
    return {"titles": titles, "owners": owners, "body": body, "problems": problems}


def _number_info(match, token=None):
    if match is None:
        return None
    number = match["number"]
    parts = re.split(r"[.．\-－]", number)
    prefix = parts[0]
    independent = token is not None and len(parts) == 2 and prefix == token
    # Raw separators and width are retained; no source normalization occurs.
    if independent:
        shape = "{appendix}" + re.sub(r"[0-9]+", "#", number[len(prefix):])
    else:
        shape = re.sub(r"[A-Za-z]+|[0-9]+", "#", number)
    if "open" in match.re.groupindex:
        shape = (match["label"] or "") + match["open"] + match["left"] + shape + match["right"] + match["close"]
    else:
        shape = match["label"] + match["gap"] + shape + match["end"]
    return {"text": match[0], "number": number, "prefix": prefix, "shape": shape,
            "independent": independent, "parts": parts,
            "sequence": int(parts[-1]) if parts[-1].isdigit() else None}


def _objects(ctx, inventory):
    records, problems = [], []
    body_nodes = {p._p for p in inventory["body"]}
    for source in collect_caption_records(ctx):
        record = dict(source)
        p = record["paragraph"]
        owner = inventory["owners"].get(p._p)
        if owner is None and p._p not in body_nodes:
            continue
        match = _CAPTION.fullmatch(_visible(p._p).strip())
        record.update(owner=owner, number_info=_number_info(match, owner["token"] if owner else None),
                      type="caption", unsafe=record["dynamic"] or record["revision"])
        if match:
            record.update(number_end=match.end("number"), title_start=match.start("title"))
        records.append(record)
    formulas, errors = _read_records(ctx)
    problems.extend(errors)
    known_formula_nodes = {r["paragraph"]._p for r in formulas}
    for p in ctx.paragraphs:
        if p._p in known_formula_nodes:
            continue
        if p._p.xpath(".//w:drawing|.//w:pict|.//w:object") and _FORMULA.fullmatch(_visible(p._p).strip()):
            formulas.append({"paragraph": p, "number": None, "number_paragraph": None, "kind": "display",
                             "opaque": True, "dynamic": False, "revision": _revised_style(p),
                             "number_unknown": False, "maths": []})
    for source in formulas:
        record = dict(source)
        p = record["paragraph"]
        owner = inventory["owners"].get(p._p)
        if owner is None and p._p not in body_nodes:
            continue
        candidates = [(p, m) for m in _FORMULA.finditer(_visible(p._p))]
        cell = next((a for a in p._p.iterancestors() if a.tag == qn("w:tc")), None)
        if not candidates and cell is not None:
            for other in ctx.paragraphs:
                other_cell = next((a for a in other._p.iterancestors() if a.tag == qn("w:tc")), None)
                if other_cell is not None and other_cell is not cell and other_cell.getparent() is cell.getparent():
                    match = _FORMULA.fullmatch(_visible(other._p).strip())
                    if match:
                        candidates.append((other, match))
        number_p, match = candidates[0] if len(candidates) == 1 else (None, None)
        number = _number_info(match, owner["token"] if owner else None)
        record.update(owner=owner, number_info=number, type="formula", number_paragraph=number_p,
                      unsafe=record["opaque"] or record["dynamic"] or record["revision"] or record["number_unknown"])
        if number_p is not None:
            record["unsafe"] = record["unsafe"] or _unsafe(number_p)
            record["number"] = {"value": number["number"], "pattern": number["shape"]}
        if cell is not None:
            record["kind"] = "table"
        elif number_p is not None and record["opaque"]:
            record["kind"] = "display"
        records.append(record)
    return records, problems


def _title_format(record):
    p = record["paragraph"]
    signature = _text_signature(p)
    # 157 enumerates font, size, alignment and number syntax. Bold/color are
    # not added as independent appendix-title errors under this identifier.
    return {"alignment": effective_paragraph(p)["alignment"],
            "text": {key: value for key, value in signature.items()
                     if key in {"east_asia_font", "ascii_font", "size_pt"}}}


def _inheritance(text):
    result, denied = set(), set()
    for sentence in re.split(r"[。；;\n]", text):
        if not re.search(r"附录.{0,60}(?:格式|排版).{0,30}正文", sentence):
            continue
        kinds = set()
        if re.search(r"图表|图|插图", sentence):
            kinds.add("figure")
        if re.search(r"图表|表格|表", sentence):
            kinds.add("table")
        if "公式" in sentence:
            kinds.add("formula")
        if re.search(r"不必|无需|不要求|不需要|不应|不得|不要|不能|不再|不与|不同于|不一致|不相同|"
                     r"除外|例外|除.{0,30}外|可不同|可以不|允许不", sentence):
            denied.update(kinds)
        elif re.search(r"附录.{0,60}(?:格式|排版).{0,10}(?:与|同).{0,4}正文.{0,10}(?:一致|相同)", sentence):
            result.update(kinds)
    # A later contradictory declaration is not overridden by a convenient
    # earlier positive sample. This is narrow literal policy extraction.
    return result - denied


def _sequence_declarations(text):
    positive, denied = False, False
    for sentence in re.split(r"[。；;\n]", text):
        if "附录" not in sentence or not re.search(r"连续编号|顺序编号|不连续|允许跳号", sentence):
            continue
        if re.search(r"不必|无需|不要求|不需要|不应|不得|不要|不能|不再|不连续|允许跳号|可以不|允许不", sentence):
            denied = True
        elif re.search(r"附录.{0,20}(?:连续编号|顺序编号)", sentence):
            positive = True
    return positive and not denied, denied


def extract_appendix_baseline(ctx, mapping=None):
    inventory = collect_appendix_inventory(ctx, mapping)
    sequence_explicit, sequence_denied = _sequence_declarations(ctx.text)
    records, problems = _objects(ctx, inventory)
    for title in inventory["titles"]:
        try:
            title["format"] = _title_format(title)
        except Exception as exc:
            title["format_error"] = type(exc).__name__
    for record in records:
        try:
            if record["unsafe"]:
                raise ValueError("模板对象含修订、域或不可读内部属性")
            if record["type"] == "caption":
                record["format"] = _caption_format(record)
                if record["kind"] == "table" and record["object"] is not None:
                    record["table_snapshot"] = _table_snapshot(Table(record["object"], ctx.document._body), record)
            else:
                record["layout"] = _layout(record, ctx)
                record["spacing"] = {k: effective_paragraph(record["paragraph"])[k]
                                     for k in ("space_before", "space_after", "line_spacing", "line_spacing_rule")}
                record["fonts"], record["font_errors"] = _font_samples(record, ctx)
        except Exception as exc:
            record["error"] = str(exc)
    return {"appendices": {"inventory": inventory, "records": records, "problems": problems,
                                  "inheritance": _inheritance(ctx.text),
                                  "sequence_explicit": sequence_explicit, "sequence_denied": sequence_denied,
                                  "number_required": bool(re.search(r"附录.{0,24}(?:图表|图、表|图和表).{0,10}公式.{0,12}(?:必须|均应|均须).{0,10}(?:独立编号|编号)", ctx.text)),
                                  "reference_optional": bool(re.search(r"附录.{0,12}(?:可以|可|允许).{0,8}(?:不必|不在|无需).{0,8}(?:正文|引用)", ctx.text))}}


def _reference_checks(base, ctx, inventory):
    findings, cited, ambiguous = [], set(), False
    targets = defaultdict(list)
    for title in inventory["titles"]:
        if title["role"] == "section" and title["token"] and not title["unsafe"]:
            targets[title["token"]].append(title)
    opaque = bool(inventory["problems"]) or any(
        p._p.xpath(".//w:drawing|.//w:pict|.//w:object") or _revised_style(p) for p in ctx.paragraphs)
    for p in inventory["body"]:
        text = _visible(p._p)
        if p._p.xpath(".//w:drawing|.//w:pict|.//w:object") or _revised_style(p):
            opaque = True
        matches = list(_LABEL.finditer(text))
        fields = scan_fields(p._p)
        if matches and (fields or fields.errors or _unsafe(p)):
            ambiguous = True
            findings.append(_notice("155", "正文附录引用含域、自动编号、修订或对象，缓存不代表当前引用目标", p))
            continue
        for match in matches:
            lead = text[max(0, match.start() - 30):match.start()]
            tail = text[match.end():]
            if re.search(r"文献|外部|引自|他人|其他论文|另一篇|附件", lead) and not re.search(r"本文|本论文|本研究", lead):
                continue
            if re.match(r"\s*(?:[-—~～至、,，]|和)\s*(?:附录\s*)?[A-Za-z0-9]", tail):
                ambiguous = True
                findings.append(_notice("155", "范围或合并附录引用无法唯一拆分", p))
                continue
            if not re.search(r"本文|本论文|本研究|参见|见|详见|见于", lead):
                ambiguous = True
                findings.append(_notice("155", "未明确指向本文的附录引用，不能排除外部材料", p))
                continue
            token = match["token"]
            cited.add(token)
            if len(targets[token]) == 1:
                continue
            if not targets[token] and not opaque:
                findings.append(_emit("155.1", "正文明确引用的附录应有唯一文末目标", "引用“" + match[0] + "”，未找到对应附录标题", p))
            else:
                findings.append(_notice("155", "引用目标存在重号或标题清单不完整，无法确认是否缺失", p))
        # A REF without a readable label may represent an appendix citation.
        if any(f["kind"] in {"REF", "PAGEREF"} for f in fields) and not matches:
            ambiguous = True
    if not base.get("reference_optional"):
        for token, titles in targets.items():
            if token in cited:
                continue
            for title in titles:
                if opaque or ambiguous:
                    findings.append(_notice("155", "正文引用清单含不可读或指向不明内容，不能断言该附录未被引用", title["paragraph"]))
                else:
                    findings.append(_emit("155.2", "文末附录应在正文出现可确认的引用", "未发现正文引用“" + title["text"] + "”", title["paragraph"]))
    for p, reason in inventory["problems"]:
        findings.append(_notice("155", reason, p))
    return findings


def _compare_tree(expected, actual, path=""):
    differences, limits = [], []
    if isinstance(actual, dict):
        for key, value in actual.items():
            if key not in expected:
                limits.append(path + str(key) + "缺少模板对应属性")
                continue
            bad, unknowns = _compare_tree(expected[key], value, path + str(key) + "/")
            differences.extend(bad)
            limits.extend(unknowns)
    elif isinstance(actual, set):
        if not isinstance(expected, set) or len(expected) != 1 or None in expected or None in actual:
            limits.append(path + "有效值不唯一或不可读")
        elif actual != expected:
            differences.append((path.rstrip("/"), str(expected), str(actual)))
    elif expected is None or actual is None:
        limits.append(path + "有效值不可读")
    elif expected != actual:
        differences.append((path.rstrip("/"), str(expected), str(actual)))
    return differences, limits


def _compare_format(rule, candidates, actual, paragraph):
    if not candidates or any(value != candidates[0] for value in candidates):
        return [_notice(rule, "模板对应角色格式缺失或不唯一", paragraph)]
    differences, limits = _compare_tree(candidates[0], actual)
    findings = []
    if differences:
        findings.append(_emit(rule, "；".join(k + "=" + e for k, e, _ in differences),
                              "；".join(k + "=" + a for k, _, a in differences), paragraph))
    if limits:
        findings.append(_notice(rule, "；".join(sorted(set(limits))), paragraph))
    return findings


def _title_checks(base, inventory):
    findings = []
    template = base.get("inventory", {}).get("titles", [])
    sections = [t for t in template if t["role"] == "section"]
    numbers = [t["number"] for t in sections if t["number"] and not t["unsafe"]]
    sequence = (len(numbers) == len(sections) and len(numbers) >= 2 and
                len({(n[0], n[2]) for n in numbers}) == 1 and
                all(b[1] == a[1] + 1 for a, b in zip(numbers, numbers[1:])))
    sequence = not base.get("sequence_denied") and (sequence or (base.get("sequence_explicit") and len(numbers) == len(sections) and bool(numbers)))
    previous = None
    for title in inventory["titles"]:
        p, role = title["paragraph"], title["role"]
        candidates = [t for t in template if t["role"] == role]
        if title["unsafe"] or any(t["unsafe"] or t.get("format_error") for t in candidates):
            findings.append(_notice("157", "标题“" + title["text"] + "”的编号边界无法解析，或含域、自动列表、未决修订、不可读有效格式", p))
            if role == "section":
                findings.append(_notice("156", "附录标题编号不可读，未由域缓存推定序号", p))
                previous = None
            continue
        findings.extend(_compare_format("157", [t["format"] for t in candidates if "format" in t], _title_format(title), p))
        if role != "section":
            continue
        shapes = {t["shape"] for t in candidates}
        if len(shapes) != 1 or None in shapes or title["shape"] is None:
            findings.extend(_notice(r, "附录编号格式基准缺失、不唯一或使用未支持编号体系", p) for r in ("156", "157"))
        elif title["shape"] not in shapes:
            # 编号形式不符：156.2（附录编号格式错误）/ 157.4（附录标题编号格式）
            findings.append(_emit("156.2", "附录编号形式=" + str(next(iter(shapes))), "当前编号形式=" + str(title["shape"]), p))
            findings.append(_emit("157.4", "附录编号形式=" + str(next(iter(shapes))), "当前编号形式=" + str(title["shape"]), p))
        if sequence and title["number"]:
            current = title["number"]
            if previous and (current[0] == previous[0]) and current[1] != previous[1] + 1:
                findings.append(_emit("156.1", "模板附录连续编号；前项序值=" + str(previous[1]) + "，下一项应为=" + str(previous[1] + 1),
                                      "当前附录=" + title["token"], p))
            previous = current
        else:
            reason = "模板连续编号声明含否定或冲突，不能据样例推定连续要求" if base.get("sequence_denied") else "模板缺少足以确认的连续编号示例或明确连续规则"
            findings.append(_notice("156", reason, p))
    return findings


def _object_key(record):
    return (record["type"], record["kind"], record.get("language"))


def _number_checks(base, records):
    findings, previous = [], {}
    for record in (r for r in records if r["owner"]):
        p = record.get("number_paragraph") or record["paragraph"]
        owner, key = record["owner"], _object_key(record)
        candidates = [r for r in base.get("records", []) if r["owner"] and _object_key(r) == key]
        if record["type"] == "formula" and record["kind"] == "inline":
            continue
        if record["unsafe"] or owner["unsafe"] or any(r["unsafe"] or r["owner"]["unsafe"] for r in candidates):
            findings.append(_notice("158", "编号含域、修订、OLE/图片公式或附录归属不可确认", p))
            previous.pop((owner["token"], key), None)
            continue
        expected = [r["number_info"] for r in candidates]
        shapes = {n["shape"] for n in expected if n}
        if not expected or any(n is None for n in expected) or len(shapes) != 1:
            findings.append(_notice("158", "模板缺少唯一的对应附录对象编号示例，不能默认要求A.1或默认必须编号", p))
            continue
        actual = record["number_info"]
        if actual is None:
            if base.get("number_required") and not record["unsafe"]:
                findings.append(_emit("158", "模板明确规定附录图表公式必须编号；对应编号形式=" + next(iter(shapes)), "可编辑对象/题注未包含可解析编号", p))
            else:
                findings.append(_notice("158", "对象未读到可解析编号，无法区分合法无编号公式或编号保存在对象内部", p))
            continue
        if actual["shape"] not in shapes:
            findings.append(_emit("158", "模板附录编号形式=" + next(iter(shapes)), "当前编号=" + actual["text"], p))
        required_independent = all(n["independent"] for n in expected)
        if required_independent and not actual["independent"]:
            # A separate failure retains the real appendix identifier.
            findings.append(_emit("158", "模板编号包含当前附录标识“" + str(owner["token"]) + "”", "当前编号=" + actual["number"], p))
        grouped = defaultdict(list)
        for candidate in candidates:
            grouped[candidate["owner"]["token"]].append(candidate["number_info"])
        sequences = [group for group in grouped.values() if len(group) >= 2]
        continuous = bool(sequences) and all(all(b["sequence"] == a["sequence"] + 1 for a, b in zip(group, group[1:])) for group in sequences)
        state_key = (owner["token"], key)
        if required_independent and continuous and actual["independent"]:
            prior = previous.get(state_key)
            if prior is not None and actual["sequence"] != prior + 1:
                findings.append(_emit("158", "模板在同一附录内对同类对象连续编号，下一序号=" + str(prior + 1), "当前编号=" + actual["number"], p))
            previous[state_key] = actual["sequence"]
        # Reset is inferred only when at least two template appendices show
        # the same first sequence number. One sample does not prove a reset.
        starts = {group[0]["sequence"] for group in grouped.values()}
        first_for_owner = not any(r is not record and r["owner"] is owner and _object_key(r) == key
                                  for r in records[:records.index(record)])
        if required_independent and actual["independent"] and first_for_owner and len(grouped) >= 2 and len(starts) == 1:
            start = next(iter(starts))
            if actual["sequence"] != start:
                findings.append(_emit("158", "模板在不同附录重新从序号" + str(start) + "开始", "当前附录首个编号=" + actual["number"], p))
    return findings


def _remap(findings):
    return [make_finding("159", TITLES["159"], f.expected, f.actual, getattr(f, "anchor", None), f.status) for f in findings]


def _object_formats(record, candidates, ctx):
    p = record["paragraph"]
    if not candidates or record["unsafe"] or any(r["unsafe"] or r.get("error") for r in candidates):
        return [_notice("159", "正文共同基准缺失或对象内部格式不可读", p)]
    if record["type"] == "formula":
        return _remap(_check_layout(record, candidates, ctx) + _check_fonts(record, candidates, ctx) + _check_spacing(record, candidates))
    findings = _compare_format("159", [r["format"] for r in candidates], _caption_format(record), p)
    if record["kind"] == "table":
        if record["object"] is None or any("table_snapshot" not in r for r in candidates):
            findings.append(_notice("159", "题注不能唯一关联相邻表格对象", p))
        else:
            actual = _table_snapshot(Table(record["object"], ctx.document._body), record)
            # Adapt detached record metadata only, not the document or old detector.
            tables = []
            for candidate in candidates:
                table = dict(candidate["table_snapshot"])
                table["caption"] = dict(table["caption"], role=record["role"])
                tables.append(table)
            findings.extend(_remap(_table_checks({"tables": tables}, [actual])))
    return findings


def _format_checks(base, records, ctx):
    findings = []
    for record in (r for r in records if r["owner"]):
        kind = "formula" if record["type"] == "formula" else record["kind"]
        p = record["paragraph"]
        if kind not in base.get("inheritance", set()):
            findings.append(_notice("159", "模板未明确该类附录对象与正文共同适用的格式关系，不能把正文实际格式强加给附录", p))
            continue
        candidates = [r for r in base.get("records", []) if not r["owner"] and _object_key(r) == _object_key(record)]
        appendix_samples = [r for r in base.get("records", []) if r["owner"] and _object_key(r) == _object_key(record)]
        conflict, template_unreadable = False, False
        for sample in appendix_samples:
            # Template samples can legitimately carry an exception. An explicit
            # contradiction means that no single inherited standard is safe.
            comparison = _object_formats(sample, candidates, SimpleNamespace(document=sample["paragraph"].part.document))
            if any(f.status == Status.UNKNOWN and "不能确认公式与编号排版后的同一行或实际坐标" not in f.actual
                   for f in comparison):
                template_unreadable = True
            if any(f.status == Status.FAIL for f in comparison):
                conflict = True
                break
        if conflict:
            findings.append(_notice("159", "模板附录示例与正文共同要求冲突，需确认例外后再比较", p))
        elif template_unreadable:
            findings.append(_notice("159", "模板附录示例含不可读格式，不能确认是否存在正文共同规则的附录例外", p))
        else:
            findings.extend(_object_formats(record, candidates, ctx))
    return findings


def check_appendix_errors(baseline, ctx, mapping=None):
    base = baseline.get("appendices", {})
    findings = []
    try:
        inventory = collect_appendix_inventory(ctx, mapping)
    except Exception as exc:
        return [_notice(rule, "附录结构读取失败：" + type(exc).__name__, ctx.paragraphs[0] if ctx.paragraphs else None) for rule in TITLES]
    for check in (lambda: _reference_checks(base, ctx, inventory), lambda: _title_checks(base, inventory)):
        try:
            findings.extend(check())
        except Exception as exc:
            findings.extend(_notice(rule, "附录标题/引用读取失败：" + type(exc).__name__, inventory["titles"][0]["paragraph"] if inventory["titles"] else None)
                            for rule in ("155", "156", "157"))
    try:
        records, problems = _objects(ctx, inventory)
        findings.extend(_number_checks(base, records))
        findings.extend(_format_checks(base, records, ctx))
        for p, reason in problems + inventory["problems"]:
            findings.extend(_notice(rule, "附录对象清单不完整：" + reason, p) for rule in ("158", "159"))
    except Exception as exc:
        findings.extend(_notice(rule, "附录对象读取失败：" + type(exc).__name__, inventory["titles"][0]["paragraph"] if inventory["titles"] else None)
                        for rule in ("158", "159"))
    return findings
