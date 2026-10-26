"""Template-driven, read-only checks for accessible equation objects.

089、094、095、098、099：本条依赖大模型语义，本阶段暂不实现。
参考文献引用：待开会讨论，本版本暂不开发。093仅核对明确的本文公式目标。
Saved SEQ/REF results are caches, never freshly evaluated numbers.
"""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
import re

from docx.oxml import OxmlElement
from thesis_format_checker.core.xmlutil import qn
from docx.text.paragraph import Paragraph

from thesis_format_checker.checks.catalog import FINE_RULES
from thesis_format_checker.core.docx_context import UNREADABLE, effective_paragraph, make_finding, unknown, memo
from thesis_format_checker.detectors.fields import scan_fields
from thesis_format_checker.detectors.headings import collect_heading_records

RULES = {"090": "公式及编号对齐不符合模板", "091": "公式编号格式不符合模板",
         "092": "公式编号序列错误", "093": "公式引用目标不一致或不存在",
         "096": "公式字体或字号不符合模板", "097": "公式段落间距或行距不符合模板"}
_NUMBER = re.compile(r"(?P<prefix>式\s*|Equation\s*|Eq\.?\s*)?(?P<open>[(（])\s*(?P<num>\d+(?:[-.．－]\d+)*[a-z]?)\s*(?P<close>[)）])", re.I)
_REFERENCE = re.compile(r"(?P<lead>本文|本章|本节)?\s*(?:公式|式)\s*(?P<label>[(（]\s*\d+(?:[-.．－]\d+)*[a-z]?\s*[)）])", re.I)
_SPACING = ("space_before", "space_after", "line_spacing", "line_spacing_rule")
_LABELS = {"space_before": "段前间距", "space_after": "段后间距", "line_spacing": "行距", "line_spacing_rule": "行距类型"}


def _unknown(rule, reason, p=None):
    return unknown(rule, _title(rule), reason, p)


#: 公式布局属性 → 细粒度规则编号（090.1 居中 / 090.2 编号右对齐）
_LAYOUT_CODES = {"math_alignment": "090.1", "paragraph_alignment": "090.1",
                 "number_alignment": "090.2", "number_tabs": "090.2", "tabs": "090.2",
                 "table": "090.2", "table_columns": "090.2"}
#: 公式段落间距属性 → 细粒度规则编号
_SPACING_CODES = {"space_before": "097.1", "space_after": "097.2",
                  "line_spacing": "097.3", "line_spacing_rule": "097.3"}


def _title(code: str) -> str:
    """细粒度规则标题取自规则目录；传入族编号时用本模块的族标题。"""

    item = FINE_RULES.get(code)
    return item.title if item else RULES[code]


def _text(node):
    return "".join(x.text or "" for x in node.iter(qn("w:t")))


def _number(text):
    matches = list(_NUMBER.finditer(text))
    if len(matches) == 1:
        match = matches[0]
    elif not matches:
        match = re.fullmatch(r"(?:式\s*|Equation\s*|Eq\.?\s*)?(?P<num>\d+(?:[-.．－]\d+)*[a-z]?)", text.strip(), re.I)
        if match is None:
            return None
    else:
        return None
    body = match.group("num")
    return {"text": match.group(), "value": re.sub(r"[．－]", lambda m: "." if m[0] == "．" else "-", body),
            "parts": tuple(int(v) for v in re.findall(r"\d+", body)),
            "simple": re.fullmatch(r"\d+(?:[-.．－]\d+)?", body) is not None,
            "pattern": re.sub(r"\d+", "#", match.group()), "start": match.start(), "end": match.end()}


def _raw_tabs(paragraph):
    props, style, seen = [paragraph._p.pPr], paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        props.append(style.element.pPr)
        style = style.base_style
    tabs = {}
    for prop in reversed(props):
        if prop is None:
            continue
        for tab in prop.findall("./" + qn("w:tabs") + "/" + qn("w:tab")):
            pos, val = tab.get(qn("w:pos")), tab.get(qn("w:val"))
            if pos is None or val is None:
                raise ValueError("制表位缺少位置或对齐属性")
            if val == "clear":
                tabs.pop(pos, None)
            else:
                tabs[pos] = (val, tab.get(qn("w:leader"), "none"))
    return tuple(sorted((int(k), *v) for k, v in tabs.items()))


def _layout(record, ctx):
    p = record["paragraph"]
    maths = p._p.xpath("./m:oMathPara")
    align = None
    if maths:
        values = [node.get(qn("m:val")) for m in maths for node in m.findall("./" + qn("m:oMathParaPr") + "/" + qn("m:jc"))]
        if len(set(values)) > 1:
            raise ValueError("同段独立数学对象有不同对齐设置")
        if values:
            align = values[0]
        else:
            defaults = ctx.document.settings.element.xpath("./m:mathPr/m:defJc")
            # No school convention is substituted for an absent math default.
            align = defaults[0].get(qn("m:val")) if defaults else None
    cell = next((a for a in p._p.iterancestors() if a.tag == qn("w:tc")), None)
    number_p = record["number_paragraph"]
    result = {"math_alignment": align, "paragraph_alignment": effective_paragraph(p)["alignment"],
              "tabs": _raw_tabs(p), "number_alignment": effective_paragraph(number_p)["alignment"] if number_p is not None else None,
              "number_tabs": _raw_tabs(number_p) if number_p is not None else (), "table": cell is not None}
    if cell is not None:
        row = cell.getparent()
        cells = list(row.findall(qn("w:tc")))
        num_cell = next((a for a in number_p._p.iterancestors() if a.tag == qn("w:tc")), None) if number_p is not None else None
        result["table_columns"] = (cells.index(cell), cells.index(num_cell) if num_cell in cells else None, len(cells))
    return result


def __read_records_impl(ctx):
    records, problems = [], []
    try:
        headings = collect_heading_records(ctx)
    except Exception:
        headings = []
    chapter = None
    region = "body"
    heading_map = {h["paragraph"]._p: h for h in headings if h["level"] == 1}
    for p in ctx.paragraphs:
        if p._p in heading_map:
            h = heading_map[p._p]
            chapter = h.get("number") if h.get("number_known") else None
            if re.match(r"\s*附录", h.get("text", "")):
                region = "appendix"
        if any(a.tag in (qn("w:del"), qn("w:moveFrom"), qn("w:txbxContent")) for a in p._p.iterancestors()):
            continue
        maths = p._p.xpath(".//m:oMath")
        objects = p._p.xpath(".//w:object")
        pictures = p._p.xpath(".//w:drawing|.//w:pict")
        normal_text = _text(p._p).strip()
        number = _number(normal_text)
        equation_ole = any("equation" in str(node.attrib).lower() or "mathtype" in str(node.attrib).lower()
                           for obj in objects for node in obj.iter())
        opaque = bool(equation_ole or ((objects or pictures) and number))
        if not maths and not opaque:
            continue
        try:
            number_p = p if number else None
            cell = next((a for a in p._p.iterancestors() if a.tag == qn("w:tc")), None)
            # Only a clearly labelled equation number in another cell of the same row.
            if cell is not None and number is None:
                candidates = []
                for candidate in ctx.paragraphs:
                    candidate_cell = next((a for a in candidate._p.iterancestors() if a.tag == qn("w:tc")), None)
                    if candidate_cell is not None and candidate_cell is not cell and candidate_cell.getparent() is cell.getparent():
                        candidate_text = _text(candidate._p).strip()
                        found = _number(candidate_text)
                        if found and found["text"] == candidate_text:
                            candidates.append((candidate, found))
                if len(candidates) == 1:
                    number_p, number = candidates[0]
            residual = normal_text
            if number_p is p and number:
                residual = normal_text[:number["start"]] + normal_text[number["end"]:]
            display = bool(p._p.xpath("./m:oMathPara")) or (not residual.strip() and bool(maths))
            kind = "table" if cell is not None and display else "display" if display else "inline"
            fields = scan_fields(number_p._p if number_p is not None else p._p)
            record = {"paragraph": p, "number_paragraph": number_p, "number": number, "kind": kind,
                      "opaque": opaque, "chapter": chapter, "region": region, "section": ctx.section_index(p),
                      "fields": fields, "dynamic": any(f["kind"] in ("SEQ", "STYLEREF", "REF") for f in fields),
                      "maths": maths}
            record["revision"] = bool(p._p.xpath(".//w:ins|.//w:del|.//w:moveFrom|.//w:moveTo|.//w:rPrChange|.//w:pPrChange")) or any(
                a.tag in (qn("w:ins"), qn("w:moveTo")) for a in p._p.iterancestors())
            record["number_unknown"] = bool(fields.errors or any(f["malformed"] for f in fields))
            if record["opaque"] and number:
                record["kind"] = "table" if cell is not None else "display"
            if record["dynamic"]:
                record["kind"] = "table" if cell is not None else "display" if display else kind
            records.append(record)
        except Exception as exc:
            problems.append((p, str(exc)))
    return records, problems


def _roles(run, formula):
    roles = []
    for node in run.iterancestors():
        if node is formula:
            break
        if node.tag in {qn("m:sub"), qn("m:sup"), qn("m:num"), qn("m:den"), qn("m:deg")}:
            roles.append(node.tag.rsplit("}", 1)[-1])
    return "/".join(reversed(roles)) or "base"


def _font_samples(record, ctx):
    from thesis_format_checker.detectors.text_runs import effective_text_runs
    p = record["paragraph"]
    samples, errors = [], []
    math_fonts = ctx.document.settings.element.xpath("./m:mathPr/m:mathFont")
    math_font = math_fonts[0].get(qn("m:val")) if math_fonts else None
    for formula in record["maths"]:
        runs = [(run, "".join(node.text or "" for node in run.iter(qn("m:t"))), _roles(run, formula))
                for run in formula.iter(qn("m:r"))]
        for control in formula.iter(qn("m:ctrlPr")):
            character = control.getparent().find(qn("m:chr"))
            text = character.get(qn("m:val")) if character is not None else None
            if text:
                runs.append((control, text, _roles(control, formula) + "/control"))
            else:
                errors.append("m:ctrlPr控制符未显式给出字符，自动生成符号的最终字体/缩放无法完整还原")
        for math_run, text, role in runs:
            if not text:
                continue
            shell = OxmlElement("w:p")
            if p._p.pPr is not None:
                shell.append(deepcopy(p._p.pPr))
            r = OxmlElement("w:r")
            rpr = math_run.find(qn("w:rPr"))
            if rpr is not None:
                r.append(deepcopy(rpr))
            t = OxmlElement("w:t")
            t.text = text
            r.append(t)
            shell.append(r)
            # Preserve a detached table ancestor chain, including style and
            # conditional-format hints, for conservative table inheritance.
            child = shell
            for ancestor in p._p.iterancestors():
                if ancestor.tag not in (qn("w:tc"), qn("w:tr"), qn("w:tbl")):
                    continue
                wrapper = OxmlElement("w:" + ancestor.tag.rsplit("}", 1)[-1])
                prop_name = {qn("w:tc"): "w:tcPr", qn("w:tr"): "w:trPr", qn("w:tbl"): "w:tblPr"}[ancestor.tag]
                source_prop = ancestor.find(qn(prop_name))
                if source_prop is not None:
                    wrapper.append(deepcopy(source_prop))
                wrapper.append(child)
                child = wrapper
            proxy = Paragraph(shell, p._parent)
            effective = effective_text_runs(proxy)
            if not effective:
                errors.append("数学运行有效属性为空")
                continue
            # The helper returns text/property pairs and resolves paragraph,
            # character style, docDefaults, theme and direct run overrides.
            props = effective[0][1]
            errors.extend(props.get("limitations", []))
            normal = math_run.find("./" + qn("m:rPr") + "/" + qn("m:nor"))
            normal_text = normal is not None and normal.get(qn("m:val"), "1") not in ("0", "false", "off")
            for category, fragment in (("中文", "".join(c for c in text if '\u4e00' <= c <= '\u9fff')),
                                       ("英文", "".join(c for c in text if c.isascii() and c.isalpha())),
                                       ("数字", "".join(c for c in text if c.isdigit())),
                                       ("符号", "".join(c for c in text if not c.isspace() and not ('\u4e00' <= c <= '\u9fff')
                                                     and not (c.isascii() and c.isalpha()) and not c.isdigit()))):
                if not fragment:
                    continue
                font_key = "east_asia_font" if category == "中文" else "hansi_font" if normal_text and any(not c.isascii() for c in fragment) else "ascii_font"
                # Normal text has ordinary Word font semantics. Math text is
                # governed by mathFont; absent math defaults are not guessed.
                font = props.get(font_key) if normal_text else math_font
                explicit_fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
                if explicit_fonts is not None and not normal_text:
                    # Word may substitute unsupported math fonts. Compare only
                    # readable declarations, with a rendering limitation below.
                    font = props.get(font_key) or math_font
                samples.append(((role, category, normal_text), fragment, {"font": font, "size_pt": props.get("size_pt")}))
    return samples, errors


def _sequence_policy(ctx):
    text = ctx.text
    global_sequence = bool(re.search(r"公式.{0,12}(?:全文|全篇).{0,6}(?:连续|统一)编号", text))
    chapter_sequence = bool(re.search(r"公式.{0,15}(?:按章编号|每章重新编号|按章节编号)", text))
    starts = re.findall(r"公式.{0,20}(?:从|起始(?:值)?[：:]?)\s*(\d+)\s*(?:开始|起|编号)", text)
    if len(set(starts)) > 1:
        return {"scope": None, "reason": "模板给出了互相冲突的公式起始编号"}
    if global_sequence and chapter_sequence:
        return {"scope": None, "reason": "模板公式全文续编与按章重置说明冲突"}
    return {"scope": "global" if global_sequence else "chapter" if chapter_sequence else None,
            "start": int(starts[0]) if starts and len(set(starts)) == 1 else None,
            "reason": "模板未明确公式编号的连续/重置范围"}


def extract_equation_baseline(ctx, mapping=None):
    records, problems = _read_records(ctx)
    for record in records:
        try:
            if record["revision"]:
                raise ValueError("模板公式含未接受的修订，无法确认采用的视图")
            record["layout"] = _layout(record, ctx)
            record["spacing"] = {key: value for key, value in effective_paragraph(record["paragraph"]).items() if key in _SPACING}
            record["fonts"], record["font_errors"] = _font_samples(record, ctx) if not record["opaque"] else ([], ["模板公式为不透明对象"])
        except Exception as exc:
            record["error"] = str(exc)
    return {"equations": {"records": records, "problems": problems, "policy": _sequence_policy(ctx),
                                "number_required": bool(re.search(r"公式.{0,8}(?:必须|均须|均应|应当).{0,6}编号", ctx.text)),
                                "reference_patterns": tuple(sorted({m["pattern"] for p in ctx.paragraphs for ref in _REFERENCE.finditer(p.text)
                                                                    if (m := _number(ref.group("label")))}))}}


def _candidates(base, record, mapping):
    candidates = [r for r in base["records"] if r["region"] == record["region"] and r["kind"] == record["kind"]]
    section_map = (mapping or {}).get("section_map", {})
    mapped = section_map.get(str(record["section"] + 1), section_map.get(record["section"] + 1)) if isinstance(section_map, dict) else None
    if mapped is not None:
        try:
            candidates = [r for r in candidates if r["section"] == int(mapped) - 1]
        except (ValueError, TypeError):
            return []
    return candidates


def _consensus(candidates, key):
    values = [r.get(key) if not r.get("opaque") and not r.get("error") else None for r in candidates]
    if not values or any(value != values[0] for value in values):
        return None
    return values[0]


def _check_layout(record, candidates, ctx):
    p, findings = record["paragraph"], []
    if record["kind"] == "inline":
        return findings
    expected, actual = _consensus(candidates, "layout"), _layout(record, ctx)
    if expected is None:
        findings.append(_unknown("090", "模板对应公式布局缺失、读取失败或存在多套无法匹配的配置", p))
    else:
        differences = {key: (value, actual.get(key)) for key, value in expected.items()
                       if value is not None and actual.get(key) is not None and value != actual.get(key)}
        if differences:
            labels = {"math_alignment": "数学对象对齐", "paragraph_alignment": "段落对齐", "tabs": "公式段制表位",
                      "number_alignment": "编号段对齐", "number_tabs": "编号段制表位", "table": "表格布局", "table_columns": "公式列/编号列/总列数"}
            # 公式居中（090.1）与编号右对齐（090.2）分别成条
            for key, (want, got) in differences.items():
                code = _LAYOUT_CODES.get(key, "090.1")
                findings.append(make_finding(code, _title(code),
                                             labels.get(key, key) + "：" + str(want),
                                             labels.get(key, key) + "：" + str(got), p))
        if expected.get("math_alignment") is None or actual.get("math_alignment") is None:
            findings.append(_unknown("090", UNREADABLE + "；独立数学对象有效对齐值缺失", p))
    if record["number"] is not None or record["dynamic"]:
        findings.append(_unknown("090", UNREADABLE + "；仅核对对齐和制表/布局表格配置，不能确认公式与编号排版后的同一行或实际坐标", record["number_paragraph"] or p))
    return findings


def _check_format(record, candidates, base):
    p, num, findings = record["number_paragraph"] or record["paragraph"], record["number"], []
    if any(r.get("opaque") or r.get("error") or r.get("number_unknown") for r in candidates):
        return [_unknown("091", UNREADABLE + "；模板对应编号存在不可读取的样例或损坏域，不能选择剩余样例代替完整基准", p)]
    fields = record["fields"]
    if fields.errors or any(f["malformed"] for f in fields):
        return [_unknown("091", UNREADABLE + "；公式编号域结构损坏", p)]
    if record["dynamic"]:
        def signatures(rec):
            # Sequence identifiers are document-local names, not school
            # formatting requirements. Compare switches, not arbitrary IDs.
            return tuple(re.sub(r'(?i)^SEQ\s+(?:"[^"]+"|[^\s\\]+)', "SEQ <公式序列>",
                                re.sub(r"\s+", " ", f["instruction"]).strip()).upper()
                         for f in rec["fields"] if f["kind"] == "SEQ")
        expected = {signatures(r) for r in candidates if r["dynamic"] and not r["fields"].errors}
        actual = signatures(record)
        if len(expected) == 1 and actual and actual not in expected:
            findings.append(make_finding("091", RULES["091"], "公式SEQ域结构 " + str(next(iter(expected))), "公式SEQ域结构 " + str(actual), p))
        findings.append(_unknown("091", UNREADABLE + "；SEQ/引用域显示值是保存缓存，未执行Word域计算，编号显示格式暂无法确认", p))
        return findings
    if num is None:
        if record["kind"] != "inline" and base["number_required"]:
            findings.append(make_finding("091", RULES["091"], "模板明确要求公式编号", "未读到可定位的公式编号", p))
        return findings
    patterns = {r["number"]["pattern"] for r in candidates if r["number"] and not r["dynamic"]}
    if len(patterns) == 1:
        expected = next(iter(patterns))
        if num["pattern"] != expected:
            findings.append(make_finding("091", RULES["091"], "公式编号形式 " + expected, num["text"], p))
    else:
        findings.append(_unknown("091", "模板对应编号形式缺失或存在多种形式，无法唯一选择", p))
    return findings


def _check_sequence(records, policy):
    numbered = [r for r in records if r["number"] or r["dynamic"]]
    if not numbered:
        return []
    if policy.get("scope") is None:
        return [_unknown("092", policy["reason"], numbered[0]["paragraph"])]
    findings, previous, seen = [], {}, defaultdict(set)
    poisoned = set()
    for rec in numbered:
        p, number = rec["number_paragraph"] or rec["paragraph"], rec["number"]
        scope = (rec["region"], rec["chapter"] if policy["scope"] == "chapter" else "global")
        if rec["opaque"] or rec["dynamic"] or rec["revision"] or rec["number_unknown"] or number is None or not number["simple"]:
            findings.append(_unknown("092", UNREADABLE + "；动态域、不可读对象或子公式编号不能作为可靠的连续序列值", p))
            poisoned.add(scope)
            continue
        if policy["scope"] == "chapter" and rec["chapter"] is None:
            findings.append(_unknown("092", "公式所在章编号无法可靠读取，不能确认章号及重置范围", p))
            continue
        if scope in poisoned:
            findings.append(_unknown("092", "同一序列前面存在无法计算的公式编号，不能推断此处连续编号", p))
            continue
        parts = number["parts"]
        expected_chapter = rec["chapter"]
        if policy["scope"] == "chapter" and len(parts) == 2 and re.fullmatch(r"\d+", str(expected_chapter or "")) and parts[0] != int(expected_chapter):
            findings.append(make_finding("092.3", _title("092.3"), "所属章节号 " + str(expected_chapter), "公式章前缀 " + str(parts[0]), p))
        seq = parts[-1]
        target = previous[scope] + 1 if scope in previous else policy.get("start")
        if number["value"] in seen[scope]:
            findings.append(make_finding("092.2", _title("092.2"), "模板规定范围内编号不得重复", "重复编号 " + number["text"], p))
        elif target is not None and seq != target:
            findings.append(make_finding("092.1", _title("092.1"), "按模板连续编号，此处序号应为 " + str(target), number["text"], p))
        seen[scope].add(number["value"])
        previous[scope] = seq
    return findings


def _check_fonts(record, candidates, ctx):
    if any(candidate.get("opaque") or candidate.get("error") for candidate in candidates):
        return [_unknown("096", UNREADABLE + "；对应模板样例包含不可读取的公式或修订，无法确定完整字体基准", record["paragraph"])]
    expected = defaultdict(lambda: {"font": set(), "size_pt": set()})
    for candidate in candidates:
        for role, _, values in candidate.get("fonts", []):
            for key, value in values.items():
                expected[role][key].add(value)
    actual, errors = _font_samples(record, ctx)
    findings, unknowns = [], set(errors)
    for role, fragment, properties in actual:
        for key, value in properties.items():
            possibilities = expected.get(role, {}).get(key, set())
            if len(possibilities) != 1 or None in possibilities or value is None:
                unknowns.add("同结构角色/字符类别的字体或基础字号基准不唯一或不可读")
                continue
            target = next(iter(possibilities))
            if target != value:
                label = "字体" if key == "font" else "基础字号（磅）"
                code = "096.1" if key == "font" else "096.2"
                findings.append(make_finding(code, _title(code), f"{role[0]}层{role[1]}{label}：{target}", f"字符“{fragment}”的{label}：{value}", record["paragraph"]))
    if not actual:
        unknowns.add("未读到可比较的OMML数学运行")
    if any(candidate.get("font_errors") for candidate in candidates):
        unknowns.add("模板存在不能完整还原的数学控制符属性")
    if unknowns:
        findings.append(_unknown("096", UNREADABLE + "；" + "；".join(sorted(unknowns)), record["paragraph"]))
    return findings


def _check_spacing(record, candidates):
    target = _consensus(candidates, "spacing")
    p = record["paragraph"]
    if target is None:
        return [_unknown("097", "模板对应独立/行内/表格公式段落基准缺失或不唯一", p)]
    actual = effective_paragraph(p)
    findings, missing = [], []
    for key in _SPACING:
        expected, current = target.get(key), actual.get(key)
        if expected is None or current is None:
            missing.append(_LABELS[key])
        elif expected != current:
            code = _SPACING_CODES.get(key, "097.1")
            findings.append(make_finding(code, _title(code), _LABELS[key] + " " + str(expected), _LABELS[key] + " " + str(current), p))
    if missing:
        findings.append(_unknown("097", UNREADABLE + "；不能读取有效" + "、".join(missing) + "，未比较公式排版后的可见高度", p))
    return findings


def _check_references(ctx, records, base, problems=()):
    findings, targets = [], defaultdict(list)
    formula_ps = {r["paragraph"]._p for r in records}
    formula_ps.update(r["number_paragraph"]._p for r in records if r["number_paragraph"] is not None)
    for record in records:
        if record["number"] and not record["dynamic"] and not record["opaque"] and not record["revision"] and not record["number_unknown"]:
            targets[record["number"]["value"]].append(record)
    opaque = any(r["opaque"] or r["dynamic"] or r["revision"] or r["number_unknown"] for r in records) or ctx.has_textboxes or bool(problems)
    chapter_by_paragraph = {}
    try:
        headings = {h["paragraph"]._p: h for h in collect_heading_records(ctx) if h["level"] == 1}
        chapter = None
        for para in ctx.paragraphs:
            if para._p in headings:
                item = headings[para._p]
                chapter = item.get("number") if item.get("number_known") else None
            chapter_by_paragraph[para._p] = chapter
    except Exception:
        pass
    # Bookmark ranges are inspected in document order so start/end need not
    # occur in the same paragraph. No bookmark or field is updated.
    active, bookmark_targets, bookmarks = {}, defaultdict(set), set()
    for node in ctx.document.element.body.iter():
        if node.tag == qn("w:bookmarkStart"):
            name, ident = node.get(qn("w:name")), node.get(qn("w:id"))
            if name:
                active[ident] = name
                bookmarks.add(name)
        elif node.tag == qn("w:bookmarkEnd"):
            active.pop(node.get(qn("w:id")), None)
        elif node.tag in (qn("w:r"), qn("m:oMath"), qn("w:p")):
            pnode = node if node.tag == qn("w:p") else next((a for a in node.iterancestors() if a.tag == qn("w:p")), None)
            if pnode in formula_ps:
                for name in active.values():
                    bookmark_targets[name].add(pnode)
    for p in ctx.paragraphs:
        if p._p in formula_ps:
            continue
        fields = scan_fields(p._p)
        ref_fields = [f for f in fields if f["kind"] == "REF"]
        for field in ref_fields:
            match = re.match(r'REF\s+(?:"([^"]+)"|([^\s\\]+))', field["instruction"], re.I)
            name = (match.group(1) or match.group(2)) if match else None
            explicit = bool(_REFERENCE.search(p.text) or re.search(r"(?:公式|式)\s*$", p.text))
            equation_target = bool(bookmark_targets.get(name))
            if not explicit and not equation_target:
                continue  # References to literature/other objects are not developed here.
            if field["malformed"] or not name:
                findings.append(_unknown("093", UNREADABLE + "；公式REF域结构或目标名不可读取", p))
            elif name not in bookmarks:
                findings.append(make_finding("093.2", _title("093.2"), "公式交叉引用指向存在的书签", "REF目标书签不存在：" + name, p))
            elif not equation_target:
                findings.append(_unknown("093", "REF书签存在，但不能确认其范围对应公式编号", p))
            else:
                findings.append(_unknown("093", UNREADABLE + "；REF目标公式存在，但保存显示值未经域更新，不能确认实时显示编号", p))
        if ref_fields:
            continue
        for match in _REFERENCE.finditer(p.text):
            number = _number(match.group("label"))
            vicinity = p.text[max(0, match.start() - 28): min(len(p.text), match.end() + 20)]
            if re.search(r"文献|参考资料|引用自|引自|\[\s*\d", vicinity) and not match.group("lead"):
                continue  # 待开会讨论，本版本暂不开发：文献中的公式引用。
            if re.match(r"\s*(?:[-—~～至、,，]|和)\s*[(（]\s*\d", p.text[match.end():]):
                findings.append(_unknown("093", "范围或多编号公式引用的边界无法唯一解析，需人工核对全部目标", p))
                continue
            if not base["reference_patterns"] or number["pattern"] not in base["reference_patterns"]:
                findings.append(_unknown("093", "模板没有可唯一对应的公式引用形式", p))
                continue
            if not match.group("lead"):
                findings.append(_unknown("093", "文字引用未明确指向本文/本章/本节，不能排除外部公式", p))
                continue
            matches = targets.get(number["value"], [])
            if match.group("lead") in ("本章", "本节"):
                chapter = chapter_by_paragraph.get(p._p)
                if chapter is None or match.group("lead") == "本节":
                    findings.append(_unknown("093", "文字引用的章/节编号作用域无法唯一确定", p))
                    continue
                matches = [r for r in matches if r["chapter"] == chapter]
            if len(matches) == 1:
                continue
            if not matches and not opaque:
                    findings.append(make_finding("093.2", _title("093.2"), "本文公式引用编号对应实际公式", "引用“" + match.group().strip() + "”，未找到对应公式", p))
            else:
                findings.append(_unknown("093", UNREADABLE + "；目标编号存在重号歧义或不可读取的公式/动态编号，不能确定是否缺失", p))
    return findings


def check_equation_errors(baseline, ctx, mapping=None):
    base = baseline.get("equations", {})
    records, problems = _read_records(ctx)
    if not base:
        return [_unknown(rule, "模板公式基准读取失败", records[0]["paragraph"] if records else None) for rule in RULES] if records else []
    findings = []
    for p, reason in problems:
        findings.extend(_unknown(rule, UNREADABLE + "；公式对象读取异常：" + reason, p) for rule in RULES)
    for record in records:
        p = record["paragraph"]
        if record["revision"]:
            findings.extend(_unknown(rule, UNREADABLE + "；公式含未接受的修订，不能确认采用的公式视图", p) for rule in RULES)
            continue
        if record["opaque"]:
            findings.extend(_unknown(rule, UNREADABLE + "；OLE/MathType或图片公式内部属性不能访问", p) for rule in ("090", "091", "096", "097"))
            continue
        candidates = _candidates(base, record, mapping)
        for rule, check in (("090", lambda: _check_layout(record, candidates, ctx)),
                            ("091", lambda: _check_format(record, candidates, base)),
                            ("096", lambda: _check_fonts(record, candidates, ctx)),
                            ("097", lambda: _check_spacing(record, candidates))):
            try:
                findings.extend(check())
            except Exception as exc:
                findings.append(_unknown(rule, UNREADABLE + "；公式检测读取异常：" + str(exc), p))
    try:
        findings.extend(_check_sequence(records, base["policy"]))
        findings.extend(_check_references(ctx, records, base, problems))
    except Exception as exc:
        findings.extend(_unknown(rule, UNREADABLE + "；公式编号关联读取异常：" + str(exc)) for rule in ("092", "093"))
    if ctx.has_textboxes and base.get("records"):
        findings.extend(_unknown(rule, UNREADABLE + "；文本框内可能包含不可定位的公式", None) for rule in RULES)
    return findings

def _read_records(ctx):
    """带缓存的入口：同一份文档只扫描一次（见 core.docx_context.memo）。"""

    return memo(ctx, "equation_records", lambda: __read_records_impl(ctx))
