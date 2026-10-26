"""页面设置检测（023–033）的取证实现（只读，以模板为准）。

023 and 025 remain exclusively in the existing paper-size and margin checkers.
028: 本条依赖大模型语义，本阶段暂不实现
032: 本条依赖大模型语义，本阶段暂不实现
"""
from __future__ import annotations

import re
from collections import defaultdict

from thesis_format_checker.core.xmlutil import qn

from thesis_format_checker.checks.catalog import FINE_RULES
from thesis_format_checker.core.docx_context import (
    UNREADABLE, DocxContext, effective_paragraph, make_finding, unknown,
)
from thesis_format_checker.core.pairing import sample_at


_TITLES = {
    "022": "前置部分页眉页脚或页码不符合模板",
    "024": "页面方向不符合模板",
    "026": "装订线位置或宽度不符合模板",
    "027": "页眉页脚距纸张边缘距离不符合模板",
    "029": "分节后页面设置或页眉页脚异常变化",
    "030": "章节分页方式需核对",
    "031": "首页或奇偶页版式选项不符合模板",
    "033": "内容超出模板内容边界或无法确定裁切风险",
}
_GROUPS = {"024": ("orientation",), "026": ("gutter", "gutter_position"),
           "027": ("header_distance", "footer_distance"),
           "031": ("different_first", "different_odd_even")}
#: 页面设置项 → 细粒度规则编号（一个设置项一条结论）
_GROUP_CODES = {"orientation": "024", "gutter": "026.2", "gutter_position": "026.1",
                "header_distance": "027.1", "footer_distance": "027.2",
                "different_first": "031.1", "different_odd_even": "031.2"}
#: 前置部分：页眉 / 页脚 各自成条
_FRONT_CODES = {"header": "022.1", "footer": "022.2"}
#: 内容超出可打印区域（033）
_WIDTH_RULE = "033.1"
_LABELS = {"orientation": "页面方向", "gutter": "装订线宽度", "gutter_position": "装订线位置",
           "header_distance": "页眉距顶端", "footer_distance": "页脚距底端",
           "different_first": "首页不同", "different_odd_even": "奇偶页不同"}


def _flag(element, name):
    node = element.find(qn("w:" + name))
    return node is not None and node.get(qn("w:val"), "true") not in {"0", "false", "off"}


def _page_number_format(section) -> str:
    """该节的有效页码编号格式；未显式设置时按 Word 默认的阿拉伯数字处理。"""

    # page_number_settings 里存的是 w:pgNumType 的属性，名字带命名空间。
    settings = {str(name).split("}")[-1]: value for name, value in (section.get("page_number_settings") or ())}
    return (settings.get("fmt") or "decimal").lower()


_NUMBER_FORMAT_LABELS = {"decimal": "阿拉伯数字", "upperroman": "大写罗马数字", "lowerroman": "小写罗马数字",
                         "upperletter": "大写字母", "lowerletter": "小写字母", "chinesecounting": "中文数字"}


def _roles(paragraphs):
    result = set()
    for p in paragraphs:
        text = re.sub(r"\s+", "", p.text).casefold()
        if text in {"摘要", "中文摘要", "内容摘要"}:
            result.add("中文摘要")
        elif text in {"abstract", "英文摘要"}:
            result.add("英文摘要")
        elif text in {"目录", "contents", "tableofcontents"}:
            result.add("目录")
        elif re.fullmatch(r".{0,20}(原创性声明|独创性声明|授权使用声明|使用授权声明).{0,10}", text):
            result.add("声明")
        elif _heading_level(p) == 1:
            result.add("正文")
    # Field labels locate a cover only when no explicit body/abstract marker is present.
    if not result and any(re.search(r"(论文题目|论文题名|学号|指导教师)\s*[:：]", p.text) for p in paragraphs):
        result.add("封面")
    return tuple(sorted(result))


def _heading_level(paragraph):
    try:
        direct = paragraph._p.find(qn("w:pPr"))
        direct_outline = None if direct is None else direct.find(qn("w:outlineLvl"))
        if direct_outline is not None:
            level = int(direct_outline.get(qn("w:val")))
            return level + 1 if 0 <= level <= 8 else None
        style = paragraph.style
        seen = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            m = re.fullmatch(r"(?:heading|标题)\s*([1-9])", style.name, re.I)
            if m:
                return int(m.group(1))
            props = style.element.find(qn("w:pPr"))
            outline = None if props is None else props.find(qn("w:outlineLvl"))
            if outline is not None:
                level = int(outline.get(qn("w:val")))
                return level + 1 if 0 <= level <= 8 else None
            style = style.base_style
    except Exception:
        return None
    return None


def _section_paragraphs(ctx):
    groups = defaultdict(list)
    for p in ctx.paragraphs:
        try:
            groups[ctx.section_index(p)].append(p)
        except Exception:
            # Section lookup uncertainty is reported by the caller instead of guessed.
            groups[None].append(p)
    return groups


def _section_anchor(paragraphs, fallback=None):
    """节级结论的锚点：优先落在**有文字**的段落上。

    封面节的第一个段落往往是校徽/图片（没有文字），把页眉页脚结论锚到那里，
    使用者只会看到批注贴在图片上、不知道说的是页眉还是页脚；这里优先取有文字的
    段落（例如校名、题目），没有文字时才退回图片段落。
    """

    return next((p for p in paragraphs if p.text.strip()),
                next((p for p in paragraphs
                      if p._p.xpath(".//w:drawing | .//w:pict | .//w:object")),
                     next(iter(paragraphs), fallback)))


def _header_info(document, sections, index, tag, variant):
    """Resolve linked stories from XML without creating any header/footer part."""
    for previous in range(index, -1, -1):
        refs = sections[previous]._sectPr.findall(qn("w:" + tag + "Reference"))
        ref = next((n for n in refs if n.get(qn("w:type")) == variant), None)
        if ref is None:
            continue
        part = document.part.related_parts[ref.get(qn("r:id"))]
        root = part.element
        words = "".join(n.text or "" for n in root.iter(qn("w:t")))
        has_object = any(n.tag in {qn("w:drawing"), qn("w:pict"), qn("w:object")} for n in root.iter())
        fields = any(n.tag in {qn("w:fldChar"), qn("w:fldSimple"), qn("w:instrText")} for n in root.iter())
        return {"present": bool(words.strip() or has_object or fields), "fields": fields,
                "text": words, "objects": has_object}
    return {"present": False, "fields": False, "text": "", "objects": False}


def _read_section(ctx, sections, index, paragraphs):
    section = sections[index]
    result = {"roles": _roles(paragraphs), "errors": {}, "index": index}
    properties = ("orientation", "gutter", "header_distance", "footer_distance", "page_width",
                  "page_height", "left_margin", "right_margin", "start_type")
    for name in properties:
        try:
            value = getattr(section, name)
            if value is None:
                if name == "gutter":
                    value = 0  # OOXML absence means no separate gutter, not a school requirement.
                else:
                    raise ValueError("属性未声明且无法解析有效值")
            result[name] = int(value)
        except Exception as exc:
            result[name] = None
            result["errors"][name] = str(exc)
    try:
        result["different_first"] = bool(section.different_first_page_header_footer)
    except Exception as exc:
        result["different_first"] = None
        result["errors"]["different_first"] = str(exc)
    try:
        settings = ctx.document.settings.element
        result["different_odd_even"] = _flag(settings, "evenAndOddHeaders")
        result["gutter_position"] = ("上" if _flag(settings, "gutterAtTop") else
                                    "右" if _flag(section._sectPr, "rtlGutter") else
                                    "内侧" if _flag(settings, "mirrorMargins") else "左")
    except Exception as exc:
        for name in ("different_odd_even", "gutter_position"):
            result[name] = None
            result["errors"][name] = str(exc)
    result["stories"] = {}
    for kind in ("header", "footer"):
        for variant in ("default", "first", "even"):
            key = kind + ":" + variant
            try:
                result["stories"][key] = _header_info(ctx.document, sections, index, kind, variant)
            except Exception as exc:
                result["stories"][key] = None
                result["errors"][key] = str(exc)
    try:
        numbering = section._sectPr.find(qn("w:pgNumType"))
        result["page_number_settings"] = None if numbering is None else tuple(sorted(numbering.attrib.items()))
    except Exception as exc:
        result["page_number_settings"] = None
        result["errors"]["page_number_settings"] = str(exc)
    return result


def _has_forced_break(ctx, paragraph):
    if effective_paragraph(paragraph).get("page_break_before"):
        return True
    p_index = next((i for i, p in enumerate(ctx.paragraphs) if p._p is paragraph._p), None)
    if p_index is None:
        return None
    # A break before any text in the heading also begins a page.
    for node in paragraph._p.iter():
        if node.tag == qn("w:t") and (node.text or "").strip():
            break
        if node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
            return True
    for previous in reversed(ctx.paragraphs[:p_index]):
        # Only a page break after the previous paragraph's last visible text qualifies.
        tail_has_break = False
        for node in previous._p.iter():
            if node.tag == qn("w:t") and (node.text or "").strip():
                tail_has_break = False
            elif node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
                tail_has_break = True
        if tail_has_break:
            return True
        props = previous._p.find(qn("w:pPr"))
        sec = None if props is None else props.find(qn("w:sectPr"))
        if sec is not None:
            try:
                current = ctx.document.sections[ctx.section_index(paragraph)]
                if int(current.start_type) in {2, 3, 4}:  # next / even / odd page section breaks
                    return True
            except Exception:
                return None
        if previous.text.strip():
            break
    return False


def _usable_width(section):
    names = ("page_width", "left_margin", "right_margin", "gutter", "gutter_position")
    if any(section.get(name) is None for name in names):
        return None
    gutter = 0 if section["gutter_position"] == "上" else section["gutter"]
    return section["page_width"] - section["left_margin"] - section["right_margin"] - gutter


def _template_layout_exceptions(ctx, sections):
    """Do not invent a margin ban when the template itself demonstrates an exception."""
    for section in sections:
        section["layout_exceptions"] = set()
    for paragraph in ctx.paragraphs:
        try:
            section = sections[ctx.section_index(paragraph)]
            available = _usable_width(section)
            if available is None:
                continue
            for inline in paragraph._p.xpath(".//wp:inline"):
                extent = inline.find(qn("wp:extent"))
                if extent is not None and int(extent.get("cx")) > available:
                    section["layout_exceptions"].add("image")
        except Exception:
            for section in sections:
                section["layout_exceptions"].add("unreadable")
    try:
        for table in ctx.document.tables:
            anchor = table.cell(0, 0).paragraphs[0]
            section = sections[ctx.section_index(anchor)]
            available = _usable_width(section)
            widths = [column.width for column in table.columns]
            if available is not None and widths and all(width is not None for width in widths):
                if sum(widths) > available:
                    section["layout_exceptions"].add("table")
    except Exception:
        for section in sections:
            section["layout_exceptions"].add("unreadable")
    # Keep the extracted namespace suitable for logging/serialization.
    for section in sections:
        section["layout_exceptions"] = tuple(sorted(section["layout_exceptions"]))


def extract_page_baseline(ctx):
    """Return a namespaced addition; existing baseline keys are never replaced."""
    result = {"sections": [], "chapter_breaks": {}, "errors": []}
    try:
        sections = list(ctx.document.sections)
        grouped = _section_paragraphs(ctx)
        result["sections"] = [_read_section(ctx, sections, i, grouped[i]) for i in range(len(sections))]
        _template_layout_exceptions(ctx, result["sections"])
        if grouped.get(None):
            result["errors"].append("部分模板段落无法定位所属节")
    except Exception as exc:
        result["errors"].append("模板分节和节设置无法读取：" + str(exc))
    # 按出现顺序记录每个级别的分页机制；判定时按“同位置”取对应标题的样例。
    observations = defaultdict(list)
    for p in ctx.paragraphs:
        level = _heading_level(p)
        if level is not None:
            try:
                forced = _has_forced_break(ctx, p)
                if forced is not None:
                    observations[level].append(forced)
            except Exception as exc:
                result["errors"].append("模板章节分页设置无法读取：" + str(exc))
    result["chapter_breaks"] = {level: tuple(values) for level, values in observations.items()}
    return {"pages": result}


def _mapped_index(mapping, index):
    raw = (mapping or {}).get("section_map", {})
    if not isinstance(raw, dict):
        return None, "section_map 不是节映射字典"
    value = raw.get(str(index + 1), raw.get(index + 1))
    if value is None:
        return None, None
    try:
        return int(value) - 1, None
    except (TypeError, ValueError):
        return None, "section_map 中的模板节号不是整数"


def _candidate_sections(template_sections, actual, mapping):
    explicit, error = _mapped_index(mapping, actual["index"])
    if error:
        return [], error
    if explicit is not None:
        if 0 <= explicit < len(template_sections):
            return [template_sections[explicit]], None
        return [], "section_map 指向不存在的模板节"
    roles = set(actual["roles"])
    if roles:
        same_role = [s for s in template_sections if set(s["roles"]) == roles]
        if same_role:
            return same_role, None
        # 论文常把封面+声明、摘要+目录合并成更少的节，节角色集合不会完全相等：
        # 退一步取“角色重叠最多”的模板节（例如摘要/目录节对应模板的摘要、目录节），
        # 而不是要求用户手写 section_map。
        scored = [(len(roles & set(s["roles"])), s) for s in template_sections]
        best = max((score for score, _ in scored), default=0)
        if best:
            return [s for score, s in scored if score == best], None
    # Multiple sections may still share the same attribute; compare only that consensus.
    return template_sections, None


#: 距离类属性：同一长度用不同单位写出时会有亚毫米级误差，按容差比较。
_DISTANCE_KEYS = {"header_distance", "footer_distance", "gutter"}
_DISTANCE_TOLERANCE_EMU = 7200  # 0.2 毫米


def auto_section_map(template_ctx: DocxContext, document_ctx: DocxContext) -> dict[str, int]:
    """按“节角色”自动对应模板节与论文节，返回 ``{论文节号: 模板节号}``（均从 1 起）。

    论文与模板的分节数通常不同（样例模板 11 节、论文 5 节）：封面/声明/摘要/目录
    可能被合并或拆分。这里按节内文字推断节角色（封面、声明、摘要、目录、正文），
    取“角色完全一致，否则角色重叠最多且唯一”的模板节作为该节的基准，从而不再要求
    用户手写 ``section_map``。无法唯一确定时不下结论，交由各规则按模板同角色共识判断。
    """

    try:
        template_sections = list(template_ctx.document.sections)
        document_sections = list(document_ctx.document.sections)
        template_groups = _section_paragraphs(template_ctx)
        document_groups = _section_paragraphs(document_ctx)
        template_data = [_read_section(template_ctx, template_sections, index,
                                       template_groups.get(index, []))
                         for index in range(len(template_sections))]
        document_data = [_read_section(document_ctx, document_sections, index,
                                       document_groups.get(index, []))
                         for index in range(len(document_sections))]
    except Exception:
        return {}
    mapping: dict[str, int] = {}
    for actual in document_data:
        roles = set(actual["roles"])
        if not roles:
            continue
        scored = [(len(roles & set(section["roles"])), section) for section in template_data]
        best = max((score for score, _ in scored), default=0)
        winners = {section["index"] for score, section in scored if score == best and score > 0}
        if len(winners) == 1:
            mapping[str(actual["index"] + 1)] = winners.pop() + 1
    return mapping


def _header_signature(section) -> tuple:
    stories = section.get("stories") or {}
    values = []
    for key in ("header:default", "header:first", "header:even"):
        story = stories.get(key)
        values.append((key, None if story is None else (
            bool(story.get("present")), bool(story.get("fields")),
            (story.get("text") or "").strip(),
            tuple(story.get("objects") or ()),
        )))
    return (
        bool(section.get("different_first")),
        bool(section.get("different_odd_even")),
        tuple(values),
    )


def auto_header_section_map(template_ctx: DocxContext,
                            document_ctx: DocxContext) -> dict[str, int]:
    """为 Word 页眉文字比较补充不影响其他节属性的安全映射。"""

    try:
        template_sections = list(template_ctx.document.sections)
        document_sections = list(document_ctx.document.sections)
        template_groups = _section_paragraphs(template_ctx)
        document_groups = _section_paragraphs(document_ctx)
        template_data = [_read_section(
            template_ctx, template_sections, index, template_groups.get(index, []))
            for index in range(len(template_sections))]
        document_data = [_read_section(
            document_ctx, document_sections, index, document_groups.get(index, []))
            for index in range(len(document_sections))]
    except Exception:
        return {}
    mapping: dict[str, int] = {}
    for actual in document_data:
        roles = set(actual["roles"])
        if not roles:
            continue
        candidates = [section for section in template_data
                      if set(section["roles"]) == roles]
        if not candidates:
            scored = [(len(roles & set(section["roles"])), section)
                      for section in template_data]
            best = max((score for score, _ in scored), default=0)
            candidates = [section for score, section in scored
                          if score == best and score > 0]
        if candidates and len({_header_signature(section) for section in candidates}) == 1:
            mapping[str(actual["index"] + 1)] = candidates[0]["index"] + 1
    return mapping


def _equal(name, expected, current) -> bool:
    if name in _DISTANCE_KEYS and isinstance(expected, int) and isinstance(current, int):
        return abs(expected - current) <= _DISTANCE_TOLERANCE_EMU
    return expected == current


def _consensus(candidates, name):
    values = [s.get(name) for s in candidates]
    if not values or any(value is None for value in values):
        return None, "模板对应属性未声明或读取失败"
    if any(not _equal(name, values[0], value) for value in values[1:]):
        return None, "模板存在不同节基准且无法唯一匹配；请核对区域或提供 section_map"
    return values[0], None


def _value(name, value):
    if name == "orientation":
        return "横向" if value == 1 else "纵向" if value == 0 else str(value)
    if name in {"gutter", "header_distance", "footer_distance"}:
        return f"{value / 36000:.3f} 毫米"
    if isinstance(value, bool):
        return "开启" if value else "关闭"
    return str(value)


def _story_text(value) -> str:
    """页眉/页脚内容的可读说明：不把 (文字, 是否有对象) 这样的原始值写进批注。"""

    if value is None:
        return "含动态域或无法读取"
    text, objects = value
    if text and objects:
        return f"文字「{text}」以及图片/其他对象"
    if text:
        return f"文字「{text}」"
    if objects:
        return "图片或其他对象"
    return "无内容"


def _limited(rule, reason, paragraph):
     return unknown(rule, _title(rule), UNREADABLE + "；" + reason, paragraph=paragraph)


def _title(code: str) -> str:
    """细粒度规则标题取自规则目录；传入族编号时用本模块的族标题。"""

    item = FINE_RULES.get(code)
    return item.title if item else _TITLES[code]


def _check_front(candidates, actual, paragraph):
    from thesis_format_checker.detectors.headers import _StoryFindings
    findings = _StoryFindings()
    findings.story = "footer"
    front = {"封面", "声明", "中文摘要", "英文摘要", "目录"}
    if not set(actual["roles"]) & front:
        return findings
    if not candidates:
        return [_limited("022", "模板没有对应的前置节，无法比对页眉页脚", paragraph)]
    expected = candidates[0]
    # 022.3：模板前置部分使用罗马数字页码时，论文前置部分不得改用阿拉伯数字。
    expected_numbering = _page_number_format(expected)
    actual_numbering = _page_number_format(actual)
    if expected_numbering == "upperroman" and actual_numbering != "upperroman":
        findings.append(make_finding(
            "022.3", _title("022.3"), "前置部分页码应使用大写罗马数字",
            "当前该节页码格式为"
            + _NUMBER_FORMAT_LABELS.get(actual_numbering, actual_numbering),
            paragraph))
    # 逐变体（普通页/首页/偶数页）比较：前置部分该有的页眉页脚不会被“变体不同”打包成无法判定，
    # 而是逐个变体给出确定结论；含 PAGE 域的变体只比结构（有无动态域、静态文字），不算域值。
    for kind, label in (("header", "页眉"), ("footer", "页脚")):
        findings.story, findings.variant = kind, None
        code = _FRONT_CODES[kind]
        if kind == "footer":
            # 前置部分按模板使用罗马页码是正常的：页脚只比“区域级配置”
            # （有没有内容、是不是动态页码域、固定文字是什么），
            # 不因为模板把页码放在偶数页变体、论文放在普通页变体就报“不应出现”。
            expected_plan = _front_story_plan(expected, kind)
            actual_plan = _front_story_plan(actual, kind)
            if expected_plan != actual_plan:
                findings.append(make_finding(
                    code, _title(code),
                    "模板前置部分" + label + "：" + _story_plan_text(expected_plan),
                    "论文前置部分" + label + "：" + _story_plan_text(actual_plan),
                    paragraph))
            continue
        for variant, variant_label in (("default", "普通页"), ("first", "首页"), ("even", "偶数页")):
            findings.variant = variant
            expected_story = _story_of(expected, kind, variant)
            actual_story = _story_of(actual, kind, variant)
            where = f"前置部分{variant_label}{label}"
            if expected_story is None or actual_story is None:
                findings.append(_limited(code, f"{where}部件或继承关系读取失败", paragraph))
                continue
            expected_present = bool(expected_story["present"])
            actual_present = bool(actual_story["present"])
            if expected_present == actual_present:
                if expected_present and bool(expected_story["fields"]) != bool(actual_story["fields"]):
                    findings.append(make_finding(
                        code, _title(code),
                        f"模板{where}为" + ("动态页码/字段" if expected_story["fields"] else "固定文字"),
                        f"论文{where}为" + ("动态页码/字段" if actual_story["fields"] else "固定文字"),
                        paragraph))
                elif (expected_present and not expected_story["fields"] and not actual_story["fields"]
                      and (expected_story["text"] or "").strip() != (actual_story["text"] or "").strip()):
                    findings.append(make_finding(
                        code, _title(code),
                        f"模板{where}文字为「{(expected_story['text'] or '无内容').strip()[:30]}」",
                        f"论文{where}文字为「{(actual_story['text'] or '无内容').strip()[:30]}」",
                        paragraph))
                continue
            if actual_present:
                findings.append(make_finding(code, _title(code), f"模板{where}无内容（不应出现）",
                                             f"论文{where}出现了" + _story_text(_story_value(actual_story)),
                                             paragraph))
            else:
                findings.append(make_finding(code, _title(code),
                                             f"模板{where}为" + _story_text(_story_value(expected_story)),
                                             f"论文{where}没有内容", paragraph))
    return findings


def _story_of(section, kind: str, variant: str):
    """取某节某页眉/页脚变体的读取结果；文档未启用该变体时用普通页变体代替。"""

    stories = section.get("stories") or {}
    order = [kind + ":" + variant]
    if variant == "first" and not section.get("different_first"):
        order = [kind + ":default", kind + ":first"]
    elif variant == "even" and not section.get("different_odd_even"):
        order = [kind + ":default", kind + ":even"]
    for key in order:
        if key in stories:
            return stories[key]
    return stories.get(kind + ":default")


def _front_story_plan(section, kind: str):
    """前置部分页眉/页脚的区域级配置：``(有内容, 是否动态页码域, 固定文字)``。"""

    stories = [story for key, story in (section.get("stories") or {}).items()
               if key.startswith(kind + ":")]
    present = any(story and story["present"] for story in stories)
    fields = any(story and story["fields"] for story in stories)
    texts = tuple(sorted({(story.get("text") or "").strip() for story in stories
                          if story and story["present"] and not story["fields"]} - {""}))
    return present, fields, texts


def _story_plan_text(plan) -> str:
    present, fields, texts = plan
    if not present:
        return "没有内容"
    if fields and texts:
        return "含动态页码/字段，另有文字" + "、".join(f"「{item[:20]}」" for item in texts)
    if fields:
        return "含动态页码/字段（各页自动编号）"
    return "固定文字" + "、".join(f"「{item[:20]}」" for item in texts) if texts else "有内容但无可读文字"


def _story_value(story):
    """把页眉读取结果转成 ``_story_text`` 需要的 (文字, 是否有对象) 形式。"""

    if story is None:
        return None
    text = (story.get("text") or "").strip()
    if story.get("fields"):
        text = "动态页码/字段" if not text else f"{text}（含动态页码/字段）"
    return (text, bool(story.get("objects")))


def _transition_signature(section):
    # No paper-size/margin comparison is implemented here; those remain existing rules.
    keys = ("orientation", "gutter", "gutter_position", "header_distance", "footer_distance",
            "different_first", "different_odd_even", "page_number_settings")
    signature = {key: section.get(key) for key in keys}
    # Static stories can be compared without evaluating fields or rendering pages.
    for key, story in section.get("stories", {}).items():
        if key.endswith(":first") and not section.get("different_first"):
            story = section.get("stories", {}).get(key.split(":")[0] + ":default")
        if key.endswith(":even") and not section.get("different_odd_even"):
            # Effective even-page content is the ordinary story when disabled.
            story = section.get("stories", {}).get(key.split(":")[0] + ":default")
        signature[key] = (story["text"], story["objects"]) if story is not None and not story["fields"] else None
    return signature


def _check_transitions(template_sections, actual_sections, matched, grouped):
    result = []
    for index in range(1, len(actual_sections)):
        paragraph = _section_anchor(grouped[index])
        left, right = matched[index - 1], matched[index]
        if len(left) != 1 or len(right) != 1 or left[0]["index"] + 1 != right[0]["index"]:
            if len(template_sections) > 1 and (len(left) != 1 or len(right) != 1):
                result.append(_limited("029", "分节边界无法唯一对应模板的前后区域；不因分节符数量不同判错", paragraph))
            continue  # Extra breaks and equal section counts do not prove a correspondence.
        expected_before, expected_after = map(_transition_signature, (left[0], right[0]))
        actual_before, actual_after = map(_transition_signature, actual_sections[index - 1:index + 1])
        differences = [key for key in expected_after
                       if expected_after[key] is not None and actual_after.get(key) is not None
                       and not (key.endswith(":first") and not actual_sections[index].get("different_first"))
                       and expected_before.get(key) == actual_before.get(key)
                       and actual_before.get(key) != actual_after[key] and expected_after[key] != actual_after[key]]
        if differences:
            # 文案只写“哪一项、从什么变成什么”，不暴露 True/False、元组等内部取值。
            def describe(section, key):
                if ":" in key:
                    return _story_text(section[key])
                return _value(key, section[key])

            def label_of(key):
                if ":" not in key:
                    return _LABELS.get(key, "页码设置")
                kind = "页眉" if key.startswith("header") else "页脚"
                variant = {"default": "", "first": "（首页）", "even": "（偶数页）"}.get(key.split(":", 1)[1], "")
                return kind + variant + "内容"

            # Each story variant has its own visible pages. Never aggregate it
            # into a section-start body annotation (possibly a declaration).
            groups = [[k] for k in differences if ":" in k]
            distances = [k for k in differences if k in ("header_distance", "footer_distance")]
            groups.extend([k] for k in distances)
            other = [k for k in differences if ":" not in k and k not in distances]
            if other:
                groups.append(other)
            for keys in groups:
                finding = make_finding(
                    "029", _TITLES["029"],
                    "本处分节后应与前一节保持一致（模板此处：" +
                    "；".join(f"{label_of(k)}为{describe(expected_after, k)}" for k in keys) + "）",
                    "本处分节后实际变为：" +
                    "；".join(f"{label_of(k)}变为{describe(actual_after, k)}" for k in keys),
                    paragraph=None if ":" in keys[0] else paragraph)
                if ":" in keys[0]:
                    finding.page_story, finding.page_variant = keys[0].split(":")
                    finding.page_scope = "variant"
                    finding.section_index = index + 1
                elif keys[0] in distances:
                    finding.anchor = None
                    finding.page_story = keys[0].split("_")[0]
                    finding.section_index = index + 1
                    finding.page_issue = f"{keys[0]}:{expected_after[keys[0]]}:{actual_after[keys[0]]}"
                result.append(finding)
                if ":" in keys[0]:
                    # A linked story can first become visible in a later
                    # section. Follow unchanged content only while that
                    # section's own template correspondence requires the same.
                    from copy import copy
                    key = keys[0]
                    for later in range(index + 1, len(actual_sections)):
                        if _transition_signature(actual_sections[later]).get(key) != actual_after[key]:
                            break
                        references = matched[later]
                        if not references or any(_transition_signature(s).get(key) != expected_after[key]
                                                 for s in references):
                            break
                        inherited = copy(finding)
                        inherited.section_index = later + 1
                        result.append(inherited)
        if any(s["errors"] or any(story is None or story["fields"] for story in s.get("stories", {}).values())
               for s in (left[0], right[0], actual_sections[index - 1], actual_sections[index])):
            result.append(_limited("029", "分节相关属性读取失败或存在动态页眉页脚域，部分分节影响无法判定", paragraph))
        elif not differences and right[0].get("start_type") != actual_sections[index].get("start_type"):
            result.append(_limited("029", "对应分节符类型与模板不同；实际分页及页眉页脚影响需要 Word 排版确认，不能仅凭类型差异判错", paragraph))
    return result


def _check_chapter_breaks(baseline, ctx):
    """030：同级标题按“同一层级 + 同一位置”与模板对应标题比较分页机制。

    模板里不同章节可能采用不同机制（例如结论章强制另起页、其他章自然分页），
    所以不再要求整级的机制唯一，而是各标题比各自对应位置的模板标题。
    """

    result = []
    groups = defaultdict(list)
    for p in ctx.paragraphs:
        level = _heading_level(p)
        if level is not None:
            groups[level].append(p)
    for level, paragraphs in groups.items():
        values = baseline.get("chapter_breaks", {}).get(level)
        if not values:
            continue
        for ordinal, p in enumerate(paragraphs):
            expected = sample_at(values, ordinal, len(paragraphs))
            if not expected:
                # 模板中对应位置的同级标题不强制另起页：本条不作要求。
                continue
            try:
                if _has_forced_break(ctx, p):
                    continue
                index = next(i for i, other in enumerate(ctx.paragraphs) if other._p is p._p)
                blanks = 0
                for previous in reversed(ctx.paragraphs[:index]):
                    if previous.text.strip() or previous._p.xpath(".//w:drawing | .//w:pict | .//w:object"):
                        break
                    blanks += 1
                reason = (f"模板对应位置的 {level} 级标题设置强制另起页；当前未发现等效分页机制，"
                          f"标题前有 {blanks} 个空段。"
                          "docx 库无法确定自然分页结果，不能认定同页或认定空段一定用于模拟分页")
                # 有 Word 分页事实时，030 已由 word_effective.pagination 按实际分页给出确定结论，
                # 这里不再重复输出“无法判定”。
                if not getattr(ctx, "render_facts", None):
                    result.append(_limited("030", reason, p))
            except Exception as exc:
                result.append(_limited("030", "章节分页或分节符读取失败：" + str(exc), p))
    return result


def _table_grid_issue(table):
    if any(node.tag == qn("w:tc") for node in table._tbl.iterancestors()):
        return "嵌套表格需要外围单元格的实际布局，不能直接按页面宽度判断"
    columns = [column.width for column in table.columns]
    if not columns or any(width is None for width in columns):
        return "表格网格列宽无法完整读取"
    for row in table._tbl.findall(qn("w:tr")):
        row_props = row.find(qn("w:trPr"))
        if row_props is not None and any(row_props.find(qn("w:" + tag)) is not None for tag in ("gridBefore", "gridAfter")):
            return "表格存在跳过的网格列，无法可靠按网格总宽判断布局"
        column = 0
        for cell in row.findall(qn("w:tc")):
            props = cell.find(qn("w:tcPr"))
            span = None if props is None else props.find(qn("w:gridSpan"))
            span_value = 1 if span is None else int(span.get(qn("w:val"), "1"))
            cell_width = None if props is None else props.find(qn("w:tcW"))
            if (span_value < 1 or column + span_value > len(columns) or cell_width is None
                    or cell_width.get(qn("w:type")) != "dxa" or cell_width.get(qn("w:w")) is None):
                return "表格单元格宽度、合并范围或网格无法可靠对应"
            expected_twips = sum(int(width) for width in columns[column:column + span_value]) // 635
            if int(cell_width.get(qn("w:w"))) != expected_twips:
                return "表格单元格宽度与网格列宽冲突，需要 Word 排版确定最终宽度"
            column += span_value
        if column != len(columns):
            return "表格每行单元格与网格列数不一致，无法可靠确定固定布局宽度"
    return None


#: 封面常用字段：封面普遍用表格排版，表格文字里出现这些字段即可认定它是封面表格。
_COVER_FIELDS = re.compile(r"(论文题目|题目|题名|作者姓名|姓名|学号|指导教师|学院|专业|班级|图书分类号|UDC|密级)")
#: 封面只可能出现在文档前部；超过该段数不再按封面表格处理，避免误豁免正文表格。
_COVER_SCAN_LIMIT = 120


def _enclosing_table(paragraph):
    """返回段落所属的 ``w:tbl`` 元素（表格套表格时取最内层），不在表格内返回 None。"""

    for node in paragraph._p.iterancestors():
        if node.tag == qn("w:tbl"):
            return node
    return None


def _inside_cover_table(ctx, paragraph) -> bool:
    """判断段落是否位于封面表格内。

    封面通常由“论文题目 / 姓名 / 学号 / 学院 / 指导教师”这类字段排成表格，
    其中的校徽、校名图片和字段位置由封面基线（规则 017）按“与模板一致”比对。
    033 只能按页面宽度推断裁切风险，对单元格内的对象没有可靠基准，
    因此封面表格整体交给 017，不再重复输出“无法判定”。
    """

    if ctx.index(paragraph) >= _COVER_SCAN_LIMIT:
        return False
    table = _enclosing_table(paragraph)
    if table is None:
        return False
    text = "".join(node.text or "" for node in table.iter(qn("w:t")))
    return bool(_COVER_FIELDS.search(text))


def _check_widths(ctx, actual_sections, matched, grouped):
    result = []
    rendered_limited = False
    for p in ctx.paragraphs:
        try:
            index = ctx.section_index(p)
            candidates = matched[index]
            widths = [_usable_width(s) for s in candidates]
            available = widths[0] if widths and None not in widths and len(set(widths)) == 1 else None
            inline = p._p.xpath(".//wp:inline")
            unsupported = p._p.xpath(".//wp:anchor | .//w:pict | .//w:object")
            if unsupported:
                result.append(_limited("033", "浮动图形、文本框或 OLE 对象的可见边界需要 Word 渲染", p))
            if not inline:
                if p.text.strip():
                    rendered_limited = True
                continue
            if any(set(s.get("layout_exceptions", ())) & {"image", "unreadable"} for s in candidates):
                result.append(_limited("033", "模板含超出内容区域的图片示例或布局读取失败，无法推定所有图片必须限制在正文边界内", p))
                continue
            if available is None or available <= 0:
                result.append(_limited("033", "模板内容区域宽度无法唯一读取或匹配，未判断图片是否越界", p))
                continue
            # A paragraph inside a table needs cell layout, which is not a page-wide image.
            if any(ancestor.tag == qn("w:tc") for ancestor in p._p.iterancestors()):
                # 封面表格里的校徽、校名等图片由 017（封面图片）规则按“与模板一致”比对，
                # 这里再按页面宽度猜一次只会产生“无法判定”噪音，所以直接跳过。
                if _inside_cover_table(ctx, p):
                    continue
                result.append(_limited("033", "表格单元格中的图片需结合合并单元格、边距和实际列布局，无法据页面宽度完成判定", p))
                continue
            effective = effective_paragraph(p)
            left = effective.get("left_indent")
            right = effective.get("right_indent")
            if left is None or right is None:
                result.append(_limited("033", "图片所在段落的有效左右缩进无法换算为绝对长度，未将未知缩进当作零", p))
                continue
            # Common effective paragraph values use points; OOXML widths use EMU.
            available -= max(float(left), 0) * 12700 + max(float(right), 0) * 12700
            for drawing in inline:
                extent = drawing.find(qn("wp:extent"))
                if extent is None or extent.get("cx") is None:
                    result.append(_limited("033", "嵌入型图片的显示宽度未能读取", p))
                    continue
                width = int(extent.get("cx"))
                if width > available:
                    result.append(make_finding(_WIDTH_RULE, _title(_WIDTH_RULE), f"模板内容区域结合当前段落缩进后的可用宽度 {available / 36000:.3f} 毫米",
                                               f"嵌入型图片宽度 {width / 36000:.3f} 毫米，超出可用内容宽度；实际裁切需打印预览确认", paragraph=p))
        except Exception as exc:
            result.append(_limited("033", "对象尺寸或所在节读取失败：" + str(exc), p))
    try:
        from docx.table import Table
        tables = [Table(element, ctx.document._body) for element in ctx.document.element.body.iter(qn("w:tbl"))
                  if not any(parent.tag == qn("w:txbxContent") for parent in element.iterancestors())]
    except Exception as exc:
        tables = []
        result.append(_limited("033", "表格集合读取失败：" + str(exc), next(iter(ctx.paragraphs), None)))
    for table in tables:
        anchor = None
        try:
            anchor = next((p for cell in table.rows[0].cells for p in cell.paragraphs), None) if table.rows else None
            if anchor is None:
                continue
            # 封面表格的布局由封面基线（017）负责，不在这里输出“无法判定”。
            if _inside_cover_table(ctx, anchor):
                continue
            candidates = matched[ctx.section_index(anchor)]
            widths = [_usable_width(s) for s in candidates]
            available = widths[0] if widths and None not in widths and len(set(widths)) == 1 else None
            if any(set(s.get("layout_exceptions", ())) & {"table", "unreadable"} for s in candidates):
                result.append(_limited("033", "模板含超出内容区域的表格示例或布局读取失败，无法推定所有表格必须限制在正文边界内", anchor))
                continue
            if table.autofit or available is None:
                result.append(_limited("033", "自动调整表格的最终宽度或对应模板内容区域无法可靠确定", anchor))
                continue
            props = table._tbl.tblPr
            if props is not None and props.find(qn("w:tblpPr")) is not None:
                result.append(_limited("033", "浮动表格位置须由 Word 排版确认", anchor))
                continue
            grid_issue = _table_grid_issue(table)
            if grid_issue:
                result.append(_limited("033", grid_issue, anchor))
                continue
            columns = [column.width for column in table.columns]
            if not columns or any(width is None for width in columns):
                result.append(_limited("033", "表格网格列宽无法完整读取", anchor))
                continue
            width = sum(int(value) for value in columns)
            indent = 0
            ind_node = props.find(qn("w:tblInd")) if props is not None else None
            if ind_node is not None:
                if ind_node.get(qn("w:type"), "dxa") != "dxa":
                    result.append(_limited("033", "表格缩进不是可读取的绝对长度，无法确定横向边界", anchor))
                    continue
                indent = int(ind_node.get(qn("w:w"), "0")) * 635
            if width + max(indent, 0) > available:
                result.append(make_finding(_WIDTH_RULE, _title(_WIDTH_RULE), f"模板该区域可用宽度 {available / 36000:.3f} 毫米",
                                           f"固定布局表格网格总宽 {width / 36000:.3f} 毫米、正向缩进 {max(indent, 0) / 36000:.3f} 毫米，超出内容宽度；实际裁切需打印预览确认", paragraph=anchor))
        except Exception as exc:
            result.append(_limited("033", "表格宽度或分节归属读取失败：" + str(exc), anchor))
    if rendered_limited:
        anchor = next((p for p in ctx.paragraphs if p.text.strip()), None)
        result.append(_limited("033", "普通文字的实际行宽、纵向分页、叠放及打印机硬件可打印区域需要 Word 渲染；本条已检测可读取的对象宽度，无法据此保证无裁切", anchor))
    return result


def check_page_errors(baseline, ctx, mapping=None):
    """Collect findings without altering source paragraphs, styles, stories or fields."""
    stored = baseline.get("pages", baseline)
    expected_sections = stored.get("sections", [])
    result = []
    first = next(iter(ctx.paragraphs), None)
    if not expected_sections:
        return [_limited(rule, "模板分节和页面基准无法读取：" + "; ".join(stored.get("errors", [])), first)
                for rule in ("022", "024", "026", "027", "029", "030", "031", "033")]
    try:
        sections = list(ctx.document.sections)
        grouped = _section_paragraphs(ctx)
        actual_sections = [_read_section(ctx, sections, i, grouped[i]) for i in range(len(sections))]
    except Exception as exc:
        return [_limited(rule, "待检测文档分节和节设置读取失败：" + str(exc), first)
                for rule in ("022", "024", "026", "027", "029", "030", "031", "033")]
    matched = []
    for index, actual in enumerate(actual_sections):
        paragraph = _section_anchor(grouped[index], first)
        candidates, mapping_error = _candidate_sections(expected_sections, actual, mapping)
        matched.append(candidates)
        for rule, names in _GROUPS.items():
            errors: list[tuple[str, str]] = []
            for name in names:
                code = _GROUP_CODES.get(name, rule)
                expected, error = _consensus(candidates, name)
                current = actual.get(name)
                if mapping_error or error or current is None:
                    errors.append((code, f"{_LABELS[name]}：{mapping_error or error or actual['errors'].get(name, '有效值读取失败')}"))
                    continue
                if not _equal(name, expected, current):
                    # 一个设置项一条结论：装订线位置与宽度、页眉与页脚距离分别成条
                    result.append(make_finding(code, _title(code),
                                               f"{_LABELS[name]}应为{_value(name, expected)}",
                                               f"{_LABELS[name]}当前为{_value(name, current)}",
                                               paragraph=paragraph))
                    if name in {"header_distance", "footer_distance"}:
                        result[-1].page_story = name.split("_")[0]
                        result[-1].page_issue = f"{name}:{expected}:{current}"
                        result[-1].section_index = index + 1
                    elif name in {"different_first", "different_odd_even"}:
                        finding = result[-1]
                        finding.page_story = "header"
                        finding.section_index = index + 1
                        if name == "different_odd_even":
                            finding.page_scope = "even_pages"
                            finding.page_skip_first = bool(actual["different_first"] or
                                                           any(c.get("different_first") for c in candidates))
                            finding.page_config = f"odd_even:{int(expected)}:{int(current)}"
                        else:
                            finding.page_scope = "first_page"
                            finding.page_config = f"first:{index}:{int(expected)}:{int(current)}"
            for code, reason in errors:
                result.append(_limited(code, reason, paragraph))
        front_findings = _check_front(candidates, actual, paragraph)
        for finding in front_findings:
            finding.section_index = index + 1
        result.extend(front_findings)
    if grouped.get(None):
        result.append(_limited("029", "部分段落无法定位所属节，未猜测节对应关系", grouped[None][0]))
    result.extend(_check_transitions(expected_sections, actual_sections, matched, grouped))
    result.extend(_check_chapter_breaks(stored, ctx))
    result.extend(_check_widths(ctx, actual_sections, matched, grouped))
    return result
